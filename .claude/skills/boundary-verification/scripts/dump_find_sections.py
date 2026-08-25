#!/usr/bin/env python3
"""find_sections()가 실제 참조 PPT에서 각 섹션을 어느 슬라이드로 탐지했는지 덤프한다.

mass-template-analysis의 dump_slide_text.py로 뽑은 "실제 슬라이드 텍스트" 목록과
이 스크립트의 출력(find_sections()가 "여기다"라고 판단한 인덱스)을 나란히 놓고 비교하면
섹션 탐지 로직이 새 템플릿에서도 맞는 슬라이드를 가리키는지 바로 확인할 수 있다.
missa_to_ppt.py를 import하므로 프로젝트 루트에서 실행해야 한다.

사용법:
    python dump_find_sections.py <참조PPT경로>
"""
import sys

from pptx import Presentation

sys.path.insert(0, '.')
import missa_to_ppt as m  # noqa: E402


def main():
    if len(sys.argv) != 2:
        print('사용법: python dump_find_sections.py <참조PPT경로>')
        return 1
    path = sys.argv[1]
    prs = Presentation(path)
    sections = m.find_sections(prs)
    for key in sorted(sections.keys()):
        val = sections[key]
        if isinstance(val, int):
            slide = prs.slides[val]
            texts = [
                sh.text_frame.text.strip().replace('\n', ' / ')[:50]
                for sh in slide.shapes
                if getattr(sh, 'has_text_frame', False) and sh.text_frame.text.strip()
            ]
            print(f'{key}\t=> slide {val}\t{" | ".join(texts)[:120]}')
        else:
            print(f'{key}\t=> {val}')
    return 0


if __name__ == '__main__':
    sys.exit(main())
