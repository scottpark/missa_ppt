#!/usr/bin/env python3
"""참조 PPT의 슬라이드별 shape 텍스트를 인덱스와 함께 덤프한다.

새 미사 유형/성당의 템플릿을 처음 조사할 때, 또는 find_sections()의 섹션 탐지 키가
실제로 어떤 슬라이드를 가리키는지 확인할 때 사용한다. missa_to_ppt.py를 import하지 않는
독립 스크립트이므로 어떤 pptx 파일에도 안전하게 쓸 수 있다.

사용법:
    python dump_slide_text.py <pptx경로> [--max-chars 40] [--start N] [--end N]
"""
import argparse
import sys

from pptx import Presentation


def dump(path: str, max_chars: int, start: int, end: int):
    prs = Presentation(path)
    slides = list(prs.slides)
    n = len(slides)
    end = n if end is None else min(end, n)
    print(f'# {path}')
    print(f'# 전체 슬라이드 수: {n}, 슬라이드 크기: {prs.slide_width}x{prs.slide_height} EMU')
    print()
    for i in range(start, end):
        slide = slides[i]
        texts = []
        for sh in slide.shapes:
            if getattr(sh, 'has_text_frame', False) and sh.text_frame.text.strip():
                t = sh.text_frame.text.strip().replace('\n', ' / ')
                texts.append(t[:max_chars])
        line = ' | '.join(texts)
        print(f'{i}\t{line[:200]}')


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('pptx')
    ap.add_argument('--max-chars', type=int, default=40)
    ap.add_argument('--start', type=int, default=0)
    ap.add_argument('--end', type=int, default=None)
    args = ap.parse_args()
    dump(args.pptx, args.max_chars, args.start, args.end)


if __name__ == '__main__':
    sys.exit(main())
