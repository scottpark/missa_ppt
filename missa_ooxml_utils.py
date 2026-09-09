"""
missa_to_ppt.py 모듈 분리 리팩토링 Phase 1 — 슬라이드/도형 범용 프리미티브.

python-pptx/lxml로 OOXML(spTree, rels, p:bg, a:pPr 등)을 직접 조작하는 저수준 유틸리티.
프로젝트 내부의 다른 모듈에 의존하지 않는 leaf 모듈이다.
"""
from __future__ import annotations

import atexit
import copy
import os
import tempfile
import uuid
from pathlib import Path
from xml.sax.saxutils import escape as _xml_escape

from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.parts.presentation import PresentationPart

HYMN_TYPES = ['입당', '봉헌', '성체', '2차봉헌', '파견']
# entry(missa_to_ppt.py main())와 content_updaters(replace_성가) 양쪽에서 쓰여서
# 여기(leaf 모듈)에 둔다 — entry→content_updaters 단방향 의존만 유지하기 위함.


# python-pptx의 _next_slide_partname은 len(sldIdLst)+1을 사용하여
# 삭제 후 add_slide 시 기존 파트명과 충돌이 발생한다.
# 패키지 레벨의 next_partname을 사용하도록 패치.
@property  # type: ignore[misc]
def _safe_next_slide_partname(self):

    return self.package.next_partname('/ppt/slides/slide%d.xml')



PresentationPart._next_slide_partname = _safe_next_slide_partname


# ─────────────────────────────────────────────────────────────────────────────
# 슬라이드 XML 유틸
# ─────────────────────────────────────────────────────────────────────────────


def _slide_text(slide) -> str:

    parts = []

    for shape in slide.shapes:

        if shape.has_text_frame:

            parts.append(shape.text_frame.text)

    return ' '.join(parts)


def all_slide_texts(prs) -> list:

    return [_slide_text(s) for s in prs.slides]


def delete_slide(prs, idx: int):

    """슬라이드 삭제 (0-based 인덱스)."""

    sldIdLst = prs.slides._sldIdLst

    sldId = sldIdLst[idx]

    rId = sldId.get(qn('r:id'))

    sldIdLst.remove(sldId)

    try:

        prs.part.drop_rel(rId)

    except Exception:

        pass


def move_slide(prs, old_idx: int, new_idx: int):

    sldIdLst = prs.slides._sldIdLst

    el = sldIdLst[old_idx]

    sldIdLst.remove(el)

    sldIdLst.insert(new_idx, el)


def _blank_layout(prs):

    """빈 레이아웃 반환 (이름 우선, 없으면 마지막 레이아웃)."""

    for layout in prs.slide_layouts:

        if layout.name in ('Blank', '빈 화면', 'blank'):

            return layout

    return prs.slide_layouts[-1]


def _copy_spTree(src_slide, dst_slide):

    """src_slide의 shape 트리를 dst_slide로 복사."""

    dst_spTree = dst_slide.shapes._spTree

    # 기본 shape 제거

    for child in list(dst_spTree):

        tag = child.tag.split('}')[-1] if '}' in child.tag else child.tag

        if tag in ('sp', 'pic', 'grpSp', 'cxnSp', 'graphicFrame'):

            dst_spTree.remove(child)



    src_spTree = src_slide.shapes._spTree

    for child in src_spTree:

        tag = child.tag.split('}')[-1] if '}' in child.tag else child.tag

        if tag in ('sp', 'pic', 'grpSp', 'cxnSp', 'graphicFrame'):

            dst_spTree.append(copy.deepcopy(child))


_IMG_ID_COUNTER = [5000]  # 충돌 방지를 위해 큰 숫자부터 시작


def _copy_image_rels(src_slide, dst_slide):

    """이미지 관계를 src에서 dst로 유니크 이름으로 복사하고 rId 매핑 반환."""

    from pptx.parts.image import ImagePart

    from pptx.opc.packuri import PackURI



    rId_map = {}

    dst_pkg = dst_slide.part.package



    for rel in src_slide.part.rels.values():

        # image 타입과 hdphoto 타입 모두 복사 (a14:imgLayer는 hdphoto를 참조함)

        if 'image' not in rel.reltype and 'hdphoto' not in rel.reltype:

            continue

        orig_part = rel.target_part

        orig_name = str(orig_part.partname)

        ext = '.' + orig_name.rsplit('.', 1)[-1] if '.' in orig_name else '.png'



        # 유니크 파일명 생성 (충돌 방지)

        _IMG_ID_COUNTER[0] += 1

        new_partname = PackURI(f'/ppt/media/image{_IMG_ID_COUNTER[0]}{ext}')



        new_part = ImagePart(new_partname, orig_part.content_type, dst_pkg, orig_part.blob)

        new_rId = dst_slide.part.relate_to(new_part, rel.reltype)

        rId_map[rel.rId] = new_rId

    return rId_map


def _update_rId_in_spTree(dst_slide, rId_map: dict):

    """복사된 shape 트리의 rId를 새 rId로 업데이트."""

    for el in dst_slide.shapes._spTree.iter():

        for attr, val in list(el.attrib.items()):

            if val in rId_map:

                el.set(attr, rId_map[val])


def _effective_bg(src_slide):

    """슬라이드 → 레이아웃 → 마스터 순으로 effective 배경 요소와 소유 part를 반환.

    p:bg는 p:cSld 내부에 있으므로 재귀 탐색(.// )을 사용한다."""

    bg = src_slide.element.find('.//' + qn('p:bg'))

    if bg is not None:

        return bg, src_slide.part

    try:

        layout = src_slide.slide_layout

        bg = layout.element.find('.//' + qn('p:bg'))

        if bg is not None:

            return bg, layout.part

        master = layout.slide_master

        bg = master.element.find('.//' + qn('p:bg'))

        if bg is not None:

            return bg, master.part

    except Exception:

        pass

    return None, None


def _copy_bg_image_rels(bg_element, src_part, dst_slide) -> dict:

    """배경(p:bg) 요소 내 이미지 r:embed rId를 dst_slide로 복사하고 rId 매핑 반환."""

    from pptx.parts.image import ImagePart

    from pptx.opc.packuri import PackURI



    R_NS = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'

    rId_map = {}

    dst_pkg = dst_slide.part.package



    for el in bg_element.iter():

        for attr, val in list(el.attrib.items()):

            if attr == f'{{{R_NS}}}embed' and val not in rId_map:

                try:

                    rel = src_part.rels[val]

                    if 'image' not in rel.reltype:

                        continue

                    orig_part = rel.target_part

                    orig_name = str(orig_part.partname)

                    ext = '.' + orig_name.rsplit('.', 1)[-1] if '.' in orig_name else '.png'

                    _IMG_ID_COUNTER[0] += 1

                    new_partname = PackURI(f'/ppt/media/image{_IMG_ID_COUNTER[0]}{ext}')

                    new_part = ImagePart(new_partname, orig_part.content_type, dst_pkg, orig_part.blob)

                    rId_map[val] = dst_slide.part.relate_to(new_part, rel.reltype)

                except Exception:

                    pass

    return rId_map


def duplicate_slide(prs, src_idx: int) -> int:

    """src_idx 슬라이드를 복제해서 맨 끝에 추가. 새 인덱스 반환."""

    src_slide = prs.slides[src_idx]

    blank_layout = _blank_layout(prs)

    new_slide = prs.slides.add_slide(blank_layout)



    _copy_spTree(src_slide, new_slide)

    rId_map = _copy_image_rels(src_slide, new_slide)

    if rId_map:

        _update_rId_in_spTree(new_slide, rId_map)



    # 배경 복사 (p:bg는 p:cSld 내부에 위치)

    src_cSld = src_slide.element.find(qn('p:cSld'))

    src_bg = src_cSld.find(qn('p:bg')) if src_cSld is not None else None

    if src_bg is not None:

        new_cSld = new_slide.element.find(qn('p:cSld'))

        if new_cSld is not None:

            existing = new_cSld.find(qn('p:bg'))

            if existing is not None:

                new_cSld.remove(existing)

            new_cSld.insert(0, copy.deepcopy(src_bg))



    # showMasterSp 속성 복사 (<p:sld> 최상위 unqualified attribute).
    # layout/배경(p:bg)이 같아도 이 속성이 다르면 마스터 배치 요소(장식·로고 등)의 노출
    # 여부가 달라져 색이 다르게 보인다. p:bg 복사만으로는 커버되지 않으므로 별도로 맞춘다.
    # 원본에 속성이 없으면(add_slide 기본 상태) 새 슬라이드에도 넣지 않는다.
    src_show_master = src_slide.element.get('showMasterSp')

    if src_show_master is not None:

        new_slide.element.set('showMasterSp', src_show_master)



    return len(prs.slides) - 1


def insert_slide_copy(prs, position: int, src_idx: int):

    """src_idx 슬라이드를 복제해서 position 위치에 삽입."""

    new_idx = duplicate_slide(prs, src_idx)

    move_slide(prs, new_idx, position)


def copy_slide_from_prs(target_prs, position: int, source_prs, source_idx: int):

    """다른 Presentation에서 슬라이드를 원본 서식(배경 포함)으로 복사해서 position에 삽입.

    소스 슬라이드의 레이아웃 이름으로 타겟에서 매칭 레이아웃을 찾아 사용한다 (원본 서식 유지).

    """

    src_slide = source_prs.slides[source_idx]



    # 소스 레이아웃 이름으로 타겟에서 매칭 레이아웃 탐색

    src_layout_name = src_slide.slide_layout.name

    target_layout = None

    for layout in target_prs.slide_layouts:

        if layout.name == src_layout_name:

            target_layout = layout

            break

    if target_layout is None:

        target_layout = _blank_layout(target_prs)



    new_slide = target_prs.slides.add_slide(target_layout)



    _copy_spTree(src_slide, new_slide)

    rId_map = _copy_image_rels(src_slide, new_slide)

    if rId_map:

        _update_rId_in_spTree(new_slide, rId_map)



    # 소스와 타겟의 슬라이드 크기가 다르면 도형이 절대 EMU 좌표 그대로 복사되어

    # 확대/축소되어 보인다 (예: 12192000x6858000 소스를 9144000x5143500 타겟에

    # 복사하면 그림이 타겟 슬라이드 폭의 127%로 넘침). 비율만큼 위치/크기를 보정한다.

    scale_x = target_prs.slide_width / source_prs.slide_width

    scale_y = target_prs.slide_height / source_prs.slide_height

    if scale_x != 1 or scale_y != 1:

        for shape in new_slide.shapes:

            if None in (shape.left, shape.top, shape.width, shape.height):

                continue

            shape.left = int(shape.left * scale_x)

            shape.top = int(shape.top * scale_y)

            shape.width = int(shape.width * scale_x)

            shape.height = int(shape.height * scale_y)



        # 도형 위치/크기만 스케일하고 글자 크기를 그대로 두면, 상자는 작아졌는데

        # 글자는 원본 크기 그대로라 텍스트가 넘치거나 화면 밖으로 잘려 보인다

        # (예: "화답송 시편 138(137)"이 "화답송 시편"으로 잘림). sz(폰트 크기)와

        # pPr의 EMU 단위 속성(marL/marR/indent/defTabSz)도 같은 비율로 축소한다.

        text_scale = min(scale_x, scale_y)

        for el in new_slide.shapes._spTree.iter():

            if el.tag in (qn('a:rPr'), qn('a:defRPr'), qn('a:endParaRPr')):

                sz = el.get('sz')

                if sz is not None:

                    el.set('sz', str(max(100, round(int(sz) * text_scale))))

            elif el.tag == qn('a:pPr'):

                for attr in ('marL', 'marR', 'indent', 'defTabSz'):

                    v = el.get(attr)

                    if v is not None:

                        el.set(attr, str(round(int(v) * text_scale)))



    # 배경 복사: 슬라이드 자체 p:bg가 없으면 레이아웃/마스터에서 상속된 배경을 명시적으로 삽입

    # p:bg는 p:cSld 내부에 위치하므로 cSld를 통해 접근/삽입

    src_bg, bg_src_part = _effective_bg(src_slide)

    if src_bg is not None:

        bg_copy = copy.deepcopy(src_bg)

        # 마스터/레이아웃에서 가져온 배경에 이미지가 있으면 rId도 복사

        if bg_src_part is not src_slide.part:

            bg_rId_map = _copy_bg_image_rels(src_bg, bg_src_part, new_slide)

            if bg_rId_map:

                for el in bg_copy.iter():

                    for attr, val in list(el.attrib.items()):

                        if val in bg_rId_map:

                            el.set(attr, bg_rId_map[val])

        new_cSld = new_slide.element.find(qn('p:cSld'))

        if new_cSld is not None:

            existing = new_cSld.find(qn('p:bg'))

            if existing is not None:

                new_cSld.remove(existing)

            new_cSld.insert(0, bg_copy)



    new_idx = len(target_prs.slides) - 1

    move_slide(target_prs, new_idx, position)


_COM_PROBE_PATH = [None]


def _com_probe_path() -> Path:
    """PowerPoint COM 실측용 임시 단일 슬라이드 pptx 경로.
    프로세스당 1개 경로를 재사용(호출마다 덮어쓰기)한다 — 매번 새 파일을 만들면
    COM 검증이 반복 호출되는 재조정 루프에서 임시 파일이 누적된다."""
    if _COM_PROBE_PATH[0] is None:
        token = f'{os.getpid()}_{uuid.uuid4().hex[:8]}'
        _COM_PROBE_PATH[0] = Path(tempfile.gettempdir()) / f'missa_com_probe_{token}.pptx'

        def _cleanup_probe_file(_path=_COM_PROBE_PATH[0]):
            try:
                _path.unlink(missing_ok=True)
            except OSError:
                pass

        atexit.register(_cleanup_probe_file)
    return _COM_PROBE_PATH[0]


def _build_com_probe_pptx(prs, slide) -> tuple:
    """prs 안의 한 슬라이드를, 크기가 동일한 임시 단일 슬라이드 Presentation으로
    복사해 디스크에 저장한다. PowerPoint COM으로 실제 줄 수를 실측하기 위한
    probe 파일 — <날짜>/log/ 폴더(커밋되는 회귀 테스트 산출물)는 쓰지 않는다.

    반환: (probe_pptx_path, content_shape의 1-based COM Shapes() 인덱스)
    """
    slide_idx = list(prs.slides).index(slide)

    probe = Presentation()
    probe.slide_width = prs.slide_width
    probe.slide_height = prs.slide_height
    while len(probe.slides) > 0:
        delete_slide(probe, 0)

    copy_slide_from_prs(probe, 0, prs, slide_idx)

    content_shape = _find_content_shape(probe.slides[0])
    shape_index_1based = list(probe.slides[0].shapes).index(content_shape) + 1

    path = _com_probe_path()
    probe.save(str(path))
    return path, shape_index_1based


# ─────────────────────────────────────────────────────────────────────────────
# 섹션 탐색
# ─────────────────────────────────────────────────────────────────────────────


def find_slide_with_text(prs, keyword: str, start: int = 0) -> int:

    for i in range(start, len(prs.slides)):

        if keyword in _slide_text(prs.slides[i]):

            return i

    return -1


def find_shape_exact_text(slide, exact: str) -> bool:

    for shape in slide.shapes:

        if shape.has_text_frame:

            for para in shape.text_frame.paragraphs:

                if para.text.strip() == exact:

                    return True

    return False


# ─────────────────────────────────────────────────────────────────────────────
# 콘텐츠 도형 탐색
# ─────────────────────────────────────────────────────────────────────────────


def _find_content_shape(slide):

    """슬라이드에서 본문 텍스트박스(가장 큰 것) 반환.

    '주님의 말씀입니다.' 등 고정 텍스트박스는 제외.

    """

    EXCLUDE_KW = ['전례문', 'Liturgy', 'Reading', '주님의 말씀입니다', '◎ 하느님', '◎ 그리스도님']

    best = None

    best_area = 0

    for shape in slide.shapes:

        if not shape.has_text_frame:

            continue

        t = shape.text_frame.text.strip()

        if not t or any(kw in t for kw in EXCLUDE_KW):

            continue

        area = shape.width * shape.height

        if area > best_area:

            best_area = area

            best = shape

    return best


def _has_ending_text(slide) -> bool:

    """슬라이드에 '주님의 말씀입니다.' 등 종료 텍스트가 있으면 True."""

    ENDING_KW = ('주님의 말씀입니다', '◎ 하느님', '◎ 그리스도님')

    for shape in slide.shapes:

        if not shape.has_text_frame:

            continue

        t = shape.text_frame.text.strip()

        if any(kw in t for kw in ENDING_KW):

            return True

    return False


def _clear_text_frame(tf):

    """텍스트 프레임의 내용을 비우고 빈 단락 하나를 남긴다."""

    from pptx.oxml import parse_xml as pptx_parse_xml

    A_NS = 'http://schemas.openxmlformats.org/drawingml/2006/main'

    txBody = tf._txBody

    for p in txBody.findall(qn('a:p')):

        txBody.remove(p)

    txBody.append(pptx_parse_xml(f'<a:p xmlns:a="{A_NS}"/>'))


# ─────────────────────────────────────────────────────────────────────────────
# 텍스트/run 조작 유틸
# ─────────────────────────────────────────────────────────────────────────────


def _para_append_run(p, new_r):

    """Run을 단락에 추가. <a:endParaRPr> 앞에 삽입해 OOXML 구조를 유지한다.

    endParaRPr 뒤에 run이 오면 PowerPoint가 해당 텍스트를 렌더링하지 않는다."""

    end_rpr = p.find(qn('a:endParaRPr'))

    if end_rpr is not None:

        end_rpr.addprevious(new_r)

    else:

        p.append(new_r)


def _replace_para_text_clone(para, new_text: str):

    """단락의 run을 제거하고 첫 run을 deepcopy해 new_text로 교체 (서식 보존)."""

    from pptx.oxml import parse_xml as pptx_parse_xml

    A_NS = 'http://schemas.openxmlformats.org/drawingml/2006/main'



    p = para._p

    runs = p.findall(qn('a:r'))

    template_r = runs[0] if runs else None



    for r in runs:

        p.remove(r)

    for br in p.findall(qn('a:br')):

        p.remove(br)



    if template_r is not None:

        new_r = copy.deepcopy(template_r)

        t_el = new_r.find(qn('a:t'))

        if t_el is not None:

            t_el.text = new_text

        _para_append_run(p, new_r)

    else:

        _para_append_run(p, pptx_parse_xml(f'<a:r xmlns:a="{A_NS}"><a:t>{_xml_escape(new_text)}</a:t></a:r>'))


def _set_화답송_content_text(tf, refrain_text: str, verse_text: str):
    """화답송 본문에 ◎ 후렴(para0) + ○ 절(para1) 설정.

    2단락 템플릿: 각 단락의 서식(pPr/rPr)을 보존하고 텍스트만 교체.
    1단락 템플릿(○만 있는 경우): ○ 단락은 그대로 두고 ◎ 단락을 앞에 구성.
    """
    from pptx.oxml import parse_xml as pptx_parse_xml
    A_NS = 'http://schemas.openxmlformats.org/drawingml/2006/main'
    txBody = tf._txBody
    paras = txBody.findall(qn('a:p'))

    if len(paras) >= 2:
        # 2단락 템플릿: para[0]=◎ 서식, para[1]=○ 서식 각각 보존
        for p, text in zip(paras[:2], [refrain_text, verse_text]):
            runs = p.findall(qn('a:r'))
            tmpl_r = copy.deepcopy(runs[0]) if runs else pptx_parse_xml(f'<a:r xmlns:a="{A_NS}"><a:t/></a:r>')
            for r in runs: p.remove(r)
            for br in p.findall(qn('a:br')): p.remove(br)
            tmpl_r.find(qn('a:t')).text = text
            _para_append_run(p, tmpl_r)
        for p in paras[2:]:
            txBody.remove(p)
    else:
        # 1단락 템플릿(○만 있음): 기존 단락을 ○ 절로 쓰고 ◎ 단락을 앞에 추가
        verse_p = paras[0] if paras else None
        verse_tmpl_r = None
        if verse_p is not None:
            runs = verse_p.findall(qn('a:r'))
            verse_tmpl_r = copy.deepcopy(runs[0]) if runs else pptx_parse_xml(f'<a:r xmlns:a="{A_NS}"><a:t/></a:r>')
            for r in runs: verse_p.remove(r)
            for br in verse_p.findall(qn('a:br')): verse_p.remove(br)
            verse_tmpl_r.find(qn('a:t')).text = verse_text
            _para_append_run(verse_p, verse_tmpl_r)

        # ◎ 단락: verse_p의 pPr(들여쓰기/탭) 복사 + ◎ 서식(37pt, bg2)
        refrain_p = pptx_parse_xml(f'<a:p xmlns:a="{A_NS}"/>')
        if verse_p is not None:
            pPr_src = verse_p.find(qn('a:pPr'))
            if pPr_src is not None:
                refrain_p.insert(0, copy.deepcopy(pPr_src))

        rPr_extra = 'lang="ko-KR" dirty="0" '
        if verse_tmpl_r is not None:
            verse_rPr = verse_tmpl_r.find(qn('a:rPr'))
            if verse_rPr is not None:
                lang = verse_rPr.get('lang', 'ko-KR')
                dirty = verse_rPr.get('dirty', '0')
                rPr_extra = f'lang="{lang}" dirty="{dirty}" '

        refrain_r = pptx_parse_xml(
            f'<a:r xmlns:a="{A_NS}">'
            f'<a:rPr {rPr_extra}sz="3700" b="1">'
            f'<a:solidFill><a:schemeClr val="bg2"/></a:solidFill>'
            f'</a:rPr>'
            f'<a:t/>'
            f'</a:r>'
        )
        refrain_r.find(qn('a:t')).text = refrain_text
        _para_append_run(refrain_p, refrain_r)

        if verse_p is not None:
            txBody.insert(list(txBody).index(verse_p), refrain_p)
        else:
            txBody.append(refrain_p)


def _set_single_para_text(tf, text: str):

    """TextFrame을 단일 단락으로 설정 (서식 보존)."""

    from pptx.oxml import parse_xml as pptx_parse_xml

    A_NS = 'http://schemas.openxmlformats.org/drawingml/2006/main'



    txBody = tf._txBody

    existing_paras = txBody.findall(qn('a:p'))



    template_r = None

    if existing_paras:

        for r in existing_paras[0].findall(qn('a:r')):

            template_r = r

            break



    for p in existing_paras[1:]:

        txBody.remove(p)



    if existing_paras:

        p = existing_paras[0]

        for r in p.findall(qn('a:r')): p.remove(r)

        for br in p.findall(qn('a:br')): p.remove(br)

        if template_r is not None:

            new_r = copy.deepcopy(template_r)

            t_el = new_r.find(qn('a:t'))

            if t_el is not None: t_el.text = text

            _para_append_run(p, new_r)

        else:

            _para_append_run(p, pptx_parse_xml(f'<a:r xmlns:a="{A_NS}"><a:t>{_xml_escape(text)}</a:t></a:r>'))


# ─────────────────────────────────────────────────────────────────────────────
# 성서 이름 유틸
# ─────────────────────────────────────────────────────────────────────────────


def _update_book_name_after_br(para, title: str) -> bool:
    """<a:br/> 이후 runs를 title 텍스트로 교체."""
    p = para._p

    brs = p.findall(qn('a:br'))

    if not brs:

        return False

    br = brs[0]

    after_runs = []

    br_found = False

    for child in list(p):

        tag = child.tag.split('}')[-1] if '}' in child.tag else child.tag

        if child is br:

            br_found = True

        elif br_found and tag == 'r':

            after_runs.append(child)

    if not after_runs:

        return False

    t_el = after_runs[0].find(qn('a:t'))

    if t_el is not None:

        t_el.text = title

    for extra_r in after_runs[1:]:

        t_el2 = extra_r.find(qn('a:t'))

        if t_el2 is not None:

            t_el2.text = ''

    return True


# ─────────────────────────────────────────────────────────────────────────────
# 배경 조작
# ─────────────────────────────────────────────────────────────────────────────


def _set_slide_bg_black(slide, prs=None):

    """슬라이드 배경을 검정으로 설정.

    p:bg를 black으로 설정하고, 레이아웃 shape들을 가리는

    전체 슬라이드 크기의 검정 사각형을 spTree 첫 번째로 삽입한다.

    """

    from pptx.oxml import parse_xml as pptx_parse_xml

    P_NS = 'http://schemas.openxmlformats.org/presentationml/2006/main'

    A_NS = 'http://schemas.openxmlformats.org/drawingml/2006/main'



    # 1. p:bg 검정으로 설정

    BLACK_BG_XML = (

        f'<p:bg xmlns:p="{P_NS}" xmlns:a="{A_NS}">'

        f'<p:bgPr><a:solidFill><a:srgbClr val="000000"/></a:solidFill>'

        f'<a:effectLst/></p:bgPr></p:bg>'

    )

    cSld = slide.element.find(qn('p:cSld'))

    if cSld is None:

        return

    existing = cSld.find(qn('p:bg'))

    if existing is not None:

        cSld.remove(existing)

    cSld.insert(0, pptx_parse_xml(BLACK_BG_XML))



    # 2. 전체 슬라이드를 덮는 검정 사각형을 spTree 최하단(첫 번째)에 삽입

    # 레이아웃 shape들이 위에서 배경을 덮는 경우 이 사각형이 그것을 가린다

    if prs is not None:

        cx = prs.slide_width

        cy = prs.slide_height

    else:

        cx, cy = 9144000, 5143500  # 와이드스크린 기본값 (EMU)



    BLACK_RECT_XML = (

        f'<p:sp xmlns:p="{P_NS}" xmlns:a="{A_NS}">'

        f'<p:nvSpPr>'

        f'<p:cNvPr id="9999" name="BlackBg"/>'

        f'<p:cNvSpPr><a:spLocks noGrp="1"/></p:cNvSpPr>'

        f'<p:nvPr/>'

        f'</p:nvSpPr>'

        f'<p:spPr>'

        f'<a:xfrm><a:off x="0" y="0"/><a:ext cx="{cx}" cy="{cy}"/></a:xfrm>'

        f'<a:prstGeom prst="rect"><a:avLst/></a:prstGeom>'

        f'<a:solidFill><a:srgbClr val="000000"/></a:solidFill>'

        f'<a:ln><a:noFill/></a:ln>'

        f'</p:spPr>'

        f'<p:txBody><a:bodyPr/><a:lstStyle/><a:p/></p:txBody>'

        f'</p:sp>'

    )

    spTree = cSld.find(qn('p:spTree'))

    if spTree is not None:

        # nvGrpSpPr, grpSpPr 다음 첫 번째 shape 위치에 삽입

        insert_pos = 2  # nvGrpSpPr(0), grpSpPr(1) 뒤

        spTree.insert(insert_pos, pptx_parse_xml(BLACK_RECT_XML))
