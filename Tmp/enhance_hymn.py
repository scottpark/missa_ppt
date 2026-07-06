#!/usr/bin/env python3
"""성가 악보 PPT 이미지 품질 개선"""

import argparse
import io
import os
import sys
from pathlib import Path

import numpy as np
from PIL import Image, ImageEnhance, ImageFilter
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn


def enhance_image(image_bytes: bytes) -> bytes:
    img = Image.open(io.BytesIO(image_bytes)).convert('L')

    # 2x 업스케일: 저해상도 획에 픽셀 여유 확보
    img_up = img.resize((img.width * 2, img.height * 2), Image.LANCZOS)

    # 노이즈 제거 (매우 약한 blur — 이후 선명화로 상쇄됨)
    img_dn = img_up.filter(ImageFilter.GaussianBlur(radius=0.5))

    # Unsharp Masking: 획 경계를 뚜렷하게, 회색조 그대로 보존
    # radius=2.0, percent=180, threshold=3
    img_sharp = img_dn.filter(ImageFilter.UnsharpMask(radius=2.0, percent=180, threshold=3))

    # 대비 향상: 어두운 획을 더 진하게, 밝은 배경은 더 밝게
    result = ImageEnhance.Contrast(img_sharp).enhance(1.4)

    buf = io.BytesIO()
    result.save(buf, format='PNG')
    return buf.getvalue()


def validate_image(original_bytes: bytes, enhanced_bytes: bytes):
    """Returns (is_valid, warnings_list)"""
    try:
        enh_arr_np = np.array(Image.open(io.BytesIO(enhanced_bytes)).convert('L'))
    except Exception as e:
        return False, [f"처리된 이미지를 열 수 없음: {e}"]

    enh_mean = enh_arr_np.mean()

    warns = []
    if enh_mean > 252:
        warns.append(f"이미지가 거의 흰색 (평균 픽셀값 {enh_mean:.0f})")
    elif enh_mean < 50:
        warns.append(f"이미지가 너무 어두움 (평균 픽셀값 {enh_mean:.0f})")

    try:
        orig_mean = np.array(Image.open(io.BytesIO(original_bytes)).convert('L')).mean()
        if abs(enh_mean - orig_mean) > 60:
            warns.append(
                f"처리 전후 밝기 큰 변화 감지 (평균 {orig_mean:.0f} -> {enh_mean:.0f})"
            )
    except Exception:
        pass

    return True, warns


def iter_pictures(shapes):
    """GROUP을 재귀적으로 탐색하여 모든 PICTURE shape을 yield."""
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from iter_pictures(shape.shapes)
        elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            yield shape


def process_pptx(input_path: str) -> None:
    prs = Presentation(input_path)
    success, warnings_list, failed = [], [], []

    for slide_idx, slide in enumerate(prs.slides, 1):
        img_num = 0
        for shape in iter_pictures(slide.shapes):
            img_num += 1
            label = f"슬라이드 {slide_idx}" if img_num == 1 else f"슬라이드 {slide_idx}-{img_num}"

            try:
                blip = shape._pic.find('.//' + qn('a:blip'))
                if blip is None:
                    continue
                rId = blip.get(qn('r:embed'))
                if rId is None:
                    continue

                image_part = slide.part.related_part(rId)
                original_blob = image_part.blob
                enhanced_blob = enhance_image(original_blob)

                ok, warns = validate_image(original_blob, enhanced_blob)
                if not ok:
                    msg = warns[0] if warns else "검증 실패"
                    print(f"  {label}: 실패 — {msg}")
                    failed.append((label, msg))
                    continue

                image_part._blob = enhanced_blob

                if warns:
                    for w in warns:
                        print(f"  {label}: 경고 — {w}")
                    warnings_list.append((label, '; '.join(warns)))
                else:
                    print(f"  {label}: 완료")
                    success.append(label)

            except Exception as e:
                print(f"  {label}: 실패 — {e}")
                failed.append((label, str(e)))

    stem = Path(input_path).stem
    output_dir = Path('성가')
    output_dir.mkdir(exist_ok=True)
    output_path = output_dir / f"{stem}_enhanced.pptx"
    prs.save(str(output_path))

    print()
    print("=" * 54)
    print(f"  처리 완료: {output_path.name}")
    print("=" * 54)
    print(f"  성공:  {len(success)}장")
    if warnings_list:
        print(f"  경고:  {len(warnings_list)}장")
        for label, msg in warnings_list:
            print(f"    → {label}: {msg}")
    if failed:
        print(f"  실패:  {len(failed)}장")
        for label, msg in failed:
            print(f"    → {label}: {msg}")
    print("=" * 54)
    print(f"  저장 위치: {output_path.resolve()}")
    print("=" * 54)


def select_file_popup() -> str:
    import tkinter as tk
    from tkinter import filedialog
    root = tk.Tk()
    root.withdraw()
    path = filedialog.askopenfilename(
        title="성가 PPT 파일 선택",
        filetypes=[("PowerPoint 파일", "*.pptx"), ("모든 파일", "*.*")],
    )
    root.destroy()
    return path or ""


def main():
    sys.stdout.reconfigure(encoding='utf-8')
    parser = argparse.ArgumentParser(description='성가 악보 PPT 이미지 품질 개선')
    parser.add_argument('input', nargs='?', help='입력 .pptx 파일 경로')
    args = parser.parse_args()

    input_path = args.input

    if not input_path:
        input_path = select_file_popup()
        if not input_path:
            print("파일을 선택하지 않았습니다.")
            sys.exit(0)

    if not os.path.isfile(input_path):
        print(f"오류: 파일을 찾을 수 없습니다 — {input_path}")
        sys.exit(1)

    if not input_path.lower().endswith('.pptx'):
        print(f"오류: .pptx 파일만 지원합니다 — {input_path}")
        sys.exit(1)

    print(f"처리 중: {input_path}")
    process_pptx(input_path)


if __name__ == '__main__':
    main()
