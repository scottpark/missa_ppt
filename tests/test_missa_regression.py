"""
missa_to_ppt.py 회귀 테스트 스위트

목적
----
주일미사·평일미사 처리 로직의 회귀 테스트. 향후 토요일미사 지원이 추가되거나
다른 기능이 변경됐을 때, 기존 주일/평일 로직이 깨지지 않았는지 확인하는 기준선.

실행
----
    pip install pytest
    pytest test_missa_regression.py -v

테스트 데이터
------------
- 주일: 20260712 (2026-07-12, 연중 제15주일)
- 평일: 20260624 (2026-06-24, 성 요한 세례자 탄생 대축일, 수요일)

두 폴더 모두 참조 PPT·JSON·성가 파일이 준비되어 있어야 한다.

새 케이스(예: 토요일미사) 추가 시 MASS_CASES 리스트에 항목을 추가하면 된다.
"""
from __future__ import annotations

import copy
import json
import re
import subprocess
import sys
import types
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

import missa_to_ppt as m
import missa_gui as gui
import missa_ooxml_utils as ou
import missa_reading_layout as rl
import missa_sections as sec
import _gen_helper as _gen

REPO_ROOT = Path(__file__).resolve().parent.parent  # tests/ -> 저장소 루트

MASS_CASES = [
    {
        "id": "sunday",
        "date": "20260712",
        "is_sunday": True,
        "hymns": {"입당": "329", "봉헌": "221", "성체": "156", "2차봉헌": "221", "파견": "25"},
    },
    {
        "id": "weekday",
        "date": "20260624",
        "is_sunday": False,
        "hymns": {"입당": "329", "봉헌": "221", "성체": "156", "파견": "25"},
    },
]

READING_SECTIONS = [
    ("제1독서_start", "제1독서_end"),
    ("제2독서_start", "제2독서_end"),
    ("복음_start", "복음_end"),
]

HYMN_TYPES = ["입당", "봉헌", "성체", "2차봉헌", "파견"]


# ─────────────────────────────────────────────────────────────────
# 1. 단위 테스트 — 순수 함수
# ─────────────────────────────────────────────────────────────────

class TestIsSundayMass:
    """향후 토요일미사(전야미사) 판단 로직 추가 시 이 클래스에
    '토요일을 평일미사/주일 전야 중 무엇으로 볼 것인가' 케이스를 추가해서
    회귀를 잡는다."""

    def test_sunday(self):
        assert m.is_sunday_mass("20260712") is True  # 일요일

    def test_weekday_friday(self):
        assert m.is_sunday_mass("20260612") is False  # 금요일

    def test_weekday_wednesday(self):
        assert m.is_sunday_mass("20260624") is False  # 수요일

    def test_saturday_currently_treated_as_weekday(self):
        # 2026-07-11은 토요일. 현재 로직(주일 여부만 판단)에서는 '평일미사'로 처리된다.
        # 토요일미사 지원이 추가되면 이 테스트는 새 기대값으로 업데이트해야 한다.
        assert m.is_sunday_mass("20260711") is False


class TestWrapLineCount:
    def test_empty_text(self):
        assert rl._wrap_line_count("") == 0

    def test_short_line_is_one_line(self):
        assert rl._wrap_line_count("짧은 문장입니다.") == 1

    def test_long_text_wraps_to_multiple_lines(self):
        text = "가" * 60  # CHARS_PER_LINE(27) 초과
        assert rl._wrap_line_count(text) >= 2


class TestParseIntoVerseUnits:
    def test_verse_numbers_are_detected(self):
        content = (
            "1 태초에 하느님께서 하늘과 땅을 창조하셨다. "
            "2 땅은 아직 모양을 갖추지 않고 비어 있었다."
        )
        units = rl.parse_into_verse_units(content)
        assert len(units) >= 2
        assert units[0]["verse_num"] == "1"


class TestSplitParaPreservesVerseColors:
    """2026-07-26 발견(규칙 13·14·15): 여러 절이 하나의 논리 단락으로 병합된 경우(continuation 절) 문단을
    두 슬라이드로 분리할 때 단락 중간 절 번호의 오렌지색이 사라지는 회귀를 방지한다. 2026-10-03 슬라이드
    단위 채우기부터 문단 분할은 `_split_paras_at`(오렌지 구간을 문단 오프셋으로 보관해 잘라 옮김)가 맡는다.
    구 TestSplitAndAdjustViaCom(_split_and_adjust_via_com의 keep ±1 재시도·롤백)은 그 함수와 함께 삭제됐다 —
    "오버플로를 성공으로 보고하지 않는다"·"재시도 후에도 절 번호 색 유지" 요구사항은 정확히 9줄만 확정하는
    구조(test_missa_progression.py SF그룹)와 아래 테스트가 지킨다."""

    @staticmethod
    def _two_verse_paras():
        return rl._reading_paras_from_units([
            {'text': '51 many words here for wrapping test purposes today', 'verse_num': '51',
             'is_continuation': False},
            {'text': '52 more words after the second verse marker appear here', 'verse_num': '52',
             'is_continuation': True},
        ])

    @staticmethod
    def _render_orange(paras):
        """paras를 실제 텍스트 프레임에 기록(_set_reading_text)해 오렌지 run의 텍스트를 읽는다."""
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        box = slide.shapes.add_textbox(0, 0, 5_000_000, 3_000_000)
        box.text_frame.text = '샘플'
        rl._set_reading_text(box.text_frame, rl._paras_to_units(paras))
        out = []
        for p in box.text_frame.paragraphs:
            for r in p.runs:
                try:
                    if r.font.color.rgb == rl.ORANGE:
                        out.append(r.text.strip())
                except AttributeError:
                    continue
        return out

    def test_verse_numbers_survive_paragraph_split_at_every_word_boundary(self):
        paras = self._two_verse_paras()
        assert len(paras) == 1
        text = paras[0]['text']
        cuts = [m.end() for m in re.finditer(' ', text)]
        assert len(cuts) > 8
        for cut_at in cuts:
            front, rest = rl._split_paras_at(paras, 0, cut_at)
            combined = self._render_orange(front) + self._render_orange(rest)
            assert combined == ['51', '52'], f"절 번호 오렌지색 유실(cut={cut_at}): {combined}"
            assert (front[-1]['text'] + ' ' + rest[0]['text']) == text

    def test_split_at_paragraph_start_keeps_whole_paragraphs_and_orange(self):
        paras = rl._reading_paras_from_units(rl.parse_into_verse_units('1 가나다.\n2 라마바.\n3 사아자.'))
        front, rest = rl._split_paras_at(paras, 1, 0)
        assert self._render_orange(front) + self._render_orange(rest) == ['1', '2', '3']


class TestComVerificationEnabledConfig:
    """config.json의 com_verification_enabled 플래그를 읽는
    _com_verification_enabled()가 true/false/키 없음(기본값 True)을 올바르게
    반영하는지 확인한다. PowerPoint COM 검증을 끄고 싶은 머신을 위한 opt-out이므로
    실패 시에도 항상 안전하게(Pillow 전용으로) 동작해야 한다."""

    def _reset_cache_and_point_config(self, monkeypatch, tmp_path, config_dict):
        config_path = tmp_path / "config.json"
        config_path.write_text(
            json.dumps(config_dict, ensure_ascii=False), encoding="utf-8"
        )
        monkeypatch.setattr(gui, "CONFIG_FILE", config_path)
        # 전역 리스트를 직접 고치면(`[0] = None`) 이 테스트가 채운 False가 뒤따르는 테스트로
        # 새어, COM 재조정이 꺼진 채로 도는 test_y2_G2가 전체 실행에서만 실패한다.
        monkeypatch.setattr(gui, "_COM_VERIFY_ENABLED_CACHE", [None])

    def test_true_when_explicitly_enabled(self, monkeypatch, tmp_path):
        self._reset_cache_and_point_config(
            monkeypatch, tmp_path, {"com_verification_enabled": True}
        )
        assert gui._com_verification_enabled() is True

    def test_false_when_explicitly_disabled(self, monkeypatch, tmp_path):
        self._reset_cache_and_point_config(
            monkeypatch, tmp_path, {"com_verification_enabled": False}
        )
        assert gui._com_verification_enabled() is False

    def test_defaults_true_when_key_absent(self, monkeypatch, tmp_path):
        self._reset_cache_and_point_config(monkeypatch, tmp_path, {})
        assert gui._com_verification_enabled() is True

    def test_result_is_memoized(self, monkeypatch, tmp_path):
        self._reset_cache_and_point_config(
            monkeypatch, tmp_path, {"com_verification_enabled": False}
        )
        assert gui._com_verification_enabled() is False
        # 캐시된 이후에는 config.json 내용이 바뀌어도 재확인하지 않는다
        (tmp_path / "config.json").write_text(
            json.dumps({"com_verification_enabled": True}), encoding="utf-8"
        )
        assert gui._com_verification_enabled() is False


# ─────────────────────────────────────────────────────────────────
# 2. 통합 테스트 — 실제 PPT 생성 + 구조 검증
# ─────────────────────────────────────────────────────────────────

def _find_output(date_str: str) -> Path:
    """임시 폴더에 생성된 결과 PPT(tests/_gen_helper.py) — 저장소 output/은 건드리지 않는다."""
    return _gen.generated_path(date_str)


def _generate(case: dict) -> str:
    """case를 임시 폴더에 생성(같은 날짜는 프로세스당 1회)하고 stdout을 반환한다. 생성 실패·내부 검증
    실패는 helper가 AssertionError로 알린다."""
    return _gen.generate_case(case["date"], case["hymns"])[1]


@pytest.fixture(scope="module", params=MASS_CASES, ids=lambda c: c["id"])
def generated_case(request):
    case = request.param
    log = _generate(case)
    path = _find_output(case["date"])
    prs = Presentation(str(path))
    sections = m.find_sections(prs)
    json_data = m.get_json_data(case["date"])
    return {
        "case": case, "prs": prs, "sections": sections, "log": log,
        "path": path, "json_data": json_data,
    }


def test_output_opens_without_corruption(generated_case):
    # Presentation()이 예외 없이 로드되면 OOXML 구조 자체는 유효함(fixture에서 이미 검증됨).
    assert len(generated_case["prs"].slides) > 0


def test_liturgy_type_detected_correctly(generated_case):
    log = generated_case["log"]
    case = generated_case["case"]
    expected = "미사 유형: 주일미사" if case["is_sunday"] else "미사 유형: 평일미사"
    assert expected in log


def test_required_sections_present(generated_case):
    sections = generated_case["sections"]
    required = [
        "입당송", "제1독서_start", "제1독서_end",
        "화답송_start", "화답송_end",
        "복음환호송", "복음_start", "복음_end",
        "영성체송",
    ]
    missing = [k for k in required if k not in sections]
    assert not missing, f"누락된 섹션: {missing}"


class TestBuildComProbePptx:
    """_build_com_probe_pptx()가 실제 독서 슬라이드를 손실 없이 임시 단일 슬라이드
    pptx로 복사하는지 확인한다 (PowerPoint COM 실측을 위한 probe 생성 단계)."""

    def test_probe_slide_dimensions_match_original(self, generated_case):
        prs, sections = generated_case["prs"], generated_case["sections"]
        slide = prs.slides[sections["제1독서_start"]]
        probe_path, _ = ou._build_com_probe_pptx(prs, slide)
        probe = Presentation(str(probe_path))
        assert probe.slide_width == prs.slide_width
        assert probe.slide_height == prs.slide_height

    def test_probe_content_text_matches_original(self, generated_case):
        prs, sections = generated_case["prs"], generated_case["sections"]
        slide = prs.slides[sections["제1독서_start"]]
        original_text = ou._find_content_shape(slide).text_frame.text
        probe_path, _ = ou._build_com_probe_pptx(prs, slide)
        probe = Presentation(str(probe_path))
        probe_text = ou._find_content_shape(probe.slides[0]).text_frame.text
        assert probe_text == original_text

    def test_probe_shape_index_is_1_based_and_points_to_content_shape(self, generated_case):
        prs, sections = generated_case["prs"], generated_case["sections"]
        slide = prs.slides[sections["제1독서_start"]]
        original_text = ou._find_content_shape(slide).text_frame.text
        probe_path, shape_index = ou._build_com_probe_pptx(prs, slide)
        probe = Presentation(str(probe_path))
        shapes = list(probe.slides[0].shapes)
        assert 1 <= shape_index <= len(shapes)
        pointed_shape = shapes[shape_index - 1]  # COM Shapes()는 1-based
        assert pointed_shape.has_text_frame
        assert pointed_shape.text_frame.text == original_text

    def test_probe_path_reused_across_calls(self, generated_case):
        prs, sections = generated_case["prs"], generated_case["sections"]
        slide_a = prs.slides[sections["제1독서_start"]]
        slide_b = prs.slides[sections["복음_start"]]
        path_a, _ = ou._build_com_probe_pptx(prs, slide_a)
        path_b, _ = ou._build_com_probe_pptx(prs, slide_b)
        assert path_a == path_b, "프로세스당 하나의 임시 경로를 재사용해야 함(누적 생성 금지)"
        # 두 번째 호출 후 파일 내용은 slide_b 기준으로 덮어써져 있어야 한다
        probe = Presentation(str(path_b))
        probe_text = ou._find_content_shape(probe.slides[0]).text_frame.text
        assert probe_text == ou._find_content_shape(slide_b).text_frame.text


def test_no_reading_slide_line_overflow(generated_case):
    """독서·복음 본문 슬라이드가 LINES_PER_SLIDE(9줄)를 넘으면 안 된다.
    2026-07-16 발견된 post-write 재조정 실패(연쇄 초과) 회귀 방지용."""
    prs, sections = generated_case["prs"], generated_case["sections"]
    problems = []
    for start_key, end_key in READING_SECTIONS:
        if start_key not in sections or end_key not in sections:
            continue
        s, e = sections[start_key], sections[end_key]
        for idx in range(s, e):
            shape = ou._find_content_shape(prs.slides[idx])
            font, box_px = rl._get_slide_render_params(prs.slides[idx])
            if shape is None or font is None:
                continue
            paras = [{'text': para.text, 'orange': []}
                     for para in shape.text_frame.paragraphs if para.text.strip()]
            lines = len(rl._pil_line_starts(paras, font, box_px))
            if lines > rl.LINES_PER_SLIDE:
                problems.append(f"{start_key} 슬라이드 {idx + 1}: {lines}줄 (>{rl.LINES_PER_SLIDE})")
    assert not problems, "\n".join(problems)


def test_no_missing_orange_verse_numbers(generated_case):
    """독서·복음의 모든 절 번호가 실제로 오렌지색 run으로 렌더링됐는지 확인한다.
    텍스트 자체는 남아 있어도 색상만 사라지는 경우(2026-07-26)가 있어, 슬라이드
    텍스트 존재 여부가 아니라 절 번호 단위로 오렌지색 run을 직접 대조한다."""
    prs, sections, json_data = (
        generated_case["prs"], generated_case["sections"], generated_case["json_data"]
    )
    problems = []
    for label, start_key, end_key in [
        ("제1독서", "제1독서_start", "제1독서_end"),
        ("제2독서", "제2독서_start", "제2독서_end"),
        ("복음", "복음_start", "복음_end"),
    ]:
        reading = json_data.get(label)
        if not reading or not reading.get("content"):
            continue
        if start_key not in sections or end_key not in sections:
            continue
        missing = sec._missing_orange_verse_numbers(
            prs, sections[start_key], sections[end_key], reading["content"]
        )
        if missing:
            problems.append(f"{label} 절 번호 오렌지색 누락: {', '.join(missing)}")
    assert not problems, "\n".join(problems)


def test_reading_verse_numbers_appear_in_ascending_order(generated_case):
    """2026-08-04 발견 회귀 방지: (구)사후 재조정의 give-up 롤백 경로가 되돌릴 위치를 잘못 계산해
    (맨 앞이 아니라 append) 다음 슬라이드에 이미 남아 있던 뒤쪽 단락보다 앞에 있어야 할
    단락이 뒤로 밀려 절 순서가 뒤바뀌는 버그가 있었다. 슬라이드 단위 채우기(2026-10-03)는 앞에서부터
    순서대로 확정하므로 구조적으로 재발하지 않지만 실제 산출물로 계속 지킨다. 오렌지색 절 번호가
    슬라이드/단락 순서대로 오름차순으로 나타나는지 확인해 이런 단락 순서
    뒤바뀜을 감지한다(존재 여부만 확인하는 test_no_missing_orange_verse_numbers
    와 달리 순서 자체를 확인)."""
    prs, sections, json_data = (
        generated_case["prs"], generated_case["sections"], generated_case["json_data"]
    )
    problems = []
    for label, start_key, end_key in [
        ("제1독서", "제1독서_start", "제1독서_end"),
        ("제2독서", "제2독서_start", "제2독서_end"),
        ("복음", "복음_start", "복음_end"),
    ]:
        reading = json_data.get(label)
        if not reading or not reading.get("content"):
            continue
        if start_key not in sections or end_key not in sections:
            continue
        s, e = sections[start_key], sections[end_key]
        seq = []
        for idx in range(s, e):
            shape = ou._find_content_shape(prs.slides[idx])
            if shape is None:
                continue
            for para in shape.text_frame.paragraphs:
                for run in para.runs:
                    try:
                        if run.font.color.rgb != rl.ORANGE:
                            continue
                    except Exception:
                        continue
                    match = re.match(r'^(\d+)', run.text.strip())
                    if match:
                        seq.append(int(match.group(1)))
        for a, b in zip(seq, seq[1:]):
            if b < a:
                problems.append(f"{label}: 절 번호 순서 역전 {a} → {b} (전체 순서: {seq})")
    assert not problems, "\n".join(problems)


def test_blackbg_covers_full_slide(generated_case):
    """BlackBg 도형이 존재하면 항상 슬라이드 전체 크기를 덮어야 한다.
    2026-07-16 발견된 _align_ending_slides_to_제2독서의 BlackBg 오분류 회귀 방지용."""
    prs = generated_case["prs"]
    problems = []
    for i, slide in enumerate(prs.slides):
        for shape in slide.shapes:
            if shape.name == "BlackBg":
                if not (
                    shape.left == 0 and shape.top == 0
                    and shape.width == prs.slide_width
                    and shape.height == prs.slide_height
                ):
                    problems.append(
                        f"슬라이드 {i + 1}: BlackBg 크기 이상 "
                        f"({shape.left},{shape.top},{shape.width},{shape.height})"
                    )
    assert not problems, "\n".join(problems)


def test_화답송_pattern_matches_mass_type(generated_case):
    """주일: 악보/텍스트 슬라이드 교대(짝수 인덱스=악보). 평일: 텍스트만, 악보 이미지 없음."""
    prs, sections, case = (
        generated_case["prs"], generated_case["sections"], generated_case["case"]
    )
    if "화답송_start" not in sections:
        pytest.skip("화답송 섹션 없음")
    s, e = sections["화답송_start"], sections["화답송_end"]
    problems = []
    for i in range(s, e):
        slide = prs.slides[i]
        has_picture = any(sh.shape_type == MSO_SHAPE_TYPE.PICTURE for sh in slide.shapes)
        rel_i = i - s
        if case["is_sunday"]:
            expected_picture = (rel_i % 2 == 0)
            if has_picture != expected_picture:
                problems.append(
                    f"화답송 슬라이드 {i + 1}: 악보 유무 불일치 "
                    f"(기대={expected_picture}, 실제={has_picture})"
                )
        else:
            if has_picture:
                problems.append(f"화답송 슬라이드 {i + 1}: 평일미사인데 악보 이미지 존재")
    assert not problems, "\n".join(problems)


def test_hymn_score_copy_matches_mass_type(generated_case):
    """주일: 성가 악보 슬라이드 복사됨. 평일: 번호만 갱신, 악보 슬라이드 없음."""
    sections, case = generated_case["sections"], generated_case["case"]
    problems = []
    for hymn in HYMN_TYPES:
        s_key, e_key = f"{hymn}_content_start", f"{hymn}_content_end"
        if s_key not in sections or e_key not in sections:
            continue
        n_content = sections[e_key] - sections[s_key]
        if case["is_sunday"]:
            if n_content == 0:
                problems.append(f"{hymn}: 주일미사인데 악보 슬라이드 없음")
        else:
            if n_content != 0:
                problems.append(f"{hymn}: 평일미사인데 악보 슬라이드 {n_content}장 존재")
    assert not problems, "\n".join(problems)


# ─────────────────────────────────────────────────────────────────
# 3. PowerPoint COM 실측 통합 테스트 (환경 조건부 skip)
# ─────────────────────────────────────────────────────────────────

def _com_available() -> bool:
    try:
        import ppt_com_verify as com
    except ImportError:
        return False
    return com.is_available()


class TestComVerify:
    """PowerPoint COM 실측 경로 전용 테스트. pywin32 미설치 또는 PowerPoint COM
    연결 불가 환경에서는 skip한다(이 저장소의 기존 pytest.skip 전례를 따름)."""

    def test_com_available_or_skips(self):
        if not _com_available():
            pytest.skip("pywin32 미설치 또는 PowerPoint COM 연결 불가")
        assert True

    def test_com_verification_resolves_known_overflow_case(self):
        """2026-06-24 케이스는 과거 Pillow 추정으로는 9줄이지만 실제 PowerPoint
        렌더링으로는 10줄인 슬라이드가 존재했던 회귀였다(제1독서 슬라이드 19,
        복음 슬라이드 38). com_verification_enabled가 기본 활성화된 상태로
        재생성한 뒤, Pillow가 아니라 실제 COM 실측으로 모든 독서/복음 슬라이드가
        9줄 이하인지 확인한다 — 원래 버그를 실제로 잡아낼 수 있는 유일한 테스트."""
        if not _com_available():
            pytest.skip("pywin32 미설치 또는 PowerPoint COM 연결 불가")
        import ppt_com_verify as com

        case = next(c for c in MASS_CASES if c["date"] == "20260624")
        _generate(case)
        path = _find_output(case["date"])
        prs = Presentation(str(path))
        sections = m.find_sections(prs)

        problems = []
        for start_key, end_key in READING_SECTIONS:
            if start_key not in sections or end_key not in sections:
                continue
            s, e = sections[start_key], sections[end_key]
            for idx in range(s, e):
                slide = prs.slides[idx]
                if ou._find_content_shape(slide) is None:
                    continue  # 종료 텍스트 병합 슬라이드 등 본문 텍스트박스가 없는 슬라이드는 대상 제외
                probe_path, shape_idx = ou._build_com_probe_pptx(prs, slide)
                real_lines = com.count_slide_lines(str(probe_path), shape_idx)
                if real_lines > rl.LINES_PER_SLIDE:
                    problems.append(
                        f"{start_key} 슬라이드 {idx + 1}: COM 실측 {real_lines}줄 (>{rl.LINES_PER_SLIDE})"
                    )
        assert not problems, "\n".join(problems)


# ---------------------------------------------------------------------------
# 종료 텍스트박스 위치 재계산 (본문+종료 통합 슬라이드)
#   test_missa_progression.py에서 승격 (2026-09-02, regression-qa).
#   버그: _reposition_merged_ending_shapes()가 line_height를
#   content_shape.height // LINES_PER_SLIDE로 역산 → 짧은 본문 박스에서 과소 산출 →
#   종료 텍스트박스가 본문과 겹침(20260906 idx 61, 약 640,000 EMU 겹침).
#   수정: line_height를 폰트 실측(ascent+descent × lnSpc)으로 직접 계산 + 한 슬라이드에
#   안 들어가면 도형 통째로 다음 슬라이드 이동(분리 금지). 요청:
#   _workspace/bugfix_ending_textbox_position/00_request.md
# ---------------------------------------------------------------------------

_ENDING_TEXT = "주님의 말씀입니다.\n◎ 하느님, 감사합니다."


def _rep_mk_prs():
    from pptx import Presentation as _P
    return _P()  # 기본 9144000×6858000


def _rep_add_reading_box(slide, lines, top, height,
                         left=1_000_000, width=7_000_000, ending=False):
    """32pt BatangChe 텍스트박스 추가. _get_slide_render_params가 폰트 메트릭을
    읽을 수 있도록 run rPr에 sz/latin이 실제로 들어가게 한다."""
    from pptx.util import Emu, Pt
    box = slide.shapes.add_textbox(Emu(left), Emu(top), Emu(width), Emu(height))
    tf = box.text_frame
    tf.word_wrap = True
    for i, text in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        run = p.add_run()
        run.text = text
        run.font.size = Pt(32)
        run.font.name = "BatangChe"
    return box


def _rep_find_ending_shapes(slide):
    return [
        sh for sh in slide.shapes
        if sh.has_text_frame and "주님의 말씀입니다" in sh.text_frame.text
    ]


def _rep_font_metrics_available(slide):
    pil, _ = rl._get_slide_render_params(slide)
    return pil is not None


class TestMergedEndingBodyMetrics:
    """_merged_ending_body_metrics()가 char-count/Pillow 추정이 실제 PowerPoint
    줄바꿈과 다를 때 COM 실측값(줄 수 + BoundHeight)을 우선 채택하는지 PowerPoint
    없이 monkeypatch로 결정적으로 검증한다.

    실사용 버그(2026-10-03, 청년미사 제2독서 종료 통합 슬라이드): char-count 근사
    (`_wrap_line_count`)와 Pillow word-wrap 둘 다 실제 5줄인 문단을 4줄로 오산해,
    종료 텍스트박스가 본문 마지막 줄과 거의 겹치는 자리(GAP≈0.09인치)에 배치됐다
    (상세: docs/ooxml-pitfalls-log.md "여러 곳에서 재사용하는 fit-보정 함수" 절
    근방, `_merged_ending_body_metrics` docstring)."""

    @pytest.fixture(autouse=True)
    def _reset_state(self, monkeypatch):
        monkeypatch.setattr(rl, "_COM_DISABLED", [False], raising=False)
        monkeypatch.setattr(gui, "_COM_VERIFY_ENABLED_CACHE", [True], raising=False)
        monkeypatch.setattr(rl, "_build_com_probe_pptx", lambda prs, slide: (Path("dummy.pptx"), 1), raising=False)
        monkeypatch.setattr(rl, "_COM_ATEXIT_REGISTERED", [False], raising=False)

    @staticmethod
    def _install_fake_com(monkeypatch, lines=None, bound_height_pt=None, raises=False):
        fake = types.ModuleType("ppt_com_verify")

        class _FakeComUnavailable(Exception):
            pass

        fake.ComVerificationUnavailable = _FakeComUnavailable
        calls = []

        def _measure(path, shape_index):
            calls.append((path, shape_index))
            if raises:
                raise _FakeComUnavailable("simulated COM failure")
            return {"lines": lines, "bound_height_pt": bound_height_pt}

        fake.measure_text_metrics = _measure
        fake.shutdown = lambda: None
        monkeypatch.setitem(sys.modules, "ppt_com_verify", fake)
        return calls

    def test_com_value_overrides_undercounted_estimate(self, monkeypatch):
        prs = _rep_mk_prs()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        content = _rep_add_reading_box(slide, ["본문 문장"], top=0, height=2_000_000)
        calls = self._install_fake_com(monkeypatch, lines=5, bound_height_pt=42.0)

        line_count, line_height = rl._merged_ending_body_metrics(prs, slide, content)

        assert len(calls) == 1
        assert line_count == 5
        assert line_height == round(42.0 * 12700 / 5)

    def test_falls_back_to_estimate_when_com_unavailable(self, monkeypatch):
        prs = _rep_mk_prs()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        content = _rep_add_reading_box(slide, ["본문 문장"], top=0, height=2_000_000)
        calls = self._install_fake_com(monkeypatch, raises=True)

        line_count, line_height = rl._merged_ending_body_metrics(prs, slide, content)

        assert len(calls) == 1
        assert line_count == rl._wrap_line_count("본문 문장")
        assert rl._COM_DISABLED[0] is True

    def test_end_to_end_reposition_uses_com_verified_gap(self, monkeypatch):
        """COM이 '추정과 다른' 줄 수/높이를 돌려주면, 종료 텍스트박스는 추정값이
        아니라 COM 값 기준(line_count+GAP)×line_height 위치에 배치된다."""
        prs = _rep_mk_prs()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        content = _rep_add_reading_box(slide, ["본문 문장"], top=0, height=3_000_000)
        ending = _rep_add_reading_box(
            slide, _ENDING_TEXT.split("\n"), top=500_000, height=1_300_000, ending=True)

        bound_height_pt = 5 * 100_000 / 12700  # 5줄 × 100,000 EMU/줄이 되도록 역산
        self._install_fake_com(monkeypatch, lines=5, bound_height_pt=bound_height_pt)

        sections = {"제1독서_start": 0, "제1독서_end": 1}
        rl._reposition_merged_ending_shapes(prs, sections)

        line_height = round(bound_height_pt * 12700 / 5)
        GAP = 2  # _reposition_merged_ending_shapes()의 로컬 상수와 동일 값
        assert ending.top == content.top + (5 + GAP) * line_height


class TestRepositionMergedEndingShapes:
    """본문+종료 통합 슬라이드의 종료 텍스트박스 재배치 회귀."""

    def test_short_body_no_overlap(self):
        """본문 5줄 + 실제보다 큰 본문 박스(20260906 idx 61 재현)에서 재계산된 종료
        텍스트박스 top이 본문 텍스트 실제 영역(5줄 분량) 아래로 내려가 겹치지 않는다.
        박스 bottom이 아니라 구현이 배치에 쓰는 실측값(line_count × line_height)과
        비교한다."""
        prs = _rep_mk_prs()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        content = _rep_add_reading_box(
            slide, [f"본문 문장 {i}" for i in range(5)], top=0, height=2_876_621)
        ending = _rep_add_reading_box(
            slide, _ENDING_TEXT.split("\n"), top=2_227_981, height=1_300_000, ending=True)

        if not _rep_font_metrics_available(slide):
            pytest.skip("Pillow/한글 폰트 미존재 — 폴백 경로는 통합 회귀가 커버")

        line_height = rl._content_line_height_emu(slide, content)
        line_count = sum(
            rl._wrap_line_count(p.text.strip())
            for p in content.text_frame.paragraphs if p.text.strip()
        )
        assert line_count == 5
        text_bottom = content.top + line_count * line_height

        sections = {"제1독서_start": 0, "제1독서_end": 1}
        rl._reposition_merged_ending_shapes(prs, sections)

        assert ending.top >= text_bottom, (
            "종료 텍스트박스가 본문 텍스트 실제 영역과 겹침", ending.top, text_bottom
        )

    def test_overflow_moves_whole_ending_to_next_slide(self):
        """본문 줄 수가 많아 종료 텍스트박스가 온전히 못 들어가면 도형을 통째로 새
        다음 슬라이드로 옮긴다(텍스트 분리 없이 단일 도형). 뒤의 무관한 슬라이드를
        덮지 않고 그 앞에 새 슬라이드를 만든다."""
        prs = _rep_mk_prs()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        _rep_add_reading_box(
            slide, [f"본문 문장 {i}" for i in range(8)], top=0, height=4_600_000)
        _rep_add_reading_box(
            slide, _ENDING_TEXT.split("\n"), top=2_227_981, height=1_300_000, ending=True)
        from pptx.util import Emu
        other = prs.slides.add_slide(prs.slide_layouts[6])
        ob = other.shapes.add_textbox(Emu(0), Emu(0), Emu(3_000_000), Emu(1_000_000))
        ob.text_frame.text = "다음 섹션 슬라이드"

        if not _rep_font_metrics_available(slide):
            pytest.skip("Pillow/한글 폰트 미존재 — 폴백 경로는 통합 회귀가 커버")

        n_before = len(prs.slides._sldIdLst)
        sections = {"제1독서_start": 0, "제1독서_end": 1}
        rl._reposition_merged_ending_shapes(prs, sections)

        assert len(prs.slides._sldIdLst) == n_before + 1
        assert not rl._has_ending_text(prs.slides[0]), "본문 슬라이드에 종료 텍스트가 남음"
        assert rl._find_content_shape(prs.slides[0]) is not None
        moved = _rep_find_ending_shapes(prs.slides[1])
        assert len(moved) == 1, "종료 텍스트가 단일 도형이 아님(분리됨)"
        assert moved[0].text_frame.text.strip() == _ENDING_TEXT, "종료 텍스트 일부 손실/분리"
        assert any(
            sh.has_text_frame and "다음 섹션" in sh.text_frame.text
            for sh in prs.slides[2].shapes
        ), "무관한 다음 슬라이드가 종료 텍스트로 덮이거나 사라짐"

    def test_line_height_is_font_measured_not_box_division(self):
        """line_height가 '박스 height ÷ 9'가 아니라 폰트 실측으로 산출된다는 것을,
        서로 다른 박스 height를 준 두 슬라이드가 같은 종료 top을 내는지로 확인한다."""
        def _make(height):
            prs = _rep_mk_prs()
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            _rep_add_reading_box(
                slide, [f"본문 {i}" for i in range(5)], top=0, height=height)
            ending = _rep_add_reading_box(
                slide, _ENDING_TEXT.split("\n"), top=2_227_981, height=1_300_000, ending=True)
            return prs, slide, ending

        prs_a, slide_a, ending_a = _make(2_876_621)
        prs_b, slide_b, ending_b = _make(4_200_000)

        if not (_rep_font_metrics_available(slide_a) and _rep_font_metrics_available(slide_b)):
            pytest.skip("Pillow/한글 폰트 미존재 — 폴백 경로는 통합 회귀가 커버")

        rl._reposition_merged_ending_shapes(prs_a, {"제1독서_start": 0, "제1독서_end": 1})
        rl._reposition_merged_ending_shapes(prs_b, {"제1독서_start": 0, "제1독서_end": 1})

        assert ending_a.top == ending_b.top, (
            "종료 top이 박스 height에 의존 — line_height가 폰트 실측이 아님",
            ending_a.top, ending_b.top,
        )


# ---------------------------------------------------------------------------
# 성가 헤더 라벨 재작성 시 run 색상 오염 버그 (bugfix_성가_label_run_color)
#
# test_missa_progression.py에서 승격(2026-09-02, regression-qa).
# 원인: _update_성가_header가 라벨을 정규화하며 라벨 글자 수가 바뀌면(예: '2차봉헌'(4)
# → '2차 봉헌'(5)), 옛 _update_prefix_in_runs가 라벨/구분자/숫자를 하나의 문자 스트림으로
# 취급해 옛 run 길이 기준으로 통짜 재배치하며 라벨 마지막 글자('헌')가 구분자 공백 run
# (우연히 회색) 슬롯으로 밀려 그 색을 물려받았다.
# 수정: 정규식 그룹 경계(라벨 vs 구분자+숫자)를 재배치의 하드 경계로 삼아 각 구간을
# 독립적으로 재배치 → 한쪽의 길이 변화가 다른 쪽 run 색상을 침범하지 않는다.
# 설계 계열: CLAUDE.md "단락 분리·병합 함수는 run 개수를 2개로 가정하지 않는다".
# ---------------------------------------------------------------------------

import missa_content_updaters as cu  # noqa: E402
from pptx.oxml.ns import qn as _cu_qn  # noqa: E402
from pptx.oxml import parse_xml as _cu_parse_xml  # noqa: E402
from pptx.util import Emu as _cu_Emu  # noqa: E402
from xml.sax.saxutils import escape as _cu_esc  # noqa: E402

# 소스 456 PPT의 실측 색상(원본 run 구성): bg1=흰색, bg2 lumMod75000=회색, FFC000=주황
_CU_FILL_XML = {
    "white": '<a:schemeClr val="bg1"/>',
    "gray": '<a:schemeClr val="bg2"><a:lumMod val="75000"/></a:schemeClr>',
    "orange": '<a:srgbClr val="FFC000"/>',
}


def _cu_make_run_xml(text: str, color: str):
    xml = (
        '<a:r xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
        '<a:rPr lang="ko-KR"><a:solidFill>{fill}</a:solidFill></a:rPr>'
        "<a:t>{t}</a:t></a:r>"
    ).format(fill=_CU_FILL_XML[color], t=_cu_esc(text))
    return _cu_parse_xml(xml)


def _cu_build_slide_with_runs(runs):
    """runs: list[(text, color_key)] → (slide, para) 한 텍스트박스 단락에 그 run들을 넣는다."""
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank
    tb = slide.shapes.add_textbox(_cu_Emu(0), _cu_Emu(0), _cu_Emu(6000000), _cu_Emu(1000000))
    para = tb.text_frame.paragraphs[0]
    p = para._p
    for r in p.findall(_cu_qn("a:r")):
        p.remove(r)
    for text, color in runs:
        p.append(_cu_make_run_xml(text, color))
    return slide, para


def _cu_run_color_key(run):
    rpr = run._r.find(_cu_qn("a:rPr"))
    if rpr is None:
        return None
    sf = rpr.find(_cu_qn("a:solidFill"))
    if sf is None:
        return None
    scheme = sf.find(_cu_qn("a:schemeClr"))
    if scheme is not None:
        val = scheme.get("val")
        return {"bg1": "white", "bg2": "gray"}.get(val, val)
    srgb = sf.find(_cu_qn("a:srgbClr"))
    if srgb is not None:
        v = srgb.get("val")
        return "orange" if v == "FFC000" else v
    return None


def _cu_char_colors(para):
    """단락을 (글자, 색상키) 리스트로 펼친다 — run 경계를 넘어 글자 단위 색상을 본다."""
    out = []
    for run in para.runs:
        at = run._r.find(_cu_qn("a:t"))
        txt = (at.text if at is not None else "") or ""
        c = _cu_run_color_key(run)
        for ch in txt:
            out.append((ch, c))
    return out


def _cu_assert_label_number_colors(para, expected_text, label_char_idx, number_char_idx):
    cc = _cu_char_colors(para)
    text = "".join(ch for ch, _ in cc)
    assert text.startswith(expected_text), (text, expected_text)
    for i in label_char_idx:
        assert cc[i][1] == "white", ("라벨 글자가 흰색이 아님", i, cc[i], cc)
    for i in number_char_idx:
        assert cc[i][1] == "orange", ("숫자 글자가 주황색이 아님", i, cc[i], cc)


def test_성가헤더_라벨확장_헌글자_색상오염_없음_synthetic():
    """라벨이 늘어나며(2차봉헌→2차 봉헌) 구분자 회색 run 슬롯으로 '헌'이 밀려 회색이 되던
    버그의 재현·수정 확인. 수정 전엔 '헌'이 회색, 수정 후엔 라벨 전체가 흰색, 숫자는 주황."""
    slide, _ = _cu_build_slide_with_runs([
        ("2", "white"),
        ("차봉헌", "white"),
        (" ", "gray"),        # 원본 제작자 실수로 이 공백만 회색
        ("456", "orange"),
        ("  ", "orange"),
        ("둘이나 셋이 모인 곳에", "orange"),
    ])
    cu._update_성가_header(slide, "2차봉헌", 456)
    para = slide.shapes[0].text_frame.paragraphs[0]
    # "2차 봉헌 456": idx 0='2',1='차',2=' ',3='봉',4='헌',5=' ',6='4',7='5',8='6'
    _cu_assert_label_number_colors(para, "2차 봉헌 456", [0, 1, 3, 4], [6, 7, 8])


def test_성가헤더_라벨축소_숫자색상_침범없음_synthetic():
    """라벨이 줄어드는 반대 방향(2차 봉헌→봉헌). 통짜 재배치라면 숫자 '45'가 라벨 run(흰색)에
    흡수되고 '6'이 구분자 회색 run으로 새지만, 그룹 경계 재배치는 숫자를 전부 주황으로 유지."""
    slide, _ = _cu_build_slide_with_runs([
        ("2차 봉헌", "white"),
        (" ", "gray"),
        ("456", "orange"),
        (" 둘이나 셋이", "orange"),
    ])
    cu._update_성가_header(slide, "봉헌", 456)
    para = slide.shapes[0].text_frame.paragraphs[0]
    # "봉헌 456": idx 0='봉',1='헌',2=' ',3='4',4='5',5='6'
    _cu_assert_label_number_colors(para, "봉헌 456", [0, 1], [3, 4, 5])


def test_성가헤더_라벨동일_길이변화없음_회귀없음_synthetic():
    """기존 정상 케이스(라벨 길이 동일 '봉헌'→'봉헌', 숫자 62→100). 라벨은 흰색, 숫자는
    주황으로 유지되고 구분자 회색이 라벨/숫자로 새지 않는다."""
    slide, _ = _cu_build_slide_with_runs([
        ("봉헌", "white"),
        (" ", "gray"),
        ("62", "orange"),
        ("  둘이나", "orange"),
    ])
    cu._update_성가_header(slide, "봉헌", 100)
    para = slide.shapes[0].text_frame.paragraphs[0]
    # "봉헌 100": idx 0='봉',1='헌',2=' ',3='1',4='0',5='0'
    _cu_assert_label_number_colors(para, "봉헌 100", [0, 1], [3, 4, 5])


def test_성가헤더_run경계가_그룹경계와_어긋남_방어적_synthetic():
    """방어적 커버리지(regression-qa 승격 시 추가): 한 원본 run이 라벨 끝 글자+구분자+숫자
    첫 글자를 함께 담아 run 경계 != 정규식 그룹 경계인 배치. 라벨이 확장(2차봉헌→2차 봉헌)돼도
    (1) 텍스트가 완전 보존되고 (2) 숫자 전용 주황 run('56')이 라벨 길이 변화에 침범당하지
    않고 주황을 유지함을 확인한다. (straddle run 내부 글자색은 소스 rPr이 결정하므로 그 run에
    속한 '4'의 색은 단언하지 않는다 — 이 함수는 텍스트만 재배치하고 rPr을 보존하기 때문.)"""
    slide, _ = _cu_build_slide_with_runs([
        ("2차봉", "white"),
        ("헌 4", "white"),        # 라벨 끝 '헌' + 구분자 ' ' + 숫자 '4'가 한 run에 (straddle)
        ("56", "orange"),          # 숫자 전용 주황 run
        (" 둘이나", "orange"),
    ])
    cu._update_성가_header(slide, "2차봉헌", 456)
    para = slide.shapes[0].text_frame.paragraphs[0]
    cc = _cu_char_colors(para)
    text = "".join(ch for ch, _ in cc)
    # (1) 텍스트 완전 보존 (라벨 1글자 확장 반영)
    assert text == "2차 봉헌 456 둘이나", text
    # (2) 라벨 글자는 흰색 유지
    for i in (0, 1, 3, 4):  # '2','차','봉','헌'
        assert cc[i][1] == "white", ("라벨 글자 색 오염", i, cc[i], cc)
    # (3) 숫자 전용 주황 run('5','6')이 라벨 확장에 침범당하지 않고 주황 유지
    assert cc[7] == ("5", "orange"), (cc[7], cc)
    assert cc[8] == ("6", "orange"), (cc[8], cc)


def test_성가헤더_실제456소스_헌글자_흰색_realdata():
    """실측 소스 PPT(성가 456)로 end-to-end 확인 — '헌' 글자가 더 이상 별도 회색으로
    분리되지 않고 라벨 전체가 흰색, 숫자 456이 주황으로 유지된다."""
    src = Path(
        r"C:\Users\Scott\OneDrive\PPT 문서\09.가톨릭 성가\성가-악보버전"
        r"\성가 456 둘이나 셋이 모인 곳에.pptx"
    )
    if not src.exists():
        pytest.skip("소스 456 PPT 없음 — 실측 테스트 건너뜀")
    prs = Presentation(str(src))
    slide = prs.slides[0]
    cu._update_성가_header(slide, "2차봉헌", 456)
    rect = next(
        sh for sh in slide.shapes
        if sh.has_text_frame and sh.name == "Rectangle 11"
    )
    para = rect.text_frame.paragraphs[0]
    _cu_assert_label_number_colors(para, "2차 봉헌 456", [0, 1, 3, 4], [6, 7, 8])


# ═══════════════════════════════════════════════════════════════════════════
# 청년미사 2단계 — Track A(OOXML 통합): find_sections(mass_type)·is_sunday 4축 분리·
# 성가 5종×4출처 통합·복음 영문 레이아웃·조합 검증.
#
# 설계: _workspace/청년미사_2단계/01_architect_design.md §1~§6,§8(요약:
# docs/청년미사 2단계 구현 계획.md §5~§10). 슬라이드 인덱스는 참조 PPT
# reference/청년미사/Template_토요일 저녁 청년 주일미사_20260822_연중 제21주일.pptx(127슬라이드,
# 날짜 2026-08-22=토요일) 실측 고정 기대값이다.
#
# 이름 충돌 방지: 접두사 `y2_`를 유지한다(예: test_y2_F1_...).
#
# 2026-09-17 regression-qa 승격: 62 regression + 25 y2 progression 재확인(독립 재실행)
# 후 test_missa_progression.py에서 이 블록을 제거하고 이 파일로 옮겼다. 이 시점부터
# 이 테스트들은 "새 동작의 명세"가 아니라 "지켜야 할 기존 동작"이다.
# ═══════════════════════════════════════════════════════════════════════════
import missa_sections as sec2
import missa_content_updaters as cu2
import missa_reading_layout as rl2
import missa_to_ppt as mtp2
import missa_gui as gui2
import missa_youth_hymn_pdf as hp
from pptx import Presentation as _Y2Presentation
from pptx.oxml.ns import qn as _y2_qn

Y2_BASE = Path(__file__).resolve().parent.parent  # tests/ -> 저장소 루트
Y2_TEMPLATE = (
    Y2_BASE / "reference" / "청년미사"
    / "Template_토요일 저녁 청년 주일미사_20260822_연중 제21주일.pptx"
)
Y2_SAMPLE_HYMN = Y2_BASE / "reference" / "청년미사" / "나주노 성가 447 영원도하시어라 그 사랑이여.pptx"
Y2_DATE = "20260822"  # 실제 토요일 — is_calendar_sunday=False 검증용


def _y2_skip_if_missing():
    if not Y2_TEMPLATE.is_file():
        pytest.skip(f"청년미사 참조 PPT가 로컬에 없음(reference/청년미사/, git 미포함): {Y2_TEMPLATE}")


@pytest.fixture
def y2_prs():
    _y2_skip_if_missing()
    return _Y2Presentation(str(Y2_TEMPLATE))


# ---------------------------------------------------------------------------
# F1-F7: find_sections(mass_type='youth')
# ---------------------------------------------------------------------------

def test_y2_F1_입당_songs(y2_prs):
    sections = sec2.find_sections(y2_prs, mass_type='youth')
    assert sections['입당_songs'] == [
        {'title_idx': 12, 'content_start': 13, 'content_end': 17},
    ]


def test_y2_F2_봉헌_songs(y2_prs):
    sections = sec2.find_sections(y2_prs, mass_type='youth')
    assert sections['봉헌_songs'] == [
        {'title_idx': 65, 'content_start': 66, 'content_end': 70},
    ]


def test_y2_F3_성체_songs_two_entries(y2_prs):
    sections = sec2.find_sections(y2_prs, mass_type='youth')
    assert sections['성체_songs'] == [
        {'title_idx': 103, 'content_start': 104, 'content_end': 106},
        {'title_idx': 106, 'content_start': 107, 'content_end': 108},
    ]


def test_y2_F4_2차봉헌_songs_and_flat_alias(y2_prs):
    sections = sec2.find_sections(y2_prs, mass_type='youth')
    assert sections['2차봉헌_songs'] == [
        {'title_idx': 109, 'content_start': 110, 'content_end': 112},
    ]
    # insert_공지사항()이 참조하는 기존 flat 키가 첫 곡을 정확히 가리켜야 한다(하위 호환 alias).
    assert sections['2차봉헌_content_end'] == 112
    assert sections['2차봉헌_divider'] == 109
    assert sections['2차봉헌_content_start'] == 110


def test_y2_F5_파견_songs_not_confused_with_greeting_slide(y2_prs):
    """114번('파 견' 인사말, PLACEHOLDER)이 116번(진짜 title, AUTO_SHAPE)보다 먼저 잡히면 안 됨."""
    sections = sec2.find_sections(y2_prs, mass_type='youth')
    assert sections['파견_songs'] == [
        {'title_idx': 116, 'content_start': 117, 'content_end': 120},
    ]


def test_y2_F6_복음환호송_middle_slide_not_the_fixed_front(y2_prs):
    sections = sec2.find_sections(y2_prs, mass_type='youth')
    assert sections['복음환호송'] == 49


def test_y2_F7_adult_path_unaffected(y2_prs):
    """mass_type 생략(기본값 'adult')은 성가 _songs 키를 만들지 않는다 — 기존 flat 전용 계약 보존."""
    sections = sec2.find_sections(y2_prs)
    assert '입당_songs' not in sections
    assert '성체_songs' not in sections


# ---------------------------------------------------------------------------
# S1-S2: is_sunday 4축 분리 — resolve_mass_flags()
# ---------------------------------------------------------------------------

def test_y2_S1_adult_matches_legacy_single_flag():
    """adult 경로는 4개 값 전부 옛 단일 is_sunday(=is_calendar_sunday)와 일치해야 한다."""
    for date_str, expected_is_sunday in [("20260712", True), ("20260624", False)]:
        flags = mtp2.resolve_mass_flags('adult', date_str)
        assert flags['is_calendar_sunday'] == expected_is_sunday
        assert flags['use_weekday_display'] == (not expected_is_sunday)
        assert flags['include_2차봉헌'] == expected_is_sunday
        assert flags['copy_hymn_scores'] == expected_is_sunday


def test_y2_S2_youth_values_do_not_derive_from_calendar(monkeypatch):
    """youth는 ②③④ 전부 고정값이며, is_sunday_mass()가 우연히 True를 반환해도 안 바뀐다."""
    flags = mtp2.resolve_mass_flags('youth', Y2_DATE)
    assert flags['is_calendar_sunday'] is False  # 실제 토요일
    assert flags['use_weekday_display'] is True
    assert flags['include_2차봉헌'] is True
    assert flags['copy_hymn_scores'] is True

    # is_sunday_mass가 (버그 등으로) True를 반환해도 ②③④는 그대로여야 한다 —
    # "우연히 값이 같다"에 암묵적으로 기대지 않는다는 설계 원칙의 실제 회귀 가드.
    monkeypatch.setattr(mtp2, "is_sunday_mass", lambda d: True)
    flags2 = mtp2.resolve_mass_flags('youth', Y2_DATE)
    assert flags2['is_calendar_sunday'] is True
    assert flags2['use_weekday_display'] is True
    assert flags2['include_2차봉헌'] is True
    assert flags2['copy_hymn_scores'] is True


# ---------------------------------------------------------------------------
# H1-H7: 성가 5종×4출처 통합
# ---------------------------------------------------------------------------

def _y2_skip_if_sample_missing():
    if not Y2_SAMPLE_HYMN.is_file():
        pytest.skip(f"나주노 447 샘플 PPT가 로컬에 없음(reference/청년미사/, git 미포함): {Y2_SAMPLE_HYMN}")


def test_y2_H1_naju_resolves_and_header_reflects_htype(monkeypatch, tmp_path):
    """H1: resolve_youth_hymn_pptx가 (Presentation, 제목)을 반환하고, 반환된 각 슬라이드
    헤더가 '{타입} 나주노 {번호} {제목}' 형식이 된다."""
    _y2_skip_if_sample_missing()
    calls = []

    def fake_build(htype, 출처, 번호, title):
        calls.append((htype, 출처, 번호, title))
        return _Y2Presentation(str(Y2_SAMPLE_HYMN))

    monkeypatch.setattr(hp, "build_hymn_pptx", fake_build)
    monkeypatch.setattr(hp, "find_song_title", lambda source, number: "영원도 하시어라 그 사랑이여")
    # find_youth_onedrive_hymn_file()은 2026-09-27(2차)부터 로컬 동기화 폴더 검색을 완전히
    # 빼고 (1) 로컬 캐시(_SCRIPT_DIR/cache/...) → (2) OneDrive 실시간 검색만 한다.
    # _SCRIPT_DIR을 tmp_path로 돌려 캐시를 격리하고(초기엔 빈 캐시라 미스), OneDrive 검색은
    # 네트워크/로그인 없이 "못 찾음"으로 스텁한다 — 그래야 캐시-미스 → 신규 생성 경로를 탄다.
    # 생성된 파일의 저장 위치(_youth_onedrive_local_sync_dir, 검색과는 별개의 판단 지점)도
    # None으로 돌려 캐시 폴더에 저장되게 하고, 그 뒤의 OneDrive 업로드 시도도 스텁한다.
    monkeypatch.setattr(gui2, "_SCRIPT_DIR", tmp_path)
    monkeypatch.setattr(gui2, "_find_onedrive_file", lambda remote_path, pattern: None)
    monkeypatch.setattr(gui2, "_youth_onedrive_local_sync_dir", lambda local_key: None)
    import missa_onedrive as od_h1
    monkeypatch.setattr(od_h1, "upload_file", lambda *a, **k: None)

    entry = {'출처': '나주노', '번호': 447, '제목': None}
    src_prs, resolved_title = cu2.resolve_youth_hymn_pptx(entry, '입당')

    assert resolved_title == "영원도 하시어라 그 사랑이여"
    assert len(calls) == 1  # 캐시에 없으므로 1회 생성
    for slide in src_prs.slides:
        hdr_text = hp._find_header(slide).text_frame.text
        assert hdr_text == "입당 나주노 447 영원도 하시어라 그 사랑이여"


def test_y2_H2_cache_reuse_skips_regeneration(monkeypatch, tmp_path):
    """H2: 같은 곡이 OneDrive 폴더에 이미 저장돼 있으면 build_hymn_pptx()가 다시 호출되지 않음."""
    _y2_skip_if_sample_missing()
    calls = []

    def fake_build(htype, 출처, 번호, title):
        calls.append((htype, 출처, 번호, title))
        return _Y2Presentation(str(Y2_SAMPLE_HYMN))

    monkeypatch.setattr(hp, "build_hymn_pptx", fake_build)
    monkeypatch.setattr(hp, "find_song_title", lambda source, number: "영원도 하시어라 그 사랑이여")
    # H1과 동일한 격리. 저장 위치와 검색 캐시 위치가 (local_sync_dir=None이라) 둘 다
    # _SCRIPT_DIR/cache/청년미사_성가/나주노 성가로 일치해야 두 번째 호출이 캐시를 찾는다.
    monkeypatch.setattr(gui2, "_SCRIPT_DIR", tmp_path)
    monkeypatch.setattr(gui2, "_find_onedrive_file", lambda remote_path, pattern: None)
    monkeypatch.setattr(gui2, "_youth_onedrive_local_sync_dir", lambda local_key: None)
    import missa_onedrive as od_h2
    monkeypatch.setattr(od_h2, "upload_file", lambda *a, **k: None)

    entry = {'출처': '나주노', '번호': 447, '제목': None}
    cu2.resolve_youth_hymn_pptx(entry, '입당')
    assert len(calls) == 1

    cu2.resolve_youth_hymn_pptx(entry, '봉헌')  # 같은 곡, 다른 용도로 재호출
    assert len(calls) == 1  # 캐시 재사용 — 재생성 없음


def test_y2_H3_catholic_hymn_header_untouched(monkeypatch, tmp_path):
    """H3: 가톨릭성가 경로는 헤더를 건드리지 않은 원본 Presentation을 반환한다
    (헤더 갱신은 호출부의 _update_성가_header() 책임)."""
    fake_prs = _Y2Presentation()
    slide = fake_prs.slides.add_slide(fake_prs.slide_layouts[6])
    box = slide.shapes.add_textbox(0, 0, 1000000, 500000)
    box.text_frame.text = "성체 161       어떤 제목"

    # find_youth_onedrive_hymn_file(subfolder=None)은 이제 로컬 동기화 폴더를 아예 보지
    # 않고 로컬 캐시(_SCRIPT_DIR/cache/청년_가톨릭성가/)부터 뒤진다 — 파일을 거기 바로
    # 둔다. OneDrive 검색까지 가면(버그가 있을 때만) 즉시 실패하도록 스텁해 조용히
    # 네트워크를 타는 회귀를 막는다.
    dest_dir = tmp_path / "cache" / "청년_가톨릭성가"
    dest_dir.mkdir(parents=True)
    fake_prs.save(str(dest_dir / "성가 161 어떤 제목.pptx"))

    monkeypatch.setattr(gui2, "_SCRIPT_DIR", tmp_path)

    def _no_network(remote_path, pattern):
        raise AssertionError("캐시 히트를 기대하는 테스트에서 OneDrive 검색이 호출됨")

    monkeypatch.setattr(gui2, "_find_onedrive_file", _no_network)

    entry = {'출처': '가톨릭성가', '번호': 161, '제목': None}
    src_prs, _resolved_title = cu2.resolve_youth_hymn_pptx(entry, '성체')

    assert src_prs.slides[0].shapes[0].text_frame.text == "성체 161       어떤 제목"


def _y2_shape_text(slide, name):
    for shape in slide.shapes:
        if shape.name == name:
            return shape.text_frame.text
    return None


def test_y2_H4_기타_returns_none_and_skips_content(y2_prs):
    """H4: 출처='기타'는 (None, 제목)을 반환하고, _replace_one_youth_song()이 콘텐츠
    슬라이드를 전혀 삭제/삽입하지 않는다(placeholder 텍스트·슬라이드 개수 불변)."""
    entry = {'출처': '기타', '번호': None, '제목': '새 제목'}
    src_prs, resolved_title = cu2.resolve_youth_hymn_pptx(entry, '성체')
    assert src_prs is None
    assert resolved_title == '새 제목'

    n_before = len(y2_prs.slides)
    placeholder_before = _y2_shape_text(y2_prs.slides[107], 'Text Placeholder 2')
    assert placeholder_before is not None and '성가 가사를 여기에 입력' in placeholder_before

    cu2._replace_one_youth_song(
        y2_prs, '성체',
        {'title_idx': 106, 'content_start': 107, 'content_end': 108},
        {'출처': '기타', '번호': None, '제목': '새 제목'},
    )

    assert len(y2_prs.slides) == n_before
    placeholder_after = _y2_shape_text(y2_prs.slides[107], 'Text Placeholder 2')
    assert placeholder_after is not None and '성가 가사를 여기에 입력' in placeholder_after


def test_y2_H5_기타_title_and_header_updated_placeholder_preserved(y2_prs):
    """H5: _replace_one_youth_song 실행 후 title 슬라이드 텍스트가 새 제목으로 바뀌고,
    콘텐츠 헤더도 갱신되지만 가사 placeholder는 그대로."""
    cu2._replace_one_youth_song(
        y2_prs, '성체',
        {'title_idx': 106, 'content_start': 107, 'content_end': 108},
        {'출처': '기타', '번호': None, '제목': '그 사랑 노래'},
    )
    title_text = _y2_shape_text(y2_prs.slides[106], 'Rectangle 7')
    assert title_text is not None and '그 사랑 노래' in title_text

    header = _y2_shape_text(y2_prs.slides[107], 'Rectangle 11')
    placeholder = _y2_shape_text(y2_prs.slides[107], 'Text Placeholder 2')
    assert header == '성체 그 사랑 노래'
    assert placeholder == '성가 가사를 여기에 입력'


def test_y2_H6_title_rebuild_no_soft_linebreak_leftover(y2_prs):
    """H6: 109번(2차봉헌, 원래 \x0b 내부 줄바꿈 2단락)에 적용 후 독립된 <a:p> 3개로
    재구성되고 \x0b가 남지 않는다."""
    slide = y2_prs.slides[109]
    cu2._update_youth_title_slide(slide, '2차봉헌', '야훼 이레', 810, '주의 자비가 내려와')

    shape = sec2._find_youth_title_shape(slide, '2차 봉헌')
    paras = shape.text_frame._txBody.findall(_y2_qn('a:p'))
    assert len(paras) == 3
    assert '\x0b' not in shape.text_frame.text
    texts = [p_.text for p_ in shape.text_frame.paragraphs]
    assert texts[0].strip().replace(' ', '') == '2차봉헌'
    assert texts[1] == '주의 자비가 내려와'
    assert '810' in texts[2]


def test_y2_H7_title_rebuild_기타_two_paragraphs_only(y2_prs):
    """H7: 106번(기타, 번호 없음)에 적용 후 정확히 2개 단락만 존재(출처문장 단락 없음)."""
    slide = y2_prs.slides[106]
    cu2._update_youth_title_slide(slide, '성체', '기타', None, '그대곁에 주님이')

    shape = sec2._find_youth_title_shape(slide, '성 체')
    paras = shape.text_frame._txBody.findall(_y2_qn('a:p'))
    assert len(paras) == 2
    texts = [p_.text for p_ in shape.text_frame.paragraphs]
    assert texts[1] == '그대곁에 주님이'


# ---------------------------------------------------------------------------
# G1-G6: 복음 영문 레이아웃
# ---------------------------------------------------------------------------

Y2_EN_GOSPEL_JSON = (
    Y2_BASE / "_workspace" / "청년미사_1단계" / "verify_20260913" / "missa_en_20260913.json"
)


def _y2_load_en_gospel_content():
    if not Y2_EN_GOSPEL_JSON.is_file():
        pytest.skip(f"20260913 영문 복음 JSON이 로컬에 없음: {Y2_EN_GOSPEL_JSON}")
    import json as _json
    data = _json.loads(Y2_EN_GOSPEL_JSON.read_text(encoding="utf-8"))
    return data["Gospel"]["content"]


_G1B_GOSPEL_20260926 = (
    'Jesus said to the chief priests and elders of the people, ‘What is your opinion? '
    'A man had two sons. He went and said to the first, “My boy, you go and work in the '
    'vineyard today.” He answered, “I will not go,” but afterwards thought better '
    'of it and went. The man then went and said the same thing to the second who answered, '
    '“Certainly, sir,” but did not go. Which of the two did the father’s will?’ '
    '‘The first’ they said. Jesus said to them, ‘I tell you solemnly, tax collectors '
    'and prostitutes are making their way into the kingdom of God before you. For John came to '
    'you, a pattern of true righteousness, but you did not believe him, and yet the tax '
    'collectors and prostitutes did. Even after seeing that, you refused to think better of it '
    'and believe in him.’'
)


def test_y2_G1b_20260926_father_fits_first_page(y2_prs):
    """G1b(2026-10-03 회귀 승격, 같은 날 슬라이드 단위 채우기에 맞춰 재작성): 실사용자 20260926 청년미사
    슬라이드 54 재현. 분할(줄 시작 계산)이 검증과 같은 박스 폭을 써야 한다 — 과거 `_PIL_WRAP_SAFETY=0.97`이
    페이지네이션만 3% 좁혀 실제로는 그 줄에 들어가는 "father's"를 다음 슬라이드로 밀어냈다. 그 실행이 실제로
    만든 영문 복음 상자 폭 949.98px는 실제 PowerPoint COM 실측으로 "...did the father's will?'"까지 들어간다고
    확인된 사실이다(docs/ooxml-pitfalls-log.md). gitignore된 output/ 산출물에 기대지 않도록 본문과 폭을
    상수로 고정하고, Pillow 폰트만 청년 참조 템플릿의 복음 슬라이드에서 얻는다(실 COM 대응 테스트:
    test_missa_progression.py test_sf_real4)."""
    sections = sec2.find_sections(y2_prs, mass_type='youth')
    pil_font, _template_box_px = rl2._get_slide_render_params(y2_prs.slides[sections['복음_start']])
    if pil_font is None:
        pytest.skip("Pillow 또는 폰트 파일(batang.ttc)을 찾을 수 없는 환경")

    paras = rl2._reading_paras_from_units(rl2.parse_into_verse_units(_G1B_GOSPEL_20260926))
    starts = rl2._pil_line_starts(paras, pil_font, 949.98)
    assert len(starts) > rl2.LINES_PER_SLIDE, "본문이 한 슬라이드에 다 들어가 절단점이 없음"
    ninth_pi, ninth_off = starts[rl2.LINES_PER_SLIDE - 1]
    tenth_pi, tenth_off = starts[rl2.LINES_PER_SLIDE]
    assert ninth_pi == tenth_pi == 0
    last_line = paras[0]['text'][ninth_off:tenth_off]
    assert 'father' in last_line, f"첫 슬라이드 마지막 줄이 'father's' 앞에서 끊김: {last_line!r}"


def test_y2_G1_pil_estimate_not_worse_than_char_based(y2_prs):
    """G1: 영문 복음 콘텐츠에 대해 Pillow 폭 기반 줄 수가 CHARS_PER_LINE 기반 줄 수보다 적거나 같아야 한다
    (영문에 한글 전각 기준 27자/줄을 쓰면 슬라이드 수를 과대추정하는 문제의 회귀 방지 — 구
    layout_units_on_slides_pil vs layout_units_on_slides 비교를 줄 시작 계산 비교로 이전)."""
    content = _y2_load_en_gospel_content()
    units = rl2.parse_into_verse_units(content)

    sections = sec2.find_sections(y2_prs, mass_type='youth')
    content_slide = y2_prs.slides[sections['복음_start']]
    pil_font, box_px = rl2._get_slide_render_params(content_slide)
    if pil_font is None:
        pytest.skip("Pillow 또는 폰트 파일(batang.ttc)을 찾을 수 없는 환경")

    paras = rl2._reading_paras_from_units(units)
    assert len(rl2._pil_line_starts(paras, pil_font, box_px)) <= len(rl2._char_line_starts(paras))


def test_y2_G3_G4_no_orange_runs_and_korean_closing_preserved(y2_prs):
    """G3: 영문 콘텐츠에는 절 번호가 없어 오렌지 run이 생성되지 않는다.
    G4: 56번(한글 클로징 TYPE B 종료 슬라이드)이 그대로 보존된다."""
    content = _y2_load_en_gospel_content()
    units = rl2.parse_into_verse_units(content)

    sections = sec2.find_sections(y2_prs, mass_type='youth')
    content_slide = y2_prs.slides[sections['복음_start']]
    pil_font, box_px = rl2._get_slide_render_params(content_slide)
    if pil_font is None:
        pytest.skip("Pillow 또는 폰트 파일(batang.ttc)을 찾을 수 없는 환경")

    closing_text_before = _y2_shape_text(y2_prs.slides[56], 'TextBox 3')
    assert closing_text_before is not None and '주님의 말씀입니다' in closing_text_before

    rl2.replace_reading_slides(
        y2_prs, sections['복음_start'], sections['복음_end'],
        units, sections['복음_start'], label='복음_영문',
    )

    sections2 = sec2.find_sections(y2_prs, mass_type='youth')
    # G3: 새 복음 범위 안에 오렌지색 run이 하나도 없어야 한다.
    from missa_reading_layout import ORANGE as _Y2_ORANGE
    orange_found = []
    for idx in range(sections2['복음_start'], sections2['복음_end']):
        for shape in y2_prs.slides[idx].shapes:
            if not shape.has_text_frame:
                continue
            for para in shape.text_frame.paragraphs:
                for run in para.runs:
                    try:
                        if run.font.color.rgb == _Y2_ORANGE:
                            orange_found.append((idx, run.text))
                    except Exception:
                        continue
    assert not orange_found, f"영문 복음 슬라이드에 오렌지 run 발견: {orange_found}"

    # G4: 한글 클로징 슬라이드가 TYPE B 종료 슬라이드로 보존됐는지(텍스트 그대로) 확인.
    # 병합/재배치로 인덱스가 이동할 수 있으므로 텍스트로 재탐색한다.
    closing_found = any(
        _y2_shape_text(y2_prs.slides[i], 'TextBox 3') is not None
        and '주님의 말씀입니다' in (_y2_shape_text(y2_prs.slides[i], 'TextBox 3') or '')
        for i in range(sections2['복음_start'], sections2['복음_end'] + 2)
    )
    assert closing_found, "한글 클로징(TYPE B 종료 슬라이드)을 찾지 못함"


def test_y2_G6_title_stays_korean_after_content_merge():
    """G6: merge_youth_gospel_content()는 title을 한글 JSON 값 그대로 두고 content만 교체한다."""
    json_data = {'복음': {'title': '마태오가 전한 거룩한 복음입니다.', 'content': '(한글 본문)'}}
    en_mass = {'Gospel': {'reference': 'Matthew 18:21-35', 'content': '(English body)'}}
    merged = mtp2.merge_youth_gospel_content(json_data, en_mass)
    assert merged['복음']['title'] == '마태오가 전한 거룩한 복음입니다.'
    assert merged['복음']['content'] == '(English body)'
    # 원본은 변경되지 않아야 한다(호출부가 다른 곳에서 json_data를 그대로 쓸 수 있으므로).
    assert json_data['복음']['content'] == '(한글 본문)'


@pytest.mark.skipif(not gui2._com_verification_enabled(), reason="config.json: com_verification_enabled=false")
def test_y2_G2_com_confirms_convergence_to_lines_per_slide(y2_prs):
    """G2: 슬라이드 단위 채우기 후 실제 PowerPoint COM 실측으로 마지막을 제외한 모든 복음 슬라이드가 정확히
    LINES_PER_SLIDE(9)줄이고 마지막은 그 이하인지 확인한다(구 구조는 "9줄 이하"까지만 보장 — 추정 오차를 이웃
    이동으로 메우던 사후 재조정의 한계였다)."""
    try:
        import ppt_com_verify as com2
    except ImportError:
        pytest.skip("pywin32 미설치")
    if not com2.is_available():
        pytest.skip("PowerPoint COM 연결 불가")

    content = _y2_load_en_gospel_content()
    units = rl2.parse_into_verse_units(content)
    sections = sec2.find_sections(y2_prs, mass_type='youth')
    rl2.replace_reading_slides(
        y2_prs, sections['복음_start'], sections['복음_end'],
        units, sections['복음_start'], label='복음_영문',
    )

    sections2 = sec2.find_sections(y2_prs, mass_type='youth')
    s, e = sections2['복음_start'], sections2['복음_end']
    from missa_ooxml_utils import _find_content_shape as _y2_find_content_shape
    from missa_ooxml_utils import _build_com_probe_pptx as _y2_build_probe
    counts = []
    for idx in range(s, e):
        slide = y2_prs.slides[idx]
        if _y2_find_content_shape(slide) is None:
            continue
        probe_path, shape_idx = _y2_build_probe(y2_prs, slide)
        counts.append(com2.count_slide_lines(str(probe_path), shape_idx))
    assert counts, "복음 본문 슬라이드가 없음"
    problems = [f"본문 슬라이드 {i + 1}: COM 실측 {c}줄 (기대 {rl2.LINES_PER_SLIDE}줄)"
                for i, c in enumerate(counts[:-1]) if c != rl2.LINES_PER_SLIDE]
    if counts[-1] > rl2.LINES_PER_SLIDE:
        problems.append(f"마지막 슬라이드 COM 실측 {counts[-1]}줄 (>{rl2.LINES_PER_SLIDE})")
    assert not problems, "\n".join(problems)


# ---------------------------------------------------------------------------
# C1-C3: 5종×4출처 조합 검증
# ---------------------------------------------------------------------------

def test_y2_C1_real_combo_full_pipeline_structure_valid(y2_prs, monkeypatch, tmp_path):
    """C1: 이번 템플릿의 실사용 5개 조합(입당=나주노447/봉헌=나주노173/성체1=가톨릭성가161/
    성체2=기타/2차봉헌=야훼이레810/파견=나주노267)으로 성가 통합+복음 영문 레이아웃을 실행해
    최종 PPT가 validate_pptx_structure() 구조 검증을 통과하는지 확인한다.

    find_youth_onedrive_hymn_file()은 2026-09-27(2차)부터 로컬 동기화 폴더 검색을 완전히
    빼고 로컬 캐시 → OneDrive 실시간 검색만 한다 — 실제 OneDrive 계정에 매주 쌓여있는
    161/173/267/810은(§읽기 전용 실측 유지) 네트워크 없이 실제 파일 내용 그대로 캐시에
    미리 복사해 두고, 원래 로컬에 없던 447만 실제 build_hymn_pptx() PDF 파이프라인을
    그대로 거치게 해 "실사용 5개 조합"의 실측 취지를 그대로 보존한다. OneDrive 원격 검색과
    실패 시 업로드는 네트워크/로그인 없이 끝나도록 스텁한다."""
    config = gui2._load_config()
    real_hymn_dir = config.get('onedrive_hymn_folder')
    real_youth_dir = config.get('onedrive_youth_hymn_folder')
    if not real_hymn_dir or not Path(real_hymn_dir).is_dir() \
            or not real_youth_dir or not Path(real_youth_dir).is_dir():
        pytest.skip("OneDrive 로컬 동기화 폴더가 이 환경에 설정돼 있지 않음")

    import shutil as _c1_shutil

    def _seed_cache(cache_subdir, src_dir, pattern):
        matches = [f for f in Path(src_dir).rglob('*.pptx')
                   if not f.name.startswith('~$') and re.search(pattern, f.name)]
        if not matches:
            pytest.skip(f"실측 대상 파일을 찾지 못함: {src_dir} / {pattern}")
        dest_dir = tmp_path / 'cache' / cache_subdir
        dest_dir.mkdir(parents=True, exist_ok=True)
        _c1_shutil.copy(str(matches[0]), str(dest_dir / matches[0].name))

    나주노_sub = hp.SOURCES['나주노']['onedrive_subfolder']
    야훼이레_sub = hp.SOURCES['야훼 이레']['onedrive_subfolder']
    _seed_cache('청년_가톨릭성가', real_hymn_dir, r'성가 161(?!\d)')
    _seed_cache(f'청년미사_성가/{나주노_sub}', Path(real_youth_dir) / 나주노_sub, r'^나주노 성가 173 ')
    _seed_cache(f'청년미사_성가/{나주노_sub}', Path(real_youth_dir) / 나주노_sub, r'^나주노 성가 267 ')
    _seed_cache(f'청년미사_성가/{야훼이레_sub}', Path(real_youth_dir) / 야훼이레_sub, r'^야훼 이레 성가 810 ')

    monkeypatch.setattr(gui2, "_SCRIPT_DIR", tmp_path)
    monkeypatch.setattr(gui2, "_find_onedrive_file", lambda remote_path, pattern: None)
    monkeypatch.setattr(gui2, "_youth_onedrive_local_sync_dir", lambda local_key: None)
    import missa_onedrive as od_c1
    monkeypatch.setattr(od_c1, "upload_file", lambda *a, **k: None)

    성가_선택 = {
        '입당': [{'출처': '나주노', '번호': 447, '제목': None}],
        '봉헌': [{'출처': '나주노', '번호': 173, '제목': None}],
        '성체': [
            {'출처': '가톨릭성가', '번호': 161, '제목': None},
            {'출처': '기타', '번호': None, '제목': '그대곁에 주님이'},
        ],
        '2차봉헌': [{'출처': '야훼 이레', '번호': 810, '제목': None}],
        '파견': [{'출처': '나주노', '번호': 267, '제목': None}],
    }
    cu2.replace_성가_youth(y2_prs, 성가_선택)

    content = _y2_load_en_gospel_content()
    units = rl2.parse_into_verse_units(content)
    sections = sec2.find_sections(y2_prs, mass_type='youth')
    content_slide = y2_prs.slides[sections['복음_start']]
    rl2.replace_reading_slides(
        y2_prs, sections['복음_start'], sections['복음_end'],
        units, sections['복음_start'], label='복음_영문',
    )

    out_path = tmp_path / "y2_c1_check.pptx"
    y2_prs.save(str(out_path))
    issues = sec2.validate_pptx_structure(str(out_path))
    assert not issues, issues

    # 헤더가 실제로 이번 주 구분/출처로 갱신됐는지도 가볍게 확인(무결성 검증에 그치지 않기 위함).
    sections2 = sec2.find_sections(y2_prs, mass_type='youth')
    header = _y2_shape_text(y2_prs.slides[sections2['입당_songs'][0]['content_start']], 'Rectangle 11')
    assert header is not None and header.startswith('입당 나주노 447')


def _y2_make_fake_catholic_hymn_pptx(dest_path, header_text):
    """실제 가톨릭성가 콘텐츠 슬라이드 구조(헤더 'Rectangle 11' + 악보 Picture)를 최소한으로
    흉내 낸 1슬라이드 PPTX를 만든다. `_find_hymn_songs_youth()`의 콘텐츠 경계 판정이 PICTURE
    도형 유무로 이뤄지므로(§1.4), Picture 없이 텍스트박스만 두면 재탐색 시 콘텐츠 범위가
    0으로 판정돼 후속 로직(2번째 곡 탐지 등)이 조용히 깨진다 — 실측(C2 최초 실패)으로 발견."""
    from PIL import Image as _Y2Image
    fake_prs = _Y2Presentation()
    slide = fake_prs.slides.add_slide(fake_prs.slide_layouts[6])
    box = slide.shapes.add_textbox(0, 0, 3000000, 500000)
    box.name = 'Rectangle 11'
    box.text_frame.text = header_text
    img_path = dest_path.parent / '_y2_dummy.png'
    if not img_path.exists():
        _Y2Image.new('RGB', (10, 10), color='black').save(img_path)
    slide.shapes.add_picture(str(img_path), 0, 600000, 1000000, 1000000)
    fake_prs.save(str(dest_path))


def test_y2_C2_성체_second_song_omitted_deletes_extra_slides(monkeypatch, tmp_path):
    """C2: 성체 2번째 곡 입력을 생략하면 106~107(title2+content2)이 삭제되고, 108번(구분)
    바로 뒤에 2차봉헌 title이 오도록 슬라이드가 재배치된다(기존 성인미사 '성체 1곡'과 동일
    동작 — 요구사항 §5.1)."""
    _y2_skip_if_missing()
    prs = _Y2Presentation(str(Y2_TEMPLATE))

    # find_youth_onedrive_hymn_file()은 로컬 동기화 폴더를 보지 않고 로컬 캐시부터
    # 뒤진다(§H3와 동일 이유) — 파일을 그 캐시 경로에 둔다.
    dest_dir = tmp_path / "cache" / "청년_가톨릭성가"
    dest_dir.mkdir(parents=True)
    _y2_make_fake_catholic_hymn_pptx(
        dest_dir / "성가 161 성체를 찬송하세.pptx", "성체 161       성체를 찬송하세",
    )
    monkeypatch.setattr(gui2, "_SCRIPT_DIR", tmp_path)

    n_before = len(prs.slides)
    sec_before = sec2.find_sections(prs, mass_type='youth')
    assert len(sec_before['성체_songs']) == 2

    원래_song1_content = sec_before['성체_songs'][0]['content_end'] - sec_before['성체_songs'][0]['content_start']
    원래_song2_slides = (
        sec_before['성체_songs'][1]['content_end'] - sec_before['성체_songs'][1]['title_idx']
    )

    성가_선택 = {'성체': [{'출처': '가톨릭성가', '번호': 161, '제목': None}]}
    cu2.replace_성가_youth(prs, 성가_선택)

    # song1 콘텐츠는 기존 슬라이드 수와 무관하게 새 소스(가짜 1슬라이드)로 교체되고,
    # song2(title+content)는 통째로 삭제된다 — 순가감 = (새 소스 1장 - 기존 song1 콘텐츠 수)
    # - song2 전체 슬라이드 수.
    expected_delta = 1 - 원래_song1_content - 원래_song2_slides
    assert len(prs.slides) == n_before + expected_delta

    sec_after = sec2.find_sections(prs, mass_type='youth')
    assert len(sec_after['성체_songs']) == 1
    # 구분 슬라이드 바로 뒤에 2차봉헌 title이 와야 한다.
    divider_idx = sec_after['성체_songs'][0]['content_end']
    assert sec_after['2차봉헌_songs'][0]['title_idx'] == divider_idx + 1


def test_y2_C3_synthetic_combo_no_constraint(monkeypatch, tmp_path):
    """C3: 조합표에 명시적 제약이 없다는 가정을 검증 — 입당=가톨릭성가, 파견=기타처럼
    설계서 실사례와 다른 조합도 예외 없이 처리된다."""
    _y2_skip_if_missing()
    prs = _Y2Presentation(str(Y2_TEMPLATE))

    dest_dir = tmp_path / "cache" / "청년_가톨릭성가"
    dest_dir.mkdir(parents=True)
    _y2_make_fake_catholic_hymn_pptx(
        dest_dir / "성가 447 영원도 하시어라 그 사랑이여.pptx",
        "입당 447       영원도 하시어라 그 사랑이여",
    )
    monkeypatch.setattr(gui2, "_SCRIPT_DIR", tmp_path)

    성가_선택 = {
        '입당': [{'출처': '가톨릭성가', '번호': 447, '제목': None}],
        '파견': [{'출처': '기타', '번호': None, '제목': '자유곡'}],
    }
    cu2.replace_성가_youth(prs, 성가_선택)  # 예외 없이 완료되어야 한다

    sections = sec2.find_sections(prs, mass_type='youth')
    header = _y2_shape_text(prs.slides[sections['입당_songs'][0]['content_start']], 'Rectangle 11')
    assert header is not None and header.startswith('입당 447')
    파견_title = _y2_shape_text(prs.slides[sections['파견_songs'][0]['title_idx']], 'Rectangle 7')
    assert 파견_title is not None and '자유곡' in 파견_title


# ---------------------------------------------------------------------------
# 리뷰 발견 사항 회귀 가드 (02b_review_report.md)
# ---------------------------------------------------------------------------

def test_y2_R1_2차봉헌_가톨릭성가_header_label_no_space(monkeypatch, tmp_path):
    """독립 리뷰 확정 버그: 2차봉헌에 가톨릭성가를 선택하면 _update_성가_header()가 성인미사
    관례(공백 있는 '2차 봉헌')로 라벨을 정규화해, 같은 슬라이드에서 나주노/야훼이레 경로가
    쓰는 공백 없는 '2차봉헌' 표기와 어긋났다. label_override로 청년 라벨을 강제해야 한다."""
    _y2_skip_if_missing()
    prs = _Y2Presentation(str(Y2_TEMPLATE))

    dest_dir = tmp_path / "cache" / "청년_가톨릭성가"
    dest_dir.mkdir(parents=True)
    _y2_make_fake_catholic_hymn_pptx(
        dest_dir / "성가 810 주의 자비가 내려와.pptx", "2차 봉헌 810       주의 자비가 내려와",
    )
    monkeypatch.setattr(gui2, "_SCRIPT_DIR", tmp_path)

    성가_선택 = {'2차봉헌': [{'출처': '가톨릭성가', '번호': 810, '제목': None}]}
    cu2.replace_성가_youth(prs, 성가_선택)

    sections = sec2.find_sections(prs, mass_type='youth')
    header = _y2_shape_text(
        prs.slides[sections['2차봉헌_songs'][0]['content_start']], 'Rectangle 11',
    )
    assert header is not None
    assert header.startswith('2차봉헌 810'), f"공백 없는 '2차봉헌' 표기여야 함: {header!r}"
    assert '2차 봉헌' not in header, f"성인미사 공백 표기가 섞이면 안 됨: {header!r}"


def test_y2_R2_성체_second_slot_non_기타_source_line_font_size(y2_prs):
    """리뷰 권장 사항: 성체 2번째 title 슬롯(106번)은 sz=2000(출처문장) run 템플릿이 원래
    없다('기타' 전용 슬라이드였으므로). 다른 출처(나주노 등)를 이 슬롯에 선택하면 출처문장이
    제목과 같은 큰 글자로 렌더링되던 것을, sz 강제 교정으로 2000pt를 유지하는지 확인."""
    cu2._update_youth_title_slide(y2_prs.slides[106], '성체', '나주노', 146, '꽃')

    shape = sec2._find_youth_title_shape(y2_prs.slides[106], '성 체')
    paras = shape.text_frame.paragraphs
    assert len(paras) == 3
    source_para_el = paras[2]._p
    run_el = source_para_el.find(_y2_qn('a:r'))
    rPr = run_el.find(_y2_qn('a:rPr'))
    assert rPr.get('sz') == '2000', f"출처문장 폰트 크기가 2000이어야 함: {rPr.get('sz')!r}"
    assert '146' in paras[2].text


# ═══════════════════════════════════════════════════════════════════════════
# 청년미사 2단계 — main() 통합 배선 (Track A 이후 마지막 통합 작업)
#
# find_sections(mass_type='youth')/resolve_mass_flags()/replace_성가_youth()/
# layout_units_on_slides_pil() 등 개별 함수는 위 Track A 섹션에서 이미 구현·검증됐지만,
# missa_to_ppt.py의 main()이 이 중 어느 것도 실제로 호출하지 않던(mass_type='adult' 고정)
# 마지막 연결점을 배선한 라운드. main()이 mass_type에 따라 올바른 함수를 올바른 인자로
# 호출하는지(스파이) + 청년 경로 전체가 실제로 유효한 pptx를 만들어내는지(구조 검증)를
# 검증한다. 성가 입력 팝업을 확인 없이 닫으면(X/Alt+F4) 무경고 빈 성가_선택이 흘러들어가
# 이번 주 성가가 하나도 갱신되지 않은 PPT가 "완료"로 저장되던 확정 버그(독립 리뷰가 실제
# 참조 PPT로 재현, `_workspace/청년미사_2단계/05_review_report.md`)에 대한 이중 방어(팝업
# RuntimeError + main() 방어 게이트) 회귀 가드도 포함한다.
#
# 이름 충돌 방지: 접두사 `y2m_`("y2 main integration")를 유지한다.
#
# 2026-09-17 regression-qa 승격: 87 regression + 9 y2m progression 재확인(독립 재실행) 후
# test_missa_progression.py에서 이 블록을 제거하고 이 파일로 옮겼다. 이 시점부터 이
# 테스트들은 "새 동작의 명세"가 아니라 "지켜야 할 기존 동작"이다.
# ═══════════════════════════════════════════════════════════════════════════
import shutil
import sys as _y2m_sys
import missa_to_ppt as mtp

Y2M_BASE = Path(__file__).resolve().parent.parent  # tests/ -> 저장소 루트
Y2M_YOUTH_TEMPLATE = (
    Y2M_BASE / "reference" / "청년미사"
    / "Template_토요일 저녁 청년 주일미사_20260822_연중 제21주일.pptx"
)
Y2M_EN_GOSPEL_FIXTURE = (
    Y2M_BASE / "_workspace" / "청년미사_1단계" / "verify_20260913" / "missa_en_20260913.json"
)
Y2M_KO_JSON_20260913 = Y2M_BASE / "output" / "20260913" / "missa_20260913.json"


def _y2m_skip_if_missing():
    if not Y2M_YOUTH_TEMPLATE.is_file():
        pytest.skip(f"청년미사 참조 PPT가 로컬에 없음: {Y2M_YOUTH_TEMPLATE}")
    if not Y2M_EN_GOSPEL_FIXTURE.is_file():
        pytest.skip(f"영문 복음 픽스처가 로컬에 없음: {Y2M_EN_GOSPEL_FIXTURE}")
    if not Y2M_KO_JSON_20260913.is_file():
        pytest.skip(f"20260913 한글 JSON이 로컬에 없음: {Y2M_KO_JSON_20260913}")


def test_y2m_1_resolve_mass_type_kr_to_en():
    """한글 팝업 반환값('성인'/'청년')이 파이프라인 전체가 쓰는 영문 키로 정확히 변환된다.
    이 변환이 main()에 배선돼 있지 않았던 것이 이번 라운드 전 유일한 미완성 연결점이었다
    (구현 계획 §13-3)."""
    assert mtp._resolve_mass_type('성인') == 'adult'
    assert mtp._resolve_mass_type('청년') == 'youth'
    # 알 수 없는 값은 KeyError로 파이프라인 전체를 죽이는 대신 안전하게 'adult'로 폴백한다.
    assert mtp._resolve_mass_type('???') == 'adult'


def test_y2m_2_youth_content_date_str_is_plus_one_day():
    """요구사항 §6.1: 사용자가 입력한 토요일 date_str에서 +1일(일요일)이 콘텐츠 조회일이다.
    월/연 경계도 정확히 넘어가야 한다(단순 정수 +1이 아니라 실제 달력 연산)."""
    assert mtp._youth_content_date_str('20260912') == '20260913'  # 토→일, 같은 달
    assert mtp._youth_content_date_str('20260830') == '20260831'  # 월말이 아닌 경우
    assert mtp._youth_content_date_str('20260228') == '20260301'  # 월 경계(2026은 평년)


def _y2m_stub_all_mass_calls(monkeypatch, calls):
    """main()의 슬라이드 조작 단계 전체를 기록용 스텁으로 치환한다. 실제 슬라이드 XML을
    조작하지 않고 "무엇이 몇 번째로 어떤 인자와 함께 호출됐는지"만 기록해, main()의 배선
    로직(mass_type에 따라 어떤 함수를 부르는가) 자체를 pptx 없이 빠르게 검증하기 위함이다."""
    for name in [
        'update_title_slide', 'update_입당송', 'update_reading_title_slide',
        'update_복음_title_slide', 'update_화답송', 'update_복음환호송', 'update_영성체송',
        'replace_시작기도문', 'replace_미사후기도', 'insert_공지사항',
        '_align_ending_slides_to_제2독서', '_reposition_merged_ending_shapes',
        'strip_ppt2007_incompatible', 'validate',
    ]:
        def _make(n):
            def _f(*a, **k):
                calls.append((n, a, k))
                return None
            return _f
        monkeypatch.setattr(mtp, name, _make(name))

    def _replace_reading_slides(*a, **k):
        calls.append(('replace_reading_slides', a, k))
        return 0
    monkeypatch.setattr(mtp, 'replace_reading_slides', _replace_reading_slides)

    def _validate_pptx_structure(*a, **k):
        calls.append(('validate_pptx_structure', a, k))
        return []
    monkeypatch.setattr(mtp, 'validate_pptx_structure', _validate_pptx_structure)

    def _replace_성가(*a, **k):
        calls.append(('replace_성가', a, k))
    monkeypatch.setattr(mtp, 'replace_성가', _replace_성가)

    def _replace_성가_youth(*a, **k):
        calls.append(('replace_성가_youth', a, k))
    monkeypatch.setattr(mtp, 'replace_성가_youth', _replace_성가_youth)

    orig_find_sections = mtp.find_sections

    def _find_sections(prs, mass_type='adult'):
        calls.append(('find_sections', (mass_type,), {}))
        return orig_find_sections(prs, mass_type=mass_type)
    monkeypatch.setattr(mtp, 'find_sections', _find_sections)


def test_y2m_3_youth_preload_path_calls_youth_functions_not_adult(monkeypatch, tmp_path):
    """청년 EXE 프리로드 경로: replace_성가_youth가 호출되고 replace_성가는 호출되지 않으며,
    find_sections가 mass_type='youth'로 호출되고, 저장 파일명이 청년 전용 포맷을 따른다."""
    _y2m_skip_if_missing()

    date_str = '20260912'          # 토요일
    content_date_str = '20260913'  # +1일 = 일요일(요구사항 §6.1)

    (tmp_path / date_str).mkdir()
    (tmp_path / content_date_str).mkdir()
    ref_dst = tmp_path / date_str / 'reference.pptx'
    shutil.copy(str(Y2M_YOUTH_TEMPLATE), str(ref_dst))
    shutil.copy(str(Y2M_KO_JSON_20260913), str(tmp_path / content_date_str / f'missa_{content_date_str}.json'))

    monkeypatch.setattr(mtp, 'OUTPUT_ROOT', str(tmp_path))

    import json as _json
    en_fixture = _json.loads(Y2M_EN_GOSPEL_FIXTURE.read_text(encoding='utf-8'))
    monkeypatch.setattr(mtp.missa_youth_gospel, 'get_youth_mass', lambda d: en_fixture)

    calls = []
    _y2m_stub_all_mass_calls(monkeypatch, calls)

    files = {
        'ref_pptx': ref_dst, '시작기도': None, '화답송_pptx': None, '화답송_img': None,
        '미사후기도': None, '공지사항': None, '성가': {},
    }
    # main()의 방어적 이중 체크(y2m_8/9, 05_review_report.md 확정 버그 대응)가 5종 전부를
    # 요구하므로, 여기서도 실제 성가 선택처럼 5종 전부를 채운다(순수 배선 스파이 목적이라
    # 전부 '기타' — 내용 검증은 y2m_5가 담당).
    성가_선택 = {
        '입당':   [{'출처': '기타', '번호': None, '제목': '테스트 입당'}],
        '봉헌':   [{'출처': '기타', '번호': None, '제목': '테스트 봉헌'}],
        '성체':   [{'출처': '기타', '번호': None, '제목': '테스트 성체'}],
        '2차봉헌': [{'출처': '기타', '번호': None, '제목': '테스트 2차봉헌'}],
        '파견':   [{'출처': '기타', '번호': None, '제목': '테스트 파견'}],
    }
    mtp._preloaded_inputs[0] = {
        'mass_type': 'youth', 'date_str': date_str, 'files': files,
        'hymn_numbers': {}, '성가_선택': 성가_선택,
    }
    monkeypatch.setattr(_y2m_sys, 'argv', ['missa_to_ppt.py'])

    mtp.main()

    names = [c[0] for c in calls]
    assert 'replace_성가_youth' in names
    assert 'replace_성가' not in names
    # replace_성가_youth가 실제로 우리가 넘긴 성가_선택으로 호출됐는지(스파이가 배선만 확인하고
    # 끝나지 않도록 — "호출됐다"만으로는 인자가 뒤바뀐 배선도 통과시킨다).
    youth_call = next(c for c in calls if c[0] == 'replace_성가_youth')
    assert youth_call[1][1] == 성가_선택

    find_sections_mass_types = {c[1][0] for c in calls if c[0] == 'find_sections'}
    assert find_sections_mass_types == {'youth'}

    expected_liturgy = _json.loads(Y2M_KO_JSON_20260913.read_text(encoding='utf-8'))['liturgy']
    expected_name = f"토요일 저녁 청년 주일미사_{date_str}_{expected_liturgy}.pptx"
    # B그룹(2026-09-24): 청년미사 출력 폴더는 output/{date_str}_youth/(mtp.output_folder_key).
    youth_folder = tmp_path / f'{date_str}_youth'
    assert (youth_folder / expected_name).exists(), list(
        youth_folder.iterdir() if youth_folder.is_dir() else []
    )


def test_y2m_4_adult_preload_path_calls_adult_functions_not_youth(monkeypatch, tmp_path):
    """성인(기존) EXE 프리로드 경로: mass_type 키가 없는(구형) preload dict도 'adult'로
    안전하게 기본 처리되고, replace_성가가 호출되며 replace_성가_youth는 호출되지 않고,
    저장 파일명은 기존 포맷 그대로다(무회귀)."""
    date_str = '20260624'  # 회귀 스위트가 쓰는 평일 고정 픽스처(git 추적 대상)
    src_folder = Y2M_BASE / 'output' / date_str
    if not src_folder.is_dir():
        pytest.skip(f'평일 고정 픽스처가 없음: {src_folder}')

    dst_folder = tmp_path / date_str
    dst_folder.mkdir()
    for f in src_folder.iterdir():
        if f.is_file():
            shutil.copy(str(f), str(dst_folder / f.name))

    monkeypatch.setattr(mtp, 'OUTPUT_ROOT', str(tmp_path))

    hymn_numbers = {'입당': 329, '봉헌': 221, '성체': 156, '2차봉헌': None, '파견': 25}
    files = mtp.find_files(date_str, hymn_numbers, copy_hymn_scores=False)

    calls = []
    _y2m_stub_all_mass_calls(monkeypatch, calls)

    # 'mass_type' 키가 아예 없는 구형 preload dict 모양으로 준다 — main()의
    # `_inp.get('mass_type', 'adult')` 폴백이 실제로 동작하는지 확인.
    mtp._preloaded_inputs[0] = {
        'date_str': date_str, 'files': files, 'hymn_numbers': hymn_numbers,
    }
    monkeypatch.setattr(_y2m_sys, 'argv', ['missa_to_ppt.py'])
    # 평일미사라 OneDrive 조회 분기(mass_type=='adult' and is_sunday_mass(...))는 자연히
    # 타지 않는다 — get_onedrive_hymn_folder()를 스텁할 필요 없음.

    mtp.main()

    names = [c[0] for c in calls]
    assert 'replace_성가' in names
    assert 'replace_성가_youth' not in names

    find_sections_mass_types = {c[1][0] for c in calls if c[0] == 'find_sections'}
    assert find_sections_mass_types == {'adult'}

    import json as _json
    expected_liturgy = _json.loads((dst_folder / f'missa_{date_str}.json').read_text(encoding='utf-8'))['liturgy']
    expected_name = f"{date_str}_{expected_liturgy}.pptx"
    assert (dst_folder / expected_name).exists(), list(dst_folder.iterdir())


def test_y2m_5_youth_end_to_end_produces_valid_pptx(monkeypatch, tmp_path):
    """C1급 종단 검증(main() 경유): 청년 mass_type을 EXE 프리로드 경로로 선택하면, OneDrive/
    로컬 리소스 없이도('기타' 출처만 사용) main()이 끝까지 돌아 validate_pptx_structure()를
    통과하는 실제 pptx를 만든다. 이번 라운드가 완성해야 하는 '실제로 끝까지 도는 파이프라인'
    그 자체를 검증하는 테스트라 슬라이드 조작 함수를 스텁하지 않는다(y2m_3과 달리 실제 실행)."""
    _y2m_skip_if_missing()

    date_str = '20260912'
    content_date_str = '20260913'

    (tmp_path / date_str).mkdir()
    (tmp_path / content_date_str).mkdir()
    ref_dst = tmp_path / date_str / 'reference.pptx'
    shutil.copy(str(Y2M_YOUTH_TEMPLATE), str(ref_dst))
    shutil.copy(str(Y2M_KO_JSON_20260913), str(tmp_path / content_date_str / f'missa_{content_date_str}.json'))

    monkeypatch.setattr(mtp, 'OUTPUT_ROOT', str(tmp_path))

    import json as _json
    en_fixture = _json.loads(Y2M_EN_GOSPEL_FIXTURE.read_text(encoding='utf-8'))
    monkeypatch.setattr(mtp.missa_youth_gospel, 'get_youth_mass', lambda d: en_fixture)

    files = {
        'ref_pptx': ref_dst, '시작기도': None, '화답송_pptx': None, '화답송_img': None,
        '미사후기도': None, '공지사항': None, '성가': {},
    }
    성가_선택 = {
        '입당':   [{'출처': '기타', '번호': None, '제목': '테스트 입당곡'}],
        '봉헌':   [{'출처': '기타', '번호': None, '제목': '테스트 봉헌곡'}],
        '성체':   [{'출처': '기타', '번호': None, '제목': '테스트 성체곡'}],
        '2차봉헌': [{'출처': '기타', '번호': None, '제목': '테스트 2차봉헌곡'}],
        '파견':   [{'출처': '기타', '번호': None, '제목': '테스트 파견곡'}],
    }
    mtp._preloaded_inputs[0] = {
        'mass_type': 'youth', 'date_str': date_str, 'files': files,
        'hymn_numbers': {}, '성가_선택': 성가_선택,
    }
    monkeypatch.setattr(_y2m_sys, 'argv', ['missa_to_ppt.py'])

    mtp.main()

    expected_liturgy = _json.loads(Y2M_KO_JSON_20260913.read_text(encoding='utf-8'))['liturgy']
    # B그룹(2026-09-24): 청년미사 출력 폴더는 output/{date_str}_youth/(mtp.output_folder_key).
    out_path = tmp_path / f'{date_str}_youth' / f"토요일 저녁 청년 주일미사_{date_str}_{expected_liturgy}.pptx"
    assert out_path.exists()

    issues = mtp.validate_pptx_structure(str(out_path))
    assert issues == [], issues

    prs = mtp.Presentation(str(out_path))
    from missa_ooxml_utils import all_slide_texts as _y2m_all_texts
    texts = _y2m_all_texts(prs)
    joined = '\n'.join(texts)
    assert expected_liturgy in joined
    # '기타' 성가 헤더가 실제로 이번 주 제목으로 갱신됐는지(§9.1 헤더 포맷 `{타입} {제목}`).
    assert '입당 테스트 입당곡' in joined
    assert '파견 테스트 파견곡' in joined


# ---------------------------------------------------------------------------
# 05_review_report.md 확정 버그 회귀 가드 —
# "성가 입력 팝업을 취소(X/Alt+F4)하면 무경고로 빈 성가_선택이 흘러 들어가, 이번 주 성가가
# 하나도 갱신되지 않은 PPT가 '완료'로 저장된다" (참조 PPT로 재현됨: replace_성가_youth(prs, {})
# 호출 시 HYMN_TYPES 5종 전부 entries=[]가 되어 조용히 스킵). 두 겹으로 막는다:
# (1) _ask_youth_hymn_popup()/_ask_mass_type_popup() 자체가 취소 시 RuntimeError를 던지도록
#     수정(다른 입력 팝업과 동일한 관례로 통일).
# (2) main()에 방어적 이중 체크 추가 — 프로그래밍적 직접 호출(테스트 등) 경로에서 빈/부분
#     성가_선택이 들어와도 sys.exit(1)로 명시적으로 중단시킨다.
# ---------------------------------------------------------------------------

def _y2m_run_popup_cancel_probe(tmp_path, func_name: str):
    """`func_name` 팝업을 별도 프로세스에서 '확인 없이 창 닫기'로 취소해보고, RuntimeError가
    실제로 발생하는지 exit code로 보고한다. 같은 프로세스에서 tk.Tk()를 연달아 새로 만들고
    destroy()하면(y2m_6/y2m_7이 연속 실행될 때) Tcl 인터프리터 상태가 간헐적으로 깨지는
    tkinter 자체의 알려진 불안정성이 관측됐다(이 프로젝트 코드의 결함이 아님) — 각 팝업을
    독립 프로세스에서 실행해 이 간섭을 원천적으로 피한다."""
    import subprocess as _y2m_subprocess

    script = f'''
import sys
sys.path.insert(0, {str(Y2M_BASE)!r})
import tkinter as tk
import missa_to_ppt as mtp

orig_mainloop = tk.Tk.mainloop

def _fake_mainloop(self, *a, **k):
    self.after(10, self.destroy)
    return orig_mainloop(self, *a, **k)

tk.Tk.mainloop = _fake_mainloop

try:
    mtp.{func_name}()
except RuntimeError:
    sys.exit(0)
sys.exit(1)
'''
    script_path = tmp_path / f'probe_{func_name}.py'
    script_path.write_text(script, encoding='utf-8')
    return _y2m_subprocess.run(
        [_y2m_sys.executable, str(script_path)],
        capture_output=True, text=True, encoding='utf-8', errors='replace',
    )


def test_y2m_6_youth_hymn_popup_cancel_raises_runtime_error(tmp_path):
    """확인 버튼이 아니라 창 닫기(X/Alt+F4)로 취소하면 무경고 {} 대신 RuntimeError를 던진다.
    실제 창 close를 tk.Tk.mainloop 자체를 스케줄된 root.destroy()로 대체해 시뮬레이션한다 —
    on_ok()를 거치지 않고 mainloop가 빠져나가는 것이 '확인 없이 닫힘'과 동일한 상태다."""
    result = _y2m_run_popup_cancel_probe(tmp_path, '_ask_youth_hymn_popup')
    assert result.returncode == 0, (
        f"RuntimeError가 발생하지 않음(exit={result.returncode})\n"
        f"stdout:\n{result.stdout}\nstderr:\n{result.stderr}"
    )


# test_y2m_7_mass_type_popup_cancel_raises_runtime_error — 2026-09-26 §J3에서 제거.
# `_ask_mass_type_popup()` 자체가 삭제(미사 유형 선택 팝업 전면 제거, CLI --미사유형로
# 직접 받음)돼 이 취소 동작을 테스트할 대상이 더 이상 없다(상세: `_workspace/
# 02_specialist_impl_notes.md` "J그룹 배치3" 절).


def test_y2m_8_main_aborts_when_성가_선택_missing_htypes(monkeypatch, tmp_path):
    """main()을 프로그래밍적으로 호출할 때(팝업을 거치지 않는 경로) 성가_선택에서 htype이
    하나라도 빠지면, replace_성가_youth가 조용히 스킵하기 전에 main()이 명시적 오류로
    중단해야 한다(sys.exit(1)) — 리뷰가 실제 참조 PPT로 재현한 '무경고 미갱신' 시나리오의
    두 번째 방어선."""
    _y2m_skip_if_missing()

    date_str = '20260912'
    content_date_str = '20260913'
    (tmp_path / date_str).mkdir()
    (tmp_path / content_date_str).mkdir()
    ref_dst = tmp_path / date_str / 'reference.pptx'
    shutil.copy(str(Y2M_YOUTH_TEMPLATE), str(ref_dst))
    shutil.copy(str(Y2M_KO_JSON_20260913), str(tmp_path / content_date_str / f'missa_{content_date_str}.json'))

    monkeypatch.setattr(mtp, 'OUTPUT_ROOT', str(tmp_path))

    import json as _json
    en_fixture = _json.loads(Y2M_EN_GOSPEL_FIXTURE.read_text(encoding='utf-8'))
    monkeypatch.setattr(mtp.missa_youth_gospel, 'get_youth_mass', lambda d: en_fixture)

    files = {
        'ref_pptx': ref_dst, '시작기도': None, '화답송_pptx': None, '화답송_img': None,
        '미사후기도': None, '공지사항': None, '성가': {},
    }
    # 완전히 빈 성가_선택 — 취소된 팝업을 흉내낸다(리뷰 재현 스크립트와 동일 시나리오).
    mtp._preloaded_inputs[0] = {
        'mass_type': 'youth', 'date_str': date_str, 'files': files,
        'hymn_numbers': {}, '성가_선택': {},
    }
    monkeypatch.setattr(_y2m_sys, 'argv', ['missa_to_ppt.py'])

    with pytest.raises(SystemExit) as exc_info:
        mtp.main()
    assert exc_info.value.code == 1
    # 성가 교체 단계(replace_성가_youth)까지 도달하지 않고 그 전에 중단돼야 한다 —
    # 저장 파일이 아예 생기지 않아야 "완료로 조용히 저장"되는 원래 버그를 확실히 막는다.
    assert not (tmp_path / f'{date_str}_youth' / f"토요일 저녁 청년 주일미사_{date_str}_연중 제24주일.pptx").exists()


def test_y2m_9_main_aborts_when_성가_선택_partially_missing(monkeypatch, tmp_path):
    """일부 htype만 채워진 성가_선택(예: 입당만 있고 나머지 4종 누락)도 동일하게 중단된다 —
    '완전히 빈 dict'만 막고 '부분적으로 빈 dict'는 놓치는 좁은 수정이 아님을 확인."""
    _y2m_skip_if_missing()

    date_str = '20260912'
    content_date_str = '20260913'
    (tmp_path / date_str).mkdir()
    (tmp_path / content_date_str).mkdir()
    ref_dst = tmp_path / date_str / 'reference.pptx'
    shutil.copy(str(Y2M_YOUTH_TEMPLATE), str(ref_dst))
    shutil.copy(str(Y2M_KO_JSON_20260913), str(tmp_path / content_date_str / f'missa_{content_date_str}.json'))

    monkeypatch.setattr(mtp, 'OUTPUT_ROOT', str(tmp_path))

    import json as _json
    en_fixture = _json.loads(Y2M_EN_GOSPEL_FIXTURE.read_text(encoding='utf-8'))
    monkeypatch.setattr(mtp.missa_youth_gospel, 'get_youth_mass', lambda d: en_fixture)

    files = {
        'ref_pptx': ref_dst, '시작기도': None, '화답송_pptx': None, '화답송_img': None,
        '미사후기도': None, '공지사항': None, '성가': {},
    }
    성가_선택 = {'입당': [{'출처': '기타', '번호': None, '제목': '테스트 입당곡'}]}
    mtp._preloaded_inputs[0] = {
        'mass_type': 'youth', 'date_str': date_str, 'files': files,
        'hymn_numbers': {}, '성가_선택': 성가_선택,
    }
    monkeypatch.setattr(_y2m_sys, 'argv', ['missa_to_ppt.py'])

    with pytest.raises(SystemExit) as exc_info:
        mtp.main()
    assert exc_info.value.code == 1

# ═══════════════════════════════════════════════════════════════════════════
# 2026-09-24 regression-qa 승격: B/A/C/D 그룹(청년미사 확장 + 성인/공통 버그 수정,
# 00_request.md) 전체 — pytest 전체 스위트 322개 all green(0 failed/skipped) 재확인,
# ooxml-code-reviewer 2라운드 리뷰 통과, D2/C1 재검증(D2: 성인 20260712 픽스처 geometry
# 원상복구 실측 확인 / C1: 나주노 151·267 원본 PDF 직접 렌더링 독립 재검증 — "버그 아님"
# 결론 재확인) 후 test_missa_progression.py에서 이 블록을 그대로 옮겼다. 이 시점부터 이
# 테스트들은 "새 동작의 명세"가 아니라 "지켜야 할 기존 동작"이다.
#
# 아래 3개는 원본 파일(옮겨오기 전 위치, line 23/375/1193 부근)에 있던 모듈 전역 의존성을
# 이 새 위치에서도 쓸 수 있도록 다시 선언한다(파일 상단부에 이미 동일 이름의 mtp/Y2M_BASE
# 등이 있지만, 아래 블록 코드 자체는 최소 변경 원칙에 따라 원본 그대로 두고 필요한 이름만
# 보강한다).
# ═══════════════════════════════════════════════════════════════════════════
from pptx.oxml.ns import qn as _qn  # noqa: E402
BASE = Path(__file__).resolve().parent.parent  # noqa: E402  (tests/ -> 저장소 루트)
import missa_youth_hymn_pdf as hp  # noqa: E402  (C1/C2 skipif 조건 계산용)
import fitz as _fitz  # noqa: E402  (C1 테스트 본문이 _fitz.open(...)을 직접 참조)
import os as _os  # noqa: E402
_YOUTH_PDFS_AVAILABLE = _os.path.isfile(hp.SOURCES["나주노"]["pdf"]) and _os.path.isfile(
    hp.SOURCES["야훼 이레"]["pdf"]
)  # noqa: E402
_YOUTH_SKIP_REASON = "나주노/야훼이레 PDF 원본이 로컬에 없음(reference/청년미사/, git 미포함)"  # noqa: E402

# ═══════════════════════════════════════════════════════════════════════════
# B그룹: output 폴더 명명 규칙 (2026-09-24 청년미사 확장 — 00_request.md)
#
# 성인미사 output/{date_str}/(무변경) vs 청년미사 output/{date_str}_youth/. 향후 어린이미사
# output/{date_str}_children/까지 하드코딩 없이 mass_type→접미사 매핑 한 줄 추가로 확장
# 가능해야 한다. "사용자 입력 날짜"(date_str)/"콘텐츠 조회 날짜"(content_date_str, 기존
# _youth_content_date_str)/"출력 폴더 키"(output_key, 이번에 신설) 3개 역할을 분리한다.
#
# 회귀 픽스처 20260624/20260712/20260705(전부 성인)이 가리키는 실제 폴더 경로는 절대 바뀌면
# 안 된다 — output_folder_key('아무날짜', 'adult') == '아무날짜'(접미사 '')로 보장한다.
#
# 안정화되면 regression-qa가 test_missa_regression.py로 승격한다.
# ═══════════════════════════════════════════════════════════════════════════
import missa_gui as gui_b  # noqa: E402


def test_output_folder_key_adult_has_no_suffix():
    """성인미사는 접미사 없음 — 회귀 픽스처 폴더 경로 보존의 핵심 전제."""
    assert gui_b.output_folder_key('20260712', 'adult') == '20260712'


def test_output_folder_key_youth_has_youth_suffix():
    assert gui_b.output_folder_key('20260919', 'youth') == '20260919_youth'


def test_output_folder_key_unknown_mass_type_falls_back_to_no_suffix():
    """미래에 매핑에 없는 mass_type이 들어와도 KeyError 대신 접미사 없음으로 안전 폴백."""
    assert gui_b.output_folder_key('20260919', 'nonexistent') == '20260919'


def test_output_folder_key_default_mass_type_is_adult():
    """mass_type을 생략하면 성인(접미사 없음) 기본값 — 기존 date_str 단일 인자 호출부 보존."""
    assert gui_b.output_folder_key('20260712') == '20260712'


def test_mass_type_output_suffix_mapping_extensible_for_children():
    """향후 어린이미사 추가 시 이 매핑에 한 줄만 추가하면 되는 구조인지 — 매핑이 dict이고
    'adult'/'youth' 두 항목만 있는지 확인(어린이미사가 아직 없다는 전제 자체를 명세)."""
    assert gui_b.MASS_TYPE_OUTPUT_SUFFIX == {'adult': '', 'youth': '_youth'}


def test_find_files_youth_scans_suffixed_folder(tmp_path, monkeypatch):
    """find_files(mass_type='youth')는 output/{date}_youth/를 스캔해야 한다 — 접미사 없는
    폴더에 같은 이름 파일이 있어도 무시한다(폴더 키 분리 확인)."""
    unsuffixed = tmp_path / '20260919'
    unsuffixed.mkdir()
    (unsuffixed / '참조 미사.pptx').write_bytes(b'')

    suffixed = tmp_path / '20260919_youth'
    suffixed.mkdir()
    (suffixed / '참조 미사.pptx').write_bytes(b'')

    monkeypatch.setattr(mtp, 'OUTPUT_ROOT', str(tmp_path))
    files = mtp.find_files('20260919', {}, copy_hymn_scores=False, mass_type='youth')
    assert files['ref_pptx'] is not None
    assert files['ref_pptx'].parent == suffixed


def test_find_files_adult_default_scans_unsuffixed_folder(tmp_path, monkeypatch):
    """mass_type을 생략(기존 호출부 100% 보존)하면 여전히 접미사 없는 폴더를 스캔한다."""
    unsuffixed = tmp_path / '20260712'
    unsuffixed.mkdir()
    (unsuffixed / '참조 미사.pptx').write_bytes(b'')

    monkeypatch.setattr(mtp, 'OUTPUT_ROOT', str(tmp_path))
    files = mtp.find_files('20260712', {}, copy_hymn_scores=False)
    assert files['ref_pptx'] is not None
    assert files['ref_pptx'].parent == unsuffixed


def test_find_files_missing_suffixed_folder_raises(tmp_path, monkeypatch):
    """청년미사인데 접미사 폴더가 없으면(접미사 없는 폴더만 있어도) FileNotFoundError —
    조용히 엉뚱한 폴더로 폴백하면 안 된다."""
    (tmp_path / '20260919').mkdir()  # 접미사 없는 폴더만 존재
    monkeypatch.setattr(mtp, 'OUTPUT_ROOT', str(tmp_path))
    with pytest.raises(FileNotFoundError):
        mtp.find_files('20260919', {}, copy_hymn_scores=False, mass_type='youth')


def test_log_dir_for_uses_output_key_when_given():
    """_log_dir_for()는 output_key가 주어지면(청년미사 등) 그 폴더를, 없으면(구버전 호출
    등) date_str 그대로 폴백한다 — 순수 함수라 tkinter 없이 검증 가능."""
    import missa_gui as gui
    assert gui._log_dir_for('20260919', '20260919_youth') == (
        Path(gui.OUTPUT_ROOT) / '20260919_youth' / 'log'
    )
    assert gui._log_dir_for('20260712', None) == (
        Path(gui.OUTPUT_ROOT) / '20260712' / 'log'
    )


def test_b_youth_preload_pipeline_saves_under_suffixed_folder(monkeypatch, tmp_path):
    """B그룹 종단 확인: 청년미사 EXE 프리로드 경로로 main()을 실행하면 최종 저장 PPT가
    output/{date}_youth/ 아래에 있어야 하고, 접미사 없는 폴더에는 아무것도 생기지 않는다.
    슬라이드 조작 함수는 전부 스텁해 배선(경로 계산)만 빠르게 검증한다."""
    youth_template = (
        BASE / "reference" / "청년미사"
        / "Template_토요일 저녁 청년 주일미사_20260822_연중 제21주일.pptx"
    )
    en_gospel_fixture = (
        BASE / "_workspace" / "청년미사_1단계" / "verify_20260913" / "missa_en_20260913.json"
    )
    ko_json_20260913 = BASE / "output" / "20260913" / "missa_20260913.json"
    if not youth_template.is_file():
        pytest.skip(f"청년미사 참조 PPT가 로컬에 없음: {youth_template}")
    if not en_gospel_fixture.is_file():
        pytest.skip(f"영문 복음 픽스처가 로컬에 없음: {en_gospel_fixture}")
    if not ko_json_20260913.is_file():
        pytest.skip(f"20260913 한글 JSON이 로컬에 없음: {ko_json_20260913}")

    import shutil
    import sys as _sys
    import json as _json

    date_str = '20260912'          # 토요일(입력)
    content_date_str = '20260913'  # +1일(콘텐츠 조회, 접미사 없음 — 어느 mass_type이든 공유 캐시)

    # 입력 파일은 임의 위치(files 딕셔너리에 직접 경로를 채우는 preload 경로라 find_files()를
    # 거치지 않는다 — 실제로 검증하려는 것은 "출력" 경로이므로 입력 위치는 무방).
    (tmp_path / date_str).mkdir()
    (tmp_path / content_date_str).mkdir()
    ref_dst = tmp_path / date_str / 'reference.pptx'
    shutil.copy(str(youth_template), str(ref_dst))
    shutil.copy(str(ko_json_20260913), str(tmp_path / content_date_str / f'missa_{content_date_str}.json'))

    monkeypatch.setattr(mtp, 'OUTPUT_ROOT', str(tmp_path))

    en_fixture = _json.loads(en_gospel_fixture.read_text(encoding='utf-8'))
    monkeypatch.setattr(mtp.missa_youth_gospel, 'get_youth_mass', lambda d: en_fixture)

    for name in [
        'update_title_slide', 'update_입당송', 'update_reading_title_slide',
        'update_복음_title_slide', 'update_화답송', 'update_복음환호송', 'update_영성체송',
        'replace_시작기도문', 'replace_미사후기도', 'insert_공지사항', 'replace_성가_youth',
        '_align_ending_slides_to_제2독서', '_reposition_merged_ending_shapes',
        'strip_ppt2007_incompatible', 'validate',
    ]:
        def _make(n):
            def _f(*a, **k):
                return None
            return _f
        monkeypatch.setattr(mtp, name, _make(name))
    monkeypatch.setattr(mtp, 'replace_reading_slides', lambda *a, **k: 0)
    monkeypatch.setattr(mtp, 'validate_pptx_structure', lambda *a, **k: [])

    files = {
        'ref_pptx': ref_dst, '시작기도': None, '화답송_pptx': None, '화답송_img': None,
        '미사후기도': None, '공지사항': None, '성가': {},
    }
    성가_선택 = {
        '입당':   [{'출처': '기타', '번호': None, '제목': '테스트 입당'}],
        '봉헌':   [{'출처': '기타', '번호': None, '제목': '테스트 봉헌'}],
        '성체':   [{'출처': '기타', '번호': None, '제목': '테스트 성체'}],
        '2차봉헌': [{'출처': '기타', '번호': None, '제목': '테스트 2차봉헌'}],
        '파견':   [{'출처': '기타', '번호': None, '제목': '테스트 파견'}],
    }
    mtp._preloaded_inputs[0] = {
        'mass_type': 'youth', 'date_str': date_str, 'files': files,
        'hymn_numbers': {}, '성가_선택': 성가_선택,
    }
    monkeypatch.setattr(_sys, 'argv', ['missa_to_ppt.py'])

    mtp.main()

    expected_liturgy = _json.loads(ko_json_20260913.read_text(encoding='utf-8'))['liturgy']
    expected_name = f"토요일 저녁 청년 주일미사_{date_str}_{expected_liturgy}.pptx"
    suffixed_folder = tmp_path / f'{date_str}_youth'
    assert (suffixed_folder / expected_name).exists(), list(
        suffixed_folder.iterdir() if suffixed_folder.is_dir() else []
    )
    # 접미사 없는 폴더(입력 파일을 임시로 넣어둔 곳)에는 결과 PPT가 없어야 한다 — 출력 폴더
    # 키가 실제로 분리됐는지의 핵심 단언.
    assert not (tmp_path / date_str / expected_name).exists()
    assert mtp._last_output_key[0] == f'{date_str}_youth'


# ═══════════════════════════════════════════════════════════════════════════
# A그룹: 입력 UI 7건 (2026-09-24 청년미사 확장 — 00_request.md, 사용자 컨펌 완료 설계)
#
# tkinter 팝업은 root.mainloop()로 블로킹되므로, 위젯을 실제로 구성한 직후(사용자가 아직
# 아무 것도 클릭하지 않은 시점) 상태를 스냅샷해야 UI 구조(행 배치/폰트/숨김 여부 등)를
# 검증할 수 있다. tk.Tk.mainloop를 "위젯 트리를 덤프하고 destroy" 콜백으로 대체해, 같은
# 프로세스에서 tk.Tk()를 연달아 만들 때의 Tcl 인터프리터 불안정(y2m_6/7 주석 참고)을 피하기
# 위해 항상 독립 서브프로세스에서 실행한다.
# ═══════════════════════════════════════════════════════════════════════════
import json as _a_json
import subprocess as _a_subprocess
import sys as _a_sys

_A_WIDGET_DUMP_HELPER = '''
def _a_widget_dump(w, out):
    info = {"class": w.winfo_class()}
    for opt in ("text", "font", "fg", "bg", "state", "value", "width"):
        try:
            info[opt] = str(w.cget(opt))
        except Exception:
            pass
    try:
        varname = w.cget("variable")
        if varname:
            info["variable_value"] = str(w.tk.globalgetvar(varname))
    except Exception:
        pass
    try:
        tvarname = w.cget("textvariable")
        if tvarname:
            info["textvariable_value"] = str(w.tk.globalgetvar(tvarname))
    except Exception:
        pass
    try:
        # 최상위 tk.Tk() 창 자체는 Grid mixin을 상속하지 않아 grid_info()가 없다(속성
        # 접근이 내부 Tcl 인터프리터로 포워딩되며 AttributeError) — Tkinter는 콜백 안의
        # 예외를 삼키고 무시하므로, 이 try/except가 없으면 root.destroy()가 영영 실행되지
        # 않아 mainloop()가 멈춘 채로 남는다(실측: 이 가드 없이 재현한 무한 행).
        gi = w.grid_info()
    except Exception:
        gi = {}
    info["grid_row"] = int(gi["row"]) if gi else None
    info["grid_column"] = int(gi["column"]) if gi else None
    info["gridded"] = bool(gi)
    out.append(info)
    for c in w.winfo_children():
        _a_widget_dump(c, out)
'''


def _a_run_widget_probe(tmp_path, popup_call: str, pre_probe: str = '', extra_vars: str = '') -> list:
    """`popup_call`(예: "gui._ask_mass_type_popup()")을 서브프로세스에서 실행하고, 그
    팝업이 mainloop에 진입한 직후(위젯 구성 완료, 사용자 조작 전) 위젯 트리 전체를 JSON으로
    덤프해 반환한다. `pre_probe`는 mainloop 진입 콜백 안에서 덤프 직전에 실행할 추가 코드
    (예: 특정 콤보박스 값을 바꾸고 <<ComboboxSelected>> 이벤트를 발생시키는 시뮬레이션)."""
    out_file = tmp_path / 'a_probe_result.json'
    script = f'''
import sys, json
sys.path.insert(0, {str(BASE)!r})
import tkinter as tk
import missa_gui as gui

{_A_WIDGET_DUMP_HELPER}

_widgets = []
_orig_mainloop = tk.Tk.mainloop

def _fake_mainloop(self, *a, **k):
    def _probe():
        root = self
        try:
{_a_indent(extra_vars, 12)}
{_a_indent(pre_probe, 12)}
            _a_widget_dump(root, _widgets)
        finally:
            # try 블록 안 어디서 예외가 나든 destroy()는 반드시 실행돼야 한다 — 안 그러면
            # mainloop()가 영영 반환하지 않아 서브프로세스가 무한 행(실측 사례, 위 가드
            # 추가 계기)한다.
            root.destroy()
    self.after(60, _probe)
    return _orig_mainloop(self, *a, **k)

tk.Tk.mainloop = _fake_mainloop

try:
    {popup_call}
except RuntimeError:
    pass

with open({str(out_file)!r}, 'w', encoding='utf-8') as f:
    json.dump(_widgets, f)
'''
    script_path = tmp_path / 'a_probe_script.py'
    script_path.write_text(script, encoding='utf-8')
    try:
        result = _a_subprocess.run(
            [_a_sys.executable, str(script_path)],
            cwd=str(BASE), capture_output=True, text=True, encoding='utf-8', errors='replace',
            timeout=30,
        )
    except _a_subprocess.TimeoutExpired as e:
        # mainloop()가 반환하지 않아 서브프로세스가 멈추는 버그 클래스(2026-09-24 test_a2
        # 최초 작성 때 실제로 겪음 — root.grid_info() AttributeError가 Tkinter 콜백 안에서
        # 조용히 삼켜져 destroy()가 실행되지 않았다) 재발 시, pytest 전체가 영원히 걸리는
        # 대신 여기서 명확한 타임아웃 실패로 끝낸다.
        raise AssertionError(
            f"위젯 프로브 서브프로세스가 30초 내에 끝나지 않음(mainloop 행 의심)\n"
            f"stdout:\n{e.stdout}\nstderr:\n{e.stderr}"
        )
    assert out_file.is_file(), (
        f"위젯 덤프 실패(exit={result.returncode})\nstdout:\n{result.stdout}\nstderr:\n{result.stderr}"
    )
    return _a_json.loads(out_file.read_text(encoding='utf-8'))


def _a_indent(code: str, spaces: int) -> str:
    if not code:
        return ''
    pad = ' ' * spaces
    return '\n'.join(pad + line for line in code.splitlines())


def test_a1_adult_label_is_plain_성인미사():
    """MASS_TYPES['성인']['label']이 '성인미사(교중미사)'가 아니라 '성인미사'여야 한다."""
    import missa_gui as gui
    assert gui.MASS_TYPES['성인']['label'] == '성인미사'


def test_a7_성체2_label_updated():
    import missa_gui as gui
    import inspect
    src = inspect.getsource(gui._ask_youth_hymn_popup)
    assert "'성체2 (선택)'" in src
    assert '성체(2번째, 선택)' not in src


# test_a2_update_button_moved_below_confirm_with_icon_and_smaller_font,
# test_a3_mass_type_popup_defaults_from_config_last_mass_type,
# test_a3_mass_type_popup_defaults_to_성인_when_config_missing_key,
# test_a3_load_config_last_mass_type_key_roundtrip — 2026-09-26 §J3에서 제거. 이 4개
# 테스트는 전부 `_ask_mass_type_popup()`(미사 유형 선택 팝업)이나 그 팝업이 읽고 쓰던
# config.json의 last_mass_type 키를 검증했는데, §J3에서 그 팝업 자체와 last_mass_type
# 읽기/쓰기가 전부 삭제(미사 유형은 이제 `--미사유형` CLI 인자로 직접 받음)돼 대상이
# 사라졌다. 업데이트 확인 버튼(test_a2)은 §J10에서 `_ask_combined_input_popup()`의 버전
# 표시 링크로 재배치되며, 그 배치의 새 테스트는 J10 배치에서 새로 작성한다(상세:
# `_workspace/02_specialist_impl_notes.md` "J그룹 배치3" 절).


def test_a4_combined_input_popup_youth_excludes_화답송_row(tmp_path):
    """청년미사(mass_type='youth')는 화답송 악보 PPT/사진 행 자체가 없어야 한다."""
    widgets = _a_run_widget_probe(tmp_path, "gui._ask_combined_input_popup(mass_type='youth')")
    labels = [w.get('text', '') for w in widgets if w['class'] == 'Label']
    assert not any('화답송' in t for t in labels), labels


def test_a4_combined_input_popup_adult_keeps_화답송_row(tmp_path):
    """성인미사(기본값)는 기존과 동일하게 화답송 악보 PPT/사진 행이 있어야 한다(무회귀)."""
    widgets = _a_run_widget_probe(tmp_path, "gui._ask_combined_input_popup(mass_type='adult')")
    labels = [w.get('text', '') for w in widgets if w['class'] == 'Label']
    assert any('화답송' in t for t in labels), labels


def test_a4_combined_input_popup_default_mass_type_is_adult(tmp_path):
    """mass_type 인자를 생략하면 기존 호출부(성인 전용 경로)와 100% 동일하게 동작한다."""
    widgets = _a_run_widget_probe(tmp_path, "gui._ask_combined_input_popup()")
    labels = [w.get('text', '') for w in widgets if w['class'] == 'Label']
    assert any('화답송' in t for t in labels), labels


# ---------------------------------------------------------------------------
# A5: 청년미사 OneDrive 커스텀 파일 브라우저 — 순수 함수(필터링/캐싱/다운로드)는 GUI 없이,
# Treeview 렌더링은 전용 서브프로세스 프로브로 검증한다.
# ---------------------------------------------------------------------------

def test_a5_onedrive_visible_children_filters_and_sorts_folders_first():
    import missa_gui as gui
    items = [
        {'name': 'b.pptx'},
        {'name': '~$temp.pptx'},
        {'name': 'notes.docx'},
        {'name': 'Folder2', 'folder': {}},
        {'name': 'a.pptx'},
        {'name': 'Folder1', 'folder': {}},
        {'name': ''},
    ]
    visible = gui._onedrive_visible_children(items)
    assert [v['name'] for v in visible] == ['Folder1', 'Folder2', 'a.pptx', 'b.pptx']


def test_a5_cached_onedrive_children_reuses_cache(monkeypatch):
    import sys
    import missa_gui as gui

    calls = []

    class _StubOD:
        @staticmethod
        def list_children(path):
            calls.append(path)
            return [{'name': 'x.pptx'}]

    monkeypatch.setitem(sys.modules, 'missa_onedrive', _StubOD)
    cache = {}
    r1 = gui._cached_onedrive_children('PPT 문서', cache)
    r2 = gui._cached_onedrive_children('PPT 문서', cache)
    assert r1 == r2 == [{'name': 'x.pptx'}]
    assert calls == ['PPT 문서']  # 두 번째 호출은 캐시 재사용 — 실제 조회는 1회만


def test_a5_download_onedrive_selection_uses_basename(tmp_path, monkeypatch):
    import sys
    import missa_gui as gui

    calls = []

    class _StubOD:
        @staticmethod
        def download_file(path, dest):
            calls.append((path, dest))
            return dest

    monkeypatch.setitem(sys.modules, 'missa_onedrive', _StubOD)
    dest_dir = tmp_path / 'out'
    result = gui._download_onedrive_selection('PPT 문서/2026/참조 미사.pptx', dest_dir)
    assert result == dest_dir / '참조 미사.pptx'
    assert calls == [('PPT 문서/2026/참조 미사.pptx', dest_dir / '참조 미사.pptx')]


def _a_run_onedrive_browser_probe(tmp_path, stub_children_code: str,
                                   open_folder_substr: str = None) -> list:
    """`_ask_onedrive_file_browser_popup()`이 렌더링한 Treeview 내용을 서브프로세스에서
    덤프해 중첩 리스트로 반환한다. ttk.Treeview 항목은 `winfo_children()`으로 보이는
    일반 위젯이 아니라 Tcl 내부 데이터라, `_A_WIDGET_DUMP_HELPER`(범용 위젯 트리 덤프)로는
    볼 수 없다 — `tree.get_children()`/`tree.item()`으로 직접 읽어야 한다.

    `open_folder_substr`(2026-09-26 J2 후속 확장): 주어지면 최상위에서 그 텍스트를
    포함하는 첫 폴더 노드에 `<<TreeviewOpen>>`을 발생시켜 실제 자식을 로드한 뒤 덤프한다
    — J2로 최상위가 3개 폴더로 고정되면서, 필터 없는 하위 레벨(파일 노드가 자식 없이
    렌더링되는지 등)을 검증하려면 최상위 파일 항목이 아니라 그 아래를 열어봐야 한다."""
    out_file = tmp_path / 'a5_probe_result.json'
    dest_dir = tmp_path / 'dest'
    script = f'''
import sys, json, types
sys.path.insert(0, {str(BASE)!r})
import tkinter as tk
import missa_gui as gui

_od_stub = types.ModuleType("missa_onedrive")
_od_stub.list_children = {stub_children_code}
_od_stub.download_file = lambda path, dest: dest
sys.modules["missa_onedrive"] = _od_stub

def _a_find_treeview(w):
    if w.winfo_class() == "Treeview":
        return w
    for c in w.winfo_children():
        found = _a_find_treeview(c)
        if found:
            return found
    return None

def _a_dump_tree(tree, parent=""):
    out = []
    for iid in tree.get_children(parent):
        item = tree.item(iid)
        out.append({{"iid": iid, "text": item["text"], "children": _a_dump_tree(tree, iid)}})
    return out

_dump = []
_orig_wait_window = tk.Toplevel.wait_window
_open_substr = {open_folder_substr!r}

def _fake_wait_window(self, *a, **k):
    def _probe():
        try:
            tv = _a_find_treeview(self)
            if tv is not None:
                if _open_substr is not None:
                    top = tv.get_children("")
                    target = next((iid for iid in top if _open_substr in tv.item(iid, "text")), None)
                    if target is not None:
                        tv.focus(target)
                        tv.event_generate("<<TreeviewOpen>>")
                        self.update()
                _dump.extend(_a_dump_tree(tv))
        finally:
            self.destroy()
    self.after(60, _probe)
    return _orig_wait_window(self, *a, **k)

tk.Toplevel.wait_window = _fake_wait_window

root = tk.Tk()
root.withdraw()
gui._ask_onedrive_file_browser_popup(root, {str(dest_dir)!r}, cache={{}})
root.destroy()

with open({str(out_file)!r}, 'w', encoding='utf-8') as f:
    json.dump(_dump, f)
'''
    script_path = tmp_path / 'a5_probe_script.py'
    script_path.write_text(script, encoding='utf-8')
    try:
        result = _a_subprocess.run(
            [_a_sys.executable, str(script_path)],
            cwd=str(BASE), capture_output=True, text=True, encoding='utf-8', errors='replace',
            timeout=30,
        )
    except _a_subprocess.TimeoutExpired as e:
        raise AssertionError(
            f"OneDrive 브라우저 프로브가 30초 내에 끝나지 않음\nstdout:\n{e.stdout}\nstderr:\n{e.stderr}"
        )
    assert out_file.is_file(), (
        f"트리 덤프 실패(exit={result.returncode})\nstdout:\n{result.stdout}\nstderr:\n{result.stderr}"
    )
    return _a_json.loads(out_file.read_text(encoding='utf-8'))


def test_a5_browser_root_level_renders_ppt_문서_children_directly(tmp_path):
    """'PPT 문서' 자체가 트리 노드로 보이지 않고, list_children('PPT 문서') 결과가 트리
    최상위 레벨에 바로 렌더링돼야 한다(사용자 명시 지시 — 00_request.md §A5).

    2026-09-26(J2) 갱신: 최상위는 실측된 3개 폴더로 필터링되므로(§J2), 임의의
    '09.가톨릭 성가'/'notes.pptx'로는 이 필터에 걸려 더 이상 재현되지 않는다 — 허용
    목록에 있는 이름으로 교체했다(J2 자체의 필터 검증은
    test_j2*(test_missa_progression.py)가 전담)."""
    stub = (
        "lambda path: {'PPT 문서': ["
        "{'name': '20.청년 미사', 'folder': {}},"
        "{'name': '13. 기도문', 'folder': {}},"
        "]}.get(path, [])"
    )
    top_level = _a_run_onedrive_browser_probe(tmp_path, stub)
    top_texts = [n['text'] for n in top_level]
    assert not any('PPT 문서' in t for t in top_texts), (
        "'PPT 문서' 폴더 자체가 트리 노드로 노출됨 — list_children() 결과가 바로 최상위여야 함",
        top_texts,
    )
    assert any('20.청년 미사' in t for t in top_texts), top_texts
    assert any('13. 기도문' in t for t in top_texts), top_texts


def test_a5_browser_folder_nodes_get_lazy_load_placeholder(tmp_path):
    """폴더 노드는 더블클릭 전에도 자리표시자 자식이 있어야(펼침 화살표가 보이도록) 하고,
    파일 노드는 자식이 없어야 한다.

    2026-09-26(J2) 갱신: 최상위 파일('notes.pptx')은 이제 3개 허용 폴더 목록에 없어
    필터링돼 최상위에서 재현할 수 없다 — 필터가 적용되지 않는 하위 레벨(허용 폴더 하나를
    열어 그 안의 파일/폴더)에서 같은 성질을 검증한다."""
    stub = (
        "lambda path: {"
        "'PPT 문서': [{'name': '20.청년 미사', 'folder': {}}],"
        "'PPT 문서/20.청년 미사': ["
        "{'name': '2.성가', 'folder': {}},"
        "{'name': 'notes.pptx'},"
        "],"
        "}.get(path, [])"
    )
    top_level = _a_run_onedrive_browser_probe(tmp_path, stub, open_folder_substr='청년 미사')
    opened = next(n for n in top_level if '청년 미사' in n['text'])
    folder_node = next(n for n in opened['children'] if '2.성가' in n['text'])
    file_node = next(n for n in opened['children'] if 'notes.pptx' in n['text'])
    assert len(folder_node['children']) == 1, folder_node
    assert file_node['children'] == [], file_node


# ---------------------------------------------------------------------------
# A6: 청년미사 성가 입력 팝업 — 헤더 행 + 제목칸 기본 숨김 + 출처='기타' 선택 시 토글
# ---------------------------------------------------------------------------

def test_a6_header_row_labels_present(tmp_path):
    widgets = _a_run_widget_probe(tmp_path, "gui._ask_youth_hymn_popup()")
    labels = [w.get('text', '') for w in widgets if w['class'] == 'Label']
    assert '성가 구분' in labels, labels
    assert '성가 번호' in labels, labels
    assert '성가 제목' in labels, labels


def test_a6_title_entry_hidden_by_default_for_all_rows(tmp_path):
    """출처 기본값('나주노')에서는 6개 행 전부 제목 입력창이 숨겨져 있어야 한다."""
    # width로 num_entry(6)/title_entry(22)를 구분한다 — grid_column/grid_row는 grid_remove()
    # 직후 grid_info()가 빈 dict가 되어 None이 되므로(숨김 위젯을 찾는 이 테스트의 핵심
    # 대상 자체를 못 찾게 됨) 쓸 수 없다. Entry 위젯은 각 htype 행마다 num_entry→title_entry
    # 순서로 생성되므로(코드 순서 그대로), width로 걸러낸 리스트의 상대 순서가 곧 행 순서다.
    widgets = _a_run_widget_probe(tmp_path, "gui._ask_youth_hymn_popup()")
    entries = [w for w in widgets if w['class'] == 'Entry']
    title_entries = [w for w in entries if w.get('width') == '22']
    assert len(title_entries) == 6, title_entries
    assert all(not w['gridded'] for w in title_entries), title_entries


def test_a6_selecting_기타_source_disables_number_and_reveals_title(tmp_path):
    """첫 행(입당)의 출처를 '기타'로 바꾸면: 번호칸은 disabled, 제목칸은 grid()로 표시돼야
    한다. 다른 행(둘째, 여전히 기본 출처)은 영향받지 않아야 한다."""
    pre_probe = '''
def _a_find_all(w, cls, out):
    if w.winfo_class() == cls:
        out.append(w)
    for c in w.winfo_children():
        _a_find_all(c, cls, out)

_combos = []
_a_find_all(root, "TCombobox", _combos)
_combos[0].set("기타")
_combos[0].event_generate("<<ComboboxSelected>>")
root.update()
'''
    widgets = _a_run_widget_probe(tmp_path, "gui._ask_youth_hymn_popup()", pre_probe=pre_probe)
    entries = [w for w in widgets if w['class'] == 'Entry']
    num_entries = [w for w in entries if w.get('width') == '6']
    title_entries = [w for w in entries if w.get('width') == '22']

    assert num_entries[0]['state'] == 'disabled', num_entries[0]
    assert title_entries[0]['gridded'] is True, title_entries[0]

    # 둘째 행(봉헌)은 여전히 기본 출처('나주노')이므로 영향받지 않아야 한다.
    assert num_entries[1]['state'] != 'disabled', num_entries[1]
    assert title_entries[1]['gridded'] is False, title_entries[1]


def test_a6_switching_away_from_기타_re_enables_number_and_hides_title(tmp_path):
    """'기타'로 바꿨다가 다시 다른 출처로 되돌리면 번호칸 재활성화 + 제목칸 재숨김."""
    pre_probe = '''
def _a_find_all(w, cls, out):
    if w.winfo_class() == cls:
        out.append(w)
    for c in w.winfo_children():
        _a_find_all(c, cls, out)

_combos = []
_a_find_all(root, "TCombobox", _combos)
_combos[0].set("기타")
_combos[0].event_generate("<<ComboboxSelected>>")
root.update()
_combos[0].set("나주노")
_combos[0].event_generate("<<ComboboxSelected>>")
root.update()
'''
    widgets = _a_run_widget_probe(tmp_path, "gui._ask_youth_hymn_popup()", pre_probe=pre_probe)
    entries = [w for w in widgets if w['class'] == 'Entry']
    num_entries = [w for w in entries if w.get('width') == '6']
    title_entries = [w for w in entries if w.get('width') == '22']

    assert num_entries[0]['state'] != 'disabled', num_entries[0]
    assert title_entries[0]['gridded'] is False, title_entries[0]


# ═══════════════════════════════════════════════════════════════════════════
# C그룹: 성가 버그 4건 (2026-09-24 청년미사 확장 — 00_request.md)
#
# 실측 근거: output/20260919/토요일 저녁 청년 주일미사_20260919_....pptx(실사용자가 실제로
# 생성해 본 산출물)를 직접 열어 슬라이드 14~17(나주노 151), 108(가톨릭성가), 117/118(2차봉헌
# 기타)을 대조했다.
# ═══════════════════════════════════════════════════════════════════════════

def test_c1_naju_151_system_band_excludes_copyright_footer():
    """detect_system_bands()가 151번 페이지 하단 저작권 줄(높이 53px)을 오선 시스템으로
    오검출하면 안 된다. 원인: 저작권 텍스트 행의 row_fill이 0.40~0.438로 드물게
    _STAFF_ROW_FILL(0.40)을 넘는 행이 5개뿐인데, 원래 _MIN_STAFF_ROWS=3이라 '시스템'으로
    잘못 인정됐다(실측: output/20260919 산출물의 슬라이드 17이 악보 대신 저작권 텍스트를
    보여줌 — 스퓨리어스 5번째 그룹이 그대로 슬라이드가 됨). 실제 오선 시스템은 항상 최소
    12개 이상의 그 rows>0.4를 가진다(7개 곡 40개 밴드 실측) — margin이 넓어(5 vs 12)
    _MIN_STAFF_ROWS를 8로 올리는 것이 _STAFF_ROW_FILL 자체를 올리는 것보다 안전하다
    (0.5로 올리면 362/447의 약한 시스템이 사라지는 회귀가 실측됨)."""
    import numpy as _np
    pi = hp.find_song_page("나주노", 151)
    d = _fitz.open(hp.SOURCES["나주노"]["pdf"])
    try:
        page = d[pi]
        img = hp.render_region(page, None, hp.RENDER_DPI)
        gray = _np.asarray(img.convert("L"))
        bands = hp.detect_system_bands(gray)
    finally:
        d.close()
    assert len(bands) == 5, bands
    last_top, last_bot = bands[-1]
    assert (last_bot - last_top) > 100, (
        "마지막 밴드가 비정상적으로 얇음(저작권 줄이 섞였을 가능성)", bands
    )


@pytest.mark.skipif(not _YOUTH_PDFS_AVAILABLE, reason=_YOUTH_SKIP_REASON)
def test_c1_known_good_songs_system_counts_unaffected_by_fix():
    """151 수정(_MIN_STAFF_ROWS 상향)이 기존에 정상이던 곡들의 시스템 수를 바꾸면 안 된다
    (test_S1_S5_system_and_slide_counts와 동일 실측표, 회귀 가드로 별도 보존)."""
    import numpy as _np
    expected = {362: 4, 173: 7, 146: 6, 810: 4, 267: 5}
    sources = {362: "나주노", 173: "나주노", 146: "나주노", 810: "야훼 이레", 267: "나주노"}
    for num, exp in expected.items():
        src = sources[num]
        pi = hp.find_song_page(src, num)
        d = _fitz.open(hp.SOURCES[src]["pdf"])
        try:
            page = d[pi]
            yr = None if src == "나주노" else hp.find_song_y_range(page, src, num)
            img = hp.render_region(page, yr, hp.RENDER_DPI)
            gray = _np.asarray(img.convert("L"))
            bands = hp.detect_system_bands(gray)
        finally:
            d.close()
        assert len(bands) == exp, (num, bands)


@pytest.mark.skipif(not _YOUTH_PDFS_AVAILABLE, reason=_YOUTH_SKIP_REASON)
def test_c1_naju_172_system_band_splits_merged_systems():
    """detect_system_bands()가 172번("나의 고백") 페이지에서 공백행 연속 길이 기반 1차
    분할이 시스템 2~9를 세그먼트 하나로 합쳐버리는 버그를 가진다 — 시스템 2~9 사이
    여백 구간에 가사 꼬리 등 미세한 잉크가 산발적으로 섞여 있어, 공백행(row_fill<0.005)
    연속 런이 전부 _SYSTEM_GAP_MIN 기반 임계값(이 페이지 기준 45px) 밑으로 끊긴다(실측
    최대 36px) — 그 결과 수정 전에는 총 2개 밴드(피크업 시스템 1개 + 시스템 2~9가
    뭉친 거대 밴드 1개)만 나와, 생성된 PPT가 슬라이드 1장에 시스템 1개(정상), 슬라이드
    2장째에 8개를 욱여넣은(가독 불가) 결과를 냈다.

    _split_merged_systems()가 오선행(row_fill>_STAFF_ROW_FILL) 밀도 군집으로 2차 분할을
    적용해 바로잡는다: 시스템 내부 오선행 간격(최대 21px, 151 실측)과 시스템 사이 간격
    (최소 60px, 362 실측 — 단 이 60px 자체는 362의 1px짜리 잉크 노이즈였고 실제 안전
    여백은 더 크다) 사이 _SYSTEM_LINE_MERGE_GAP=40을 기준으로 군집을 나누고,
    _MIN_STAFF_ROWS(8, 151 저작권 줄 오검출 방지용으로 이미 확립된 값)로 노이즈 군집을
    버린다. 172 페이지 전체를 오선행 군집으로 직접 재측정하면 9개 군집 사이 8개의 큰
    간격(167~190px, 전부 40보다 월등히 큼)이 뚜렷이 분리되어 나타난다."""
    import numpy as _np
    pi = hp.find_song_page("나주노", 172)
    d = _fitz.open(hp.SOURCES["나주노"]["pdf"])
    try:
        page = d[pi]
        img = hp.render_region(page, None, hp.RENDER_DPI)
        gray = _np.asarray(img.convert("L"))
        bands = hp.detect_system_bands(gray)
    finally:
        d.close()
    assert len(bands) == 9, bands
    for top, bot in bands:
        assert (bot - top) > 20, (
            "시스템 밴드가 비정상적으로 얇음(분할 경계 계산 오류 가능성)", bands
        )
    # 밴드는 페이지 순서대로 겹치지 않고 나열돼야 한다(분할 후 재조립 과정에서
    # 중간점 계산이 역전되거나 트림이 다음 밴드를 침범하지 않았는지 확인). 두 밴드가
    # 같은 ink 행에서 맞닿는 것(next_top == prev_bot)은 트림 경계가 정확히 인접
    # 시스템의 마지막/첫 ink row를 공유하는 정상 케이스라 겹침으로 보지 않는다.
    for (_, prev_bot), (next_top, _) in zip(bands, bands[1:]):
        assert next_top >= prev_bot, bands


def test_c1_naju_151_copyright_crop_resolves():
    """캐시에 151 항목을 추가한 뒤 resolve_copyright_crop()이 None이 아닌 이미지를
    반환해야 한다(요구사항: 정상은 매 슬라이드 우하단 반복 표시).

    02b_review_report.md가 151의 bbox가 실제 저작권 2줄 중 1줄만(그것도 잘려서) 담는다고
    지적했으나, 독립 재검증 결과(PDF 196페이지를 직접 렌더링 — 아래 주석) 이는 버그가
    아니라 267/447과 동일한 기존 확립 규칙("Copyright © {저작권자}" 한 줄만 담고
    `COPYRIGHT_TRIGGERS`="Administered by"/"Adm. by" 직전에서 자른다, `_MIN_STAFF_ROWS`
    문서화 주석의 "트리거 직전까지" 원칙)이 이 곡에도 동일하게 적용된 것이다:
    - 267 원문 한 줄 전체: "Copyright ©천태혁. 진경. Adm. by KCMCA. All rights reserved.
      Used by permission." — 캐시된 크롭은 "진경"에서 잘려 "Adm. by..." 이후 전부 제외.
    - 151 원문 1행: "Copyright © 1984 ... Praise Inc. Administered by CopyCare Asia."
      (2행: "All rights reserved. Used by permission. ...") — 캐시된 크롭은 "Praise Inc"에서
      잘려 "Administered by..."와 2행 전체가 제외됨. 267과 정확히 같은 패턴(트리거 직전
      cut)이며, 151의 "All rights reserved..." 문구는 267에서도 이미 트리거 이후로 제외되던
      바로 그 문구다 — 151만 특별 취급할 근거가 없다.
    아래 단언은 이 "1줄만, 트리거 직전까지"라는 확립된 형태를 회귀 가드로 굳힌다(폭이
    넓고 높이가 낮은 단일 줄 비율 — 향후 실수로 2줄을 담게 크롭이 바뀌면 이 비율이
    깨져 잡힌다)."""
    d = _fitz.open(hp.SOURCES["나주노"]["pdf"])
    try:
        pi = hp.find_song_page("나주노", 151)
        page = d[pi]
        img = hp.resolve_copyright_crop("나주노", 151, page)
    finally:
        d.close()
    assert img is not None
    assert img.width > 0 and img.height > 0
    # 단일 줄 크롭이면 폭이 높이의 여러 배는 되어야 한다(267/447과 동일 형태 — 2줄을
    # 담으면 높이가 거의 2배가 되어 이 비율이 깨진다).
    assert img.width > img.height * 5, (img.width, img.height)


def test_c1_naju_151_every_slide_has_score_and_copyright_pictures():
    """build_hymn_pptx() 종단 검증: 151번 곡의 모든 슬라이드가 악보 1장 + 저작권 1장,
    정확히 2개의 그림을 가져야 한다(저작권이 마지막 슬라이드에만 나오거나 아예 안 나오는
    회귀를 모두 잡는다)."""
    prs = hp.build_hymn_pptx('입당', '나주노', 151)
    assert len(prs.slides) >= 1
    for i, slide in enumerate(prs.slides):
        pics = [sh for sh in slide.shapes if sh.shape_type == 13]
        assert len(pics) == 2, (i, len(pics), [p.name for p in pics])


def test_c2_duplicate_slide_preserves_source_slide_layout(tmp_path):
    """duplicate_slide()는 '가장 먼저 찾은 Blank류 레이아웃'(_blank_layout())을 무조건 쓰는
    대신, 원본 슬라이드가 실제로 쓰던 레이아웃을 그대로 물려받아야 한다.

    실측 원인: assets/청년미사_성가_template.pptx의 유일한 슬라이드는 레이아웃 '빈 화면'을
    쓰는데, missa_ooxml_utils._blank_layout()은 slide_layouts를 순회하며 'Blank'/'빈 화면'/
    'blank' 중 처음 매칭되는 것을 반환한다 — 이 템플릿의 레이아웃 목록은 'Blank'가 '빈 화면'
    보다 먼저 나와(Title Slide, Title and Content, Two Content, Title Only, Blank, 빈 화면)
    duplicate_slide()로 복제된 슬라이드는 전부 'Blank'를 쓰게 된다. 원본(슬라이드 0)만
    '빈 화면'으로 남아 슬라이드 간 레이아웃이 갈라진다(2026-09-24 실사용자 보고 재현:
    output/20260919 산출물에서 나주노 151 슬라이드 14='빈 화면', 15~17='Blank')."""
    import missa_ooxml_utils as ou
    from pptx import Presentation as _CPresentation

    prs = _CPresentation(str(hp.TEMPLATE))
    original_layout_name = prs.slides[0].slide_layout.name
    assert original_layout_name != 'Blank', (
        "전제 확인: 템플릿 슬라이드가 'Blank'가 아닌 레이아웃을 써야 이 테스트가 의미 있음"
    )
    new_idx = ou.duplicate_slide(prs, 0)
    assert prs.slides[new_idx].slide_layout.name == original_layout_name, (
        prs.slides[new_idx].slide_layout.name, original_layout_name,
    )


def test_c2_naju_151_fresh_build_all_slides_share_same_layout():
    """C2 수정 후 build_hymn_pptx()로 새로 만든 151번 곡은 슬라이드 전부가 원본(템플릿)과
    동일한 레이아웃을 써야 한다(층별로 갈라지지 않음)."""
    prs = hp.build_hymn_pptx('입당', '나주노', 151)
    layout_names = {slide.slide_layout.name for slide in prs.slides}
    assert len(layout_names) == 1, layout_names


def test_c3_catholic_hymn_title_derived_from_filename_when_not_provided(tmp_path, monkeypatch):
    """가톨릭성가는 UI에 제목 입력칸이 없다(A6 — 번호만 입력). resolve_youth_hymn_pptx()가
    찾은 파일명('성가 {번호} {제목}.pptx' 관례, 실측: OneDrive '성가 96 하느님 약속하신.pptx')
    에서 제목을 자동으로 뽑아야 한다 — 그렇지 않으면 `제목 or ''`가 빈 문자열이 되어 title
    슬라이드가 '성 체 / (빈 줄) / 가톨릭 성가 96'처럼 제목이 누락된 채 저장된다(2026-09-24
    실사용자 보고, output/20260919 슬라이드 108 재현)."""
    import missa_content_updaters as cu3
    import missa_gui as gui3
    from pptx import Presentation as _C3Presentation

    # find_youth_onedrive_hymn_file()은 로컬 동기화 폴더를 보지 않고 로컬 캐시부터
    # 뒤진다(§H3와 동일 이유) — 파일을 그 캐시 경로에 둔다.
    fake_prs = _C3Presentation()
    fake_prs.slides.add_slide(fake_prs.slide_layouts[6])
    dest_dir = tmp_path / 'cache' / '청년_가톨릭성가'
    dest_dir.mkdir(parents=True)
    fake_path = dest_dir / '성가 96 하느님 약속하신 분.pptx'
    fake_prs.save(str(fake_path))

    monkeypatch.setattr(gui3, '_SCRIPT_DIR', tmp_path)

    entry = {'출처': '가톨릭성가', '번호': 96, '제목': None}
    _src_prs, resolved_title = cu3.resolve_youth_hymn_pptx(entry, '성체')

    assert resolved_title == '하느님 약속하신 분', repr(resolved_title)


def test_c3_catholic_hymn_explicit_title_overrides_filename(tmp_path, monkeypatch):
    """제목이 명시적으로 주어지면(향후 UI 변경 대비) 파일명 추정보다 우선한다."""
    import missa_content_updaters as cu3
    import missa_gui as gui3
    from pptx import Presentation as _C3Presentation

    fake_prs = _C3Presentation()
    fake_prs.slides.add_slide(fake_prs.slide_layouts[6])
    dest_dir = tmp_path / 'cache' / '청년_가톨릭성가'
    dest_dir.mkdir(parents=True)
    fake_path = dest_dir / '성가 96 하느님 약속하신 분.pptx'
    fake_prs.save(str(fake_path))

    monkeypatch.setattr(gui3, '_SCRIPT_DIR', tmp_path)

    entry = {'출처': '가톨릭성가', '번호': 96, '제목': '사용자가 직접 입력한 제목'}
    _src_prs, resolved_title = cu3.resolve_youth_hymn_pptx(entry, '성체')

    assert resolved_title == '사용자가 직접 입력한 제목'


_Y2_TEMPLATE_C4 = (
    BASE / "reference" / "청년미사"
    / "Template_토요일 저녁 청년 주일미사_20260822_연중 제21주일.pptx"
)


@pytest.fixture
def y2_prs2():
    if not _Y2_TEMPLATE_C4.is_file():
        pytest.skip(f"청년미사 참조 PPT가 로컬에 없음(reference/청년미사/, git 미포함): {_Y2_TEMPLATE_C4}")
    from pptx import Presentation as _C4Presentation
    return _C4Presentation(str(_Y2_TEMPLATE_C4))


# ---------------------------------------------------------------------------
# C4: 2차봉헌(출처='기타')에 엉뚱한 성가(템플릿 기본 예시)가 남는 버그
#
# 실측(F3/F4, test_missa_regression.py): 참조 템플릿의 성체2 슬롯(title_idx=106,
# content=107)은 '기타' 예시로 만들어져 있어 콘텐츠가 항상 가사 placeholder만 있고 실제
# 악보가 없다 — 이것이 사용자가 "정상"이라 부른 114번 패턴이다. 반면 2차봉헌 슬롯
# (title_idx=109, content=110~111)은 야훼이레 810 "실제 곡"을 기본 예시로 담고 있다.
# 기존 `_replace_one_youth_song()`의 '기타' 분기는 "콘텐츠를 건드리지 않는다"는 전제였는데,
# 이 전제는 성체2에만 우연히 맞고 2차봉헌에는 틀렸다 — 그 결과 2차봉헌에 '기타'를 선택하면
# 헤더 텍스트만(그것도 content_start 슬라이드 1장에만) 갱신되고 실제 야훼이레 810 악보
# 그림은 그대로 남는다(2026-09-24 실사용자 보고, output/20260919 슬라이드 117/118 재현).
# ---------------------------------------------------------------------------

def test_c4_기타_2차봉헌_content_replaced_not_left_as_template_default(y2_prs2):
    """2차봉헌에 출처='기타'를 선택하면 템플릿 기본 예시(야훼이레 810 실제 악보 그림)가
    남아있으면 안 되고, 성체2 슬롯과 동일한 '빈' 콘텐츠(가사 placeholder만)로 교체돼야
    하며, 콘텐츠 슬라이드 수도 성체2처럼 1장이어야 한다(원래 2차봉헌은 2장)."""
    import missa_content_updaters as cu4
    import missa_sections as sec4

    성가_선택 = {'2차봉헌': [{'출처': '기타', '번호': None, '제목': '테스트 2차봉헌 제목'}]}
    cu4.replace_성가_youth(y2_prs2, 성가_선택)

    sections = sec4.find_sections(y2_prs2, mass_type='youth')
    slot = sections['2차봉헌_songs'][0]

    # 성체2처럼 콘텐츠 슬라이드 1장으로 축소돼야 한다(원래 템플릿 기본값은 2장).
    assert slot['content_end'] - slot['content_start'] == 1, slot

    content_slide = y2_prs2.slides[slot['content_start']]
    pic_names = [sh.name for sh in content_slide.shapes if sh.shape_type == 13]
    # 야훼이레 810의 실제 악보 그림(F4 실측 이름) 흔적이 전혀 없어야 한다.
    assert 'Picture 4' not in pic_names and 'Picture 5' not in pic_names, pic_names

    header = hp._find_header(content_slide)
    assert '테스트 2차봉헌 제목' in header.text_frame.text, header.text_frame.text


def test_c4_기타_2차봉헌_leaves_no_leftover_scratch_slide(y2_prs2):
    """빈 콘텐츠 복제용으로 임시로 만든 scratch 슬라이드가 최종 결과물에 남아있으면
    안 된다(전체 슬라이드 수가 순수하게 2차봉헌 콘텐츠 1장 축소분(-1)만 반영해야 함)."""
    import missa_content_updaters as cu4

    before = len(y2_prs2.slides)
    성가_선택 = {'2차봉헌': [{'출처': '기타', '번호': None, '제목': '테스트'}]}
    cu4.replace_성가_youth(y2_prs2, 성가_선택)
    after = len(y2_prs2.slides)
    assert after == before - 1, (before, after)


def test_c4_성체2_own_기타_slot_unaffected_by_self_replacement(y2_prs2):
    """성체2 자신이 '기타'로 처리될 때도(자기 자신의 콘텐츠로 자기 자신을 교체하는 셈) 헤더가
    정확히 갱신되고 구조가 깨지지 않아야 한다(회귀 가드 — '정상'이라던 114 패턴 자체가
    이번 수정으로 깨지면 안 됨)."""
    import missa_content_updaters as cu4
    import missa_sections as sec4

    성가_선택 = {
        '성체': [
            {'출처': '기타', '번호': None, '제목': '첫째 성체'},
            {'출처': '기타', '번호': None, '제목': '둘째 성체'},
        ],
    }
    cu4.replace_성가_youth(y2_prs2, 성가_선택)

    sections = sec4.find_sections(y2_prs2, mass_type='youth')
    songs = sections['성체_songs']
    assert len(songs) == 2, songs
    slot2 = songs[1]
    assert slot2['content_end'] - slot2['content_start'] == 1, slot2
    header = hp._find_header(y2_prs2.slides[slot2['content_start']])
    assert '둘째 성체' in header.text_frame.text, header.text_frame.text


# ---------------------------------------------------------------------------
# D그룹: 나머지 버그 3건 (00_request.md D1~D3)
#
# D1 — 슬라이드 53(청년 템플릿 '복음환호송' 가운데 슬라이드): 실사용 산출물
# (output/20260919/...pptx) 실측 확인 결과, JSON 복음환호송(예:
# {'content': '◎ 알렐루야.\n○ 문구\n◎ 알렐루야.'})의 ◎ 줄까지 그대로 렌더링돼
# "○ \t알렐루야."가 문구 앞뒤에 남았다. 청년 템플릿의 '복음환호송' 가운데 슬라이드는
# 앞뒤에 별도의 고정 'ALLELUIA'/'알렐루야' 장식 슬라이드(51/53 등)가 이미 있어, 이
# Rectangle 3 도형 자체는 원래 ○ 구절 문단 1개짜리다(성인 템플릿처럼 ◎/○/◎ 3문단
# 구조가 아님). `update_복음환호송()`의 기존 로직은 성인 템플릿의 3문단(para[0]=◎,
# para[1]=○,para[2]=◎) 전제로 짜여 있어, 청년처럼 문단이 1개뿐이면 그 하나의 문단이
# tmpl_allel0/tmpl_verse/tmpl_allel2 세 역할에 전부 별칭(alias)되고, ◎ 줄 처리 시
# "○ " 라벨 run을 그대로 물려받아 "○ \t알렐루야."로 오염된다. 근본 수정: mass_type='youth'
# 일 때는 ◎ 줄(비-verse)을 애초에 처리 대상에서 제외하고 ○ 구절 문단 하나만 갱신한다 —
# 이렇게 하면 원래 템플릿 문단 개수(1개)가 유지돼 폰트도 강제 축소(_adjust_fit_if_needed)
# 없이 원본 그대로 남는다(D1의 "가운데 문구 폰트 불일치" 증상도 같은 원인의 다른 증상).
# ---------------------------------------------------------------------------

def test_d1_복음환호송_youth_drops_alleluia_bookend_lines(y2_prs2):
    import missa_content_updaters as cu_d
    import missa_sections as sec_d

    sections = sec_d.find_sections(y2_prs2, mass_type='youth')
    idx = sections['복음환호송']
    json_data = {'복음환호송': {
        'title': '테스트 참조',
        'content': '◎ 알렐루야.\n○ 테스트 검증 문구입니다.\n◎ 알렐루야.',
    }}

    cu_d.update_복음환호송(y2_prs2, json_data, sections, mass_type='youth')

    rect = next(
        sh for sh in y2_prs2.slides[idx].shapes
        if sh.has_text_frame and sh.name == 'Rectangle 3'
    )
    paras = rect.text_frame._txBody.findall(_qn('a:p'))
    # 청년 템플릿 원본과 동일하게 문단 1개(○ 구절)만 남아야 한다 — ◎ 알렐루야 줄이
    # 별도 문단으로 추가되면 안 된다.
    assert len(paras) == 1, [p.findtext(_qn('a:r') + '/' + _qn('a:t')) for p in paras]
    text = rect.text_frame.text
    assert '알렐루야' not in text, text
    assert '◎' not in text, text
    assert '테스트 검증 문구입니다' in text, text


def test_d1_복음환호송_youth_preserves_template_font_on_verse_line(y2_prs2):
    """가운데 구절 문단의 폰트(크기/서체)가 템플릿 원본 그대로 유지돼야 한다 —
    ◎ 줄이 남아 문단이 3개가 되면 _adjust_fit_if_needed()가 오버플로 방지를 위해
    폰트를 강제로 줄여 템플릿과 달라진다(D1 '가운데 문구 폰트 불일치')."""
    import missa_content_updaters as cu_d
    import missa_sections as sec_d

    sections = sec_d.find_sections(y2_prs2, mass_type='youth')
    idx = sections['복음환호송']
    rect_before = next(
        sh for sh in y2_prs2.slides[idx].shapes
        if sh.has_text_frame and sh.name == 'Rectangle 3'
    )
    r_before = rect_before.text_frame._txBody.findall(_qn('a:p'))[0].findall(_qn('a:r'))[-1]
    sz_before = r_before.find(_qn('a:rPr')).get('sz')
    assert sz_before is not None

    json_data = {'복음환호송': {
        'title': '테스트 참조',
        'content': '◎ 알렐루야.\n○ 테스트 검증 문구입니다.\n◎ 알렐루야.',
    }}
    cu_d.update_복음환호송(y2_prs2, json_data, sections, mass_type='youth')

    rect_after = next(
        sh for sh in y2_prs2.slides[idx].shapes
        if sh.has_text_frame and sh.name == 'Rectangle 3'
    )
    r_after = rect_after.text_frame._txBody.findall(_qn('a:p'))[0].findall(_qn('a:r'))[-1]
    sz_after = r_after.find(_qn('a:rPr')).get('sz')
    assert sz_after == sz_before, (sz_before, sz_after)


def test_d1_복음환호송_adult_default_mass_type_signature_backward_compatible(y2_prs2):
    """mass_type 파라미터는 기본값 'adult'로 기존 호출부(위치 인자 3개만 넘기는 호출)를
    깨면 안 된다 — 시그니처 변경 자체가 회귀 위험이므로 최소한 예외 없이 호출되는지만
    확인한다(실제 성인 3문단 렌더링 결과는 기존 회귀 스위트가 20260624/20260712
    픽스처로 이미 보호)."""
    import missa_content_updaters as cu_d
    import missa_sections as sec_d

    sections = sec_d.find_sections(y2_prs2, mass_type='youth')
    json_data = {'복음환호송': {
        'title': '테스트 참조',
        'content': '◎ 알렐루야.\n○ 테스트 검증 문구입니다.\n◎ 알렐루야.',
    }}
    # mass_type 인자 없이 호출 — 기본값 'adult' 경로로 예외 없이 동작해야 한다.
    cu_d.update_복음환호송(y2_prs2, json_data, sections)


# ---------------------------------------------------------------------------
# D2 — 슬라이드 57-60(청년 템플릿 '영문 복음' 본문): 실측(reference/청년미사/
# Template_..._20260822_...pptx, find_sections(mass_type='youth')['복음_start':'복음_end']
# = 53..57) 결과, 템플릿이 원래 담고 있던 영문 샘플 콘텐츠 자체가 페이지마다 다른 박스
# 높이를 쓴다(53/54=4809458, 55=2101024 — 템플릿 제작자가 자신의 샘플 텍스트 분량에 맞춰
# 마지막 페이지 박스를 손으로 줄여놓은 것으로 보인다). `replace_reading_slides()`가 새
# 페이지를 필요한 만큼만 삽입/삭제하고(needed==n_usable이면 삽입/삭제 없음) 기존 슬라이드를
# 그대로 재사용하는데, 그 "그대로"에 이 이질적인 박스 높이가 섞여 있다. 반면 페이지 분배
# (`layout_units_on_slides_pil`)는 `sec['복음_start']`(53번, 큰 박스) 하나만 기준으로 각
# 페이지의 word-wrap 줄 수를 계산한다 — 그 결과 55번처럼 원래 작은 박스를 물려받는 페이지는
# "큰 박스 기준으로 계산된 줄 수"가 실제 작은 박스에 안 들어가 마지막 줄이 잘린다(58/59
# 워드랩 잘림). 추가로 이 영문 텍스트박스는 `algn` 속성이 아예 없어(None) 기본 정렬(좌측)로
# 나오는데, 요구사항은 중앙 정렬이다.
# ---------------------------------------------------------------------------

def _d2_units(text, n=3):
    """9줄 분량 텍스트 n개를 이어 붙인 units(parse_into_verse_units 결과) — 슬라이드 단위 채우기가
    (9줄 + 9줄 + 나머지)로 최소 n장을 채우도록 한다."""
    import missa_reading_layout as rl_d
    return rl_d.parse_into_verse_units('\n'.join([text] * n))


def _d2_nine_line_text(pil_font, box_px):
    """LINES_PER_SLIDE(9)에 딱 맞는 워드랩 텍스트를 실측으로 만든다 — 너무 짧은 텍스트를 쓰면 필요한
    슬라이드 수가 3장 미만이 되어 3개 슬라이드 박스를 검증하는 이 테스트들의 전제가 어긋난다."""
    import missa_reading_layout as rl_d
    words = ('lorem ipsum dolor sit amet consectetur adipiscing elit sed do '
              'eiusmod tempor incididunt ut labore et dolore magna aliqua ' * 10).split()
    line = ''
    lines = []
    for w in words:
        cand = (line + ' ' + w) if line else w
        if pil_font.getlength(cand) <= box_px:
            line = cand
        else:
            lines.append(line)
            line = w
        if len(lines) >= 9:
            break
    if len(lines) < 9 and line:
        lines.append(line)
    text = ' '.join(lines)
    assert len(rl_d._pil_line_starts([{'text': text, 'orange': []}], pil_font, box_px)) == 9
    return text


def test_d2_replace_reading_slides_normalizes_page_height_to_template(y2_prs2):
    """needed==n_usable(3페이지, 삽입/삭제 없음)일 때도, normalize_page_size=True를 넘기면
    재사용되는 각 페이지 콘텐츠 박스의 height가 template_idx(53번, 4809458)와 같아져야
    한다 — 55번처럼 템플릿에 원래 박혀 있던 더 작은 박스(2101024)를 그대로 물려받으면
    안 된다. 리뷰(02b_review_report.md)에서 이 정규화가 성인 경로에서도 항상 실행되면
    기존 슬라이드 지오메트리를 조용히 바꾸는 부작용이 실측 확인돼(20260712 제2독서
    슬라이드 62/63), 기본값을 False로 바꾸고 호출부가 명시적으로 opt-in하도록 했다."""
    import missa_reading_layout as rl_d
    import missa_sections as sec_d

    sections = sec_d.find_sections(y2_prs2, mass_type='youth')
    cs, ce = sections['복음_start'], sections['복음_end']
    template_shape = rl_d._find_content_shape(y2_prs2.slides[cs])
    template_height = template_shape.height
    assert template_height == 4809458  # 실측 고정값(위 주석 근거)
    pil_font, box_px = rl_d._get_slide_render_params(y2_prs2.slides[cs])

    units = _d2_units(_d2_nine_line_text(pil_font, box_px))
    rl_d.replace_reading_slides(y2_prs2, cs, ce, units, cs, label='복음',
                                 normalize_page_size=True)

    for page_i in range(3):
        shape = rl_d._find_content_shape(y2_prs2.slides[cs + page_i])
        assert shape.height == template_height, (page_i, shape.height, template_height)


def test_d2_replace_reading_slides_normalize_page_size_default_off(y2_prs2):
    """normalize_page_size 기본값(False)이면 청년 템플릿에서도 페이지 박스 크기를 건드리지
    않아야 한다 — 명시적으로 켜지 않는 한 이 부작용이 발생하면 안 된다(회귀 가드,
    02b_review_report.md 대응)."""
    import missa_reading_layout as rl_d
    import missa_sections as sec_d

    sections = sec_d.find_sections(y2_prs2, mass_type='youth')
    cs, ce = sections['복음_start'], sections['복음_end']
    pil_font, box_px = rl_d._get_slide_render_params(y2_prs2.slides[cs])
    before_heights = [
        rl_d._find_content_shape(y2_prs2.slides[cs + i]).height for i in range(3)
    ]
    assert len(set(before_heights)) > 1, before_heights  # 전제: 실제로 서로 다름(53/54 vs 55)

    units = _d2_units(_d2_nine_line_text(pil_font, box_px))
    rl_d.replace_reading_slides(y2_prs2, cs, ce, units, cs, label='복음')

    after_heights = [
        rl_d._find_content_shape(y2_prs2.slides[cs + i]).height for i in range(3)
    ]
    assert after_heights == before_heights, (before_heights, after_heights)


def test_d2_replace_reading_slides_adult_fixture_geometry_untouched_by_default():
    """리뷰가 실측 확인한 실제 성인 회귀 픽스처(20260712 Template, 제2독서 콘텐츠
    슬라이드 62/63)로, normalize_page_size 기본값이 이 픽스처의 기존 박스 크기 차이를
    절대 건드리지 않음을 직접 검증한다. 실제 파이프라인과 동일하게(단편 필러 텍스트가
    아니라) 그날 실제 JSON 제2독서 콘텐츠로 슬라이드 단위 채우기를 돌려 실제 경로 그대로 재현한다."""
    import json as _json
    import missa_reading_layout as rl_d
    import missa_sections as sec_d
    from pptx import Presentation as _AdultPresentation

    path = (
        BASE / 'output' / '20260712'
        / 'Template_20260628_연중 제13주일 (교황주일).pptx'
    )
    json_path = BASE / 'output' / '20260712' / 'missa_20260712.json'
    if not path.is_file() or not json_path.is_file():
        pytest.skip(f'성인 회귀 픽스처 없음: {path}')
    prs = _AdultPresentation(str(path))
    sections = sec_d.find_sections(prs)
    cs, ce = sections['제2독서_start'], sections['제2독서_end']
    # n_usable=2(62,63) + n_ending=1(64, "주님의 말씀입니다." 종료 슬라이드) — 62/63만
    # 본문 페이지 대상이다. before/after는 이 2장만 비교한다.
    before = [rl_d._find_content_shape(prs.slides[cs + i]).height for i in range(2)]
    assert len(set(before)) > 1, before  # 실측 전제: 62/63 박스 크기가 원래부터 다름

    content = _json.loads(json_path.read_text(encoding='utf-8'))['제2독서']['content']
    units = rl_d.parse_into_verse_units(content)
    rl_d.replace_reading_slides(prs, cs, ce, units, cs, line_spacing=1.1, label='제2독서')

    after = [rl_d._find_content_shape(prs.slides[cs + i]).height for i in range(2)]
    assert after == before, (before, after)


def test_d2_replace_reading_slides_center_aligns_when_requested(y2_prs2):
    """align='ctr'을 넘기면 새로 채워진 각 단락의 algn이 'ctr'이 돼야 한다(영문 복음
    전용 요구사항 — 기존 template pPr에는 algn 자체가 없어 상속만으로는 중앙 정렬이
    안 된다)."""
    import missa_reading_layout as rl_d
    import missa_sections as sec_d

    sections = sec_d.find_sections(y2_prs2, mass_type='youth')
    cs, ce = sections['복음_start'], sections['복음_end']
    pil_font, box_px = rl_d._get_slide_render_params(y2_prs2.slides[cs])
    units = _d2_units(_d2_nine_line_text(pil_font, box_px))
    rl_d.replace_reading_slides(y2_prs2, cs, ce, units, cs, label='복음', align='ctr')

    for page_i in range(3):
        shape = rl_d._find_content_shape(y2_prs2.slides[cs + page_i])
        for para in shape.text_frame.paragraphs:
            if not para.text.strip():
                continue
            pPr = para._pPr
            assert pPr is not None and pPr.get('algn') == 'ctr', (page_i, para.text)


def test_d2_replace_reading_slides_align_default_none_backward_compatible(y2_prs2):
    """align 기본값(None)은 기존 호출부(adult 독서/복음, 성인 회귀 픽스처)를 절대
    건드리면 안 된다 — 명시적으로 넘기지 않으면 algn을 새로 주입하지 않는다."""
    import missa_reading_layout as rl_d
    import missa_sections as sec_d

    sections = sec_d.find_sections(y2_prs2, mass_type='youth')
    cs, ce = sections['복음_start'], sections['복음_end']
    pil_font, box_px = rl_d._get_slide_render_params(y2_prs2.slides[cs])
    units = _d2_units(_d2_nine_line_text(pil_font, box_px))
    rl_d.replace_reading_slides(y2_prs2, cs, ce, units, cs, label='복음')

    shape = rl_d._find_content_shape(y2_prs2.slides[cs])
    for para in shape.text_frame.paragraphs:
        if not para.text.strip():
            continue
        pPr = para._pPr
        assert pPr is None or pPr.get('algn') is None


# ---------------------------------------------------------------------------
# D3 — 슬라이드 106(청년 템플릿 '영성체송'): 실사용 산출물(output/20260919/...pptx)
# 실측 결과, 이 슬라이드의 본문 도형('직사각형 2')은 <a:spAutoFit/>(텍스트에 맞춰 도형
# 높이 자동 조절)이 걸려 있어 실제 PowerPoint에서 문구 길이에 따라 상자가 아래로
# 늘어난다. `_adjust_fit_if_needed()`가 이미 "겹치면 줄간격→폰트 축소" 로직을 갖고
# 있고 `update_영성체송()`도 이미 이 함수를 호출하지만(D3가 요구하는 "성인 경로의 유사
# 로직" 재사용 확인 결과, 정확히 이 함수가 그 로직이었다), char-count 근사
# (`_estimate_text_lines`)와 실제 필요 높이 사이의 안전 마진이 실측 기준 약 2.7%밖에
# 안 돼(전례문 상단까지 여유 4177044 EMU, 추정 높이 4064000 EMU) 이 청년미사
# 영성체송 문구처럼 경계에 딱 걸치는 경우 조정이 아예 트리거되지 않는다. CLAUDE.md에
# 이미 기록된 "Pillow ascent+descent가 실제 PowerPoint 줄 높이보다 14~18% 작다"는
# 실측 보정치를 재사용해(신규 임의 값 도입 대신) 안전 마진을 추가한다 — 이러면 이
# 근소한 초과 케이스가 정확히 (a) 줄간격 축소 단계에서 잡힌다(폰트 축소까지는 필요
# 없음, 요구사항 "자동 줄간격 축소로 한 슬라이드에 맞춤"과 정확히 일치).
# ---------------------------------------------------------------------------

def test_d3_영성체송_youth_triggers_line_spacing_reduction_for_near_miss_overflow(y2_prs2):
    import missa_content_updaters as cu_d
    import missa_sections as sec_d

    sections = sec_d.find_sections(y2_prs2, mass_type='youth')
    json_data = {'영성체송': {
        'title': '테스트 참조',
        'content': (
            '주님이 말씀하신다. 누구든지 사람들 앞에서 나를 안다고 증언하면, '
            '나도 하늘에 계신 내 아버지 앞에서 그를 안다고 증언하리라.'
        ),
    }}
    cu_d.update_영성체송(y2_prs2, json_data, sections)

    slide = y2_prs2.slides[sections['영성체송']]
    shape = next(
        sh for sh in slide.shapes
        if sh.has_text_frame and '누구든지' in sh.text_frame.text
    )
    spc_pct = cu_d._shape_first_para_line_spacing_pct(shape)
    assert spc_pct == 100000, spc_pct  # 줄간격이 100%로 축소돼야 한다(템플릿 원본은 160%)

    font_emu = cu_d._shape_first_run_font_size_emu(shape)
    assert font_emu == int(40 * 12700), font_emu  # 줄간격 축소만으로 충분 — 폰트까지 줄이면 과교정


def test_d3_영성체송_youth_short_content_unaffected(y2_prs2):
    """충분히 짧은 문구는 안전 마진을 추가해도 여전히 조정이 트리거되면 안 된다
    (과도하게 보수적인 마진으로 인한 불필요한 축소 방지 회귀 가드)."""
    import missa_content_updaters as cu_d
    import missa_sections as sec_d

    sections = sec_d.find_sections(y2_prs2, mass_type='youth')
    json_data = {'영성체송': {'title': '테스트 참조', 'content': '주님이 말씀하신다.'}}
    cu_d.update_영성체송(y2_prs2, json_data, sections)

    slide = y2_prs2.slides[sections['영성체송']]
    shape = next(
        sh for sh in slide.shapes
        if sh.has_text_frame and '주님이 말씀하신다' in sh.text_frame.text
    )
    spc_pct = cu_d._shape_first_para_line_spacing_pct(shape)
    assert spc_pct == 160000, spc_pct  # 템플릿 원본 줄간격 그대로 유지


# ═══════════════════════════════════════════════════════════════════════════
# [승격 2026-10-03] SF그룹 — 독서/복음 "슬라이드 단위 정확히 9줄 채우기"(replace_reading_slides) 회귀
#
# test_missa_progression.py에서 승격. 가짜 COM(결정적) 테스트 sf1~sf31과 실제 PowerPoint COM 테스트
# sf_real1~6(pywin32/PowerPoint 부재 시 자체 skip). 아래 상수·헬퍼는 SF그룹 전용이다.
# 의존 산출물: output/20260705·20260712 템플릿은 git 추적됨. 청년 템플릿(reference/, gitignore)과
# 영문 복음 픽스처(_workspace/, gitignore)는 없으면 해당 테스트가 skip 된다.
# ═══════════════════════════════════════════════════════════════════════════
BASE = REPO_ROOT
_E_BASE = REPO_ROOT
_E_YOUTH_TEMPLATE = (
    _E_BASE / "reference" / "청년미사"
    / "Template_토요일 저녁 청년 주일미사_20260822_연중 제21주일.pptx"
)
_E_EN_GOSPEL_FIXTURE = (
    _E_BASE / "_workspace" / "청년미사_1단계" / "verify_20260913" / "missa_en_20260913.json"
)

# 로마서 8,31-39 제2독서 재현 상수 — SF-real3(test_sf_real3_*)이 사용한다.
# 구 TestSplitAndAdjustViaComKeepsBestPartialResult / test_second_reading_rebalance_converges_without_
# large_deficit(2026-09-20 제2독서 6줄/4줄 미달 버그)는 2026-10-03 슬라이드 단위 채우기로 사후 재조정이 사라져
# 삭제됐고, 그 요구사항(대량 미달 금지)은 비마지막 슬라이드 "정확히 9줄"로 강화돼 SF1/SF-real3이 지킨다.
import types as _b_types

import missa_reading_layout as rl_b
import missa_sections as sec_b
from pptx import Presentation as _BPresentation

_B_TEMPLATE = (
    BASE / "output" / "20260712"
    / "Template_20260628_연중 제13주일 (교황주일).pptx"
)

# 2026-08-23자 실사용자 로마서 8,31-39 본문 그대로(20260920 주일 제2독서 재현에 쓰인 실측값).
_B_CONTENT_ROM_8_31_39 = (
    "형제 여러분, 31 하느님께서 우리 편이신데 누가 우리를 대적하겠습니까? "
    "32 당신의 친아드님마저 아끼지 않으시고 우리 모두를 위하여 내어 주신 분께서, "
    "어찌 그 아드님과 함께 모든 것을 우리에게 베풀어 주지 않으시겠습니까? "
    "33 하느님께 선택된 이들을 누가 고발할 수 있겠습니까? 그들을 의롭게 해 주시는 분은 "
    "하느님이십니다. 34 누가 그들을 단죄할 수 있겠습니까? 돌아가셨다가 참으로 되살아나신 분, "
    "또 하느님의 오른쪽에 앉아 계신 분, 그리고 우리를 위하여 간구해 주시는 분이 바로 "
    "그리스도 예수님이십니다. 35 무엇이 우리를 그리스도의 사랑에서 갈라놓을 수 있겠습니까? "
    "환난입니까? 역경입니까? 박해입니까? 굶주림입니까? 헐벗음입니까? 위험입니까? 칼입니까? "
    "36 이는 성경에 기록된 그대로입니다. \"저희는 온종일 당신 때문에 살해되며 도살될 양처럼 "
    "여겨집니다.\" 37 그러나 우리는 우리를 사랑해 주신 분의 도움에 힘입어 이 모든 것을 이겨 "
    "내고도 남습니다. 38 나는 확신합니다. 죽음도, 삶도, 천사도, 권세도, 현재의 것도, "
    "미래의 것도, 권능도, 39 저 높은 곳도, 저 깊은 곳도, 그 밖의 어떠한 피조물도 우리 주 "
    "그리스도 예수님에게서 드러난 하느님의 사랑에서 우리를 떼어 놓을 수 없습니다."
)


# ═══════════════════════════════════════════════════════════════════════════
# W그룹(이력) — 2026-10-03 복음 슬라이드 꼬리 줄 widow('짓는' 두 글자만 마지막 줄) 방지
#
# 구 _prevent_widow_tails 후처리 테스트(w0~w11)는 2026-10-03 슬라이드 단위 채우기로 후처리 자체가 사라져
# 삭제됐다. 요구사항 이전: 줄 경계에서 자르므로 꼬리가 구조적으로 안 생김 → test_sf21(결정적)·
# test_sf_real2(20261004 실 COM, 비마지막 정확히 9줄). 아래는 SF그룹이 쓰는 헬퍼/상수만 남긴 것이다.
# ═══════════════════════════════════════════════════════════════════════════

import sys as _w_sys
from pptx.dml.color import RGBColor as _WRGB
from pptx.util import Emu as _WEmu, Pt as _WPt


def _w_mk_prs():
    from pptx import Presentation as _P
    return _P()


def _w_add_slide(prs, paras):
    """paras: [[(text, orange_bool), ...], ...] — 한 단락 = run 목록."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(_WEmu(1_000_000), _WEmu(500_000), _WEmu(7_000_000), _WEmu(5_000_000))
    tf = box.text_frame
    tf.word_wrap = True
    for i, runs in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        for text, orange in runs:
            r = p.add_run()
            r.text = text
            r.font.size = _WPt(32)
            if orange:
                r.font.color.rgb = _WRGB(0xFF, 0xC0, 0x00)
    return slide


_W_GOSPEL_20261004_TITLE = '마태오가 전한 거룩한 복음입니다.'
_W_GOSPEL_20261004 = (
    '그때에 예수님께서 수석 사제들과 백성의 원로들에게 말씀하셨다. 33 “다른 비유를 들어 보아라. 어떤 밭 임자가 ‘포도밭을 '
    '일구어 울타리를 둘러치고 포도 확을 파고 탑을 세웠다.’ 그리고 소작인들에게 내주고 멀리 떠났다. 34 포도 철이 가까워지자 '
    '그는 자기 몫의 소출을 받아 오라고 소작인들에게 종들을 보냈다. 35 그런데 소작인들은 그들을 붙잡아 하나는 매질하고 하나는 '
    '죽이고 하나는 돌을 던져 죽이기까지 하였다. 36 주인이 다시 처음보다 더 많은 종을 보냈지만, 소작인들은 그들에게도 같은 '
    '짓을 하였다. 37 주인은 마침내 ‘내 아들이야 존중해 주겠지.’ 하며 그들에게 아들을 보냈다. 38 그러나 소작인들은 아들을'
    ' 보자, ‘저자가 상속자다. 자, 저자를 죽여 버리고 우리가 그의 상속 재산을 차지하자.’ 하고 저희끼리 말하면서, 39 그를'
    ' 붙잡아 포도밭 밖으로 던져 죽여 버렸다. 40 그러니 포도밭 주인이 와서 그 소작인들을 어떻게 하겠느냐?” 41 “그렇게 '
    '악한 자들은 가차 없이 없애 버리고, 제때에 소출을 바치는 다른 소작인들에게 포도밭을 내줄 것입니다.” 하고 그들이 대답하자,'
    ' 42 예수님께서 그들에게 말씀하셨다. “너희는 성경에서 이 말씀을 읽어 본 적이 없느냐? ‘집 짓는 이들이 내버린 돌, 그 '
    '돌이 모퉁이의 머릿돌이 되었네. 이는 주님께서 이루신 일, 우리 눈에 놀랍기만 하네.’ 43 그러므로 내가 너희에게 말한다. '
    '하느님께서는 너희에게서 하느님의 나라를 빼앗아, 그 소출을 내는 민족에게 주실 것이다.”'
)
_W_GOSPEL_TEMPLATE = BASE / 'output' / '20260705' / 'Template_성수축복_20260614_연중 제11주일.pptx'


# ═══════════════════════════════════════════════════════════════════════════
# SF그룹 — 2026-10-03 독서/복음 "슬라이드 바이 슬라이드 정확히 채우기" 구조 교체
#
# 설계: _workspace/refactor_슬라이드단위_채우기/design.md. 전체 추정 분할(layout_units_on_slides*)
# 후 사후 보정(_rebalance_reading_slides_post_write 등)하던 구조를, 각 슬라이드의 실제 본문 상자에서
# 줄 시작 오프셋을 실측(COM, 불가 시 Pillow)해 정확히 첫 9줄만 확정하고 나머지로 다음 슬라이드를
# 같은 방식으로 채우는 구조로 바꿨다. 불변조건: 비마지막 슬라이드 정확히 9줄 / 마지막 <= 9줄 /
# 텍스트 무손실 / 절 번호 오렌지 / 오버플로 없음.
#
# 테스트 도구: 가짜 COM(고정폭 글자 10px, 슬라이드 상자 폭으로 줄 수 결정 — 구현과 독립된 줄바꿈
# 구현)과 Pillow 폴백(_SFFont). 전역 상태는 monkeypatch.setattr로 새 객체로 교체한다.
# ═══════════════════════════════════════════════════════════════════════════
import re as _sf_re

_SF_PX = 10.0                      # 글자당 폭(px)
_SF_BOX_EMU = 2_400_000            # 기본 상자 폭 → 252px → 25자/줄
_SF_A = 'http://schemas.openxmlformats.org/drawingml/2006/main'


class _SFFont:
    def getlength(self, s):
        return len(s) * _SF_PX


def _sf_max_chars(shape):
    return int((shape.width / 9525.0) // _SF_PX)


def _sf_line_starts(par_texts, max_chars):
    """고정폭 그리디 word-wrap — 구현(_pil_wrap_spans)과 독립된 테스트용 구현.
    각 문단의 첫 줄은 문단 시작, 이후 줄은 단어(비공백) 시작. 반환: 전체 텍스트('\\r' 결합) 기준 줄 시작 오프셋."""
    starts, base = [], 0
    for t in par_texts:
        toks = [(m.start(), m.end()) for m in _sf_re.finditer(r'\S+', t)]
        if toks:
            line_start = 0
            starts.append(base)
            for ts, te in toks:
                if te - line_start > max_chars and ts > line_start:
                    line_start = ts
                    starts.append(base + ts)
        base += len(t) + 1
    return starts


def _sf_slide_paras(slide):
    shape = rl_b._find_content_shape(slide)
    return [p for p in shape.text_frame.paragraphs if p.text.strip()] if shape else []


def _sf_slide_lines(slide):
    shape = rl_b._find_content_shape(slide)
    return len(_sf_line_starts([p.text for p in _sf_slide_paras(slide)], _sf_max_chars(shape)))


def _sf_fake_com(raises=False, corrupt=False):
    fake = _b_types.ModuleType('ppt_com_verify')

    class _Unavail(Exception):
        pass

    fake.ComVerificationUnavailable = _Unavail
    fake.calls = []
    fake.registry = {}
    fake.shutdown = lambda: None

    def measure_line_starts(path, shape_index, max_lines=None):
        fake.calls.append(path)
        if raises:
            raise _Unavail('simulated')
        slide = fake.registry[path]
        shape = rl_b._find_content_shape(slide)
        texts = [p.text for p in shape.text_frame.paragraphs]
        full = '\r'.join(texts)
        starts = _sf_line_starts(texts, _sf_max_chars(shape))
        ends = starts[1:] + [len(full)]
        line_texts = [full[s:e] for s, e in zip(starts, ends)]
        if corrupt:
            line_texts = ['@@@' + t for t in line_texts]   # 오프셋과 텍스트가 어긋난 이상 응답
        n = len(starts)
        k = n if max_lines is None else min(n, max_lines)
        return {'lines': n, 'starts': starts[:k], 'texts': line_texts[:k]}

    fake.measure_line_starts = measure_line_starts
    return fake


def _sf_install(monkeypatch, com='fake', box_px=None):
    """com='fake'(기본): 가짜 COM / None: COM 비활성(Pillow 폴백) / 모듈 객체: 그대로 사용."""
    font = _SFFont()
    monkeypatch.setattr(
        rl_b, '_get_slide_render_params',
        lambda slide: (font, box_px if box_px is not None
                       else rl_b._find_content_shape(slide).width / 9525.0),
    )
    monkeypatch.setattr(rl_b, '_COM_DISABLED', [False], raising=False)
    monkeypatch.setattr(rl_b, '_COM_MISMATCH_COUNT', {}, raising=False)
    monkeypatch.setattr(rl_b, '_COM_ATEXIT_REGISTERED', [False], raising=False)
    if com is None:
        monkeypatch.setattr(rl_b, '_com_verification_enabled', lambda: False)
        return None
    if com == 'fake':
        com = _sf_fake_com()
    monkeypatch.setattr(rl_b, '_com_verification_enabled', lambda: True)

    def _probe(prs, slide):
        key = f'probe-{id(slide)}-{len(com.registry)}'
        com.registry[key] = slide
        return key, 1

    monkeypatch.setattr(rl_b, '_build_com_probe_pptx', _probe, raising=False)
    monkeypatch.setitem(_w_sys.modules, 'ppt_com_verify', com)
    return com


def _sf_add_body_slide(prs, width=_SF_BOX_EMU, text='샘플 본문'):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(_WEmu(500_000), _WEmu(300_000), _WEmu(width), _WEmu(5_000_000))
    tf = box.text_frame
    tf.word_wrap = True
    r = tf.paragraphs[0].add_run()
    r.text = '1 ' + text
    r.font.size = _WPt(32)
    return slide


def _sf_add_ending_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(_WEmu(500_000), _WEmu(300_000), _WEmu(2_000_000), _WEmu(500_000))
    box.text_frame.text = '주님의 말씀입니다.'
    return slide


def _sf_prs(n_body=1, n_ending=0, widths=None):
    prs = _w_mk_prs()
    for i in range(n_body):
        _sf_add_body_slide(prs, width=(widths[i] if widths else _SF_BOX_EMU))
    for _ in range(n_ending):
        _sf_add_ending_slide(prs)
    return prs


def _sf_fill(prs, text, n_body=None, n_ending=0, **kw):
    units = rl_b.parse_into_verse_units(text)
    n_body = n_body if n_body is not None else len(prs.slides) - n_ending
    return rl_b.replace_reading_slides(prs, 0, n_body + n_ending, units, 0, label='T', **kw)


def _sf_body_slides(prs):
    """본문 상자가 있는 슬라이드(종료 전용 슬라이드는 _find_content_shape가 None; 종료 병합 슬라이드는 포함)."""
    return [s for s in prs.slides if rl_b._find_content_shape(s) is not None]


def _sf_all_text(prs):
    return ''.join(p.text for s in _sf_body_slides(prs) for p in _sf_slide_paras(s))


def _sf_nospace(s):
    return _sf_re.sub(r'\s+', '', s)


def _sf_assert_invariants(prs, text):
    """비마지막 슬라이드 정확히 9줄, 마지막 <= 9줄, 텍스트 무손실, 절 번호 오렌지(규칙 15)."""
    slides = _sf_body_slides(prs)
    assert slides
    counts = [_sf_slide_lines(s) for s in slides]
    assert all(c == rl_b.LINES_PER_SLIDE for c in counts[:-1]), counts
    assert 1 <= counts[-1] <= rl_b.LINES_PER_SLIDE, counts
    units = rl_b.parse_into_verse_units(text)
    expected = ''.join(u['text'] for u in units)
    assert _sf_nospace(_sf_all_text(prs)) == _sf_nospace(expected)
    if any(u['verse_num'] for u in units):
        assert sec_b._missing_orange_verse_numbers(prs, 0, len(prs.slides), text) == []
    return counts


def _sf_run_orange(r):
    clr = r.find(f'{{{_SF_A}}}rPr/{{{_SF_A}}}solidFill/{{{_SF_A}}}srgbClr')
    return clr is not None and clr.get('val', '').upper() == 'FFC000'


def _sf_orange_texts(slide):
    return [r.find(f'{{{_SF_A}}}t').text
            for p in _sf_slide_paras(slide) for r in p._p.findall(f'{{{_SF_A}}}r') if _sf_run_orange(r)]


# 문장이 끝나지 않고 이어지는 continuation 절과 절 중간 분할을 모두 만드는 한글 본문(20261004 복음 상수 재사용)
_SF_GOSPEL = _W_GOSPEL_20261004
_SF_LINE = ' '.join(['가나다라마'] * 4) + '.'   # 24자 — 25자/줄 상자에서 정확히 1줄


def test_sf1_com_measured_every_nonlast_slide_exactly_nine_lines(monkeypatch):
    """SF1: 한 슬라이드만 있는 템플릿에서 시작해도 슬라이드를 늘려가며 비마지막 슬라이드를
    정확히 9줄로 채운다(실측=가짜 COM 기준). Pillow 추정 폭을 일부러 크게 틀리게 줘도(과거
    'Pillow=9줄, COM=10줄' 상황) 수렴은 COM 실측이 결정한다."""
    com = _sf_install(monkeypatch, box_px=400.0)   # Pillow 추정 폭 400px vs 실제(COM) 252px
    prs = _sf_prs(1)
    _sf_fill(prs, _SF_GOSPEL)
    counts = _sf_assert_invariants(prs, _SF_GOSPEL)
    assert len(counts) >= 3
    assert com.calls, 'COM 실측을 한 번도 쓰지 않음'


def test_sf2_slide_count_is_driven_by_measurement_not_template_count(monkeypatch):
    """SF2: 템플릿 본문 슬라이드가 필요한 수보다 많으면 남는 슬라이드를 삭제하고, 적으면 복제해 늘린다.
    반환값은 순삽입 슬라이드 수(종료 슬라이드 위치 계산이 의존)."""
    _sf_install(monkeypatch)
    prs_few = _sf_prs(1)
    shift_few = _sf_fill(prs_few, _SF_GOSPEL)
    n = len(prs_few.slides)
    assert shift_few == n - 1 and n >= 3

    prs_many = _sf_prs(8)
    shift_many = _sf_fill(prs_many, _SF_GOSPEL)
    assert len(prs_many.slides) == n
    assert shift_many == n - 8


def test_sf3_per_slide_geometry_is_measured_in_each_slides_own_box(monkeypatch):
    """SF3: 기존 슬라이드 박스 폭이 서로 달라도(성인 템플릿 62/63 같은 사례) 각 슬라이드를 그 슬라이드의 실제
    상자에서 측정해 정확히 9줄로 채운다."""
    _sf_install(monkeypatch)
    prs = _sf_prs(4, widths=[2_400_000, 3_000_000, 2_000_000, 2_600_000])
    _sf_fill(prs, _SF_GOSPEL)
    slides = _sf_body_slides(prs)
    counts = [_sf_slide_lines(s) for s in slides]
    assert all(c == 9 for c in counts[:-1]) and counts[-1] <= 9, counts
    widths = {rl_b._find_content_shape(s).width for s in slides}
    assert len(widths) > 1, '전제: 박스 폭이 실제로 다름'


def test_sf4_verse_numbers_orange_and_continuation_fragments_not_orange(monkeypatch):
    """SF4: 절 번호는 오렌지 run으로 렌더링되고(규칙 15), 슬라이드 경계에서 잘려 이어지는 조각의 본문
    run은 오렌지가 아니다. 오렌지 run은 절 번호 텍스트뿐이다."""
    _sf_install(monkeypatch)
    prs = _sf_prs(1)
    _sf_fill(prs, _SF_GOSPEL)
    _sf_assert_invariants(prs, _SF_GOSPEL)
    for s in _sf_body_slides(prs):
        for t in _sf_orange_texts(s):
            assert _sf_re.fullmatch(r'\d+(,\d+)?\s*', t), f'오렌지 run에 절 번호가 아닌 텍스트: {t!r}'


def test_sf5_cut_inside_paragraph_keeps_formats_and_a_p_child_order(monkeypatch):
    """SF5: 문단 중간에서 잘린 조각의 a:p 자식 순서(pPr? → r* → endParaRPr?)와 본문 run 서식(sz)이 유지된다
    (규칙 1·2·13·14). 이어지는 조각은 새 슬라이드 첫 문단이며 절 번호(오렌지)로 시작하지 않는다."""
    _sf_install(monkeypatch)
    prs = _sf_prs(1)
    _sf_fill(prs, _SF_GOSPEL, line_spacing=1.1)
    order = {'pPr': 0, 'r': 1, 'endParaRPr': 2}
    cut_found = False
    for i, s in enumerate(_sf_body_slides(prs)):
        for p in _sf_slide_paras(s):
            tags = [c.tag.split('}')[1] for c in p._p]
            assert tags == sorted(tags, key=lambda t: order[t]), tags
            for r in p._p.findall(f'{{{_SF_A}}}r'):
                assert r.find(f'{{{_SF_A}}}rPr').get('sz') == '3200'
        first = _sf_slide_paras(s)[0]
        if i > 0 and not _sf_re.match(r'^\d', first.text):
            cut_found = True
            assert not _sf_run_orange(first._p.findall(f'{{{_SF_A}}}r')[0])
    assert cut_found, '전제: 이 본문은 문단 중간에서 잘린다'


def test_sf6_no_overflow_when_com_real_wrap_differs_from_pillow_in_both_directions(monkeypatch):
    """SF6: Pillow 추정이 COM보다 넓거나 좁게 틀려도 결과는 COM 실측 기준으로 정확히 9줄이다."""
    for est in (150.0, 252.0, 900.0):
        _sf_install(monkeypatch, box_px=est)
        prs = _sf_prs(2)
        _sf_fill(prs, _SF_GOSPEL)
        _sf_assert_invariants(prs, _SF_GOSPEL)


def test_sf7_exact_multiple_of_nine_lines_makes_no_extra_empty_slide(monkeypatch):
    """SF7: 본문이 정확히 18줄이면 슬라이드 2장(9+9) — 빈 3번째 슬라이드를 만들지 않는다."""
    _sf_install(monkeypatch)
    text = '\n'.join(_SF_LINE for _ in range(18))
    prs = _sf_prs(1)
    _sf_fill(prs, text)
    counts = [_sf_slide_lines(s) for s in _sf_body_slides(prs)]
    assert counts == [9, 9], counts


def test_sf8_empty_units_leave_one_empty_slide_without_crash(monkeypatch):
    """SF8: 본문이 비어 있어도 예외 없이 슬라이드 1장(빈 본문)이 남는다(과거 `[[]]` 동작 보존)."""
    _sf_install(monkeypatch)
    prs = _sf_prs(3)
    shift = _sf_fill(prs, '')
    assert len(prs.slides) == 1 and shift == -2


def test_sf9_pillow_path_same_loop_when_com_disabled(monkeypatch):
    """SF9: com_verification_enabled=false면 COM을 한 번도 호출하지 않고 같은 슬라이드 단위 루프를 Pillow
    word-wrap(분할과 검증이 같은 박스 폭)으로 수행해 비마지막 슬라이드가 정확히 9줄이다."""
    com = _sf_fake_com()
    monkeypatch.setitem(_w_sys.modules, 'ppt_com_verify', com)
    _sf_install(monkeypatch, com=None)
    prs = _sf_prs(1)
    _sf_fill(prs, _SF_GOSPEL)
    _sf_assert_invariants(prs, _SF_GOSPEL)
    assert com.calls == []


def test_sf10_com_failure_switches_to_pillow_and_disables_com_permanently(monkeypatch, capsys):
    """SF10: COM 호출이 ComVerificationUnavailable로 실패하면 _COM_DISABLED를 켜고 그 슬라이드부터 Pillow로
    계속한다(예외 없음). 결과 불변조건은 동일."""
    com = _sf_install(monkeypatch, com=_sf_fake_com(raises=True))
    prs = _sf_prs(1)
    _sf_fill(prs, _SF_GOSPEL)
    _sf_assert_invariants(prs, _SF_GOSPEL)
    assert rl_b._COM_DISABLED[0] is True
    assert len(com.calls) == 1
    assert '[경고]' in capsys.readouterr().out


def test_sf11_inconsistent_com_answer_falls_back_for_that_slide_only(monkeypatch, capsys):
    """SF11: 줄 시작 오프셋과 줄 텍스트가 어긋나는 이상 응답이면 그 슬라이드만 Pillow로 폴백하고
    COM을 전역 비활성화하지 않는다."""
    com = _sf_install(monkeypatch, com=_sf_fake_com(corrupt=True))
    prs = _sf_prs(1)
    _sf_fill(prs, _SF_GOSPEL)
    _sf_assert_invariants(prs, _SF_GOSPEL)
    assert rl_b._COM_DISABLED[0] is False
    assert len(com.calls) >= 2          # 매 슬라이드 COM을 계속 시도
    assert '[경고]' in capsys.readouterr().out


def test_sf12_mismatch_between_pillow_and_com_is_counted_and_logged(monkeypatch, capsys):
    """SF12: 같은 텍스트에 대해 Pillow 추정 줄 수와 COM 실측 줄 수가 다르면 `_COM_MISMATCH_COUNT`를 올리고
    기존과 같은 '줄 수 불일치 감지' 경고를 출력한다(관측 가능한 동작 보존)."""
    _sf_install(monkeypatch, box_px=400.0)
    prs = _sf_prs(1)
    _sf_fill(prs, _SF_GOSPEL)
    assert sum(rl_b._COM_MISMATCH_COUNT.values()) >= 1
    assert '줄 수 불일치 감지' in capsys.readouterr().out


def test_sf13_prefix_grows_until_tenth_line_is_visible(monkeypatch):
    """SF13: 측정용 프리픽스가 너무 짧아 10번째 줄이 안 보이면 프리픽스를 늘려 재측정해 같은 결과가 나온다."""
    _sf_install(monkeypatch)
    prs_ref = _sf_prs(1)
    _sf_fill(prs_ref, _SF_GOSPEL)
    ref = [[p.text for p in _sf_slide_paras(s)] for s in _sf_body_slides(prs_ref)]

    com = _sf_install(monkeypatch)
    monkeypatch.setattr(rl_b, '_MEASURE_PREFIX_CHARS', 60)
    prs = _sf_prs(1)
    _sf_fill(prs, _SF_GOSPEL)
    got = [[p.text for p in _sf_slide_paras(s)] for s in _sf_body_slides(prs)]
    assert got == ref
    assert len(com.calls) > len(ref)      # 재측정이 일어났다


def test_sf14_ending_merge_after_last_slide_is_confirmed(monkeypatch):
    """SF14: 종료 텍스트 병합은 마지막 본문 슬라이드 확정 이후 기존 규칙(<=5줄이면 마지막 본문 슬라이드에
    통합) 그대로 동작하고, 반환값은 실제 슬라이드 수 변화와 일치한다."""
    _sf_install(monkeypatch)
    text = '\n'.join(_SF_LINE for _ in range(11))   # 11줄 → 9 + 2
    prs = _sf_prs(2, n_ending=1)
    before = len(prs.slides)
    shift = _sf_fill(prs, text, n_body=2, n_ending=1)
    assert len(prs.slides) == before + shift == 2
    last = prs.slides[1]
    assert any('주님의 말씀입니다' in sh.text_frame.text for sh in last.shapes if sh.has_text_frame)


def test_sf15_ending_slide_kept_when_last_slide_is_long(monkeypatch):
    """SF15: 마지막 본문 슬라이드가 6줄 이상이면 종료 슬라이드는 별도로 남는다(종료 슬라이드가 본문 뒤에 위치)."""
    _sf_install(monkeypatch)
    text = '\n'.join(_SF_LINE for _ in range(15))   # 15줄 → 9 + 6
    prs = _sf_prs(2, n_ending=1)
    shift = _sf_fill(prs, text, n_body=2, n_ending=1)
    assert len(prs.slides) == 3 and shift == 0
    assert rl_b._has_ending_text(prs.slides[2]) and not rl_b._has_ending_text(prs.slides[1])
    assert _sf_slide_lines(prs.slides[0]) == 9 and _sf_slide_lines(prs.slides[1]) == 6


def test_sf16_ending_slide_pushed_when_more_body_slides_are_needed(monkeypatch):
    """SF16: 필요한 본문 슬라이드 수가 기존보다 많으면 본문이 종료 슬라이드 앞에 삽입된다(종료는 항상 맨 뒤)."""
    _sf_install(monkeypatch)
    prs = _sf_prs(1, n_ending=1)
    shift = _sf_fill(prs, _SF_GOSPEL, n_body=1, n_ending=1)
    assert shift >= 2
    last = prs.slides[len(prs.slides) - 1]
    assert rl_b._has_ending_text(last)
    slides = _sf_body_slides(prs)
    assert [_sf_slide_lines(s) for s in slides[:-1]] == [9] * (len(slides) - 1)


def test_sf17_align_and_line_spacing_applied_to_every_filled_slide(monkeypatch):
    """SF17: align/line_spacing 인자는 모든 슬라이드(복제된 슬라이드 포함)의 모든 문단에 적용된다."""
    _sf_install(monkeypatch)
    prs = _sf_prs(1)
    _sf_fill(prs, _SF_GOSPEL, align='ctr', line_spacing=1.1)
    assert len(_sf_body_slides(prs)) >= 3
    for s in _sf_body_slides(prs):
        for p in _sf_slide_paras(s):
            pPr = p._p.find(f'{{{_SF_A}}}pPr')
            assert pPr.get('algn') == 'ctr'
            assert pPr.find(f'{{{_SF_A}}}lnSpc/{{{_SF_A}}}spcPct').get('val') == '110000'


def test_sf18_normalize_page_size_copies_template_box_before_measuring(monkeypatch):
    """SF18: normalize_page_size=True면 각 슬라이드 상자를 template 슬라이드 크기로 맞춘 뒤 측정한다."""
    _sf_install(monkeypatch)
    prs = _sf_prs(3, widths=[2_400_000, 1_800_000, 2_000_000])
    _sf_fill(prs, _SF_GOSPEL, normalize_page_size=True)
    widths = {rl_b._find_content_shape(s).width for s in _sf_body_slides(prs)}
    assert widths == {2_400_000}
    counts = [_sf_slide_lines(s) for s in _sf_body_slides(prs)]
    assert all(c == 9 for c in counts[:-1]) and counts[-1] <= 9


def test_sf19_split_helpers_preserve_orange_ranges_and_text():
    """SF19(규칙 13·14·15): 문단 분할 헬퍼가 run 개수를 가정하지 않고 오렌지 구간을 잘린 조각에 정확히 옮긴다.
    continuation으로 병합된 여러 절 번호가 한 문단에 있는 경우도."""
    merged = rl_b._reading_paras_from_units(
        [{'text': '51 aaa bbb ccc', 'verse_num': '51', 'is_continuation': False},
         {'text': '52 ddd eee', 'verse_num': '52', 'is_continuation': True}])
    assert len(merged) == 1
    text = merged[0]['text']
    assert text == '51 aaa bbb ccc 52 ddd eee'
    cut = text.index('ccc')
    front, rest = rl_b._split_paras_at(merged, 0, cut)
    f_units, r_units = rl_b._paras_to_units(front), rl_b._paras_to_units(rest)
    assert f_units[-1]['text'] == '51 aaa bbb' and f_units[-1]['extra_verses'] == [(0, '51')]
    assert r_units[0]['text'] == 'ccc 52 ddd eee' and r_units[0]['extra_verses'] == [(4, '52')]
    # 두 번째 절 번호 바로 뒤에서 자르기 — 앞 조각의 오렌지 구간이 trailing space 제거 후에도 유지
    cut2 = text.index('ddd')
    front2, rest2 = rl_b._split_paras_at(merged, 0, cut2)
    f2 = rl_b._paras_to_units(front2)[-1]
    assert f2['text'] == '51 aaa bbb ccc 52' and f2['extra_verses'] == [(0, '51'), (15, '52')]
    assert rl_b._paras_to_units(rest2)[0]['extra_verses'] == []


def test_sf20_cut_at_paragraph_start_moves_whole_paragraphs():
    """SF20: 절단점이 문단 시작이면 앞 조각엔 그 이전 문단만, 뒤 조각엔 그 문단부터 통째로 들어간다."""
    paras = rl_b._reading_paras_from_units(rl_b.parse_into_verse_units('1 가나다.\n2 라마바.\n3 사아자.'))
    front, rest = rl_b._split_paras_at(paras, 1, 0)
    assert [p['text'] for p in front] == ['1 가나다.']
    assert [p['text'] for p in rest] == ['2 라마바.', '3 사아자.']


def test_sf21_gospel_20261004_cut_lines_are_exact_wrap_lines_no_two_char_tail(monkeypatch):
    """SF21(구 W그룹 widow 요구사항 이전): 20261004 복음에서 문단 중간에서 잘린 비마지막 슬라이드의 마지막 줄은
    줄바꿈의 정확한 한 줄이며(앞 슬라이드는 정확히 9줄), 그 줄에 뒤 조각의 첫 단어가 더 못 들어간다 —
    '짓는' 두 글자만 남는 꼬리가 구조적으로 생기지 않는다."""
    _sf_install(monkeypatch)
    prs = _sf_prs(1)
    _sf_fill(prs, _SF_GOSPEL)
    slides = _sf_body_slides(prs)
    checked = 0
    for cur, nxt in zip(slides, slides[1:]):
        shape = rl_b._find_content_shape(cur)
        texts = [p.text for p in _sf_slide_paras(cur)]
        starts = _sf_line_starts(texts, _sf_max_chars(shape))
        assert len(starts) == 9
        last_line = '\r'.join(texts)[starts[-1]:].strip()
        nxt_first = _sf_slide_paras(nxt)[0].text
        if not _sf_re.match(r'^\d', nxt_first) and texts[-1].rstrip()[-1] not in '.!?”’':
            first_word = nxt_first.split(' ')[0]
            assert len(last_line) + 1 + len(first_word) > _sf_max_chars(shape), (last_line, first_word)
            checked += 1
    assert checked >= 1


def test_sf22_cut_never_exceeds_nine_lines_even_with_long_unbreakable_word(monkeypatch):
    """SF22: 한 줄보다 긴 단어가 있어도 오버플로 없이 슬라이드를 구성한다(비마지막 슬라이드는 정확히 9줄)."""
    _sf_install(monkeypatch)
    text = '가나다 ' + '가' * 40 + ' 라마바' * 60
    prs = _sf_prs(1)
    _sf_fill(prs, text)
    slides = _sf_body_slides(prs)
    assert len(slides) >= 2
    for s in slides[:-1]:
        assert _sf_slide_lines(s) == 9


# ---------------------------------------------------------------------------
# SF-real: 실제 PowerPoint COM으로 모든 비마지막 슬라이드가 정확히 9줄인지 확인한다.
# ---------------------------------------------------------------------------

_SF_G_20260926_EN = (
    'Jesus said to the chief priests and elders of the people, ‘What is your opinion? '
    'A man had two sons. He went and said to the first, “My boy, you go and work in the '
    'vineyard today.” He answered, “I will not go,” but afterwards thought better '
    'of it and went. The man then went and said the same thing to the second who answered, '
    '“Certainly, sir,” but did not go. Which of the two did the father’s will?’ '
    '‘The first’ they said. Jesus said to them, ‘I tell you solemnly, tax collectors '
    'and prostitutes are making their way into the kingdom of God before you. For John came to '
    'you, a pattern of true righteousness, but you did not believe him, and yet the tax '
    'collectors and prostitutes did. Even after seeing that, you refused to think better of it '
    'and believe in him.’'
)


def _sf_real_com_or_skip():
    try:
        import ppt_com_verify as com
    except ImportError:
        pytest.skip('pywin32 미설치')
    if not com.is_available():
        pytest.skip('PowerPoint COM 연결 불가')
    return com


def _sf_real_assert_nine(prs, start, end, text):
    """모든 본문 슬라이드를 실제 COM으로 재측정: 비마지막 == 9, 마지막 <= 9, 텍스트 보존, 오렌지."""
    com = _sf_real_com_or_skip()
    from missa_ooxml_utils import _build_com_probe_pptx as _probe
    body = [prs.slides[i] for i in range(start, end) if rl_b._find_content_shape(prs.slides[i]) is not None]
    counts = []
    for slide in body:
        path, idx = _probe(prs, slide)
        counts.append(com.count_slide_lines(str(path), idx))
    assert all(c == rl_b.LINES_PER_SLIDE for c in counts[:-1]), counts
    assert counts[-1] <= rl_b.LINES_PER_SLIDE, counts
    got = ''.join(p.text for s in body for p in rl_b._find_content_shape(s).text_frame.paragraphs)
    expected = ''.join(u['text'] for u in rl_b.parse_into_verse_units(text))
    assert _sf_nospace(got) == _sf_nospace(expected)
    if any(u['verse_num'] for u in rl_b.parse_into_verse_units(text)):
        assert sec_b._missing_orange_verse_numbers(prs, start, end, text) == []
    return counts


def test_sf_real1_measure_line_starts_partitions_text():
    """SF-real1: ppt_com_verify.measure_line_starts가 줄 시작 오프셋(0-based)·줄 텍스트를 서로 일관되게 돌려준다."""
    com = _sf_real_com_or_skip()
    prs = _w_mk_prs()
    _w_add_slide(prs, [[('1 ', True), ('태초에 하느님께서 하늘과 땅을 창조하셨다. 땅은 아직 모양을 갖추지 않고 비어 있었다. 어둠이 심연을 덮고 있었다.', False)],
                       [('2 ', True), ('두번째 단락입니다 짧은 것. ', False)]])
    from missa_ooxml_utils import _build_com_probe_pptx as _probe
    slide = prs.slides[0]
    path, idx = _probe(prs, slide)
    full = '\r'.join(p.text for p in rl_b._find_content_shape(slide).text_frame.paragraphs)
    n = com.count_slide_lines(str(path), idx)
    info = com.measure_line_starts(str(path), idx, 3)
    assert info['lines'] == n and n > 3
    assert info['starts'][0] == 0 and info['starts'] == sorted(set(info['starts'])) and len(info['starts']) == 3
    for s, t in zip(info['starts'], info['texts']):
        assert full.startswith(t.rstrip('\r '), s)


@pytest.mark.skipif(not _W_GOSPEL_TEMPLATE.is_file(), reason='성인 템플릿 픽스처 없음')
def test_sf_real2_20261004_adult_gospel_all_nonlast_slides_exactly_nine(tmp_path):
    """SF-real2(구 W9 이전): 20261004 성인 복음을 실제 COM으로 — 모든 비마지막 슬라이드 정확히 9줄(구 구조는 75번이 8줄)."""
    _sf_real_com_or_skip()
    from pptx import Presentation as _P
    import shutil as _sh
    import missa_content_updaters as cu
    work = tmp_path / 't.pptx'
    _sh.copy(_W_GOSPEL_TEMPLATE, work)
    prs = _P(str(work))
    sec = sec_b.find_sections(prs)
    cu.update_복음_title_slide(prs, sec['복음_title'], _W_GOSPEL_20261004_TITLE)
    rl_b.replace_reading_slides(prs, sec['복음_start'], sec['복음_end'],
                                rl_b.parse_into_verse_units(_W_GOSPEL_20261004), sec['복음_start'], label='복음')
    sec = sec_b.find_sections(prs)
    counts = _sf_real_assert_nine(prs, sec['복음_start'], sec['복음_end'], _W_GOSPEL_20261004)
    assert len(counts) >= 3


@pytest.mark.skipif(not _B_TEMPLATE.is_file(), reason='제2독서 재현용 참조 PPT 없음')
def test_sf_real3_second_reading_20260920_converges_to_nine():
    """SF-real3(구 test_second_reading_rebalance_converges_without_large_deficit 이전): 로마서 8,31-39 제2독서가
    '1줄 이내 미달 허용'이 아니라 비마지막 정확히 9줄."""
    _sf_real_com_or_skip()
    prs = _BPresentation(str(_B_TEMPLATE))
    sections = sec_b.find_sections(prs)
    rl_b.replace_reading_slides(prs, sections['제2독서_start'], sections['제2독서_end'],
                                rl_b.parse_into_verse_units(_B_CONTENT_ROM_8_31_39), sections['제2독서_start'],
                                label='제2독서')
    s2 = sec_b.find_sections(prs)
    _sf_real_assert_nine(prs, s2['제2독서_start'], s2['제2독서_end'], _B_CONTENT_ROM_8_31_39)


@pytest.mark.skipif(not _E_YOUTH_TEMPLATE.is_file(), reason='청년 템플릿 없음')
@pytest.mark.parametrize('which', ['20260913', '20260926'])
def test_sf_real4_youth_english_gospel_exactly_nine(which):
    """SF-real4(구 test_g1c / test_y2_G2 / test_y2_G1b 이전): 영문 복음이 추정 오차(Pillow=9줄, COM=10줄)에도
    비마지막 슬라이드 정확히 9줄로 수렴하고, 20260926에서는 "father's"가 첫 슬라이드에 들어간다."""
    _sf_real_com_or_skip()
    if which == '20260913':
        if not _E_EN_GOSPEL_FIXTURE.is_file():
            pytest.skip('20260913 영문 복음 JSON 없음')
        import json as _j
        content = _j.loads(_E_EN_GOSPEL_FIXTURE.read_text(encoding='utf-8'))['Gospel']['content']
    else:
        content = _SF_G_20260926_EN
    prs = _BPresentation(str(_E_YOUTH_TEMPLATE))
    sections = sec_b.find_sections(prs, mass_type='youth')
    rl_b.replace_reading_slides(prs, sections['복음_start'], sections['복음_end'],
                                rl_b.parse_into_verse_units(content), sections['복음_start'], label='복음',
                                normalize_page_size=True)
    s2 = sec_b.find_sections(prs, mass_type='youth')
    _sf_real_assert_nine(prs, s2['복음_start'], s2['복음_end'], content)
    if which == '20260926':
        first = rl_b._find_content_shape(prs.slides[s2['복음_start']]).text_frame.text
        assert "father’s" in first, first



# ---------------------------------------------------------------------------
# SF 후속(2026-10-03, 독립 리뷰 반영): 종료 병합 판정을 실측 줄 수로 / 절 번호 단독 꼬리 / 예외 / 계측 초기화
# ---------------------------------------------------------------------------

_SF_WIDE_EMU = 6_000_000                       # 629px → 62자/줄
_SF_L60 = ' '.join(['abcdefghi'] * 6) + '.'    # 60자: 넓은 상자에서 정확히 1줄, 27자/줄 추정으로는 3줄
_SF_NARROW_EMU = 1_200_000                     # 125px → 12자/줄


def _sf_merge_case(n_lines, width, line=_SF_L60, n_ending=1):
    prs = _sf_prs(1, n_ending=n_ending, widths=[width])
    text = '\n'.join([line] * n_lines)
    _sf_fill(prs, text, n_body=1, n_ending=n_ending)
    return prs


@pytest.mark.parametrize('n_lines,merged', [(4, True), (5, True), (6, False), (7, False)])
def test_sf23_ending_merge_uses_measured_line_count_com(monkeypatch, n_lines, merged):
    """SF23(사용자 결정 2026-10-03): 종료 병합 판정은 27자/줄 추정이 아니라 마지막 슬라이드를 확정할 때 얻은
    COM 실측 줄 수로 한다. 이 입력은 실측 n줄, 추정 3n줄이다 — 경계 정확히 5줄은 병합, 6줄은 미병합
    (`<=`→`<` 변이 방지)."""
    _sf_install(monkeypatch)
    prs = _sf_merge_case(n_lines, _SF_WIDE_EMU)
    assert _sf_slide_lines(prs.slides[0]) == n_lines
    assert len(prs.slides) == (1 if merged else 2)


@pytest.mark.parametrize('n_lines', [1, 2, 3, 4, 5, 6, 7])
def test_sf32_allow_ending_merge_false_keeps_ending_slide_separate(monkeypatch, n_lines):
    """SF32(사용자 결정 2026-10-05): 청년 영문 복음은 마지막 슬라이드가 몇 줄이든(병합 기준 5줄 이하여도) 종료 슬라이드
    ('주님의 말씀입니다/그리스도님, 찬미합니다')를 본문에 병합하지 않는다 — allow_ending_merge=False."""
    _sf_install(monkeypatch)
    prs = _sf_prs(1, n_ending=1, widths=[_SF_WIDE_EMU])
    _sf_fill(prs, chr(10).join([_SF_L60] * n_lines), n_body=1, n_ending=1, allow_ending_merge=False)
    assert len(prs.slides) == 2
    assert _sf_slide_lines(prs.slides[0]) == n_lines
    body_has_ending = any(sh.has_text_frame and '주님의 말씀입니다' in sh.text_frame.text for sh in prs.slides[0].shapes)
    end_has_ending = any(sh.has_text_frame and '주님의 말씀입니다' in sh.text_frame.text for sh in prs.slides[1].shapes)
    assert not body_has_ending and end_has_ending


@pytest.mark.parametrize('n_lines,merged', [(5, True), (6, False)])
def test_sf24_ending_merge_uses_pillow_total_when_com_unavailable(monkeypatch, n_lines, merged):
    """SF24: COM이 없으면 Pillow 줄 수로 판정한다(같은 경계)."""
    _sf_install(monkeypatch, com=None)
    prs = _sf_merge_case(n_lines, _SF_WIDE_EMU)
    assert len(prs.slides) == (1 if merged else 2)


@pytest.mark.parametrize('n_lines,merged', [(1, True), (2, False)])
def test_sf25_ending_merge_falls_back_to_char_estimate_without_any_font(monkeypatch, n_lines, merged):
    """SF25: 폰트도 없을 때만 27자 근사 줄 수로 판정한다(이 입력은 줄당 3줄 추정 → 1문단 3줄 병합, 2문단 6줄 미병합)."""
    _sf_install(monkeypatch, com=None)
    monkeypatch.setattr(rl_b, '_get_slide_render_params', lambda slide: (None, 0.0))
    prs = _sf_merge_case(n_lines, _SF_WIDE_EMU)
    assert len(prs.slides) == (1 if merged else 2)


def test_sf26_ending_not_merged_when_estimate_underestimates_actual_lines(monkeypatch):
    """SF26: 반대 방향 — 좁은 상자에서 각 문단이 실측 2줄(추정 1줄)이라 3문단=실측 6줄이면, 구 추정(3줄)은
    병합했지만 실측 기준으로는 병합하지 않는다."""
    _sf_install(monkeypatch)
    prs = _sf_merge_case(3, _SF_NARROW_EMU, line='abcdefghi jklmnopqr.')
    assert _sf_slide_lines(prs.slides[0]) == 6
    assert len(prs.slides) == 2


def test_sf27_mismatch_counter_is_reset_per_section_call(monkeypatch):
    """SF27(리뷰 변이검사): `_COM_MISMATCH_COUNT`는 섹션(제1독서→제2독서→복음)마다 새로 시작한다 — 이전
    섹션의 슬라이드별 누적이 다음 섹션 경고에 섞이면 안 된다."""
    _sf_install(monkeypatch)
    monkeypatch.setattr(rl_b, '_COM_MISMATCH_COUNT', {999: 5}, raising=False)
    prs = _sf_prs(1)
    _sf_fill(prs, _SF_GOSPEL)
    assert 999 not in rl_b._COM_MISMATCH_COUNT


def test_sf28_missing_body_box_on_cloned_slide_raises_with_label_and_slide(monkeypatch):
    """SF28: 복제한 슬라이드에서 본문 상자를 못 찾으면 남은 본문을 조용히 버리지 말고 RuntimeError로 중단한다
    (라벨·슬라이드 번호 포함)."""
    _sf_install(monkeypatch)
    orig = rl_b.insert_slide_copy

    def _bad_copy(prs, pos, src):
        orig(prs, pos, src)
        for sh in prs.slides[pos].shapes:
            if sh.has_text_frame:
                sh.text_frame.text = ''

    monkeypatch.setattr(rl_b, 'insert_slide_copy', _bad_copy)
    prs = _sf_prs(1)
    with pytest.raises(RuntimeError) as ei:
        rl_b.replace_reading_slides(prs, 0, 1, rl_b.parse_into_verse_units(_SF_GOSPEL), 0, label='복음')
    assert '복음' in str(ei.value) and '슬라이드 2' in str(ei.value)


def test_sf31_non_last_slides_are_exactly_nine_lines(monkeypatch):
    """SF31: 비마지막 슬라이드는 정확히 9줄이다."""
    _sf_install(monkeypatch)
    prs = _sf_prs(1)
    _sf_fill(prs, _SF_GOSPEL)
    counts = [_sf_slide_lines(s) for s in _sf_body_slides(prs)]
    assert counts[:-1] == [9] * (len(counts) - 1)


def test_sf_real5_ending_merge_decided_by_real_com_line_count():
    """SF-real5(사용자 결정): 청년 템플릿 영문 복음에서 마지막 슬라이드의 실제 COM 줄 수가 5 이하일 때만 종료
    텍스트가 본문 슬라이드에 병합된다(구 27자 추정은 라틴 본문을 2배 가까이 과대 추정해 실측 4~5줄도 미병합)."""
    com = _sf_real_com_or_skip()
    if not _E_YOUTH_TEMPLATE.is_file():
        pytest.skip('청년 템플릿 없음')
    from missa_ooxml_utils import _build_com_probe_pptx as _probe
    words = ('lorem ipsum dolor sit amet consectetur adipiscing elit sed do eiusmod tempor incididunt '
             'ut labore et dolore magna aliqua').split()
    checked = set()
    for n in (90, 110, 190, 215, 235):
        text = ' '.join(words[i % len(words)] for i in range(n))
        prs = _BPresentation(str(_E_YOUTH_TEMPLATE))
        sections = sec_b.find_sections(prs, mass_type='youth')
        rl_b.replace_reading_slides(prs, sections['복음_start'], sections['복음_end'],
                                    rl_b.parse_into_verse_units(text), sections['복음_start'], label='복음',
                                    normalize_page_size=True)
        s2 = sec_b.find_sections(prs, mass_type='youth')
        body = [prs.slides[i] for i in range(s2['복음_start'], s2['복음_end'])
                if rl_b._find_content_shape(prs.slides[i]) is not None]
        last = body[-1]
        path, idx = _probe(prs, last)
        actual = com.count_slide_lines(str(path), idx)
        merged = any(sh.has_text_frame and '주님의 말씀입니다' in sh.text_frame.text for sh in last.shapes)
        assert merged == (actual <= 5), (n, actual, merged)
        checked.add(merged)
    assert checked == {True, False}, '전제: 병합/미병합 두 경우가 모두 검증돼야 함'


@pytest.mark.skipif(not _W_GOSPEL_TEMPLATE.is_file(), reason='성인 템플릿 픽스처 없음')
def test_sf_real6_probe_line_breaks_equal_final_slide_line_breaks(tmp_path, monkeypatch):
    """SF-real6: 프로브(임시 기록) 단계에서 측정한 첫 9줄이 확정 후 실제 슬라이드의 줄바꿈과 동일하고 줄 시작이
    단어 경계다(템플릿은 한글을 단어 단위로 줄바꿈한다). 프로브 상자와 실제 상자가 같은 설정이라 절단 오프셋이
    실제 렌더링과 어긋나지 않는다는 확인."""
    com = _sf_real_com_or_skip()
    from pptx import Presentation as _P
    import shutil as _sh
    from missa_ooxml_utils import _build_com_probe_pptx as _probe
    work = tmp_path / 't.pptx'
    _sh.copy(_W_GOSPEL_TEMPLATE, work)
    prs = _P(str(work))
    sec = sec_b.find_sections(prs)
    recorded = []
    orig = com.measure_line_starts

    def _spy(path, idx, max_lines=10):
        info = orig(path, idx, max_lines)
        recorded.append(info)
        return info

    monkeypatch.setattr(com, 'measure_line_starts', _spy)
    rl_b.replace_reading_slides(prs, sec['복음_start'], sec['복음_end'],
                                rl_b.parse_into_verse_units(_W_GOSPEL_20261004), sec['복음_start'], label='복음')
    monkeypatch.setattr(com, 'measure_line_starts', orig)
    sec = sec_b.find_sections(prs)
    body = [prs.slides[i] for i in range(sec['복음_start'], sec['복음_end'])
            if rl_b._find_content_shape(prs.slides[i]) is not None]
    assert len(recorded) == len(body) >= 3
    for slide, rec in zip(body[:-1], recorded):
        shape = rl_b._find_content_shape(slide)
        path, idx = _probe(prs, slide)
        fin = com.measure_line_starts(str(path), idx, 40)
        full = '\r'.join(p.text for p in shape.text_frame.paragraphs)
        assert [t.rstrip() for t in fin['texts']][:9] == [t.rstrip() for t in rec['texts']][:9]
        assert all(full[s - 1] in ' \r' for s in fin['starts'][1:]), '단어 중간에서 시작하는 줄'
