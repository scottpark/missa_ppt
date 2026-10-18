"""OneDrive 로컬 폴더 루트(공용/성인/청년) 기반 동작 테스트 — 성가 조회/저장·결과 복사·창 구성.

루트 관리·경로 규칙·찾아보기 기억의 명세는 test_missa_progression.py의 test_fs_* 가 맡는다
(2026-10-18 폴더 구조 개편, SPEC_OneDrive_폴더구조_개편.md). 이 파일은 그 루트를 쓰는 소비 동작을 검증한다.
"""
from pathlib import Path

import pytest

import missa_gui as gui


@pytest.fixture
def ppt_root(tmp_path):
    """새 구조의 세 루트를 모두 채운 'PPT 문서'(마스터 PC). `gui._ROOTS`에 세 루트를 설정해 돌려준다."""
    parent = tmp_path / "PPT 문서"
    common = parent / "00.공용 PPT"
    adult = parent / "10.성인 미사"
    youth = parent / "20.청년 미사"
    for p in (youth / "2.성가" / "나주노 성가", youth / "2.성가" / "야훼이레 성가",
              common / "09.가톨릭 성가" / "성가-악보버전", adult / "10.기타 PPT문서"):
        p.mkdir(parents=True)
    gui._ROOTS.update(common=common, adult=adult, youth=youth)
    return parent


@pytest.fixture
def cfg(tmp_path, monkeypatch):
    path = tmp_path / "config.json"
    monkeypatch.setattr(gui, "CONFIG_FILE", path)
    gui._reset_ppt_roots()
    return path


def test_local_hymn_lookup_finds_file(cfg, ppt_root):
    f = ppt_root / "20.청년 미사" / "2.성가" / "나주노 성가" / "나주노 성가 172 나의 고백.pptx"
    f.write_bytes(b"x")
    (f.parent / "~$나주노 성가 171 임시.pptx").write_bytes(b"x")
    assert gui.find_youth_hymn_file("youth_hymn", "나주노 성가", r"^나주노 성가 172 ") == f
    assert gui.find_youth_hymn_file("youth_hymn", "나주노 성가", r"^나주노 성가 999 ") is None


def test_local_hymn_lookup_missing_folder_is_none(cfg, ppt_root):
    assert gui.find_youth_hymn_file("youth_hymn", "없는 성가 폴더", r".") is None


def test_youth_output_is_no_longer_auto_copied_or_uploaded():
    """2026-10-03: 청년 결과 PPT는 output/에만 생성한다 — 자동 복사·업로드 함수는 삭제됐다."""
    import missa_to_ppt as m
    assert not hasattr(m, "_upload_youth_output_to_onedrive")


# ── 결과창 '원드라이브로 복사'(2026-10-03, 2026-10-18 새 루트 기준) ──

def test_copy_result_youth_goes_to_year_folder_under_youth_root(cfg, ppt_root, tmp_path):
    out = tmp_path / "결과.pptx"
    out.write_bytes(b"pptx")
    dest = gui.copy_result_to_onedrive("youth", out, "20261003")
    assert Path(dest) == ppt_root / "20.청년 미사" / "2026"
    assert (ppt_root / "20.청년 미사" / "2026" / "결과.pptx").read_bytes() == b"pptx"
    assert out.exists()


def test_copy_result_adult_goes_to_year_folder_under_adult_root(cfg, ppt_root, tmp_path):
    out = tmp_path / "성인.pptx"
    out.write_bytes(b"a")
    dest = gui.copy_result_to_onedrive("adult", out, "20261004")
    assert Path(dest) == ppt_root / "10.성인 미사" / "2026"
    assert (ppt_root / "10.성인 미사" / "2026" / "성인.pptx").read_bytes() == b"a"


def test_copy_result_exists_detects_same_name(cfg, ppt_root, tmp_path):
    out = tmp_path / "x.pptx"
    out.write_bytes(b"1")
    assert not gui.result_copy_exists("youth", out, "20261003")
    gui.copy_result_to_onedrive("youth", out, "20261003")
    assert gui.result_copy_exists("youth", out, "20261003")


def test_copy_result_without_root_raises(cfg, tmp_path):
    out = tmp_path / "a.pptx"
    out.write_bytes(b"1")
    for mt in ("adult", "youth"):
        with pytest.raises(RuntimeError):
            gui.copy_result_to_onedrive(mt, out, "20261004")


def _stub_messagebox(monkeypatch, answer=True):
    from tkinter import messagebox
    shown = []
    monkeypatch.setattr(messagebox, "showinfo", lambda t, m, **k: shown.append(("info", m)))
    monkeypatch.setattr(messagebox, "showerror", lambda t, m, **k: shown.append(("error", m)))
    monkeypatch.setattr(messagebox, "showwarning", lambda t, m, **k: shown.append(("warn", m)))
    monkeypatch.setattr(messagebox, "askyesno", lambda t, m, **k: (shown.append(("ask", m)), answer)[1])
    return shown


def test_copy_ui_announces_destination_folder(cfg, ppt_root, tmp_path, monkeypatch):
    shown = _stub_messagebox(monkeypatch)
    out = tmp_path / "r.pptx"
    out.write_bytes(b"1")
    gui._copy_result_with_ui(None, "youth", out, "20261003")
    assert shown[-1][0] == "info"
    assert str(ppt_root / "20.청년 미사" / "2026") in shown[-1][1]


def test_copy_ui_asks_before_overwrite_and_respects_no(cfg, ppt_root, tmp_path, monkeypatch):
    out = tmp_path / "r.pptx"
    out.write_bytes(b"new")
    dest_file = ppt_root / "20.청년 미사" / "2026" / "r.pptx"
    dest_file.parent.mkdir(parents=True)
    dest_file.write_bytes(b"old")
    shown = _stub_messagebox(monkeypatch, answer=False)
    gui._copy_result_with_ui(None, "youth", out, "20261003")
    assert [k for k, _ in shown] == ["ask"]
    assert dest_file.read_bytes() == b"old"


def test_copy_ui_reports_failure(cfg, ppt_root, tmp_path, monkeypatch):
    shown = _stub_messagebox(monkeypatch)
    gui._copy_result_with_ui(None, "youth", tmp_path / "없는파일.pptx", "20261003")
    assert shown[-1][0] == "error"


@pytest.mark.parametrize("mass_type", ["adult", "youth"])
def test_copy_ui_without_root_asks_picker_then_cancel_warns(cfg, tmp_path, monkeypatch, mass_type):
    monkeypatch.setattr(gui, "_ask_onedrive_path_popup", lambda **k: "")
    shown = _stub_messagebox(monkeypatch)
    out = tmp_path / "a.pptx"
    out.write_bytes(b"1")
    gui._copy_result_with_ui(None, mass_type, out, "20261004")
    assert shown[-1][0] == "warn"


def test_result_window_copy_button_sits_right_after_open_file_with_same_font():
    """2026-10-04: '원드라이브로 복사'는 '파일 열기' 바로 옆(같은 btn_frame), 같은 글꼴 크기(공용 팩토리 사용)."""
    import inspect
    src = inspect.getsource(gui._show_result_window)
    assert (src.index("'open_file', '파일 열기'") < src.index("'copy', '원드라이브로 복사'")
            < src.index("'open_folder', '폴더 열기'") < src.index("'close', '닫기'"))
    assert "copy_frame" not in src and "_copy_result_with_ui" in src
    factory = inspect.getsource(gui._result_button)
    assert "_UI['font_sz'], 'bold'" in factory   # 4개 모두 같은 글꼴 크기(한 사이즈 작게 줄이지 않음)


def test_run_gui_mode_asks_roots_before_combined_popup(monkeypatch):
    """루트 확보(필요한 것만 선택창)가 통합 입력창보다 먼저 일어난다 — 성인·청년 모두."""
    import missa_to_ppt as m
    for kr, mt in (("성인", "adult"), ("청년", "youth")):
        calls = []
        monkeypatch.setattr(m, "ensure_ppt_roots",
                            lambda mass_type, ask_if_missing=False: (calls.append((mass_type, ask_if_missing)), True)[1])
        monkeypatch.setattr(m, "_ask_combined_input_popup",
                            lambda mass_type: (calls.append("combined"), ("20261004", {}))[1])
        monkeypatch.setattr(m, "_ask_numbers_popup", lambda n, s: {})
        monkeypatch.setattr(m, "_ask_youth_hymn_popup", lambda: {})
        monkeypatch.setattr(m, "_run_with_progress_window", lambda fn: ("", ""))
        monkeypatch.setattr(m, "_show_result_window", lambda *a, **k: None)
        m._preloaded_inputs[0] = None
        try:
            m._run_gui_mode(kr)
        finally:
            m._preloaded_inputs[0] = None
        assert calls[:2] == [(mt, True), "combined"]


# ── 찾아보기 연동 ──

def test_combined_popup_browse_uses_remembered_dir_for_every_row():
    """통합 입력창의 '찾아보기'는 로컬 파일 다이얼로그 하나뿐(Graph 브라우저 없음)이며 initialdir을 쓰고 고른
    폴더를 행별로 기억한다."""
    import inspect
    src = inspect.getsource(gui._ask_combined_input_popup)
    assert "browse_initial_dir(mass_type, key)" in src and "remember_browse_dir(mass_type, key, path)" in src
    assert "_ask_onedrive_file_browser_popup" not in src and "_od_cache" not in src


def test_resolve_new_hymn_saved_into_local_folder(cfg, ppt_root, monkeypatch):
    import missa_content_updaters as cu
    import missa_youth_hymn_pdf as yh
    from pptx import Presentation
    monkeypatch.setattr(yh, "find_song_title", lambda src, n: "나의 고백")
    monkeypatch.setattr(yh, "build_hymn_pptx", lambda *a, **k: Presentation())
    monkeypatch.setattr(yh, "build_header_runs", lambda *a, **k: None)
    prs, title = cu.resolve_youth_hymn_pptx({"출처": "나주노", "번호": 172, "제목": None}, "입당")
    assert title == "나의 고백"
    saved = ppt_root / "20.청년 미사" / "2.성가" / "나주노 성가" / "나주노 성가 172 나의 고백.pptx"
    assert saved.exists()


def test_existing_hymn_file_title_comes_from_filename_without_pdf_lookup(cfg, ppt_root, monkeypatch):
    """2026-10-04 성능: 이미 만든 나주노/야훼이레 성가 파일이 있으면 제목을 파일명에서 얻고 PDF 조회(야훼 이레는
    574쪽 텍스트 검색 ≈ 2.5초)를 하지 않는다. 파일이 없을 때만 PDF에서 찾는다."""
    import missa_content_updaters as cu
    import missa_youth_hymn_pdf as yh
    from pptx import Presentation
    folder = ppt_root / "20.청년 미사" / "2.성가" / "나주노 성가"
    Presentation().save(str(folder / "나주노 성가 172 나의 고백.pptx"))
    monkeypatch.setattr(yh, "find_song_title", lambda *a: pytest.fail("PDF 제목 조회가 호출되면 안 됨"))
    monkeypatch.setattr(yh, "build_header_runs", lambda *a, **k: None)
    _, title = cu.resolve_youth_hymn_pptx({"출처": "나주노", "번호": 172, "제목": None}, "입당")
    assert title == "나의 고백"


def test_find_sections_youth_songs_unchanged_by_shared_scan(monkeypatch):
    """성가 탐색은 슬라이드 1회 스캔을 5종이 공유하지만, 스캔 없이 호출한 결과와 같아야 한다."""
    import missa_sections as ms
    from pptx import Presentation
    tpl = Path(__file__).resolve().parent.parent / "reference" / "청년미사" / \
        "Template_토요일 저녁 청년 주일미사_20260822_연중 제21주일.pptx"
    if not tpl.is_file():
        pytest.skip("청년 템플릿 없음(reference/, git 미포함)")
    prs = Presentation(str(tpl))
    scan = ms._scan_youth_slides(prs)
    for kw in ms.HYMN_LABEL_KW.values():
        assert ms._find_hymn_songs_youth(prs, kw, scan) == ms._find_hymn_songs_youth(prs, kw)
        legacy = []
        n = len(prs.slides)
        # 구 구현(도형을 직접 순회)과 동일한 결과인지 별도 계산으로 확인
        i = 0
        while i < n:
            if ms._is_hymn_title_youth(prs.slides[i], kw):
                cs = ce = i + 1
                while ce < n and ms._slide_has_picture(prs.slides[ce]):
                    ce += 1
                legacy.append((i, cs, ce))
                i = ce
                if not (i < n and ms._is_hymn_title_youth(prs.slides[i], kw)):
                    break
            else:
                i += 1
        got = [(s["title_idx"], s["content_start"], s["content_end"]) for s in ms._find_hymn_songs_youth(prs, kw, scan)]
        assert got == legacy


def test_result_buttons_use_option_c_palette():
    """목업 C안: 파일 열기 #59001d · 복사 #871b24 · 폴더 열기 #3f5a73 · 닫기 흰 배경+#59001d 테두리."""
    c = gui._RESULT_BTN_COLORS
    assert c["open_file"][0] == "#59001d" and c["copy"][0] == "#871b24" and c["open_folder"][0] == "#3f5a73"
    assert c["close"][0] == "#ffffff" and c["close"][1] == "#59001d" and c["close"][3] == "#59001d"


def test_result_button_close_has_border_frame_and_others_match_background():
    import tkinter as tk
    root = tk.Tk()
    root.withdraw()
    try:
        close = gui._result_button(root, "close", "닫기", lambda: None, 14)
        copy = gui._result_button(root, "copy", "복사", lambda: None, 16)
        assert close.cget("bg") == "#59001d"          # 테두리 색 프레임
        assert copy.cget("bg") == "#871b24"           # 테두리 없음 = 배경색(높이 맞춤)
        assert close.winfo_children()[0].cget("bg") == "#ffffff"
    finally:
        root.destroy()
