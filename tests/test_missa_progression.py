# -*- coding: utf-8 -*-
"""프로그레션 테스트 — 아직 안정화되지 않은 신규 기능을 구현 전에 명세한다.

현재 진행 중: 화답송 악보 이미지(PNG/JPG) 입력 지원.
- 1차 마일스톤(missa_psalm_score_image.py — 이미지 → 슬라이드 변환): 완료·사용자 승인.
- 2차 마일스톤(missa_to_ppt.py 배선 — find_files/CLI 오버라이드/입력창/update_화답송):
  이번 라운드. 설계서 §4(통합 지점) 그대로 구현.

설계서: _workspace/01_architect_design.md §4(통합 지점), §7(1차 마일스톤 16개 테스트 대상 행동).
기대값은 3개 예제 세트(20260809/20260816/20260823) 실측 정답이다(추측 아님).

안정화되면 regression-qa가 test_missa_regression.py로 승격하고 이 파일에서 제거한다.
"""
import io
from pathlib import Path

import numpy as np
import pytest
from PIL import Image

import missa_psalm_score_image as ps

BASE = Path(__file__).resolve().parent.parent  # tests/ -> 저장소 루트(reference/·assets/ 등의 기준)
REF = BASE / "reference" / "화답송악보"
TEMPLATE = BASE / "assets" / "화답송_악보_template.pptx"

# 세트별 원본 이미지 경로와 실측 기대값 (설계서 §0.1, §2.3, §7)
CASES = {
    "20260809": {
        "img": REF / "20260809" / "20260809_043842947.jpg",
        "crop_left": 372,
        "crop_right": 3906,
        "split": 2146,
        "top_ratio": 0.502,
    },
    "20260816": {
        "img": REF / "20260816" / "20260816_043822155.png",
        "crop_left": 223,
        "crop_right": 1953,
        "split": 1201,
        "top_ratio": 0.565,
    },
    "20260823": {
        "img": REF / "20260823" / "20260823_043736869.png",
        "crop_left": 372,
        "crop_right": 3906,
        # split=2304(top_ratio 0.547)는 이전 라운드의 오검출 버그가 낳은 값이었다 — 실제로는
        # 8분음표 기둥+깃발(스템+플래그)이었고 진짜 마디선이 아니었다(사용자가 실제 생성물을
        # 육안 확대 검토해 발견, _workspace/02c 라운드). detect_barlines()의 노트헤드/스템 인접
        # 배제 필터 추가 후 실측 재검증한 참값은 split=1971(진짜 마디선, 확대 이미지로 확인:
        # 좌우가 깨끗한 공백인 순수 수직선). "값이 그럴듯하게 보인다"만으로 설계서 수치를
        # 맹신하지 않고, 매 라운드 실제 이미지 확대 검토로 재검증해야 한다는 교훈.
        "split": 1971,
        "top_ratio": 0.452,
    },
}


@pytest.fixture(scope="module", params=list(CASES))
def case(request):
    c = CASES[request.param]
    if not c["img"].is_file():
        pytest.skip(
            f"화답송 악보 참조 이미지가 로컬에 없음(reference/화답송악보/, git 미포함): {c['img']}"
        )
    c = dict(c, name=request.param)
    return c


@pytest.fixture(scope="module")
def gray(case):
    return ps.load_gray_array(case["img"])


# ---------------------------------------------------------------------------
# 이미지 처리 (PowerPoint 불필요, 순수 함수 — 빠른 회귀 후보)
# ---------------------------------------------------------------------------

def test_01_watermark_crop_start(case, gray):
    """워터마크 크롭 시작 컬럼이 분리 공백 띠 우측에 온다 (±5px)."""
    x = ps.detect_watermark_crop_x(gray)
    assert abs(x - case["crop_left"]) <= 5, (case["name"], x)


def test_02_watermark_excluded_from_crop(case, gray):
    """워터마크가 크롭 결과에서 제외된다.

    (밀도 비교는 부적합 — 실측상 오선보 시작부 clef/조표가 워터마크 글자보다 잉크가
    더 촘촘하다.) 대신 두 사실로 강하게 단언한다: (1) 제거되는 좌측 영역에 실제
    워터마크 잉크가 존재하고(제거할 대상이 있음), (2) 크롭 경계가 분리 공백 띠 위에
    정확히 놓여 크롭 직전 컬럼이 완전 백색이다 → 워터마크 컬럼은 전부 크롭 좌측에 남아
    결과에서 빠진다.
    """
    crop_x = ps.detect_watermark_crop_x(gray)
    ink = gray < 128
    assert crop_x > 0, case["name"]
    assert ink[:, 10:crop_x].sum() > 0, ("제거 영역에 워터마크 잉크가 있어야 함", case["name"])
    assert ink[:, crop_x - 1].sum() == 0, ("크롭 경계가 공백 띠 위에 놓여야 함", case["name"])


def test_03_content_right_edge(case, gray):
    """오선보 유효 우측끝 검출 — 트레일링 공백 제거 (±5px)."""
    left, right, top, bottom = ps.detect_content_bounds(gray)
    assert abs(right - case["crop_right"]) <= 5, (case["name"], right)


def test_04_split_on_barline(case, gray):
    """분할점이 마디 경계(바라인) 위에 온다 — 중앙창 안 + 탐지 바라인 집합의 원소 (±20px)."""
    crop_left = ps.detect_watermark_crop_x(gray)
    left, right, top, bottom = ps.detect_content_bounds(gray)
    band = ps.staff_band(gray)
    barlines = ps.detect_barlines(gray, band)
    split = ps.choose_split_x(barlines, crop_left, right)
    assert abs(split - case["split"]) <= 20, (case["name"], split)
    # 중앙창 [0.30, 0.70] 안
    content_w = right - crop_left
    assert crop_left + 0.30 * content_w <= split <= crop_left + 0.70 * content_w
    # 폴백이 아닌 한 실제 탐지 바라인 중 하나여야 한다
    assert any(abs(split - b) <= 20 for b in barlines), (case["name"], barlines)


def test_05_balanced_two_lines(case, gray):
    """두 줄 폭 비율이 균형: |상단폭-하단폭|/content_w ≤ 0.15, 상단비율 실측 일치(±0.03)."""
    crop_left = ps.detect_watermark_crop_x(gray)
    left, right, top, bottom = ps.detect_content_bounds(gray)
    band = ps.staff_band(gray)
    split = ps.choose_split_x(ps.detect_barlines(gray, band), crop_left, right)
    content_w = right - crop_left
    up = split - crop_left
    low = right - split
    assert abs(up - low) / content_w <= 0.15, (case["name"], up, low)
    assert abs(up / content_w - case["top_ratio"]) <= 0.03, (case["name"], up / content_w)


def test_06_lossless_split(case, gray):
    """원본 픽셀 무손실 분할: 상단폭 + 하단폭 == crop_right - crop_left (겹침·누락 없음)."""
    crop_left = ps.detect_watermark_crop_x(gray)
    left, right, top, bottom = ps.detect_content_bounds(gray)
    band = ps.staff_band(gray)
    split = ps.choose_split_x(ps.detect_barlines(gray, band), crop_left, right)
    img = ps.load_flattened_image(CASES[case["name"]]["img"])
    up_img, low_img = ps.split_and_place(img, (crop_left, top, right, bottom + 1), split)
    assert up_img.width + low_img.width == right - crop_left, (case["name"],)
    assert up_img.height == low_img.height  # 같은 높이로 잘림


def test_07_rgba_flatten_safe():
    """RGBA 투명 입력이 검정 잉크로 오인되지 않는다 (알파 평탄화 검증).

    완전 투명 RGBA 이미지는 흰 배경 합성 후 그레이스케일에서 잉크(어두운 픽셀)가
    전혀 없어야 한다. 평탄화를 빠뜨리면 투명 픽셀이 0(검정)으로 변해 전부 잉크가 된다.
    """
    transparent = Image.new("RGBA", (100, 40), (0, 0, 0, 0))
    buf = io.BytesIO()
    transparent.save(buf, format="PNG")
    buf.seek(0)
    gray = ps.load_gray_array(buf)
    ink = gray < 128
    assert ink.mean() == 0.0, "투명 픽셀이 잉크로 오인됨 — 평탄화 누락"


def test_08_fallback_no_watermark_gap():
    """폴백: 워터마크 미검출 시 미크롭(crop_left==0) + 경고, 예외 없음.

    좌측 30%에 넓은 zero-run이 없는 합성 입력. 음표 손실보다 워터마크 잔존이 안전.
    """
    # 좌측 30% 안에 완전 백색 컬럼(zero-run)이 하나도 없도록 전 폭에 가로 잉크 바를 둔다
    arr = np.full((40, 300), 255, dtype=np.uint8)
    arr[10:30, :] = 0  # 모든 컬럼에 잉크 → 분리 공백 띠 없음
    x = ps.detect_watermark_crop_x(arr)
    assert x == 0


def test_09_fallback_no_barline_geometric_center():
    """폴백: 중앙창 바라인 0개 → 예외 없이 기하 중앙 근처 분할점 반환."""
    # 바라인(고밀도 세로선)이 전혀 없는 합성 이미지
    arr = np.full((40, 300), 255, dtype=np.uint8)
    arr[18:22, 20:280] = 0  # 얇은 가로 오선만
    band = ps.staff_band(arr)
    barlines = ps.detect_barlines(arr, band)
    split = ps.choose_split_x(barlines, 20, 280)
    assert 20 < split < 280
    assert abs(split - (20 + 280) // 2) <= 30  # 기하 중앙 근처


def test_09b_barline_rejects_notehead_fused_stem():
    """detect_barlines가 노트헤드/스템에 융합된 획을 마디선 후보에서 제외한다.

    실사용자가 20260823 미리보기에서 col=2304를 확대해 8분음표 기둥+깃발임을 확인한 사례의
    재발 방지 테스트(설계서 §2.4, 02c 라운드). 합성 이미지로 두 후보를 만든다:
    - col=100: 오선 전체를 관통하는 순수 수직선(좌우 공백) — 진짜 마디선.
    - col=150: 오선 전체를 관통하는 수직선 + 한쪽 행간에 노트헤드 모양 블록이 융합됨 — 가짜.
    둘 다 col_fill>0.85(밴드 전체 관통)라 1차 임계값만으로는 구분되지 않으므로, 노트헤드 인접
    배제 필터가 실제로 두 후보를 갈라내는지 확인한다.
    """
    w, h = 220, 60
    img = np.full((h, w), 255, dtype=np.uint8)
    for r in (10, 20, 30, 40, 50):
        img[r, :] = 0  # 오선 5줄, 폭 전체 관통
    img[10:51, 100:102] = 0  # 진짜 마디선: 순수 수직선, 좌우 공백
    img[10:51, 150:152] = 0  # 가짜 후보의 기둥(스템)
    img[24:37, 143:161] = 0  # 가짜 후보에 융합된 노트헤드(한쪽 행간에 폭넓게 번짐)

    band = ps.staff_band(img)
    barlines = ps.detect_barlines(img, band)
    assert any(abs(b - 100) <= 2 for b in barlines), ("진짜 마디선이 후보에서 빠짐", barlines)
    assert not any(abs(b - 150) <= 2 for b in barlines), (
        "노트헤드에 융합된 스템이 마디선으로 오검출됨", barlines
    )


# ---------------------------------------------------------------------------
# 슬라이드 생성 (템플릿 보존)
# ---------------------------------------------------------------------------

@pytest.fixture(scope="module")
def rendered(case):
    return ps.render_화답송_score_slide(CASES[case["name"]]["img"], "시편 00(00)", TEMPLATE)


def _iter_shapes(slide):
    return list(slide.shapes)


def test_10_exactly_one_slide(rendered):
    """출력 Presentation은 슬라이드 정확히 1장."""
    assert len(rendered.slides) == 1


def test_11_slide_size_preserved(rendered):
    """슬라이드 크기 12192000×6858000 유지 (copy_slide_from_prs 스케일 1:1 보장)."""
    assert rendered.slide_width == 12192000
    assert rendered.slide_height == 6858000


def test_12_template_text_preserved(rendered):
    """템플릿 라벨/저작권 텍스트박스가 그대로 보존된다."""
    texts = [
        sh.text_frame.text
        for sh in _iter_shapes(rendered.slides[0])
        if sh.has_text_frame
    ]
    joined = "\n".join(texts)
    assert "Responsorial" in joined
    assert "한국천주교중앙협의회" in joined
    assert "2022" in joined


def test_13_title_replaced(rendered):
    """제목박스 텍스트가 '화 답 송   {title}'로 교체, 다른 텍스트박스는 불변."""
    slide = rendered.slides[0]
    title_boxes = [
        sh for sh in _iter_shapes(slide)
        if sh.has_text_frame and "화 답 송" in sh.text_frame.text
    ]
    assert len(title_boxes) == 1
    assert title_boxes[0].text_frame.text.strip() == "화 답 송   시편 00(00)"
    # 저작권 박스는 제목 교체의 영향을 받지 않는다
    copyright_boxes = [
        sh for sh in _iter_shapes(slide)
        if sh.has_text_frame and "한국천주교중앙협의회" in sh.text_frame.text
    ]
    assert len(copyright_boxes) == 1


def test_13b_title_paragraph_child_order_schema_valid(rendered):
    """제목 교체 후 <a:p> 자식 순서가 CT_TextParagraph 스키마를 지킨다.

    스키마 시퀀스: pPr?, (run|br)*, endParaRPr?. run이 endParaRPr보다 뒤에 오면
    PowerPoint가 "손상된 파일"로 취급할 수 있는 스키마 위반이다. validate_pptx_structure()는
    XML 제어문자·끊어진 rId만 검사하고 요소 순서를 검사하지 않으므로(missa_to_ppt.py:6071),
    이 순서는 lxml로 직접 확인해야 한다(02b_review_report.md 리뷰 지적 — CLAUDE.md의
    "pPr append 금지" 함정이 <a:p> 자식 레벨(run vs endParaRPr)에서 재발했었음).
    """
    from pptx.oxml.ns import qn

    slide = rendered.slides[0]
    title_boxes = [
        sh for sh in _iter_shapes(slide)
        if sh.has_text_frame and "화 답 송" in sh.text_frame.text
    ]
    assert len(title_boxes) == 1
    p = title_boxes[0].text_frame.paragraphs[0]._p
    tags = [child.tag for child in p]
    eprp_tag = qn("a:endParaRPr")
    r_tag = qn("a:r")
    assert r_tag in tags, "제목 run이 존재해야 함"
    if eprp_tag in tags:
        r_indices = [i for i, t in enumerate(tags) if t == r_tag]
        eprp_index = tags.index(eprp_tag)
        assert max(r_indices) < eprp_index, (
            "run이 endParaRPr보다 앞에 와야 한다(스키마 위반)", tags
        )


def test_14_two_score_pictures_placed(rendered):
    """악보 이미지 PICTURE 2장 추가 (width>1M), 각각 상/하 밴드 안, 종횡비 유지."""
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    slide = rendered.slides[0]
    big_pics = [
        sh for sh in _iter_shapes(slide)
        if sh.shape_type == MSO_SHAPE_TYPE.PICTURE and sh.width > 1_000_000
    ]
    assert len(big_pics) == 2, [sh.width for sh in big_pics]
    tops = sorted(sh.top for sh in big_pics)
    # 상단 밴드 top≈1.0M, 하단 밴드 top≈3.65M (밴드 내 세로 중앙정렬이라 약간 아래로 이동)
    assert 1_000_000 <= tops[0] < 3_000_000, tops
    assert 3_000_000 <= tops[1] < 6_000_000, tops
    # 밴드 폭(11392000, left여백 400000) 안에 있고 왜곡 없이 배치
    for sh in big_pics:
        assert sh.left >= 400_000 - 1
        assert sh.left + sh.width <= 400_000 + 11_392_000 + 1
        assert sh.height <= 2_500_000 + 1


def test_14b_uniform_scale_across_both_lines():
    """상/하 두 단의 폭이 크게 다르면 공통(더 작은) 배율을 적용해 음표 크기를 통일한다.

    사용자 요구(02c 라운드): "일부 악보는 첫째 단과 둘째 단의 폭이 다를 수 있다 — 그 경우
    전체적인 비율을 조정해서... 첫째 단과 둘째 단의 높이가 더 작아지는 거지." 각자 독립적으로
    자기 밴드에 맞춰 스케일링하면(예전 방식) 폭이 좁은 쪽이 상대적으로 확대돼 두 줄의 음표
    크기가 달라 보인다. 원본 픽셀 높이가 같은 두 이미지(폭만 크게 다름)를 넣었을 때, 결과
    슬라이드의 두 PICTURE 높이(EMU)가 같아야 한다 — 배율이 공통이라는 강한 증거.
    """
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    wide_img = Image.new("RGB", (2000, 100), "white")
    narrow_img = Image.new("RGB", (300, 100), "white")
    prs = ps.build_slide(TEMPLATE, "시편 00(00)", wide_img, narrow_img)
    slide = prs.slides[0]
    big_pics = [
        sh for sh in slide.shapes
        if sh.shape_type == MSO_SHAPE_TYPE.PICTURE and sh.width > 1_000_000
    ]
    assert len(big_pics) == 2
    heights = sorted(sh.height for sh in big_pics)
    # 원본 픽셀 높이가 동일(100=100)하므로 배율이 같다면 결과 EMU 높이도 같아야 한다.
    assert abs(heights[0] - heights[1]) <= 2, (
        "두 단의 배율이 달라 높이가 다름 — 공통 배율 미적용", heights
    )
    widths = sorted(sh.width for sh in big_pics)
    # 폭 비율(2000:300 ≈ 6.67:1)이 결과에도 그대로 유지돼야 한다(공통 배율은 왜곡을 만들지 않음).
    assert abs(widths[1] / widths[0] - 2000 / 300) < 0.05


def test_15_residual_shapes_preserved(rendered):
    """잔재 4도형(12×7·41×17 EMU PICTURE) 그대로 — 템플릿 무손상."""
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    slide = rendered.slides[0]
    tiny_pics = [
        sh for sh in _iter_shapes(slide)
        if sh.shape_type == MSO_SHAPE_TYPE.PICTURE and sh.width < 1_000_000
    ]
    assert len(tiny_pics) == 4, [(sh.width, sh.height) for sh in tiny_pics]


def test_16_copy_slide_from_prs_compatible(rendered, tmp_path):
    """생성 Presentation을 타깃에 copy_slide_from_prs로 복사 시 손상 없이 rel까지 복사."""
    import missa_to_ppt as mtp
    import missa_ooxml_utils as ou
    from pptx import Presentation

    target = Presentation(str(TEMPLATE))
    ou.copy_slide_from_prs(target, len(target.slides), rendered, 0)
    out = tmp_path / "copied.pptx"
    target.save(str(out))
    problems = mtp.validate_pptx_structure(str(out))
    assert problems == [], problems


# ---------------------------------------------------------------------------
# 2차 마일스톤: missa_to_ppt.py 배선 (설계서 §4 통합 지점)
# ---------------------------------------------------------------------------

import missa_to_ppt as mtp  # noqa: E402  (모듈 상단 import 관례를 따르되, 1차 섹션과 분리해 배치)


def test_find_files_prefers_pptx_over_image(tmp_path):
    """화답송_pptx가 있으면 화답송_img 탐색 자체를 건너뛰고 pptx를 사용한다(기존 동작 보존)."""
    (tmp_path / "화답송 악보 시편 1.pptx").write_bytes(b"")
    (tmp_path / "photo.jpg").write_bytes(b"")
    files = mtp.find_files(str(tmp_path), {}, copy_hymn_scores=False)
    assert files["화답송_pptx"] is not None
    assert files["화답송_pptx"].name == "화답송 악보 시편 1.pptx"
    assert files["화답송_img"] is None


def test_find_files_named_image_when_no_pptx(tmp_path):
    """pptx가 없고 파일명에 '화답송'이 들어간 이미지가 있으면 그것을 화답송_img로 찾는다."""
    (tmp_path / "화답송 사진.jpg").write_bytes(b"")
    files = mtp.find_files(str(tmp_path), {}, copy_hymn_scores=False)
    assert files["화답송_pptx"] is None
    assert files["화답송_img"] is not None
    assert files["화답송_img"].name == "화답송 사진.jpg"


def test_find_files_unique_image_fallback_when_unnamed(tmp_path):
    """파일명에 '화답송'이 없어도(촬영 앱 타임스탬프 파일명) 폴더 내 유일 이미지면 폴백한다."""
    (tmp_path / "20260809_043842947.jpg").write_bytes(b"")
    files = mtp.find_files(str(tmp_path), {}, copy_hymn_scores=False)
    assert files["화답송_img"] is not None
    assert files["화답송_img"].name == "20260809_043842947.jpg"


def test_find_files_ambiguous_multiple_unnamed_images_no_fallback(tmp_path):
    """이름 매칭이 안 되고 이미지가 여러 개면 어떤 게 화답송 악보인지 모호하므로 폴백하지
    않는다 — 잘못된 이미지를 화답송으로 오인하는 것보다 미검출이 안전하다."""
    (tmp_path / "photo1.jpg").write_bytes(b"")
    (tmp_path / "photo2.png").write_bytes(b"")
    files = mtp.find_files(str(tmp_path), {}, copy_hymn_scores=False)
    assert files["화답송_img"] is None


def test_apply_hwadapsong_override_routes_image_extension():
    """--화답송 오버라이드가 이미지 확장자면 화답송_img로 가고, 반대쪽 키는 비워진다."""
    files = {"화답송_pptx": Path("auto_found.pptx"), "화답송_img": None}
    mtp.apply_화답송_override(files, "C:/mass/photo.jpg")
    assert files["화답송_img"] == Path("C:/mass/photo.jpg")
    assert files["화답송_pptx"] is None


def test_apply_hwadapsong_override_routes_pptx_extension():
    """--화답송 오버라이드가 pptx 확장자면 화답송_pptx로 가고, 반대쪽 키는 비워진다."""
    files = {"화답송_pptx": None, "화답송_img": Path("auto_found.jpg")}
    mtp.apply_화답송_override(files, "C:/mass/manual.pptx")
    assert files["화답송_pptx"] == Path("C:/mass/manual.pptx")
    assert files["화답송_img"] is None


def _화답송_json(title: str) -> dict:
    return {
        "화답송": {
            "title": title,
            "content": "◎ 후렴 테스트\n○ 절1\n○ 절2\n○ 절3",
        }
    }


def test_update_화답송_uses_image_when_no_pptx_end_to_end(tmp_path):
    """update_화답송이 화답송_img_path만 주어졌을 때 render_화답송_score_slide()를 거쳐
    만들어진 슬라이드가 실제 결과 PPT에 copy_slide_from_prs 경로로 정상 복사되는지 확인한다
    (실제 참조 예제 이미지로 end-to-end 검증, 설계서 §4.3)."""
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    prs = mtp.Presentation(
        str(BASE / "output" / "20260712" / "Template_20260628_연중 제13주일 (교황주일).pptx")
    )
    sections = mtp.find_sections(prs)
    json_data = _화답송_json("시편 99(98)")
    img_path = CASES["20260809"]["img"]

    mtp.update_화답송(prs, json_data, sections, None, 화답송_img_path=img_path, is_sunday=True)

    start = sections["화답송_start"]
    score_slide = prs.slides[start]  # i=0(짝수) => 악보 슬라이드
    big_pics = [
        sh for sh in score_slide.shapes
        if sh.shape_type == MSO_SHAPE_TYPE.PICTURE and sh.width > 1_000_000
    ]
    assert len(big_pics) == 2, "이미지 입력 경로로 악보 슬라이드에 상/하 악보 그림 2장이 있어야 함"

    texts = [sh.text_frame.text for sh in score_slide.shapes if sh.has_text_frame]
    assert any("시편 99(98)" in t for t in texts), texts

    out = tmp_path / "e2e_img.pptx"
    prs.save(str(out))
    assert mtp.validate_pptx_structure(str(out)) == []


def test_update_화답송_existing_pptx_path_unaffected(tmp_path):
    """기존 pptx 입력 경로가 이번 배선 변경으로 전혀 달라지지 않았는지 확인(회귀 방지)."""
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    prs = mtp.Presentation(
        str(BASE / "output" / "20260712" / "Template_20260628_연중 제13주일 (교황주일).pptx")
    )
    sections = mtp.find_sections(prs)
    json_data = _화답송_json("시편 67(66)")
    pptx_path = REF / "20260816" / "시편 67(66).pptx"

    mtp.update_화답송(prs, json_data, sections, pptx_path, 화답송_img_path=None, is_sunday=True)

    start = sections["화답송_start"]
    score_slide = prs.slides[start]
    big_pics = [
        sh for sh in score_slide.shapes
        if sh.shape_type == MSO_SHAPE_TYPE.PICTURE and sh.width > 1_000_000
    ]
    assert len(big_pics) == 2

    out = tmp_path / "e2e_pptx.pptx"
    prs.save(str(out))
    assert mtp.validate_pptx_structure(str(out)) == []


def test_update_화답송_prefers_pptx_over_image(monkeypatch):
    """화답송_pptx_path와 화답송_img_path가 둘 다 주어지면 pptx가 우선한다(설계서 §4.3 우선순위).

    render_화답송_score_slide가 절대 호출되지 않아야 함을 몽키패치로 직접 확인한다.
    """
    import missa_psalm_score_image

    def _must_not_be_called(*args, **kwargs):
        raise AssertionError("이미지 경로가 호출됨 — pptx가 우선이어야 하는데 위반됨")

    monkeypatch.setattr(
        missa_psalm_score_image, "render_화답송_score_slide", _must_not_be_called
    )

    prs = mtp.Presentation(
        str(BASE / "output" / "20260712" / "Template_20260628_연중 제13주일 (교황주일).pptx")
    )
    sections = mtp.find_sections(prs)
    json_data = _화답송_json("시편 1(1)")
    pptx_path = REF / "20260816" / "시편 67(66).pptx"
    img_path = CASES["20260809"]["img"]

    # 예외 없이 끝나야 함 = render_화답송_score_slide가 호출되지 않았다는 뜻
    mtp.update_화답송(
        prs, json_data, sections, pptx_path, 화답송_img_path=img_path, is_sunday=True
    )


# ---------------------------------------------------------------------------
# 공지사항 슬라이드 삽입 (설계서 _workspace/공지사항_슬라이드/01_architect_design.md)
# §6 테스트 대상 행동 1~9. 8번(GUI tkinter 확장자 검증)만 자동 테스트 제외.
# 기대 인덱스는 세 git 추적 픽스처의 실측 find_sections() 값이다(추측 아님):
#   20260712(주일) n=147, 2차봉헌_content_end=132(blank)
#   20260705(성수축복) n=144, 2차봉헌_content_end=129(blank)
#   20260624(평일) 2차봉헌_* 키 부재
# ---------------------------------------------------------------------------

from pptx import Presentation as _Presentation  # noqa: E402
from pptx.util import Emu as _Emu  # noqa: E402
from missa_content_updaters import insert_공지사항  # noqa: E402
from missa_ooxml_utils import all_slide_texts as _all_slide_texts  # noqa: E402

import _gen_helper as _gen_fx  # noqa: E402


class _LazyFixtures:
    """`_FIXTURES["20260712"]` → 임시 폴더에 생성된 결과 PPT 경로(저장소 output/의 결과 PPT는 추적하지 않는다)."""

    def __getitem__(self, date_str):
        return _gen_fx.generated_path(date_str)


_FIXTURES = _LazyFixtures()


def _make_공지사항_pptx(path, n):
    """식별 가능한 텍스트를 가진 n장짜리 공지사항 PPT를 임시 생성한다.

    각 슬라이드에 '공지사항 슬라이드 {i}' 텍스트박스를 넣어 삽입 후 서식/순서 보존을
    강하게 단언할 수 있게 한다. n==0이면 슬라이드 없는 빈 PPT를 만든다.
    """
    prs = _Presentation()
    for i in range(n):
        slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank layout
        box = slide.shapes.add_textbox(_Emu(914400), _Emu(914400), _Emu(3657600), _Emu(914400))
        box.text_frame.text = f"공지사항 슬라이드 {i}"
    prs.save(str(path))
    return prs


def _src_texts(path):
    src = _Presentation(str(path))
    return [
        "\n".join(
            sh.text_frame.text for sh in s.shapes if sh.has_text_frame
        ).strip()
        for s in src.slides
    ]


def test_공지사항_01_insert_count_주일(tmp_path):
    """[삽입 총수] N장짜리 공지사항 → 반환 N+1, 전체 슬라이드 == 기존 + N + 1."""
    from missa_sections import find_sections

    n_src = 4
    src = tmp_path / "공지사항.pptx"
    _make_공지사항_pptx(src, n_src)

    prs = _Presentation(str(_FIXTURES["20260712"]))
    before = len(prs.slides)
    sec = find_sections(prs)
    added = insert_공지사항(prs, src, sec)

    assert added == n_src + 1
    assert len(prs.slides) == before + n_src + 1


def test_공지사항_02_one_divider_each_side(tmp_path):
    """[앞뒤 구분선 정확히 1장씩] 앞 구분선(기존 content_end)·뒤 구분선(신규) 모두 blank,
    그 뒤는 non-blank(마침 축복). 공지사항 경계에 2장 연속 blank가 없음."""
    from missa_sections import find_sections

    n_src = 3
    src = tmp_path / "공지사항.pptx"
    _make_공지사항_pptx(src, n_src)

    prs = _Presentation(str(_FIXTURES["20260712"]))
    sec = find_sections(prs)
    front = sec["2차봉헌_content_end"]  # 132 (실측)
    insert_공지사항(prs, src, sec)

    texts = _all_slide_texts(prs)
    is_blank = [not t.strip() for t in texts]

    # 앞 구분선: content_end 자리 그대로 blank
    assert is_blank[front], ("앞 구분선이 blank가 아님", front)
    # 공지사항 콘텐츠 n_src장: 모두 non-blank
    for i in range(n_src):
        assert not is_blank[front + 1 + i], ("공지사항 콘텐츠가 blank임", front + 1 + i)
    # 뒤 구분선: 콘텐츠 마지막 다음 자리 blank
    rear = front + 1 + n_src
    assert is_blank[rear], ("뒤 구분선이 blank가 아님", rear)
    # 2026-09-13 실측 버그의 end-to-end 재발 방지: 뒤 구분선도 앞 구분선과 동일하게
    # showMasterSp="0"이어야 한다(다르면 마스터 상속 요소가 노출돼 색이 달라 보임).
    front_show_master = prs.slides[front].element.get("showMasterSp")
    rear_show_master = prs.slides[rear].element.get("showMasterSp")
    assert rear_show_master == front_show_master, (
        "뒤 구분선의 showMasterSp가 앞 구분선과 다름", front_show_master, rear_show_master,
    )
    # 뒤 구분선 다음은 non-blank(마침 축복 기도)
    assert not is_blank[rear + 1], ("뒤 구분선 다음이 blank임 — 2장 연속 blank", rear + 1)
    # front-1(2차봉헌 마지막 콘텐츠)도 non-blank → front 앞뒤로 2장 연속 blank 없음
    assert not is_blank[front - 1], ("앞 구분선 앞이 blank임 — 2장 연속 blank", front - 1)


def test_공지사항_03_content_format_preserved(tmp_path):
    """[콘텐츠 서식 보존] 삽입된 공지사항 슬라이드 텍스트가 원본과 동일(순서 포함).
    뒤 구분선 슬라이드는 shapes 0개."""
    from missa_sections import find_sections

    n_src = 3
    src = tmp_path / "공지사항.pptx"
    _make_공지사항_pptx(src, n_src)
    src_texts = _src_texts(src)  # ['공지사항 슬라이드 0', '...1', '...2']

    prs = _Presentation(str(_FIXTURES["20260712"]))
    sec = find_sections(prs)
    front = sec["2차봉헌_content_end"]
    insert_공지사항(prs, src, sec)

    for i in range(n_src):
        slide = prs.slides[front + 1 + i]
        joined = "\n".join(
            sh.text_frame.text for sh in slide.shapes if sh.has_text_frame
        ).strip()
        assert joined == src_texts[i], (i, joined, src_texts[i])

    rear = front + 1 + n_src
    assert len(prs.slides[rear].shapes) == 0, "뒤 구분선은 원본 blank 복제라 shapes 0개여야 함"


def test_공지사항_04_no_2차봉헌_slot_skips_평일(tmp_path):
    """[2차봉헌 없는 평일 → 경고 후 skip] 예외 없이 반환 0, 슬라이드 수 불변."""
    from missa_sections import find_sections

    src = tmp_path / "공지사항.pptx"
    _make_공지사항_pptx(src, 2)

    prs = _Presentation(str(_FIXTURES["20260624"]))
    before = len(prs.slides)
    sec = find_sections(prs)
    assert sec.get("2차봉헌_content_end") is None  # 실측: 키 자체가 없음

    added = insert_공지사항(prs, src, sec)
    assert added == 0
    assert len(prs.slides) == before


def test_공지사항_05_insert_성수축복(tmp_path):
    """[성수축복(2차봉헌 있는 특수 주일)] content_end=129 기준 삽입 성립."""
    from missa_sections import find_sections

    n_src = 2
    src = tmp_path / "공지사항.pptx"
    _make_공지사항_pptx(src, n_src)

    prs = _Presentation(str(_FIXTURES["20260705"]))
    before = len(prs.slides)
    sec = find_sections(prs)
    front = sec["2차봉헌_content_end"]  # 129 (실측)
    added = insert_공지사항(prs, src, sec)

    assert added == n_src + 1
    assert len(prs.slides) == before + n_src + 1

    texts = _all_slide_texts(prs)
    is_blank = [not t.strip() for t in texts]
    assert is_blank[front]
    rear = front + 1 + n_src
    assert is_blank[rear]
    assert not is_blank[rear + 1]


def test_공지사항_06a_missing_path_returns_zero(tmp_path):
    """[로드 실패 방어] 존재하지 않는 경로 → 반환 0, 슬라이드 불변, 예외 없음."""
    from missa_sections import find_sections

    prs = _Presentation(str(_FIXTURES["20260712"]))
    before = len(prs.slides)
    sec = find_sections(prs)
    added = insert_공지사항(prs, tmp_path / "없는파일.pptx", sec)
    assert added == 0
    assert len(prs.slides) == before


def test_공지사항_06b_zero_slide_pptx_returns_zero(tmp_path):
    """[0장 방어] 슬라이드 0장짜리 공지사항 PPT → 반환 0, 슬라이드 불변."""
    from missa_sections import find_sections

    src = tmp_path / "empty.pptx"
    _make_공지사항_pptx(src, 0)

    prs = _Presentation(str(_FIXTURES["20260712"]))
    before = len(prs.slides)
    sec = find_sections(prs)
    added = insert_공지사항(prs, src, sec)
    assert added == 0
    assert len(prs.slides) == before


def test_공지사항_07_front_not_blank_skips(tmp_path):
    """[앞 구분선이 blank 아님 → skip] content_end가 non-blank를 가리키면 반환 0(방어 (c))."""
    from missa_sections import find_sections

    src = tmp_path / "공지사항.pptx"
    _make_공지사항_pptx(src, 2)

    prs = _Presentation(str(_FIXTURES["20260712"]))
    before = len(prs.slides)
    sec = find_sections(prs)
    # 132(blank) 대신 133(마침 축복, non-blank)을 가리키도록 조작
    sec = dict(sec, **{"2차봉헌_content_end": 133})
    added = insert_공지사항(prs, src, sec)
    assert added == 0
    assert len(prs.slides) == before


def test_공지사항_structure_valid_after_insert(tmp_path):
    """[구조 무손상] 삽입·저장 후 validate_pptx_structure()가 빈 리스트(끊어진 rId/제어문자 없음)."""
    from missa_sections import find_sections

    src = tmp_path / "공지사항.pptx"
    _make_공지사항_pptx(src, 3)

    prs = _Presentation(str(_FIXTURES["20260712"]))
    sec = find_sections(prs)
    insert_공지사항(prs, src, sec)

    out = tmp_path / "out.pptx"
    prs.save(str(out))
    assert mtp.validate_pptx_structure(str(out)) == []


def test_공지사항_08_cli_rejects_non_pptx(monkeypatch):
    """[CLI 확장자 검증] --공지사항에 .pptx 아닌 경로 → SystemExit(오류 종료)."""
    monkeypatch.setattr(
        "sys.argv", ["missa_to_ppt.py", "20260712", "--공지사항", "foo.docx", "--test"]
    )
    with pytest.raises(SystemExit):
        mtp.parse_args()


def test_공지사항_08b_cli_accepts_pptx_into_5tuple(monkeypatch):
    """[CLI 통과] .pptx면 parse_args가 7-튜플을 반환하고 5번째가 공지사항 경로다
    (청년미사 CLI 지원 추가로 mass_type/성가_선택이 6·7번째로 붙어 5-튜플→7-튜플로
    확장됐다 — 앞 5개 위치는 그대로다)."""
    monkeypatch.setattr(
        "sys.argv", ["missa_to_ppt.py", "20260712", "--공지사항", "notice.pptx", "--test"]
    )
    result = mtp.parse_args()
    assert len(result) == 7
    assert result[4] == "notice.pptx"


def test_공지사항_08c_cli_default_none(monkeypatch):
    """[CLI 기본값] --공지사항 미지정이면 5번째가 None(선택 입력)."""
    monkeypatch.setattr("sys.argv", ["missa_to_ppt.py", "20260712", "--test"])
    result = mtp.parse_args()
    assert len(result) == 7
    assert result[4] is None


def test_공지사항_09_find_files_detects_by_name(tmp_path):
    """[자동 탐색] 파일명에 '공지사항'이 든 .pptx가 files['공지사항']에 잡힌다."""
    (tmp_path / "20260712_ref.pptx").write_bytes(b"")
    (tmp_path / "공지사항.pptx").write_bytes(b"")
    files = mtp.find_files(str(tmp_path), {}, copy_hymn_scores=False)
    assert files["공지사항"] is not None
    assert files["공지사항"].name == "공지사항.pptx"


def test_공지사항_09b_find_files_none_when_absent(tmp_path):
    """[미입력 무변경 가드 근거] 공지사항 파일이 없으면 files['공지사항']==None →
    main의 `if files.get('공지사항')` 가드가 False라 insert_공지사항이 호출되지 않는다."""
    (tmp_path / "20260712_ref.pptx").write_bytes(b"")
    files = mtp.find_files(str(tmp_path), {}, copy_hymn_scores=False)
    assert files.get("공지사항") is None
    assert not files.get("공지사항")  # 가드 falsy


# ---------------------------------------------------------------------------
# duplicate_slide()의 showMasterSp 속성 보존 (2026-09-13 실측 버그)
# insert_공지사항이 만든 뒤 구분 슬라이드(142, duplicate_slide 산출물)가 앞 구분 슬라이드
# (133, hand-authored)와 색이 달라 보였다. 두 슬라이드는 layout/배경이 100% 동일했으나
# 최상위 <p:sld>의 showMasterSp 속성이 앞쪽은 "0"(마스터 상속 요소 숨김), 뒤쪽은 부재
# (python-pptx add_slide 기본값 = 마스터 요소 노출)였던 것이 원인. duplicate_slide가
# spTree/이미지 rel/p:bg는 복사하면서 <p:sld> 자체의 showMasterSp 속성은 복사하지 않았다.
# ---------------------------------------------------------------------------

from missa_ooxml_utils import duplicate_slide as _duplicate_slide  # noqa: E402


def test_showmastersp_present_copied_to_duplicate():
    """원본 <p:sld showMasterSp="0">이면 복제본의 <p:sld>도 showMasterSp="0"이어야 한다.

    이 속성이 누락되면 layout/배경이 동일해도 마스터 배치 요소(장식·로고 등)가 복제본에서만
    노출돼 색이 달라 보인다. p:bg 복사만으로는 커버되지 않는 별도 속성이다.
    """
    prs = _Presentation()
    src = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    src.element.set("showMasterSp", "0")
    new_idx = _duplicate_slide(prs, 0)
    assert prs.slides[new_idx].element.get("showMasterSp") == "0"


def test_showmastersp_absent_stays_absent():
    """원본에 showMasterSp 속성이 없으면 복제본에도 추가하지 않는다(불필요한 속성 주입 금지).

    원본이 add_slide 기본 상태(속성 부재 = 마스터 요소 노출)면 복제본도 그 상태를 그대로
    유지해야 한다 — 원본에 없던 속성을 새로 만들어 넣으면 안 된다.
    """
    prs = _Presentation()
    src = prs.slides.add_slide(prs.slide_layouts[6])
    assert src.element.get("showMasterSp") is None  # 전제: add_slide 기본은 속성 부재
    new_idx = _duplicate_slide(prs, 0)
    assert prs.slides[new_idx].element.get("showMasterSp") is None


def test_showmastersp_real_fixture_handauthored_slide():
    """실측 재현: 20260712 idx132(hand-authored, showMasterSp="0")를 복제하면 복제본도 "0".

    2026-09-13 프로덕션 파일의 앞 구분 슬라이드(idx132)와 동일한 성격의 슬라이드다(git 추적
    픽스처 20260712 idx132도 showMasterSp="0"임을 실측). 이 슬라이드를 재사용해 뒤 구분
    슬라이드를 만들 때 속성이 보존돼야 두 슬라이드가 같은 색으로 렌더링된다.
    """
    prs = _Presentation(str(_FIXTURES["20260712"]))
    assert prs.slides[132].element.get("showMasterSp") == "0"  # 픽스처 실측 전제
    new_idx = _duplicate_slide(prs, 132)
    assert prs.slides[new_idx].element.get("showMasterSp") == "0"


# ===========================================================================
# 청년미사 1단계 — 모듈 1: 영문 복음 조회 (missa_youth_gospel.py)
# 설계서 _workspace/청년미사_1단계/01_architect_design.md §2, §9(G1~G7).
#
# 파싱 의존 테스트(G1/G4~G7)는 커밋된 실측 HTML 픽스처로 결정론적으로 돈다(universalis는
# 약 열흘치만 서빙하므로 날짜 기반 라이브 테스트는 시간이 지나면 깨진다). 리다이렉트 함정
# (G2/G3)만 라이브로 검증한다 — 20200101/20270101은 현재 날짜와 무관하게 항상 범위를 벗어나
# 리다이렉트되므로 시간에 안정적이다.
# ===========================================================================
import missa_youth_gospel as yg

GOSPEL_FIX = BASE / "reference" / "청년미사" / "gospel_fixtures"


def _fix_html(date_str):
    return (GOSPEL_FIX / f"{date_str}.html").read_text(encoding="utf-8")


def test_G1_parse_reference_and_body_start():
    """G1: 범위 내 날짜 HTML → reference/본문 정확 추출."""
    g = yg.parse_gospel(_fix_html("20260913"))
    assert g["reference"] == "Matthew 18:21-35"
    assert g["content"].startswith("Peter went up to Jesus")


def test_G2_far_past_redirect_raises():
    """G2(최우선): 과거로 먼 날짜(→/mass.htm 리다이렉트) → GospelFetchError.

    라이브. 조용히 오늘자 복음이 반환되는 가장 위험한 케이스를 막는지 확인."""
    with pytest.raises(yg.GospelFetchError):
        yg.get_youth_gospel("20200101")


def test_G3_far_future_redirect_raises():
    """G3: 미래로 먼 날짜(→/n-otherdates.htm) → GospelFetchError. 라이브."""
    with pytest.raises(yg.GospelFetchError):
        yg.get_youth_gospel("20270101")


def test_G4_p_pi_paragraph_boundary():
    """G4: p/pi 형식 본문 단락 경계가 \\n\\n으로 유지."""
    g = yg.parse_gospel(_fix_html("20260913"))
    assert "\n\n" in g["content"]
    # pi 단락("And so the kingdom…")이 p 단락 뒤 \n\n 경계로 이어짐
    parts = g["content"].split("\n\n")
    assert len(parts) >= 2
    assert any("kingdom of heaven" in p for p in parts)


def test_G5_v_vi_format_collected():
    """G5: v/vi 형식(족보 20260908)도 본문 수집."""
    g = yg.parse_gospel(_fix_html("20260908"))
    assert g["reference"] == "Matthew 1:1-16,18-23"
    assert "Abraham was the father of Isaac" in g["content"]
    # v 형식은 절마다 div → 여러 단락으로 수집됨
    assert g["content"].count("\n\n") >= 5


def test_G6_html_entities_decoded():
    """G6: &#8216; 등 HTML 엔티티가 유니코드 문자로 디코딩(&# 잔존 없음)."""
    g = yg.parse_gospel(_fix_html("20260913"))
    assert "‘" in g["content"]  # &#8216; 곡따옴표
    assert "&#" not in g["content"]
    assert "&#" not in g["reference"]


def test_G7_theme_and_date_in_full_json(monkeypatch):
    """G7: get_youth_gospel JSON에 theme·date 포함, 값 정확(fetch를 픽스처로 대체)."""
    monkeypatch.setattr(yg, "fetch_gospel_html", lambda d: _fix_html("20260913"))
    out = yg.get_youth_gospel("20260913")
    assert out["date"] == "20260913"
    assert out["복음"]["theme"] == "To be forgiven, you must forgive"
    assert out["복음"]["reference"] == "Matthew 18:21-35"
    assert out["복음"]["content"].startswith("Peter went up to Jesus")


# ===========================================================================
# 청년미사 1단계 — 모듈 2: PDF 성가집 → PPT (missa_youth_hymn_pdf.py)
# 설계서 §3, §9(H1~H4, S1~S6, Hd1~Hd4, C1~C5, P1~P4).
# 실측 자산(나주노/야훼 PDF, 샘플 PPT)을 기준으로 검증한다 — 하드코딩 기대값 금지.
# ===========================================================================
import missa_youth_hymn_pdf as hp
import fitz as _fitz
import os as _os

_YOUTH_PDFS_AVAILABLE = _os.path.isfile(hp.SOURCES["나주노"]["pdf"]) and _os.path.isfile(
    hp.SOURCES["야훼 이레"]["pdf"]
)
_YOUTH_SKIP_REASON = "나주노/야훼이레 PDF 원본이 로컬에 없음(reference/청년미사/, git 미포함)"


@pytest.fixture(scope="module")
def naju_doc():
    d = _fitz.open(hp.SOURCES["나주노"]["pdf"])
    yield d
    d.close()


@pytest.fixture(scope="module")
def yahwe_doc():
    d = _fitz.open(hp.SOURCES["야훼 이레"]["pdf"])
    yield d
    d.close()


@pytest.mark.skipif(not _YOUTH_PDFS_AVAILABLE, reason=_YOUTH_SKIP_REASON)
def test_H1_naju_page_and_title():
    """H1: 나주노 362 → TOC 기반 page index 426, 제목 '삼위일체'(번호 접두 제거)."""
    assert hp.find_song_page("나주노", 362) == 426
    assert hp.find_song_title("나주노", 362) == "삼위일체"


@pytest.mark.skipif(not _YOUTH_PDFS_AVAILABLE, reason=_YOUTH_SKIP_REASON)
def test_H2_yahwe_page():
    """H2: 야훼 이레 810 → search_for 기반 page index 572."""
    assert hp.find_song_page("야훼 이레", 810) == 572


@pytest.mark.skipif(not _YOUTH_PDFS_AVAILABLE, reason=_YOUTH_SKIP_REASON)
def test_H3_yahwe_y_range_excludes_neighbors(yahwe_doc):
    """H3: 야훼810 y범위 = (제목아래≈396.6, 저작권top≈743.7). 809(y≈65)·811(다음 페이지) 미포함."""
    page = yahwe_doc[572]
    top, bottom = hp.find_song_y_range(page, "야훼 이레", 810)
    assert 390 < top < 400          # 810 제목 아래
    assert 740 < bottom < 748       # 810 저작권 줄 top
    assert top > 84                 # 809 제목(y1≈84) 아래에서 시작 → 809 혼입 없음


@pytest.mark.skipif(not _YOUTH_PDFS_AVAILABLE, reason=_YOUTH_SKIP_REASON)
def test_H4_naju_scan_no_text_layer(naju_doc):
    """H4: 나주노 스캔 페이지 get_text() 길이 0이어도 크래시 없이 이미지 경로 진행."""
    page = naju_doc[426]
    assert len(page.get_text("text").strip()) == 0  # 스캔 = 텍스트 레이어 없음(전제)
    # 렌더가 예외 없이 이미지를 반환
    img = hp.render_region(page, None, hp.RENDER_DPI)
    assert img.width > 0 and img.height > 0


# --- 시스템 밴드 / 슬라이드 수 (S1~S6) --------------------------------------
# 5곡의 시스템 밴드는 렌더가 느리므로(82MB PDF) 모듈 스코프에서 1회 계산해 공유한다.
_SONGS = {362: "나주노", 173: "나주노", 146: "나주노", 810: "야훼 이레", 267: "나주노"}
_EXPECTED_SYSTEMS = {362: 4, 173: 7, 146: 6, 810: 4, 267: 5}
_EXPECTED_SLIDES = {362: 2, 173: 4, 146: 3, 810: 2, 267: 3}


@pytest.fixture(scope="module")
def hymn_bands():
    import numpy as _np
    import missa_psalm_score_image as _ps
    result = {}
    for num, src in _SONGS.items():
        pi = hp.find_song_page(src, num)
        d = _fitz.open(hp.SOURCES[src]["pdf"])
        page = d[pi]
        yr = None if src == "나주노" else hp.find_song_y_range(page, src, num)
        img = hp.render_region(page, yr, hp.RENDER_DPI)
        gray = _np.asarray(img.convert("L"))
        bands = hp.detect_system_bands(gray)
        content_top = int(_np.where((gray < 128).any(axis=1))[0].min())
        result[num] = {"bands": bands, "content_top": content_top,
                       "shape": gray.shape, "gray": gray}
        d.close()
    return result


@pytest.fixture(scope="module")
def render_447():
    """447 전용 렌더(gray+bands) — 항목 5/6 검증용(447은 _SONGS에 없음)."""
    import numpy as _np
    pi = hp.find_song_page("나주노", 447)
    d = _fitz.open(hp.SOURCES["나주노"]["pdf"])
    page = d[pi]
    img = hp.render_region(page, None, hp.RENDER_DPI)
    gray = _np.asarray(img.convert("L"))
    bands = hp.detect_system_bands(gray)
    d.close()
    return {"gray": gray, "bands": bands}


@pytest.mark.parametrize("num", [362, 173, 146, 810, 267])
@pytest.mark.skipif(not _YOUTH_PDFS_AVAILABLE, reason=_YOUTH_SKIP_REASON)
def test_S1_S5_system_and_slide_counts(hymn_bands, num):
    """S1~S5: 각 곡의 시스템 수와 슬라이드 수(2개/장)가 요구사항 표와 일치."""
    bands = hymn_bands[num]["bands"]
    assert len(bands) == _EXPECTED_SYSTEMS[num]
    assert len(hp.group_systems(bands, 2)) == _EXPECTED_SLIDES[num]


@pytest.mark.parametrize("num", [362, 173, 146, 810, 267])
@pytest.mark.skipif(not _YOUTH_PDFS_AVAILABLE, reason=_YOUTH_SKIP_REASON)
def test_S6_bands_ordered_non_overlapping(hymn_bands, num):
    """S6: 밴드가 오름차순·비겹침이고, 상단 텍스트(제목/작곡가)가 첫 밴드에 포함되지 않는다.

    나주노는 페이지 전체를 렌더하므로 제목/작곡가가 콘텐츠 최상단에 있다 — 첫 시스템 밴드는
    반드시 그 아래에서 시작해야 한다(작곡가 병합 트림 검증). 크롭 PNG 육안 검증은 구현 노트
    참고(설계서 §3.4 실측 검증 원칙)."""
    info = hymn_bands[num]
    bands = info["bands"]
    for (t, b) in bands:
        assert b > t
    for (t0, b0), (t1, b1) in zip(bands, bands[1:]):
        assert t1 > b0  # 비겹침·오름차순
    if _SONGS[num] == "나주노":
        # 제목/작곡가가 콘텐츠 최상단에 있고, 첫 시스템은 그보다 아래에서 시작
        assert bands[0][0] > info["content_top"]


# --- 저작권 (C1~C4; C5는 슬라이드 조립 섹션) --------------------------------
@pytest.mark.parametrize("num", [362, 173, 146])
@pytest.mark.skipif(not _YOUTH_PDFS_AVAILABLE, reason=_YOUTH_SKIP_REASON)
def test_C1_naju_no_copyright(num):
    """C1: 나주노 362/173/146 → 저작권 크롭 None(이미지 요소 생성 안 함)."""
    pi = hp.find_song_page("나주노", num)
    d = _fitz.open(hp.SOURCES["나주노"]["pdf"])
    try:
        assert hp.resolve_copyright_crop("나주노", num, d[pi]) is None
    finally:
        d.close()


@pytest.mark.skipif(not _YOUTH_PDFS_AVAILABLE, reason=_YOUTH_SKIP_REASON)
def test_C2_naju_267_copyright_before_trigger():
    """C2: 나주노267 → 저작권 크롭 1개, 캐시 bbox가 'Adm. by' 트리거 전까지(가로 긴 한 줄)."""
    cache = hp._load_naju_cache()
    assert cache["267"]["has_copyright"] is True
    assert cache["267"]["trigger"] == "Adm. by"
    pi = hp.find_song_page("나주노", 267)
    d = _fitz.open(hp.SOURCES["나주노"]["pdf"])
    try:
        crop = hp.resolve_copyright_crop("나주노", 267, d[pi])
    finally:
        d.close()
    assert crop is not None
    assert crop.width > crop.height * 3  # 한 줄 텍스트(가로로 김)


@pytest.mark.skipif(not _YOUTH_PDFS_AVAILABLE, reason=_YOUTH_SKIP_REASON)
def test_C3_yahwe_810_right_edge_before_trigger(yahwe_doc):
    """C3: 야훼810 저작권 크롭 우측 경계가 'Administered by' 직전(그 이후 제외). 결정론적."""
    page = yahwe_doc[572]
    left_x, right_x, y0, y1 = hp.yahwe_copyright_bounds(page, 810)
    adm_x0 = min(h.x0 for h in page.search_for("Administered by"))
    assert left_x < right_x <= adm_x0        # 'Administered by' 시작 이전에서 끝남
    cw_x0 = min(h.x0 for h in page.search_for("Copyright"))
    assert abs(left_x - cw_x0) < 1           # 좌측은 'Copyright'에서 시작
    crop = hp.resolve_copyright_crop("야훼 이레", 810, page)
    assert crop is not None and crop.width > crop.height * 3


@pytest.mark.skipif(not _YOUTH_PDFS_AVAILABLE, reason=_YOUTH_SKIP_REASON)
def test_C4_naju_cache_miss_returns_none(caplog):
    """C4: 캐시에 없는 나주노 번호 → 경고 로그 + None(크래시/오크롭 없음)."""
    d = _fitz.open(hp.SOURCES["나주노"]["pdf"])
    try:
        import logging
        with caplog.at_level(logging.WARNING):
            crop = hp.resolve_copyright_crop("나주노", 999, d[0])
    finally:
        d.close()
    assert crop is None
    assert any("999" in r.message or "수동 확인" in r.message for r in caplog.records)


# 마침표(".")는 이 스캔·300dpi에서 폭 3~4px, 한글/라틴 글자는 ≥10px. 8px가 둘 사이 안전
# 경계다(글자로 끝나면 last-run 폭 > 8, 꼬리 마침표만 남으면 ≤ 8). test_C6 회귀 판정 기준.
_PERIOD_MAX_W_PX = 8


def _crop_col_ink_runs(pil_img):
    """크롭 이미지의 세로 잉크 열(어두운 픽셀 존재) 연속 구간 [(x0, x1, width), …]."""
    col = (np.asarray(pil_img.convert("L")) < 128).sum(axis=0)
    runs = []
    i, n = 0, len(col)
    while i < n:
        if col[i] > 0:
            s = i
            while i < n and col[i] > 0:
                i += 1
            runs.append((s, i - 1, i - s))
        else:
            i += 1
    return runs


@pytest.mark.parametrize("num", [447, 267])
@pytest.mark.skipif(not _YOUTH_PDFS_AVAILABLE, reason=_YOUTH_SKIP_REASON)
def test_C6_naju_copyright_excludes_trailing_period(num):
    """C6: 나주노 저작권 크롭이 이름 마지막 '글자'에서 끝난다(꼬리 마침표 제외).

    요구사항(§Copyright, 177~183행)은 '마침표 자체도 제외'를 명시한다. 이전 고정 4px back-off는
    마침표+공백 폭보다 작아 크롭 우측에 마침표가 남았고(리뷰 확정 이슈), 크기 비율만 보던
    기존 C2/P2는 이를 못 잡았다. 여기서는 런타임 경로(resolve_copyright_crop, 캐시 bbox·OCR
    없음)의 실제 크롭 우측 끝단 픽셀을 검사한다 — 마지막 세로 잉크 런이 글자 폭(>8px)이면
    글자로 끝난 것, 마침표 폭(≤8px)만 남으면 회귀."""
    d = _fitz.open(hp.SOURCES["나주노"]["pdf"])
    try:
        crop = hp.resolve_copyright_crop("나주노", num, d[hp.find_song_page("나주노", num)])
    finally:
        d.close()
    assert crop is not None
    runs = _crop_col_ink_runs(crop)
    assert len(runs) >= 4  # 'Copyright ©…이름' — 내용이 온전(과트림 방지)
    assert runs[-1][2] > _PERIOD_MAX_W_PX, (
        f"성가 {num}: 크롭이 글자가 아니라 마침표 폭({runs[-1][2]}px)으로 끝남 — 꼬리 마침표 잔존")


@pytest.mark.skipif(not _YOUTH_PDFS_AVAILABLE, reason=_YOUTH_SKIP_REASON)
def test_C6b_naju_267_keeps_internal_period():
    """C6b: 267은 이름 자체에 마침표가 있다('천태혁. 진경'). 꼬리 마침표만 제거하고 내부
    마침표는 보존돼야 한다 — 크롭 중간(마지막 런 이전)에 마침표 폭 런이 남아 있어야 하고,
    '진경'까지 온전히 포함돼(마지막 두 글자가 글자 폭) 과트림이 아니어야 한다."""
    d = _fitz.open(hp.SOURCES["나주노"]["pdf"])
    try:
        crop = hp.resolve_copyright_crop("나주노", 267, d[hp.find_song_page("나주노", 267)])
    finally:
        d.close()
    runs = _crop_col_ink_runs(crop)
    assert any(w <= _PERIOD_MAX_W_PX for _, _, w in runs[:-1]), "내부 마침표(천태혁.)가 사라짐"
    assert runs[-1][2] > _PERIOD_MAX_W_PX and runs[-2][2] > _PERIOD_MAX_W_PX, "'진경' 과트림"


@pytest.mark.skipif(not _YOUTH_PDFS_AVAILABLE, reason=_YOUTH_SKIP_REASON)
def test_C6c_447_matches_sample_no_trailing_period():
    """C6c: 447 런타임 크롭이 제공 샘플(SAMPLE447)과 같은 '마침표 없는 글자 끝' 속성을 가진다.

    샘플·재생성 크롭 모두 마지막 세로 잉크 런이 글자 폭(>8px)으로, 저작권자 이름 뒤 마침표가
    없다. 리뷰어가 지적한 '샘플과의 직접 대조' 갭을 이 속성 일치로 닫는다(해상도가 달라 픽셀
    완전 일치는 부적합)."""
    d = _fitz.open(hp.SOURCES["나주노"]["pdf"])
    try:
        crop = hp.resolve_copyright_crop("나주노", 447, d[hp.find_song_page("나주노", 447)])
    finally:
        d.close()
    sample = Image.open(
        Path(__file__).resolve().parent.parent / "_workspace" / "청년미사_1단계" / "crops"
        / "SAMPLE447_copyright.png")
    assert _crop_col_ink_runs(sample)[-1][2] > _PERIOD_MAX_W_PX   # 샘플도 글자로 끝남
    assert _crop_col_ink_runs(crop)[-1][2] > _PERIOD_MAX_W_PX     # 재생성도 동일 속성


# --- 헤더 / 산출물 구조 (Hd1~Hd4, C5, P1~P4) --------------------------------
from pptx.oxml.ns import qn as _qn


def _hdr_runs(slide):
    for sh in slide.shapes:
        if sh.name == "Rectangle 11":
            return sh.text_frame.paragraphs[0]._p.findall(_qn("a:r"))
    return []


def _run_text(r):
    t = r.find(_qn("a:t"))
    return t.text if t is not None else ""


def _run_color(r):
    rPr = r.find(_qn("a:rPr"))
    if rPr is None:
        return None
    fill = rPr.find(_qn("a:solidFill"))
    if fill is None:
        return None
    srgb = fill.find(_qn("a:srgbClr"))
    if srgb is not None:
        return ("srgb", srgb.get("val"))
    prst = fill.find(_qn("a:prstClr"))
    if prst is not None:
        return ("prst", prst.get("val"))
    return None


def _score_pics(slide):
    return [s for s in slide.shapes if s.shape_type == 13 and s.top < 4800000]


def _copyright_pics(slide):
    return [s for s in slide.shapes if s.shape_type == 13 and s.top >= 4800000]


@pytest.fixture(scope="module")
def built():
    return {
        362: hp.build_hymn_pptx("입당", "나주노", 362),
        267: hp.build_hymn_pptx("파견", "나주노", 267),
        447: hp.build_hymn_pptx("입당", "나주노", 447),
        810: hp.build_hymn_pptx("2차봉헌", "야훼 이레", 810),
    }


@pytest.mark.skipif(not _YOUTH_PDFS_AVAILABLE, reason=_YOUTH_SKIP_REASON)
def test_Hd1_header_on_every_slide(built):
    """Hd1: 'Rectangle 11' 헤더가 모든 슬라이드에 반복."""
    for num, prs in built.items():
        for slide in prs.slides:
            assert any(sh.name == "Rectangle 11" for sh in slide.shapes), num


@pytest.mark.skipif(not _YOUTH_PDFS_AVAILABLE, reason=_YOUTH_SKIP_REASON)
def test_Hd2_run_order_and_colors(built):
    """Hd2: run 순서·색 — 구분=흰색(prstClr white), 출처·번호·제목=FFC000."""
    runs = _hdr_runs(built[362].slides[0])
    texts = [_run_text(r) for r in runs]
    assert texts[0] == "입당 "
    assert _run_color(runs[0]) == ("prst", "white")
    # 공백 아닌 세그먼트: 출처/번호/제목
    seg = [(_run_text(r), _run_color(r)) for r in runs[1:] if _run_text(r).strip()]
    assert seg[0] == ("나주노", ("srgb", "FFC000"))
    assert seg[1] == ("362", ("srgb", "FFC000"))
    assert seg[2] == ("삼위일체", ("srgb", "FFC000"))


@pytest.mark.skipif(not _YOUTH_PDFS_AVAILABLE, reason=_YOUTH_SKIP_REASON)
def test_Hd3_font_size_and_spacing_preserved(built):
    """Hd3: 헤더 모든 run이 28pt(sz=2800)·자간 spc=300 유지(템플릿 서식 보존)."""
    runs = _hdr_runs(built[362].slides[0])
    for r in runs:
        rPr = r.find(_qn("a:rPr"))
        assert rPr.get("sz") == "2800"
        assert rPr.get("spc") == "300"


@pytest.mark.skipif(not _YOUTH_PDFS_AVAILABLE, reason=_YOUTH_SKIP_REASON)
def test_Hd4_unknown_구분_raises():
    """Hd4: 구분이 HYMN_TYPES에 없으면 에러."""
    with pytest.raises(ValueError):
        hp.build_hymn_pptx("존재하지않는구분", "나주노", 362)


@pytest.mark.skipif(not _YOUTH_PDFS_AVAILABLE, reason=_YOUTH_SKIP_REASON)
def test_C1_C5_copyright_presence_and_position(built):
    """C1/C5: 362=저작권 0개; 267/810/447=저작권 1개씩, 우하단 샘플 계열 위치·높이."""
    for slide in built[362].slides:
        assert len(_copyright_pics(slide)) == 0
    for num in (267, 810, 447):
        for slide in built[num].slides:
            pics = _copyright_pics(slide)
            assert len(pics) == 1, num
            p = pics[0]
            assert abs(p.top - 4855468) < 30000            # 샘플 top
            assert abs(p.height - 288032) < 30000          # 샘플 높이 계열
            assert abs((p.left + p.width) - 8510525) < 30000  # 샘플 우측 끝


@pytest.mark.skipif(not _YOUTH_PDFS_AVAILABLE, reason=_YOUTH_SKIP_REASON)
def test_P1_slide_size(built):
    """P1: 슬라이드 크기 9144000×5143500."""
    for prs in built.values():
        assert prs.slide_width == 9144000
        assert prs.slide_height == 5143500


@pytest.mark.skipif(not _YOUTH_PDFS_AVAILABLE, reason=_YOUTH_SKIP_REASON)
def test_P2_447_matches_sample(built):
    """P2: 나주노447 재생성 → 4장, 헤더 텍스트·색, 저작권 유무 샘플과 일치."""
    prs = built[447]
    assert len(prs.slides) == 4
    runs = _hdr_runs(prs.slides[0])
    seg = [(_run_text(r), _run_color(r)) for r in runs[1:] if _run_text(r).strip()]
    assert _run_text(runs[0]) == "입당 " and _run_color(runs[0]) == ("prst", "white")
    assert seg[0] == ("나주노", ("srgb", "FFC000"))
    assert seg[1] == ("447", ("srgb", "FFC000"))
    assert seg[2] == ("영원도 하시어라 그 사랑이여", ("srgb", "FFC000"))
    for slide in prs.slides:
        assert len(_copyright_pics(slide)) == 1  # 447은 저작권 있음


@pytest.mark.skipif(not _YOUTH_PDFS_AVAILABLE, reason=_YOUTH_SKIP_REASON)
def test_P3_validate_structure(built, tmp_path):
    """P3: validate_pptx_structure 통과(구조 손상 없음). PowerPoint 육안은 구현 노트."""
    from missa_sections import validate_pptx_structure
    for num, prs in built.items():
        f = tmp_path / f"out_{num}.pptx"
        prs.save(str(f))
        assert validate_pptx_structure(str(f)) == []


@pytest.mark.skipif(not _YOUTH_PDFS_AVAILABLE, reason=_YOUTH_SKIP_REASON)
def test_P4_one_score_per_slide_below_header(built):
    """P4: 슬라이드별 악보 크롭 1장, 헤더(523220) 아래 배치."""
    for num, prs in built.items():
        for slide in prs.slides:
            score = _score_pics(slide)
            assert len(score) == 1, num
            assert score[0].top >= 523220


@pytest.mark.skipif(not _YOUTH_PDFS_AVAILABLE, reason=_YOUTH_SKIP_REASON)
def test_Hd_header_child_order_valid(built):
    """헤더 <a:p> 자식 순서 검증 — 모든 run이 endParaRPr **앞**에 온다.

    validate_pptx_structure()는 요소 순서를 검사하지 않으므로(CLAUDE.md 2026-08-29,
    ooxml-pitfalls) run/endParaRPr 순서를 lxml로 직접 단언한다. build_header_runs가
    _insert_run_before_end_para_rpr로 삽입하지 않으면 run이 endParaRPr 뒤로 가 스키마 위반이
    되지만 검사기·다른 테스트를 통과해버린다."""
    for num, prs in built.items():
        for slide in prs.slides:
            for sh in slide.shapes:
                if sh.name != "Rectangle 11":
                    continue
                p = sh.text_frame.paragraphs[0]._p
                children = [c.tag for c in p]
                eprp_tag = _qn("a:endParaRPr")
                r_tag = _qn("a:r")
                if eprp_tag in children:
                    eprp_i = children.index(eprp_tag)
                    last_r = max((i for i, t in enumerate(children) if t == r_tag), default=-1)
                    assert last_r < eprp_i, (num, "run이 endParaRPr 뒤에 옴")


# ===========================================================================
# 청년미사 1단계 후속(v2) — 설계서 01b_architect_design_v2.md
# 항목 6(P5~P7) per-group 크롭 · 항목 5(S7~S11) 동적 패킹 · 항목 3(Hd8) 야훼이레 표시명
# 항목 4(C7 폴백 마침표 · C8 높이0.9 · C9 447 샘플 종횡비) · 항목 2(Hf1~Hf4 헤더 축소)
# 항목 1(M1~M9 다중 섹션 파싱)
# ===========================================================================


# --- 항목 6: per-group 크롭 bbox (P5~P7) -----------------------------------
def _tail_has_ink(gray, box, frac=0.98):
    """크롭 박스의 우측 끝(frac~1.0 열 범위)에 잉크가 있으면 True(우측 흰 여백 없음)."""
    left, top, right, bottom = box
    sub = (gray[top:bottom, left:right] < hp._INK)
    w = sub.shape[1]
    return bool(sub[:, int(w * frac):].any())


# (P5~P7은 test_missa_regression.py로 승격됨 — 2026-09-11 QA.)


# --- 항목 5: 동적 패킹 (S7~S11) --------------------------------------------
def _sizes(groups):
    return [len(g) for g in groups]


def _synthetic_gray(bands, width):
    """밴드 tuple 목록에 맞춰 전 폭에 잉크가 찬 합성 gray(그 외 흰색)를 만든다."""
    h = max(b for _, b in bands) + 5
    g = np.full((h, width), 255, np.uint8)
    for t, b in bands:
        g[t:b + 1, :] = 0
    return g


_PACK_EXPECTED = {362: [2, 2], 173: [2, 2, 2, 1], 146: [2, 2, 2],
                  810: [2, 2], 267: [2, 2, 1]}


@pytest.mark.parametrize("num", [362, 173, 146, 810, 267])
@pytest.mark.skipif(not _YOUTH_PDFS_AVAILABLE, reason=_YOUTH_SKIP_REASON)
def test_S7_pack_systems_no_regression_5songs(hymn_bands, num):
    """S7: min-slide 균형 패커가 5곡 각각 현재와 동일한 그룹 크기를 반환(무회귀)."""
    groups = hp.pack_systems(hymn_bands[num]["bands"], hymn_bands[num]["gray"])
    assert _sizes(groups) == _PACK_EXPECTED[num]


@pytest.mark.skipif(not _YOUTH_PDFS_AVAILABLE, reason=_YOUTH_SKIP_REASON)
def test_S7_pack_systems_447(render_447):
    """S7: 447도 현재와 동일 [2,2,2,1](무회귀)."""
    groups = hp.pack_systems(render_447["bands"], render_447["gray"])
    assert _sizes(groups) == [2, 2, 2, 1]


def test_S8_three_pack_reduces_slides():
    """S8: 짧은 시스템 5개, 첫 3개가 종횡비·safety 통과 → [3,2] 2장(현재 규칙이면 [2,2,1] 3장).
    3-pack이 슬라이드를 줄인다. (합성 데이터라 PDF 불필요.)"""
    bands = [(0, 39), (60, 99), (120, 159), (180, 219), (240, 279)]
    gray = _synthetic_gray(bands, width=2000)
    assert _sizes(hp.pack_systems(bands, gray)) == [3, 2]


def test_S9_tight_three_pack_rejected_by_safety():
    """S9: 3개가 종횡비는 통과하나 vfill이 safety 초과(810형 97%) → 3-pack 탈락, [2,2] 유지."""
    # 3-band span/width = 449/1000 = 0.449 > band_aspect*safety(0.463*0.95=0.440),
    # 하지만 ≤ band_aspect(0.463) — 즉 겹치진 않으나 너무 빡빡. 2-band=289/1000=0.289 통과.
    bands = [(0, 129), (160, 289), (320, 449), (480, 609)]
    gray = _synthetic_gray(bands, width=1000)
    assert _sizes(hp.pack_systems(bands, gray)) == [2, 2]


def test_S10_balance_prefers_2_2_over_3_1():
    """S10: 3-pack이 기하적으로 들어가도 슬라이드 수가 같으면 [3,1]이 아니라 균형 [2,2]."""
    bands = [(0, 39), (60, 99), (120, 159), (180, 219)]  # 4개 다 짧음(3-pack 여유 통과)
    gray = _synthetic_gray(bands, width=2000)
    assert _sizes(hp.pack_systems(bands, gray)) == [2, 2]


# S11(3-pack 발동 슬라이드 PNG 겹침 없음)은 육안 검증 — 구현 노트 참고.


# --- 항목 3: 야훼이레 표시명 (Hd8) -----------------------------------------
@pytest.mark.skipif(not _YOUTH_PDFS_AVAILABLE, reason=_YOUTH_SKIP_REASON)
def test_Hd8_yahwe_display_label_no_space(built):
    """Hd8: 야훼 810 헤더 출처 run == '야훼이레'(공백 없음). SOURCES 키/조회는 '야훼 이레' 유지."""
    runs = _hdr_runs(built[810].slides[0])
    seg = [(_run_text(r), _run_color(r)) for r in runs[1:] if _run_text(r).strip()]
    assert seg[0] == ("야훼이레", ("srgb", "FFC000"))
    assert hp.find_song_page("야훼 이레", 810) == 572  # 내부 조회는 원본 키 유지


def test_Hd8_source_display_mapping():
    """Hd8: SOURCE_DISPLAY는 표시명만 매핑, SOURCES 키는 원본 유지."""
    assert hp.SOURCE_DISPLAY.get("야훼 이레") == "야훼이레"
    assert "야훼 이레" in hp.SOURCES and "야훼이레" not in hp.SOURCES


# --- 항목 4-1: 폴백 꼬리 마침표 트림 일반화 (C7) ---------------------------
def test_C7_trailing_period_trim_generalized():
    """C7: _trim_trailing_period_px가 트리거 유무와 무관하게 우측 꼬리 마침표를 제거하고,
    이름 내부 마침표(267식)는 보존한다. (합성 이미지라 PDF 불필요.)"""
    from PIL import Image as _Im
    # 글자(20px) + 공백 + 꼬리 마침표(4px)
    a = np.full((30, 60), 255, np.uint8)
    a[10:20, 5:25] = 0
    a[16:20, 40:44] = 0
    out = hp._trim_trailing_period_px(_Im.fromarray(a))
    runs = _crop_col_ink_runs(out)
    assert runs[-1][2] > _PERIOD_MAX_W_PX  # 마침표 제거됨(마지막 런=글자)
    # 글자 + 내부 마침표 + 글자 → 내부 마침표 보존, 트림 없음
    b = np.full((30, 80), 255, np.uint8)
    b[10:20, 5:25] = 0
    b[16:20, 30:34] = 0
    b[10:20, 45:65] = 0
    out2 = hp._trim_trailing_period_px(_Im.fromarray(b))
    runs2 = _crop_col_ink_runs(out2)
    assert any(w <= _PERIOD_MAX_W_PX for _, _, w in runs2[:-1])  # 내부 마침표 보존
    assert runs2[-1][2] > _PERIOD_MAX_W_PX


# --- 항목 4-2: 저작권 높이 0.9배 (C8) --------------------------------------
@pytest.mark.skipif(not _YOUTH_PDFS_AVAILABLE, reason=_YOUTH_SKIP_REASON)
def test_C8_copyright_height_fixed_09(built):
    """C8: 저작권 이미지 높이 == round(288032*0.9)=259229(전곡), 폭 == round(높이×크롭종횡비),
    우측끝 8510525·바닥 5143500 앵커 유지."""
    assert hp._COPY_H == 259229
    for num in (267, 810, 447):
        for slide in built[num].slides:
            p = _copyright_pics(slide)[0]
            assert p.height == 259229
            assert abs((p.left + p.width) - 8510525) < 3        # 우측끝 앵커
            assert abs((p.top + p.height) - 5143500) < 3        # 바닥 앵커


# --- 항목 4-3: 447 캐시 크롭이 샘플 종횡비에 근접 (C9) ----------------------
@pytest.mark.skipif(not _YOUTH_PDFS_AVAILABLE, reason=_YOUTH_SKIP_REASON)
def test_C9_447_crop_aspect_matches_sample():
    """C9: 447 저작권 크롭 종횡비가 손제작 샘플(≈5.17)에 근접(±0.5). 재조정 전 6.03은 red.

    샘플보다 과도하게 넓던 447 크롭을 샘플 비율로 재조정(항목 4-3). 크롭은 여전히 마침표 없이
    글자로 끝나야 한다(C6c 속성 유지)."""
    d = _fitz.open(hp.SOURCES["나주노"]["pdf"])
    try:
        crop = hp.resolve_copyright_crop("나주노", 447, d[hp.find_song_page("나주노", 447)])
    finally:
        d.close()
    aspect = crop.width / crop.height
    sample = Image.open(
        Path(__file__).resolve().parent.parent / "_workspace" / "청년미사_1단계" / "crops"
        / "SAMPLE447_copyright.png")
    s_aspect = sample.width / sample.height
    assert abs(aspect - s_aspect) < 0.5, (aspect, s_aspect)
    assert _crop_col_ink_runs(crop)[-1][2] > _PERIOD_MAX_W_PX  # 여전히 글자로 끝남


# --- 항목 2: 헤더 자동 축소 (Hf1~Hf4) --------------------------------------
def _template_header_runs(구분, 출처, number, title):
    prs = hp.Presentation(str(hp.TEMPLATE))
    hp.build_header_runs(prs.slides[0], 구분, 출처, number, title)
    return _hdr_runs(prs.slides[0])


def _run_sz(r):
    rPr = r.find(_qn("a:rPr"))
    return rPr.get("sz") if rPr is not None else None


def _run_spc(r):
    rPr = r.find(_qn("a:rPr"))
    return rPr.get("spc") if rPr is not None else None


@pytest.mark.skipif(not _YOUTH_PDFS_AVAILABLE, reason=_YOUTH_SKIP_REASON)
def test_Hf1_short_title_no_shrink(built):
    """Hf1: 짧은 제목(362 '삼위일체')은 spc=300·28pt 그대로(축소 없음)."""
    for r in _hdr_runs(built[362].slides[0]):
        assert _run_spc(r) == "300"
        assert _run_sz(r) == "2800"


def test_Hf2_overflow_reduces_spacing_only():
    """Hf2: spc=300에서 넘치나 spc=0이면 들어가는 제목 → 전 run spc→0, 폰트는 28pt 유지."""
    runs = _template_header_runs("입당", "나주노", 447, "가" * 15)
    for r in runs:
        assert _run_spc(r) == "0"
        assert _run_sz(r) == "2800"


def test_Hf3_still_overflow_reduces_title_font_only():
    """Hf3: spc=0으로도 넘치는 제목 → 제목 run sz만 축소(<2800), 구분/출처/번호는 2800 유지."""
    runs = _template_header_runs("입당", "나주노", 447, "가" * 20)
    title_r = runs[-1]
    assert int(_run_sz(title_r)) < 2800
    assert int(_run_sz(title_r)) >= hp._MIN_TITLE_PT * 100
    # 제목 아닌 run(구분/출처/번호/공백)은 2800 유지
    for r in runs[:-1]:
        assert _run_sz(r) == "2800"


def test_Hf4_after_shrink_fits_usable():
    """Hf4: 축소 후 최종 측정폭 ≤ usable(안 넘침)."""
    prs = hp.Presentation(str(hp.TEMPLATE))
    hp.build_header_runs(prs.slides[0], "입당", "나주노", 447, "가" * 20)
    hdr = hp._find_header(prs.slides[0])
    runs = _hdr_runs(prs.slides[0])
    total = hp._header_total_emu(runs)
    assert total <= hp._usable_header_emu(hdr)


# --- 항목 1: 다중 섹션 파싱 (M1~M9) ----------------------------------------
def test_M1_parse_mass_all_sections():
    """M1: parse_mass가 5개 섹션 키 반환, Gospel reference 정확."""
    m = yg.parse_mass(_fix_html("20260913"))
    for k in ("First reading", "Responsorial Psalm", "Second reading",
              "Gospel Acclamation", "Gospel"):
        assert k in m, k
    assert m["Gospel"]["reference"] == "Matthew 18:21-35"


def test_M2_first_reading():
    m = yg.parse_mass(_fix_html("20260913"))
    assert m["First reading"]["reference"] == "Ecclesiasticus 27:33-28:9"
    assert m["First reading"]["content"]


def test_M3_psalm_reference_only():
    """M3: Responsorial Psalm은 reference만(content/theme 키 없음)."""
    psalm = yg.parse_mass(_fix_html("20260913"))["Responsorial Psalm"]
    assert psalm["reference"] == "Psalm 102(103):1-4,9-12"
    assert "content" not in psalm and "theme" not in psalm


def test_M4_acclamation_content_excludes_or_body():
    """M4: Acclamation content가 'Alleluia'로 시작, Or: 본문이 섞이지 않음(다음 섹션 종료 경계)."""
    accl = yg.parse_mass(_fix_html("20260913"))["Gospel Acclamation"]
    assert accl["content"].startswith("Alleluia")
    assert "I give you a new commandment" not in accl["content"]
    assert accl["content"].count("\n\n") >= 3


def test_M5_acclamation_or_nested():
    """M5: Gospel Acclamation의 'or' 중첩 = {reference 'Jn13:34', content '…love one another…'}."""
    accl = yg.parse_mass(_fix_html("20260913"))["Gospel Acclamation"]
    assert accl["or"]["reference"] == "Jn13:34"
    assert "love one another" in accl["or"]["content"]


def test_M6_no_korean_keys():
    """M6: 모든 키가 ASCII(한글 '복음' 부재)."""
    m = yg.parse_mass(_fix_html("20260913"))
    def _all_keys(d):
        for k, v in d.items():
            yield k
            if isinstance(v, dict):
                yield from _all_keys(v)
    for k in _all_keys(m):
        assert k.isascii(), k


def test_M7_theme_presence_by_section():
    """M7: First/Second/Gospel엔 theme, Psalm/Acclamation엔 theme 키 없음."""
    m = yg.parse_mass(_fix_html("20260913"))
    assert m["First reading"]["theme"] == "Forgive your neighbour the hurt he does you; and when you pray, your sins will be forgiven"
    assert m["Second reading"]["theme"] == "Alive or dead, we belong to the Lord"
    assert m["Gospel"]["theme"] == "To be forgiven, you must forgive"
    assert "theme" not in m["Responsorial Psalm"]
    assert "theme" not in m["Gospel Acclamation"]


def test_M8_main_default_filename(monkeypatch, tmp_path):
    """M8: main 인자 미지정 시 missa_en_YYYYMMDD.json으로 저장, 영어 섹션 포함."""
    monkeypatch.setattr(yg, "fetch_gospel_html", lambda d: _fix_html("20260913"))
    monkeypatch.chdir(tmp_path)
    rc = yg.main(["20260913"])
    assert rc == 0
    out = tmp_path / "missa_en_20260913.json"
    assert out.exists()
    import json as _json
    data = _json.loads(out.read_text(encoding="utf-8"))
    assert data["date"] == "20260913"
    assert data["Gospel"]["reference"] == "Matthew 18:21-35"
    assert "복음" not in data


def test_M9_get_youth_mass_schema(monkeypatch):
    """M9: get_youth_mass가 date + 영어 섹션 반환, Psalm 빈 본문은 오류 아님."""
    monkeypatch.setattr(yg, "fetch_gospel_html", lambda d: _fix_html("20260913"))
    out = yg.get_youth_mass("20260913")
    assert out["date"] == "20260913"
    assert out["First reading"]["reference"] == "Ecclesiasticus 27:33-28:9"
    assert "복음" not in out


# ═══════════════════════════════════════════════════════════════════════════
# 2026-09-26 실사용자 산출물 후속 요청 5건 (청년미사, output/20260926_youth/토요일 저녁
# 청년 주일미사_20260926_연중 제26주일.pptx 육안 검수로 발견) — 코디네이터 지시 직접 반영,
# 별도 설계서 없음. 항목:
#   E1. 결과 PPT를 성당 공용 OneDrive(PPT 문서/20.청년 미사/{연도})에 저장 직후 업로드(신규)
#   E2. 제목 슬라이드 날짜가 콘텐츠 조회일(+1일, 일요일)로 나옴 → 사용자가 입력한 토요일
#       date_str을 표시해야 함
#   E3. (철회, 2026-09-26) 슬라이드 17의 "▶" 단독 장식 슬라이드 삭제를 요청했으나, 직접
#       조사 결과 참조 템플릿에 원래부터 있던 것으로 확인되자 사용자가 "제거할 필요 없다"고
#       판단해 요청 자체를 취소했다 — 구현·테스트 모두 원상복구(missa_sections.py의
#       _remove_youth_arrow_divider_slides()와 그 호출부 제거).
#   E4. 슬라이드 50: 복음환호송 청년 템플릿의 ○ 구절 run 구조(라벨/탭/본문이 별도 run)를
#       성인 템플릿(라벨+탭 결합 run) 기준 인덱스로 잘못 가정해 탭 run에 본문을 덮어써
#       본문이 라벨 폰트(굴림류)로 렌더링되고 탭 정렬이 깨짐
#   E5. 슬라이드 54/55: D2가 "박스를 중앙에 배치"할 요구사항을 "문단을 가운데 정렬"로
#       잘못 구현 — 문단은 왼쪽 정렬로 되돌리고 박스 위치(normalize_page_size)는 유지
#
# 안정화되면 regression-qa가 test_missa_regression.py로 승격한다.
# ═══════════════════════════════════════════════════════════════════════════
import shutil as _e_shutil
import sys as _e_sys
import json as _e_json

_E_BASE = Path(__file__).resolve().parent.parent  # tests/ -> 저장소 루트
_E_YOUTH_TEMPLATE = (
    _E_BASE / "reference" / "청년미사"
    / "Template_토요일 저녁 청년 주일미사_20260822_연중 제21주일.pptx"
)
_E_EN_GOSPEL_FIXTURE = (
    _E_BASE / "_workspace" / "청년미사_1단계" / "verify_20260913" / "missa_en_20260913.json"
)
_E_KO_JSON_20260913 = _E_BASE / "output" / "20260913" / "missa_20260913.json"


def _e_skip_if_missing():
    for p in (_E_YOUTH_TEMPLATE, _E_EN_GOSPEL_FIXTURE, _E_KO_JSON_20260913):
        if not p.is_file():
            pytest.skip(f"청년미사 실측 픽스처가 로컬에 없음: {p}")


@pytest.fixture
def e_y2_prs():
    """C4 등 이전 라운드가 쓰던 것과 동일한 실측 참조 PPT(단독 로드, main() 미경유)."""
    _e_skip_if_missing()
    from pptx import Presentation as _EPresentation
    return _EPresentation(str(_E_YOUTH_TEMPLATE))


# E1(청년 결과 PPT 자동 OneDrive 업로드)는 2026-10-03 폐지 — 결과창 '원드라이브로 복사' 버튼으로 대체
# (tests/test_missa_youth_local_mode.py의 복사 테스트 참고).


# ---------------------------------------------------------------------------
# E2: 제목 슬라이드 날짜 — 콘텐츠 조회일이 아니라 사용자가 입력한 날짜를 보여줘야 함
# ---------------------------------------------------------------------------

def test_e2_title_slide_shows_override_date_not_json_date():
    import missa_content_updaters as cu
    from pptx.util import Emu
    from pptx import Presentation as _EPresentation2

    # 최소 제목 슬라이드 프레젠테이션(레이아웃 6=Blank)을 즉석에서 만든다.
    prs = _EPresentation2()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Emu(0), Emu(0), Emu(1000000), Emu(1000000))
    tf = box.text_frame
    tf.text = '2026년 9월 27일'
    tf.add_paragraph().text = '연중 제26주일'

    json_data = {'date': '2026-09-27', 'liturgy': '연중 제26주일'}  # +1일 콘텐츠 조회일(일요일)

    cu.update_title_slide(prs, json_data, date_str='2026-09-26')  # 사용자가 입력한 토요일

    assert prs.slides[0].shapes[0].text_frame.paragraphs[0].text == '2026년 9월 26일'


def test_e2_title_slide_backward_compatible_without_override():
    """date_str 인자를 안 넘기면 기존처럼 json_data['date']를 그대로 쓴다(성인 경로 무변경)."""
    import missa_content_updaters as cu
    from pptx.util import Emu
    from pptx import Presentation as _EPresentation3

    prs = _EPresentation3()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Emu(0), Emu(0), Emu(1000000), Emu(1000000))
    tf = box.text_frame
    tf.text = '2026년 6월 28일'
    tf.add_paragraph().text = '연중 제13주일'

    json_data = {'date': '2026-06-28', 'liturgy': '연중 제13주일'}
    cu.update_title_slide(prs, json_data)

    assert prs.slides[0].shapes[0].text_frame.paragraphs[0].text == '2026년 6월 28일'


# ---------------------------------------------------------------------------
# E4: 복음환호송 청년 ○ 구절 — 라벨/탭/본문 run 경계를 성인 템플릿 인덱스로 잘못 가정
# ---------------------------------------------------------------------------

def test_e4_복음환호송_youth_verse_keeps_tab_char_and_body_font(e_y2_prs):
    import missa_content_updaters as cu
    import missa_sections as sec_e
    from pptx.oxml.ns import qn as _e4_qn

    prs = e_y2_prs
    sections = sec_e.find_sections(prs, mass_type='youth')
    idx = sections['복음환호송']
    slide = prs.slides[idx]

    verse_text = '주님이 말씀하신다. 내 양들은 내 목소리를 알아듣는다. 나는 그들을 알고 그들은 나를 따른다.'
    json_data = {'복음환호송': {'content': f'◎ 알렐루야.\n○ {verse_text}\n◎ 알렐루야.'}}

    cu.update_복음환호송(prs, json_data, sections, mass_type='youth')

    target = None
    for shape in slide.shapes:
        if shape.has_text_frame and shape.text_frame.text.strip() and verse_text[:10] in shape.text_frame.text:
            target = shape
            break
    assert target is not None, [s.text_frame.text[:40] for s in slide.shapes if s.has_text_frame]

    para = target.text_frame.paragraphs[0]
    runs = para._pPr.getparent().findall(_e4_qn('a:r'))
    run_texts = [(r.find(_e4_qn('a:t')).text or '') for r in runs]

    # 탭 문자를 담은 run이 그대로 남아 있어야 한다(정렬용 tabLst와 짝을 이루는 실제 탭).
    assert '\t' in run_texts, run_texts

    tab_idx = run_texts.index('\t')
    body_run = runs[tab_idx + 1]
    # 폰트 확인은 rPr/a:latin의 typeface 속성으로 직접 조회한다.
    rpr = body_run.find(_e4_qn('a:rPr'))
    latin = rpr.find(_e4_qn('a:latin')) if rpr is not None else None
    assert latin is not None and latin.get('typeface') == '바탕', (
        f"본문 run 폰트가 템플릿(바탕체)과 다름: {latin.get('typeface') if latin is not None else None}"
    )
    assert verse_text in run_texts[tab_idx + 1]


# ---------------------------------------------------------------------------
# E5 + 종합: main() 실경로(youth, 복음 영문 슬라이드 포함) 1회 실행으로 E2/E3/E4/E5를
# 실제 산출물 기준으로 재확인한다. 공유 비용이 큰 main() 호출을 모듈 스코프로 1회만 수행.
# ---------------------------------------------------------------------------

@pytest.fixture(scope="module")
def e_full_output(tmp_path_factory):
    _e_skip_if_missing()
    import missa_to_ppt as mtp

    tmp_path = tmp_path_factory.mktemp("e_full_output")
    date_str = '20260912'
    content_date_str = '20260913'

    (tmp_path / date_str).mkdir()
    (tmp_path / content_date_str).mkdir()
    ref_dst = tmp_path / date_str / 'reference.pptx'
    _e_shutil.copy(str(_E_YOUTH_TEMPLATE), str(ref_dst))
    _e_shutil.copy(str(_E_KO_JSON_20260913), str(tmp_path / content_date_str / f'missa_{content_date_str}.json'))

    orig_output_root = mtp.OUTPUT_ROOT
    mtp.OUTPUT_ROOT = str(tmp_path)

    en_fixture = _e_json.loads(_E_EN_GOSPEL_FIXTURE.read_text(encoding='utf-8'))
    orig_get_youth_mass = mtp.missa_youth_gospel.get_youth_mass
    mtp.missa_youth_gospel.get_youth_mass = lambda d: en_fixture

    upload_calls = []  # missa_onedrive 폐지 — 업로드 경로 자체가 없다(e1은 빈 목록을 확인).

    files = {
        'ref_pptx': ref_dst, '시작기도': None, '화답송_pptx': None, '화답송_img': None,
        '미사후기도': None, '공지사항': None, '성가': {},
    }
    성가_선택 = {
        '입당':   [{'출처': '기타', '번호': None, '제목': '테스트 입당곡'}],
        '봉헌':   [{'출처': '기타', '번호': None, '제목': '테스트 봉헌곡'}],
        '성체':   [{'출처': '기타', '번호': None, '제목': '테스트 성체곡'}],
        '2차봉헌': [{'출처': '기타', '번호': None, '제목': '테스트 2차봉헌곡'}],
        '파견':   [{'출처': '기타', '번호': None, '제목': '테스트 파견곡'}],
    }
    mtp._preloaded_inputs[0] = {
        'mass_type': 'youth', 'date_str': date_str, 'files': files,
        'hymn_numbers': {}, '성가_선택': 성가_선택,
    }
    orig_argv = _e_sys.argv
    _e_sys.argv = ['missa_to_ppt.py']

    try:
        mtp.main()
        expected_liturgy = _e_json.loads(_E_KO_JSON_20260913.read_text(encoding='utf-8'))['liturgy']
        out_path = tmp_path / f'{date_str}_youth' / f"토요일 저녁 청년 주일미사_{date_str}_{expected_liturgy}.pptx"
        prs = mtp.Presentation(str(out_path))
        yield {'prs': prs, 'out_path': out_path, 'date_str': date_str, 'upload_calls': upload_calls}
    finally:
        _e_sys.argv = orig_argv
        mtp._preloaded_inputs[0] = None
        mtp.OUTPUT_ROOT = orig_output_root
        mtp.missa_youth_gospel.get_youth_mass = orig_get_youth_mass


def test_e1_full_pipeline_does_not_upload_result_to_onedrive(e_full_output):
    """2026-10-03: 청년 결과 PPT는 output/에만 생성한다 — 파이프라인이 OneDrive 업로드를 하면 안 된다
    (운영자가 검증 후 결과창 '원드라이브로 복사' 또는 수동으로 올린다)."""
    result_uploads = [c for c in e_full_output['upload_calls']
                      if c[0] == 'upload_file' and str(c[1]).endswith(e_full_output['out_path'].name)]
    assert result_uploads == []
    assert e_full_output['out_path'].is_file()


def _e_gospel_ending_slide(prs):
    """복음 구간에서 '주님의 말씀입니다' 종료 문구가 있는 슬라이드 (인덱스, 슬라이드)."""
    import missa_sections as sec_e
    sections = sec_e.find_sections(prs, mass_type='youth')
    for i in range(sections['복음_start'], sections['복음_end']):
        sl = prs.slides[i]
        if any(sh.has_text_frame and '주님의 말씀입니다' in sh.text_frame.text for sh in sl.shapes):
            return i, sl
    return None, None


def test_e6_full_pipeline_youth_english_gospel_ending_is_separate_slide(e_full_output):
    """2026-10-05 사용자 결정: 청년 영문 복음은 마지막 본문 슬라이드가 몇 줄이든 '주님의 말씀입니다/그리스도님, 찬미합니다'가
    그 다음 별도 슬라이드(템플릿 57쪽과 동일)에 있어야 한다 — 이 픽스처(20260913 복음)는 마지막 슬라이드가 4줄이라
    예전에는 본문에 병합됐다."""
    from pptx import Presentation as _P
    idx, ending = _e_gospel_ending_slide(e_full_output['prs'])
    assert ending is not None, '복음 종료 문구 슬라이드가 없음'
    # 종료 슬라이드에는 영문 본문이 없고(본문 상자는 비어 있음), 직전 슬라이드에는 종료 문구가 없다.
    body_texts = [sh.text_frame.text.strip() for sh in ending.shapes
                  if sh.has_text_frame and '주님의 말씀입니다' not in sh.text_frame.text
                  and '전례문 ©' not in sh.text_frame.text]
    assert all(t == '' for t in body_texts), body_texts
    prev = e_full_output['prs'].slides[idx - 1]
    assert not any(sh.has_text_frame and '주님의 말씀입니다' in sh.text_frame.text for sh in prev.shapes)
    # 위치는 템플릿의 복음 종료 슬라이드(57쪽)와 같다.
    t_prs = _P(str(_E_YOUTH_TEMPLATE))
    _t_idx, t_end = _e_gospel_ending_slide(t_prs)
    geo = lambda sl: sorted((sh.left, sh.top, sh.width, sh.height) for sh in sl.shapes
                            if sh.has_text_frame and '주님의 말씀입니다' in sh.text_frame.text)
    assert geo(ending) == geo(t_end), (geo(ending), geo(t_end))


def test_e2_full_pipeline_title_slide_shows_input_saturday_date(e_full_output):
    from missa_ooxml_utils import all_slide_texts as _e2_all_texts
    texts = _e2_all_texts(e_full_output['prs'])
    joined = '\n'.join(texts)
    assert '2026년 9월 12일' in joined  # 사용자가 입력한 토요일(date_str='20260912')
    assert '2026년 9월 13일' not in joined  # 콘텐츠 조회용 +1일(일요일)이 노출되면 안 됨


def test_e4_full_pipeline_복음환호송_body_run_uses_batang_font(e_full_output):
    import missa_sections as sec_e
    from pptx.oxml.ns import qn as _e4b_qn

    sections = sec_e.find_sections(e_full_output['prs'], mass_type='youth')
    slide = e_full_output['prs'].slides[sections['복음환호송']]
    target = None
    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue
        t = shape.text_frame.text.strip()
        if t and '복음 환호송' not in t and '환호송' not in t and '전례문' not in t and 'ALLELUIA' not in t:
            target = shape
            break
    assert target is not None

    para = target.text_frame.paragraphs[0]
    p_el = para._pPr.getparent()
    runs = p_el.findall(_e4b_qn('a:r'))
    run_texts = [(r.find(_e4b_qn('a:t')).text or '') for r in runs]
    assert '\t' in run_texts, run_texts
    tab_idx = run_texts.index('\t')
    body_run = runs[tab_idx + 1]
    rpr = body_run.find(_e4b_qn('a:rPr'))
    latin = rpr.find(_e4b_qn('a:latin')) if rpr is not None else None
    assert latin is not None and latin.get('typeface') == '바탕'


def test_e5_full_pipeline_영문_복음_paragraphs_left_aligned_not_centered(e_full_output):
    import missa_sections as sec_e
    from missa_ooxml_utils import _find_content_shape as _e5_find_content_shape

    sections = sec_e.find_sections(e_full_output['prs'], mass_type='youth')
    s, e = sections['복음_start'], sections['복음_end']
    assert e > s, "영문 복음 콘텐츠 슬라이드가 비어 있음(테스트 전제 깨짐)"

    checked_any = False
    for i in range(s, e):
        shape = _e5_find_content_shape(e_full_output['prs'].slides[i])
        if shape is None:
            continue
        for para in shape.text_frame.paragraphs:
            if not para.text.strip():
                continue
            checked_any = True
            algn = para._pPr.get('algn') if para._pPr is not None else None
            assert algn != 'ctr', f"슬라이드 {i}: 문단이 가운데 정렬됨(algn={algn}) — 왼쪽 정렬이어야 함"
    assert checked_any, "정렬을 확인할 실제 텍스트 문단을 찾지 못함(테스트 전제 깨짐)"


# ═══════════════════════════════════════════════════════════════════════════
# 2026-09-26 F그룹 — 청년미사 JSON/로그 저장 위치 버그(코디네이터가 CLI 실측으로 직접 확인
# 후 지시, 재조사 불필요). `20260926 --미사유형 청년` 실행 결과:
#   F1. 한글 미사 JSON이 output/{date_str}_youth/가 아니라 output/{content_date_str}/(콘텐츠
#       조회용 +1일 날짜, 무관한 다른 산출물과 섞이는 폴더)에 저장됨 — get_json_data()가
#       여전히 고정 경로 Path(OUTPUT_ROOT)/date_str/missa_{date_str}.json을 쓰기 때문.
#   F2. 영문 미사 JSON은 아예 파일로 저장된 적이 없음(반환 dict를 메모리에서만 쓰고 버림) —
#       missa_youth_gospel.main() CLI 진입점에만 저장 코드가 있고 파이프라인은 안 탐.
#   F3. 로그 경로는 이미 B그룹에서 올바르게 배선돼 있을 가능성이 높음(_last_output_key[0] =
#       output_folder_key(date_str, mass_type) → _log_dir_for()) — 코드 미실행 GUI 경로라
#       실측 확인만 필요.
# 이번에도 회귀 스위트는 생략(사용자 지시), 항목별 프로그레션 테스트로만 검증한다.
# ═══════════════════════════════════════════════════════════════════════════
import json as _f_json

_F_BASE = Path(__file__).resolve().parent.parent  # tests/ -> 저장소 루트


def _f_write_json(path, data: dict):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(_f_json.dumps(data, ensure_ascii=False, indent=2), encoding='utf-8')


# ---------------------------------------------------------------------------
# F1: 한글 미사 JSON 저장 폴더
# ---------------------------------------------------------------------------

def test_f1_get_json_data_adult_backward_compatible_default_path(tmp_path, monkeypatch):
    """mass_type/input_date_str을 안 넘기는 기존 호출부(성인 경로, 회귀 픽스처 포함)는
    100% 예전과 동일한 경로(output/{date_str}/missa_{date_str}.json)를 써야 한다."""
    import missa_to_ppt as mtp

    monkeypatch.setattr(mtp, 'OUTPUT_ROOT', str(tmp_path))
    date_str = '20260712'
    expected_path = tmp_path / date_str / f'missa_{date_str}.json'
    _f_write_json(expected_path, {'date': '2026-07-12', 'liturgy': '테스트'})

    result = mtp.get_json_data(date_str)  # mass_type 생략 → 'adult' 기본값

    assert result['liturgy'] == '테스트'


def test_f1_get_json_data_youth_uses_input_date_output_folder_not_content_date(tmp_path, monkeypatch):
    """청년미사는 폴더를 input_date_str(사용자가 입력한 토요일) 기준
    output_folder_key(=..._youth)로 잡아야 한다 — content_date_str(+1일) 폴더가 아니라."""
    import missa_to_ppt as mtp

    monkeypatch.setattr(mtp, 'OUTPUT_ROOT', str(tmp_path))
    input_date_str = '20260926'      # 사용자가 입력한 토요일
    content_date_str = '20260927'    # 콘텐츠 조회용 +1일(일요일)

    correct_path = tmp_path / f'{input_date_str}_youth' / f'missa_{content_date_str}.json'
    _f_write_json(correct_path, {'date': '2026-09-27', 'liturgy': '연중 제26주일'})

    # 예전 버그 경로(콘텐츠 날짜 폴더)에는 아무것도 안 둔다 — 거기서 읽으면 안 되므로.
    wrong_path = tmp_path / content_date_str / f'missa_{content_date_str}.json'
    assert not wrong_path.exists()

    result = mtp.get_json_data(content_date_str, mass_type='youth', input_date_str=input_date_str)

    assert result['liturgy'] == '연중 제26주일'
    assert not wrong_path.parent.exists() or not wrong_path.exists()


def test_f1_get_json_data_youth_cache_hit_skips_crawl(tmp_path, monkeypatch):
    import missa_to_ppt as mtp
    import missa_to_json as m2j

    monkeypatch.setattr(mtp, 'OUTPUT_ROOT', str(tmp_path))
    input_date_str, content_date_str = '20260926', '20260927'
    correct_path = tmp_path / f'{input_date_str}_youth' / f'missa_{content_date_str}.json'
    _f_write_json(correct_path, {'date': '2026-09-27', 'liturgy': '캐시됨'})

    def _boom(*a, **k):
        raise AssertionError('캐시가 있는데도 크롤링을 시도함')

    monkeypatch.setattr(m2j, 'fetch_html', _boom)

    result = mtp.get_json_data(content_date_str, mass_type='youth', input_date_str=input_date_str)
    assert result['liturgy'] == '캐시됨'


def test_f1_get_json_data_youth_crawls_into_correct_folder_when_missing(tmp_path, monkeypatch):
    """캐시 미스 시 새로 크롤링하면(HTML fetch/parse는 content_date_str 기준 그대로) 저장
    위치는 반드시 output_folder_key(input_date_str, 'youth') 폴더여야 한다."""
    import missa_to_ppt as mtp
    import missa_to_json as m2j

    monkeypatch.setattr(mtp, 'OUTPUT_ROOT', str(tmp_path))
    input_date_str, content_date_str = '20260926', '20260927'

    calls = []

    def _fake_fetch_html(date_str):
        calls.append(('fetch', date_str))
        return '<html></html>'

    def _fake_parse_missa(html, date_str):
        calls.append(('parse', date_str))
        return {'date': '2026-09-27', 'liturgy': '새로 크롤링'}

    monkeypatch.setattr(m2j, 'fetch_html', _fake_fetch_html)
    monkeypatch.setattr(m2j, 'parse_missa', _fake_parse_missa)

    result = mtp.get_json_data(content_date_str, mass_type='youth', input_date_str=input_date_str)

    assert result['liturgy'] == '새로 크롤링'
    assert calls == [('fetch', content_date_str), ('parse', content_date_str)]
    saved_path = tmp_path / f'{input_date_str}_youth' / f'missa_{content_date_str}.json'
    assert saved_path.is_file()
    assert _f_json.loads(saved_path.read_text(encoding='utf-8'))['liturgy'] == '새로 크롤링'


# ---------------------------------------------------------------------------
# F2: 영문 미사 JSON 파일 저장(신규) — 지금까지는 메모리에서만 쓰고 버려졌다
# ---------------------------------------------------------------------------

def test_f2_get_youth_english_mass_cache_hit_skips_network(tmp_path, monkeypatch):
    import missa_to_ppt as mtp

    monkeypatch.setattr(mtp, 'OUTPUT_ROOT', str(tmp_path))
    output_key = '20260926_youth'
    content_date_str = '20260927'
    cached_path = tmp_path / output_key / f'missa_en_{content_date_str}.json'
    _f_write_json(cached_path, {'date': content_date_str, 'Gospel': {'content': '캐시됨'}})

    def _boom(d):
        raise AssertionError('캐시가 있는데도 네트워크 조회를 시도함')

    monkeypatch.setattr(mtp.missa_youth_gospel, 'get_youth_mass', _boom)

    result = mtp._get_youth_english_mass(content_date_str, output_key)
    assert result['Gospel']['content'] == '캐시됨'


def test_f2_get_youth_english_mass_cache_miss_fetches_and_saves(tmp_path, monkeypatch):
    import missa_to_ppt as mtp

    monkeypatch.setattr(mtp, 'OUTPUT_ROOT', str(tmp_path))
    output_key = '20260926_youth'
    content_date_str = '20260927'

    fetched = {'date': content_date_str, 'Gospel': {'content': '새로 조회'}}
    monkeypatch.setattr(mtp.missa_youth_gospel, 'get_youth_mass', lambda d: fetched)

    result = mtp._get_youth_english_mass(content_date_str, output_key)

    assert result == fetched
    saved_path = tmp_path / output_key / f'missa_en_{content_date_str}.json'
    assert saved_path.is_file()
    assert _f_json.loads(saved_path.read_text(encoding='utf-8')) == fetched


# ---------------------------------------------------------------------------
# F3: 로그 저장 폴더 배선 검증(기존에 이미 맞게 돼 있을 가능성이 높다는 코디네이터 소견 —
# 직접 main() 1회 실행해 실측으로 확인)
# ---------------------------------------------------------------------------

def test_f3_log_dir_wiring_points_to_youth_output_folder_after_main_run(tmp_path, monkeypatch):
    import sys as _f3_sys
    import missa_to_ppt as mtp
    import missa_gui as gui

    _e_skip_if_missing()  # 이 파일 상단(E그룹)에 정의된 실측 픽스처 가드 재사용

    date_str = '20260912'
    content_date_str = '20260913'
    (tmp_path / date_str).mkdir()
    (tmp_path / content_date_str).mkdir()
    ref_dst = tmp_path / date_str / 'reference.pptx'
    _e_shutil.copy(str(_E_YOUTH_TEMPLATE), str(ref_dst))

    monkeypatch.setattr(mtp, 'OUTPUT_ROOT', str(tmp_path))
    monkeypatch.setattr(gui, 'OUTPUT_ROOT', str(tmp_path))

    en_fixture = _e_json.loads(_E_EN_GOSPEL_FIXTURE.read_text(encoding='utf-8'))
    monkeypatch.setattr(mtp.missa_youth_gospel, 'get_youth_mass', lambda d: en_fixture)

    files = {
        'ref_pptx': ref_dst, '시작기도': None, '화답송_pptx': None, '화답송_img': None,
        '미사후기도': None, '공지사항': None, '성가': {},
    }
    성가_선택 = {
        '입당': [{'출처': '기타', '번호': None, '제목': 'T'}],
        '봉헌': [{'출처': '기타', '번호': None, '제목': 'T'}],
        '성체': [{'출처': '기타', '번호': None, '제목': 'T'}],
        '2차봉헌': [{'출처': '기타', '번호': None, '제목': 'T'}],
        '파견': [{'출처': '기타', '번호': None, '제목': 'T'}],
    }
    mtp._preloaded_inputs[0] = {
        'mass_type': 'youth', 'date_str': date_str, 'files': files,
        'hymn_numbers': {}, '성가_선택': 성가_선택,
    }
    monkeypatch.setattr(_f3_sys, 'argv', ['missa_to_ppt.py'])

    try:
        mtp.main()
    finally:
        mtp._preloaded_inputs[0] = None

    assert mtp._last_date_str[0] == date_str
    assert mtp._last_output_key[0] == f'{date_str}_youth'

    log_dir = gui._log_dir_for(mtp._last_date_str[0], mtp._last_output_key[0])
    assert log_dir == tmp_path / f'{date_str}_youth' / 'log'


# ═══════════════════════════════════════════════════════════════════════════
# G그룹 — 2026-09-26 (F그룹 후속) 코디네이터 실측 버그 2건
#
# 버그1(영문 복음 조기 줄바꿈): layout_units_on_slides_pil()이 검증
# (_rendered_wrap_count/_count_slide_lines_rendered/COM)과 다른(3% 더 좁은) 박스 폭으로
# 페이지네이션해, 실제로는 그 슬라이드에 들어가는 단어까지 다음 슬라이드로 밀어냄
# (실사용자 9/26 청년미사 슬라이드 54, "...did the" 뒤 "father's"가 잘림). 원인은
# `_PIL_WRAP_SAFETY=0.97`이 페이지네이션에만 곱해지고 검증 경로에는 곱해지지 않던 것.
# 이 상수는 2026-09-17에 반대 방향 오차(Pillow가 실제보다 줄 수를 과소 예측하는 경계
# 케이스, 20260913 영문 복음 4번째 슬라이드)를 막으려고 도입됐다 — 제거하면 그 문제가
# 재발할 수 있으므로, "제거해도 안전한지"를 post-write COM 재조정 안전망으로 실측
# 확인(아래 g1c) 후 제거했다.
#
# 버그2(OneDrive 브라우저 폴더 열 때마다 3~6초): missa_onedrive.get_access_token()이
# msal.PublicClientApplication을 매 호출마다 새로 생성(약 0.9~1.0초/회, network-based
# authority의 instance discovery 재수행 추정) — 캐시된 토큰 조회(acquire_token_silent)
# 자체는 0초인데 앱 생성 자체가 병목. 프로세스 수명 모듈 전역 캐시로 앱을 1회만 생성해
# 재사용하도록 수정.
# ═══════════════════════════════════════════════════════════════════════════


# ---------------------------------------------------------------------------
# G1: 영문 복음 조기 줄바꿈 버그 수정
# ---------------------------------------------------------------------------

# (G1a/G1c 삭제 — 2026-10-03 슬라이드 단위 채우기: 분할과 검증이 같은 박스 폭을 쓴다는 요구사항은
#  test_y2_G1b(Pillow 경로, 949.98px)·test_sf9(Pillow 루프)가, "추정 오차에도 COM 실측 9줄 수렴"은
#  test_sf1/test_sf6/test_sf_real4가 이어받았다. G1b는 test_missa_regression.py로 승격됨 —
#  2026-10-03: test_y2_G1b_20260926_father_fits_first_page)


# ---------------------------------------------------------------------------
# G2: OneDrive 커스텀 파일 브라우저 폴더 열 때마다 느린 버그(MSAL 앱 매번 재생성) 수정
# ---------------------------------------------------------------------------


# ═══════════════════════════════════════════════════════════════════════════
# H그룹 — 2026-09-26 (G그룹 후속) OneDrive 파일 브라우저: 폴더 확장 화살표 클릭이
# 무반응(영구히 "(불러오는 중...)")인 버그
#
# 근본 원인: `missa_gui._ask_onedrive_file_browser_popup()`의 lazy-load가 `<Double-1>`
# (더블클릭)에만 걸려 있었다. ttk.Treeview의 확장 화살표(▶/▼) **단일 클릭**은 이 바인딩과
# 무관한 네이티브 Tk 토글이라 `_on_double_click`이 전혀 호출되지 않는다 — 그 결과
# "(불러오는 중...)" 플레이스홀더만 뜬 채 실제 fetch가 영구히 일어나지 않는다(실사용자
# 스크린샷 재현 — "20.청년 미사" 펼침 아이콘은 `[-]`로 바뀌었는데 자식은 그대로 플레이스홀더).
# `<<TreeviewOpen>>` 가상 이벤트(화살표 클릭·더블클릭·키보드 등 **어떤 방식으로 펼치든
# 항상 발생**)로 lazy-load 트리거를 옮겨 고친다. 기존 더블클릭 핸들러는 그대로 유지하고,
# 두 경로가 겹쳐 호출돼도 이미 있던 already_loaded 가드가 중복 fetch를 막는다.
# ═══════════════════════════════════════════════════════════════════════════


# ═══════════════════════════════════════════════════════════════════════════
# I그룹 — 2026-09-26 (H그룹 후속) OneDrive에서 파일 선택 후, 입력 확인 화면의 경로
# 표시가 비실용적인 버그
#
# 대상: `_ask_combined_input_popup(mass_type='youth')`의 파일 선택 행들. '찾아보기'로
# OneDrive 파일을 고르면 반환된 **로컬 다운로드 경로**(`output\{date}_youth\{파일명}`)가
# 그대로 Entry(46자 폭)에 표시돼, 실제로 어떤 파일인지 구분되는 핵심 부분(파일명)이 잘려서
# 안 보인다(실사용자 스크린샷으로 확인). 요구사항(사용자 확인 완료): (1) 표시는 로컬 경로가
# 아니라 OneDrive 원본 경로에서 'PPT 문서/' 접두사를 뺀 상대경로부터, (2) 파일명이 최소
# 절반 이상 보이도록(Entry 스크롤을 끝까지 이동).
#
# 설계 제약: 표시 문자열과 파이프라인이 실제로 쓰는 값은 반드시 분리해야 한다 —
# on_ok()가 Entry의 StringVar 값을 그대로 파일 경로로 쓰므로, 표시를 상대경로로 바꾸면
# on_ok()가 존재하지 않는 경로를 쓰게 되는 회귀가 생긴다. 그래서 실제 로컬 경로는 별도
# dict(`_od_real_paths`)에 저장하고, on_ok()의 각 필드 조회는 그 dict를 우선한다
# (`_resolved_field_value`). `_ask_onedrive_file_browser_popup()`은 이제 local_path 하나가
# 아니라 (local_path, remote_path) 튜플을 반환한다(취소 시 ('', '')) — 표시용 상대경로를
# 만들려면 remote_path가 필요하기 때문. 호출부는 missa_gui.py 한 곳(그리고 이 파일의
# 서브프로세스 프로브 스크립트들)뿐이라 시그니처 변경이 안전함을 grep으로 재확인했다.
# ═══════════════════════════════════════════════════════════════════════════


class _I_FakeEntry:
    """`entry.xview_moveto(1.0)` 호출 여부만 기록하는 대역 — 실제 Tk Entry 위젯 없이도
    "파일명이 보이도록 스크롤했는지"를 검증할 수 있게 한다."""

    def __init__(self):
        self.xview_calls = []

    def xview_moveto(self, frac):
        self.xview_calls.append(frac)


# ═══════════════════════════════════════════════════════════════════════════
# J그룹 배치1 — 2026-09-26 OneDrive 파일 선택 팝업 UI (J1 안내문구 단순화 + J2 최상위
# 폴더 3개로 제한 + J5 백그라운드 prefetch)
#
# J1: `_ask_onedrive_file_browser_popup()`의 안내 Label 문구를 "{root} 폴더 안에서
# 파일을 선택해 주세요 (더블클릭으로 폴더 열기)"에서 "파일을 선택해 주세요"로 단순화.
# J2: 최상위 트리 레벨에서는 실제 Graph API로 실측된 3개 폴더('11.공지사항 PPT문서',
# '13. 기도문', '20.청년 미사')만 보여준다 — 하위 폴더 탐색은 그대로 전부 보여줌. 공백
# 유무 차이에 안전하도록 정규화 비교.
# J5(선택 구현): 팝업이 열리는 시점에 위 3개 폴더의 자식 목록을 백그라운드 스레드로
# prefetch해 캐시를 미리 채운다 — 사용자가 실제로 펼칠 때 캐시 히트로 즉시 뜨게 함.
# ═══════════════════════════════════════════════════════════════════════════


# ═══════════════════════════════════════════════════════════════════════════
# J그룹 배치2 — 2026-09-26 파일 순서 + 로그 (J4 청년미사 파일 행 순서 변경 + J6 로그에
# 공지사항 누락 보강)
#
# J4: `_ask_combined_input_popup()`의 file_rows가 성인미사는 [ref_pptx, 시작기도,
# 화답송_pptx, 공지사항, 미사후기도]인데, 청년미사만 [ref_pptx, 시작기도, 미사후기도,
# 공지사항](화답송 행 자체는 기존처럼 없음)으로 바꾼다. 성인미사 순서는 절대 그대로.
# J6: `missa_to_ppt.py`의 "[1] 파일 확인..." 로그 블록에 참조 PPT/시작기도/화답송/
# 미사후기도는 출력하는데 공지사항이 빠져 있다 — 같은 형식으로 한 줄 추가한다.
# ═══════════════════════════════════════════════════════════════════════════


_J_WIDGET_DUMP_HELPER = '''
def _j_widget_dump(w, out):
    info = {"class": w.winfo_class()}
    for opt in ("text", "font", "fg", "bg", "state", "width", "cursor"):
        try:
            info[opt] = str(w.cget(opt))
        except Exception:
            pass
    try:
        gi = w.grid_info()
    except Exception:
        gi = {}
    info["grid_row"] = int(gi["row"]) if gi else None
    info["grid_column"] = int(gi["column"]) if gi else None
    info["gridded"] = bool(gi)
    out.append(info)
    for c in w.winfo_children():
        _j_widget_dump(c, out)
'''


def _j_run_widget_probe(tmp_path, popup_call: str) -> list:
    """`popup_call`을 서브프로세스에서 실행하고 mainloop 진입 직후 위젯 트리를 JSON으로
    덤프해 반환한다. A그룹의 `_a_run_widget_probe`(test_missa_regression.py)와 같은 기법
    이지만, 이 파일은 회귀 파일에 의존하지 않고 자기 완결적으로 유지하기 위해 별도로
    둔다(이후 J그룹 배치의 다른 위젯 검증에도 재사용 가능)."""
    out_file = tmp_path / 'j_probe_result.json'
    script = f'''
import sys, json
sys.path.insert(0, {str(BASE)!r})
import tkinter as tk
import missa_gui as gui

{_J_WIDGET_DUMP_HELPER}

_widgets = []
_orig_mainloop = tk.Tk.mainloop

def _fake_mainloop(self, *a, **k):
    def _probe():
        try:
            _j_widget_dump(self, _widgets)
        finally:
            self.destroy()
    self.after(60, _probe)
    return _orig_mainloop(self, *a, **k)

tk.Tk.mainloop = _fake_mainloop

try:
    {popup_call}
except RuntimeError:
    pass

with open({str(out_file)!r}, 'w', encoding='utf-8') as f:
    json.dump(_widgets, f)
'''
    script_path = tmp_path / 'j_probe_script.py'
    script_path.write_text(script, encoding='utf-8')
    import subprocess as _j4_subprocess
    import sys as _j4_sys
    import json as _j4_json
    try:
        result = _j4_subprocess.run(
            [_j4_sys.executable, str(script_path)],
            cwd=str(BASE), capture_output=True, text=True, encoding='utf-8', errors='replace',
            timeout=30,
        )
    except _j4_subprocess.TimeoutExpired as e:
        raise AssertionError(
            f"J그룹 위젯 프로브가 30초 내에 끝나지 않음\nstdout:\n{e.stdout}\nstderr:\n{e.stderr}"
        )
    assert out_file.is_file(), (
        f"위젯 덤프 실패(exit={result.returncode})\nstdout:\n{result.stdout}\nstderr:\n{result.stderr}"
    )
    return _j4_json.loads(out_file.read_text(encoding='utf-8'))


def test_j4_youth_file_rows_order_changed_but_adult_order_untouched(tmp_path):
    """J4: 청년미사는 [참조, 시작기도, 미사후기도, 공지사항] 순서로, 성인미사는 기존
    [참조, 시작기도, 화답송, 공지사항, 미사후기도] 순서 그대로여야 한다."""
    youth_widgets = _j_run_widget_probe(tmp_path, "gui._ask_combined_input_popup(mass_type='youth')")
    youth_labels = sorted(
        (w for w in youth_widgets
         if w['class'] == 'Label' and w.get('gridded') and w.get('grid_column') == 0
         and w.get('grid_row', 0) >= 1),
        key=lambda w: w['grid_row'],
    )
    youth_texts = [w['text'] for w in youth_labels]
    assert len(youth_texts) == 4, youth_texts
    assert '참조' in youth_texts[0], youth_texts
    assert '시작기도' in youth_texts[1], youth_texts
    assert '미사' in youth_texts[2] and '기도' in youth_texts[2], youth_texts
    assert '공지사항' in youth_texts[3], youth_texts

    adult_widgets = _j_run_widget_probe(tmp_path, "gui._ask_combined_input_popup(mass_type='adult')")
    adult_labels = sorted(
        (w for w in adult_widgets
         if w['class'] == 'Label' and w.get('gridded') and w.get('grid_column') == 0
         and w.get('grid_row', 0) >= 1),
        key=lambda w: w['grid_row'],
    )
    adult_texts = [w['text'] for w in adult_labels]
    assert len(adult_texts) == 5, adult_texts
    assert '참조' in adult_texts[0], adult_texts
    assert '시작기도' in adult_texts[1], adult_texts
    assert '화답송' in adult_texts[2], (
        "성인미사 화답송 행 위치가 J4 작업 중 잘못 바뀜(절대 건드리지 않아야 함)", adult_texts,
    )
    assert '공지사항' in adult_texts[3], adult_texts
    assert '미사' in adult_texts[4] and '기도' in adult_texts[4], adult_texts


def test_j6_file_check_log_includes_공지사항():
    """J6: '[1] 파일 확인...' 로그 블록에 참조 PPT/시작기도/화답송/미사후기도와 같은
    형식(없으면 "없음")으로 공지사항 항목이 추가돼야 한다."""
    import inspect
    import missa_to_ppt as ppt
    src = inspect.getsource(ppt.main)
    assert 'files["공지사항"].name if files.get("공지사항") else "없음"' in src, (
        "[1] 파일 확인 로그에 공지사항 항목이 참조 PPT/시작기도/화답송/미사후기도와 같은 "
        "형식(없으면 '없음')으로 추가되지 않음"
    )


# ═══════════════════════════════════════════════════════════════════════════
# J그룹 배치3 — 2026-09-26 미사 유형 선택 팝업 전면 제거 + CLI 직접 실행 (J3+J9-CLI아키텍처)
#
# `python missa_to_ppt.py --미사유형 청년`(날짜 없음)을 실행하면 미사 유형 선택 팝업 없이
# 바로 청년미사 입력 UI가 뜨게 한다. `_ask_mass_type_popup()` 자체와 호출부 2곳을 삭제하고,
# config.json의 last_mass_type 읽기/쓰기도 제거한다. `--미사유형`에 '어린이'를 choices로
# 추가(아직 미지원 안내 후 정상 종료). `_ask_combined_input_popup()`의 DIALOG_TITLE에
# mass_type별 접미사 추가.
#
# 핵심 함정(실측 발견): `main()`이 `if len(sys.argv) == 1:`로 "GUI 흐름인가"를 판단했는데,
# `--미사유형 청년`처럼 date 없이 플래그만 있는 GUI 모드 실행에서는 sys.argv 길이가 1이
# 아니라서(`_run_gui_mode()`가 채운 `_preloaded_inputs[0]`를 쓰고 있으면서도) 잘못 CLI
# 분기(`parse_args()`, date=None이라 이후 로직이 깨짐)로 빠질 뻔했다. 참 신호는
# `_preloaded_inputs[0]`의 유무이지 argv 길이가 아니다 — `_main_should_use_preloaded_or_
# fallback()`으로 수정.
# ═══════════════════════════════════════════════════════════════════════════


def test_j3a_arg_parser_date_is_optional():
    """date 없이(플래그만) 파싱해도 argparse가 사용법 오류로 죽지 않아야 한다."""
    parser = mtp._build_arg_parser()
    args = parser.parse_args(['--미사유형', '청년'])
    assert args.date is None
    assert args.mass_type_kr == '청년'


def test_j3b_arg_parser_accepts_어린이_choice():
    """--미사유형 choices에 '어린이'가 추가돼야 한다."""
    parser = mtp._build_arg_parser()
    args = parser.parse_args(['--미사유형', '어린이'])
    assert args.mass_type_kr == '어린이'


def test_j3c_arg_parser_date_still_parses_when_given():
    """date를 주면 여전히 정상적으로 그 값을 받는다(기존 CLI 전체 모드 회귀 없음)."""
    parser = mtp._build_arg_parser()
    args = parser.parse_args(['20260927', '--미사유형', '청년'])
    assert args.date == '20260927'
    assert args.mass_type_kr == '청년'


def test_j3d_arg_parser_default_mass_type_is_성인_when_omitted():
    parser = mtp._build_arg_parser()
    args = parser.parse_args(['20260927'])
    assert args.mass_type_kr == '성인'


def test_j3e_main_dispatch_prefers_preloaded_even_with_cli_flags_present(monkeypatch):
    """실측 버그 재현/회귀 가드: `_preloaded_inputs[0]`가 채워져 있으면, sys.argv에
    `--미사유형 청년`처럼 date 없는 CLI 플래그가 남아있어도 '프리로드 사용'으로 판단해야
    한다 — 예전 `len(sys.argv) == 1` 기준으로는 이 경우 잘못 CLI 분기로 빠졌다."""
    monkeypatch.setattr(mtp.sys, 'argv', ['missa_to_ppt.py', '--미사유형', '청년'])
    mtp._preloaded_inputs[0] = {'dummy': True}
    try:
        assert mtp._main_should_use_preloaded_or_fallback() is True
    finally:
        mtp._preloaded_inputs[0] = None


def test_j3f_main_dispatch_falls_back_when_no_preloaded_and_date_missing(monkeypatch):
    """프리로드도 없고 date도 없으면(프로그래밍적 직접 호출) 여전히 "프리로드/fallback"
    분기로 가야 한다(내부에서 다시 preloaded 유무로 나뉨)."""
    monkeypatch.setattr(mtp.sys, 'argv', ['missa_to_ppt.py', '--미사유형', '청년'])
    mtp._preloaded_inputs[0] = None
    assert mtp._main_should_use_preloaded_or_fallback() is True


def test_j3g_main_dispatch_uses_cli_branch_when_date_given(monkeypatch):
    """date가 실제로 주어지면 CLI 전체 모드(parse_args() 사용)로 가야 한다."""
    monkeypatch.setattr(mtp.sys, 'argv', ['missa_to_ppt.py', '20260927', '--미사유형', '청년'])
    mtp._preloaded_inputs[0] = None
    assert mtp._main_should_use_preloaded_or_fallback() is False


def test_j3h_dispatch_cli_or_gui_calls_run_gui_mode_when_date_omitted(monkeypatch):
    """`_dispatch_cli_or_gui()`(구 __main__ 블록)가 date 없이 실행되면 `_ask_mass_type_popup`
    같은 팝업 없이 바로 `_run_gui_mode(mass_type_kr)`을 호출해야 한다."""
    calls = []
    monkeypatch.setattr(mtp.sys, 'argv', ['missa_to_ppt.py', '--미사유형', '청년'])
    monkeypatch.setattr(mtp, '_run_gui_mode', lambda mass_type_kr: calls.append(mass_type_kr))
    mtp._dispatch_cli_or_gui()
    assert calls == ['청년'], calls


def test_j3i_dispatch_cli_or_gui_calls_main_when_date_given(monkeypatch):
    """date가 주어지면 `main()`을 호출해야 한다(`_run_gui_mode`는 호출되지 않음)."""
    gui_calls = []
    main_calls = []
    monkeypatch.setattr(mtp.sys, 'argv', ['missa_to_ppt.py', '20260927', '--미사유형', '청년', '--test'])
    monkeypatch.setattr(mtp, '_run_gui_mode', lambda mass_type_kr: gui_calls.append(mass_type_kr))
    monkeypatch.setattr(mtp, 'main', lambda: main_calls.append(True))
    mtp._dispatch_cli_or_gui()
    assert main_calls == [True]
    assert gui_calls == []


def test_j3j_dispatch_cli_or_gui_shows_children_message_and_skips_gui(monkeypatch):
    """'어린이'가 선택되면 입력 UI(`_run_gui_mode`)를 열지 않고 안내만 띄운 뒤 끝나야
    한다."""
    gui_calls = []
    msg_calls = []
    monkeypatch.setattr(mtp.sys, 'argv', ['missa_to_ppt.py', '--미사유형', '어린이'])
    monkeypatch.setattr(mtp, '_run_gui_mode', lambda mass_type_kr: gui_calls.append(mass_type_kr))
    monkeypatch.setattr(
        mtp, '_show_children_mass_not_supported_message', lambda: msg_calls.append(True)
    )
    mtp._dispatch_cli_or_gui()
    assert msg_calls == [True]
    assert gui_calls == []


def test_j3j2_dispatch_cli_or_gui_shows_children_message_even_with_date_given(monkeypatch):
    """리뷰 라운드11 지적: '어린이' 가드가 date 유무와 무관하게 최우선으로 걸려야 한다.
    date가 주어져도(`python missa_to_ppt.py 20260927 --미사유형 어린이`) main()이 조용히
    성인미사로 폴백 실행되면 안 되고, 안내만 뜨고 끝나야 한다."""
    gui_calls = []
    msg_calls = []
    main_calls = []
    monkeypatch.setattr(
        mtp.sys, 'argv',
        ['missa_to_ppt.py', '20260927', '--미사유형', '어린이', '--test'],
    )
    monkeypatch.setattr(mtp, '_run_gui_mode', lambda mass_type_kr: gui_calls.append(mass_type_kr))
    monkeypatch.setattr(mtp, 'main', lambda: main_calls.append(True))
    monkeypatch.setattr(
        mtp, '_show_children_mass_not_supported_message', lambda: msg_calls.append(True)
    )
    mtp._dispatch_cli_or_gui()
    assert msg_calls == [True]
    assert gui_calls == []
    assert main_calls == [], (
        "date가 주어졌다는 이유로 어린이 미가드가 우회되어 main()이 조용히 실행됨"
    )


def test_j3k_ask_mass_type_popup_function_removed():
    """`_ask_mass_type_popup()` 함수 자체가 삭제됐어야 한다."""
    import missa_gui as gui
    assert not hasattr(gui, '_ask_mass_type_popup'), (
        "_ask_mass_type_popup()이 아직 남아 있음 — §J3에서 완전히 제거해야 함"
    )


def test_j3l_last_mass_type_no_longer_written_by_gui_module():
    """config.json에 last_mass_type을 쓰는 코드가 missa_gui.py에 더 이상 없어야 한다."""
    import inspect
    import missa_gui as gui
    src = inspect.getsource(gui)
    assert 'last_mass_type' not in src, (
        "missa_gui.py에 last_mass_type 참조가 남아 있음 — §J3에서 완전히 제거해야 함"
    )


def test_j3m_run_gui_mode_skips_mass_type_popup_and_uses_given_kr(monkeypatch):
    """`_run_gui_mode('청년')`을 호출하면 (더 이상 존재하지 않을) 미사유형 팝업을 부르지
    않고, 주어진 '청년'을 그대로 mass_type으로 써서 청년 흐름(ensure_ppt_roots →
    _ask_combined_input_popup(mass_type='youth') → _ask_youth_hymn_popup)을 타야 한다."""
    calls = []
    monkeypatch.setattr(mtp, 'ensure_ppt_roots', lambda mt, ask_if_missing=False: (calls.append('login'), True)[1])
    monkeypatch.setattr(
        mtp, '_ask_combined_input_popup',
        lambda mass_type: (calls.append(('combined', mass_type)), ('20260927', {}))[1],
    )
    monkeypatch.setattr(mtp, '_ask_youth_hymn_popup', lambda: (calls.append('youth_hymn'), {})[1])
    monkeypatch.setattr(mtp, '_run_with_progress_window', lambda fn: (calls.append('progress'), ('', ''))[1])
    monkeypatch.setattr(mtp, '_show_result_window', lambda *a, **k: calls.append(('result', a, k)))
    mtp._preloaded_inputs[0] = None
    try:
        mtp._run_gui_mode('청년')
    finally:
        mtp._preloaded_inputs[0] = None
    assert calls[0] == 'login', calls
    assert calls[1] == ('combined', 'youth'), calls
    assert calls[2] == 'youth_hymn', calls
    assert calls[3] == 'progress', calls


def test_j3n_dialog_title_has_mass_type_suffix():
    """`_ask_combined_input_popup()`의 창 제목에 mass_type별 접미사가 붙어야 한다.
    root.title()은 일반 위젯 트리 덤프에 잡히지 않으므로(최상위 창 자체의 속성), 소스
    검사로 매핑이 실제로 존재하는지 확인한다."""
    import missa_gui as gui
    import inspect
    src = inspect.getsource(gui._ask_combined_input_popup)
    assert "' (성인미사)'" in src, src
    assert "' (청년미사)'" in src, src


# ═══════════════════════════════════════════════════════════════════════════
# J그룹 배치5 — 2026-09-26 liturgy 필드에 <sub> 부제 포함 (J8)
#
# `extract_liturgy_name(soup)`은 `<meta name="title">`의 content 속성에서만 파싱한다
# (예: "2026.09.27 [녹] 연중 제26주일" → "연중 제26주일"). 그런데 실제 페이지의
# `<h3 id="missa_title">` 안에는 부제가 `<sub>` 태그로 따로 들어갈 때가 있다 — 이 경우
# 최종 liturgy 값은 "연중 제26주일 (세계 이주민과 난민의 날)"이어야 한다(부제가 공백
# 하나 두고 뒤에 이어붙음). 기존 meta-tag 파싱 로직은 그대로 두고(기존 회귀 픽스처
# 20260624/20260705/20260712가 이 함수에 의존), <h3 id="missa_title"> 안에 <sub> 태그가
# 있으면 그 텍스트를 결과에 추가로 붙이는 방식으로 보강한다.
# ═══════════════════════════════════════════════════════════════════════════


def test_j8a_extract_liturgy_name_appends_sub_subtitle_from_h3():
    """코디네이터가 제시한 실제 예시 그대로 재현."""
    import missa_to_json as m2j
    from bs4 import BeautifulSoup
    html = '''
    <html><head><meta name="title" content="2026.09.27 [녹] 연중 제26주일"></head>
    <body>
    <h3 id="missa_title">
        <span style="color:green;">[녹]</span> 연중 제26주일
            <span class="float-right"><sub>(세계 이주민과 난민의 날)</sub></span>
    </h3>
    </body></html>
    '''
    soup = BeautifulSoup(html, 'html.parser')
    result = m2j.extract_liturgy_name(soup)
    assert result == '연중 제26주일 (세계 이주민과 난민의 날)', result


def test_j8b_extract_liturgy_name_without_sub_unchanged():
    """<sub>가 없으면 기존 meta-tag 파싱 결과 그대로(회귀 보호) — 대부분의 평범한
    주일/평일미사는 부제가 없다."""
    import missa_to_json as m2j
    from bs4 import BeautifulSoup
    html = '''
    <html><head><meta name="title" content="2026.06.14 [녹] 연중 제11주일"></head>
    <body><h3 id="missa_title"><span style="color:green;">[녹]</span> 연중 제11주일</h3></body>
    </html>
    '''
    soup = BeautifulSoup(html, 'html.parser')
    result = m2j.extract_liturgy_name(soup)
    assert result == '연중 제11주일', result


def test_j8c_extract_liturgy_name_no_h3_falls_back_to_meta_only():
    """h3#missa_title 자체가 없어도(예: 페이지 구조가 다른 과거 크롤 데이터) 기존
    meta-tag 파싱은 그대로 동작해야 한다 — 회귀 픽스처 20260624/20260705/20260712가
    이 함수에 의존하므로 기존 동작 보존이 최우선."""
    import missa_to_json as m2j
    from bs4 import BeautifulSoup
    html = (
        '<html><head><meta name="title" '
        'content="2026.06.24 [백] 성 요한 세례자 탄생 대축일"></head><body></body></html>'
    )
    soup = BeautifulSoup(html, 'html.parser')
    result = m2j.extract_liturgy_name(soup)
    assert result == '성 요한 세례자 탄생 대축일', result


def test_j8d_extract_liturgy_name_h3_sub_empty_text_ignored():
    """<sub> 태그가 있어도 텍스트가 공백뿐이면 덧붙이지 않는다(방어적 — 빈 꼬리 공백이
    최종 문자열에 남지 않게)."""
    import missa_to_json as m2j
    from bs4 import BeautifulSoup
    html = '''
    <html><head><meta name="title" content="2026.09.27 [녹] 연중 제26주일"></head>
    <body><h3 id="missa_title">연중 제26주일<span class="float-right"><sub>   </sub></span></h3></body>
    </html>
    '''
    soup = BeautifulSoup(html, 'html.parser')
    result = m2j.extract_liturgy_name(soup)
    assert result == '연중 제26주일', result


# ═══════════════════════════════════════════════════════════════════════════
# J그룹 배치4 — 2026-09-26 PowerPoint 백그라운드 실행 안내 + 포그라운드 조사 (J7)
#
# ppt_com_verify.py의 print()는 콘솔에만 찍혀서 GUI 모드(tkinter 진행률 창)에서는
# 실시간으로 보이지 않는다. `_run_gui_mode()`가 실제 처리(_run_with_progress_window(main))
# 를 시작하기 직전에, 성인/청년 동일하게 한 번만 messagebox 안내를 띄워야 한다.
#
# 포커스 탈취 여부는 실측(2026-09-26, EnumWindows로 PP12FrameClass 프레임 창을 찾아
# GetForegroundWindow()와 비교, 좀비 방지를 위해 매 실행 shutdown() 후 잔존 프로세스
# 없음까지 확인, 2회 반복)으로 확인했다 — Application.Visible=True로 프레임 창이
# 실제로 "보이는"(IsWindowVisible=1, 최소화 안 됨) 상태가 되지만, Dispatch 직후와
# Presentations.Open() 측정 도중 둘 다 GetForegroundWindow()는 그대로였다(포커스를
# 가져가지 않음). 그래서 코드 수정(ShowWindow 등)은 하지 않고, 안내 문구만 이 실측과
# 일치하도록 정확히 쓴다 — "화면에 창이 뜨지 않습니다"(사실과 다름, 창은 실제로
# visible해짐)가 아니라 "포커스를 가져가지 않는다"(실측으로 확인된 사실)로 표현한다.
# ═══════════════════════════════════════════════════════════════════════════


def test_j7a_no_separate_powerpoint_notice_popup_anymore():
    """2026-10-03: PowerPoint 백그라운드 안내 팝업(messagebox)은 폐지됐다 — 진행 창 첫머리 안내로 대체."""
    assert not hasattr(mtp, '_show_powerpoint_background_notice')
    import inspect
    assert '_show_powerpoint_background_notice' not in inspect.getsource(mtp._run_gui_mode)


def test_j7b_progress_window_shows_notice_for_five_seconds_then_removes_it():
    """진행 창이 안내 라벨을 만들고(첫머리), `_POWERPOINT_NOTICE_SECONDS`(=5)초 뒤 지운다.
    진행 메시지(label_var)는 그 라벨과 별개라 5초 동안 같이 보인다."""
    import inspect
    import missa_gui as gui
    assert gui._POWERPOINT_NOTICE_SECONDS == 5
    src = inspect.getsource(gui._run_with_progress_window)
    assert '_POWERPOINT_NOTICE_TEXT' in src and '_POWERPOINT_NOTICE_SECONDS * 1000' in src
    assert 'notice.destroy()' in src
    # 안내 라벨이 진행 메시지 라벨보다 먼저 pack된다(첫머리).
    assert src.index('notice.pack(') < src.index("label_var = tk.StringVar")


def test_j7c_notice_wording_does_not_falsely_claim_no_window_appears():
    """실측(§J7)과 다른 주장("화면에 창이 뜨지 않습니다")을 안내 문구에 남기면 안 된다 —
    프레임 창은 visible해지지만 포커스는 가져가지 않는다."""
    import missa_gui as gui
    assert '화면에 창이 뜨지 않습니다' not in gui._POWERPOINT_NOTICE_TEXT
    assert '포커스' in gui._POWERPOINT_NOTICE_TEXT


def test_j7d_ppt_com_verify_ensure_app_notice_wording_fixed():
    """ppt_com_verify.py `_ensure_app()`의 콘솔 안내도 같은 실측 근거로 문구를 고쳐야
    한다 — "화면에 창이 뜨지 않습니다"는 사실과 다르다(창은 visible해지고, 포커스만
    안 가져간다)."""
    import ppt_com_verify as pcv
    import inspect
    src = inspect.getsource(pcv._ensure_app)
    assert '화면에 창이 뜨지 않습니다' not in src, src


# ═══════════════════════════════════════════════════════════════════════════
# J그룹 배치7 — 2026-09-26 버전 표시 + 업데이트 확인 버튼→링크 전환 (J10)
#
# VERSION 파일(repo 루트, 형식 YYYYMMDD-NN)을 읽어 표시하고, §J3에서 옮겨둔
# `_on_check_update(parent)`를 그대로 재배선한다(재구현 금지). '업데이트 확인'을
# tk.Button에서 tk.Label(cursor='hand2' + <Button-1> bind)로 바꾼다.
# ═══════════════════════════════════════════════════════════════════════════


def test_j10a_get_app_version_reads_version_file(tmp_path, monkeypatch):
    """VERSION 파일이 있으면 그 내용을 그대로(공백 제거) 반환해야 한다."""
    import missa_gui as gui
    (tmp_path / 'VERSION').write_text('20260101-02', encoding='utf-8')
    monkeypatch.setattr(gui, '_SCRIPT_DIR', tmp_path)
    assert gui._get_app_version() == '20260101-02'


def test_j10b_get_app_version_falls_back_to_today_when_missing(tmp_path, monkeypatch):
    """VERSION 파일이 없으면 오늘 날짜 + '-01'로 폴백해야 한다(자동 증가 메커니즘은
    두지 않는다 — 코디네이터 명시 지시로 과설계 판단)."""
    import missa_gui as gui
    from datetime import datetime
    monkeypatch.setattr(gui, '_SCRIPT_DIR', tmp_path)
    expected = datetime.now().strftime('%Y%m%d') + '-01'
    assert gui._get_app_version() == expected


def _j10_run_click_probe(tmp_path) -> dict:
    """'🔄 업데이트 확인' 텍스트를 가진 Label을 찾아 실제로 `<Button-1>` 이벤트를 발생시켜,
    monkeypatch한 `gui._on_check_update`가 호출되는지 직접 확인한다(단순히 위젯이
    존재하는지가 아니라 실제로 배선됐는지를 검증 — 존재만 확인하면 클릭해도 아무 일도
    안 일어나는 회귀를 놓친다)."""
    out_file = tmp_path / 'j10_click_result.json'
    script = f'''
import sys, json
sys.path.insert(0, {str(BASE)!r})
import tkinter as tk
import missa_gui as gui

_clicked = []
gui._on_check_update = lambda parent: _clicked.append(True)

def _find_update_label(w):
    if w.winfo_class() == 'Label' and '업데이트' in str(w.cget('text')):
        return w
    for c in w.winfo_children():
        found = _find_update_label(c)
        if found is not None:
            return found
    return None

_orig_mainloop = tk.Tk.mainloop

def _fake_mainloop(self, *a, **k):
    def _probe():
        try:
            target = _find_update_label(self)
            if target is not None:
                # x/y 없이 event_generate('<Button-1>')만 호출하면 이 환경(Tk on
                # Windows)에서는 바인딩이 조용히 트리거되지 않는다(직접 최소 재현으로
                # 확인) — 위젯 내부 좌표를 명시해야 실제 클릭처럼 처리된다. 좌표를
                # 줘도 update_idletasks()로 위젯이 실제 배치(geometry)를 먼저 확정하지
                # 않으면 이벤트가 조용히 무시되는 경우가 있었다(직접 재현) — 반드시
                # 이벤트 발생 전에 호출한다.
                target.update_idletasks()
                target.event_generate('<Button-1>', x=5, y=5)
                self.update()
        finally:
            self.destroy()
    self.after(60, _probe)
    return _orig_mainloop(self, *a, **k)

tk.Tk.mainloop = _fake_mainloop

try:
    gui._ask_combined_input_popup(mass_type='adult')
except RuntimeError:
    pass

with open({str(out_file)!r}, 'w', encoding='utf-8') as f:
    json.dump({{'clicked': _clicked}}, f)
'''
    script_path = tmp_path / 'j10_click_script.py'
    script_path.write_text(script, encoding='utf-8')
    import subprocess as _j10_subprocess
    import sys as _j10_sys
    import json as _j10_json
    try:
        result = _j10_subprocess.run(
            [_j10_sys.executable, str(script_path)],
            cwd=str(BASE), capture_output=True, text=True, encoding='utf-8', errors='replace',
            timeout=30,
        )
    except _j10_subprocess.TimeoutExpired as e:
        raise AssertionError(
            f"J10 클릭 프로브가 30초 내에 끝나지 않음\nstdout:\n{e.stdout}\nstderr:\n{e.stderr}"
        )
    assert out_file.is_file(), (
        f"J10 클릭 프로브 실패(exit={result.returncode})\nstdout:\n{result.stdout}\nstderr:\n{result.stderr}"
    )
    return _j10_json.loads(out_file.read_text(encoding='utf-8'))


def test_j10c_version_row_shows_version_and_update_link(tmp_path):
    """`_ask_combined_input_popup()`에 버전 텍스트 Label과 '🔄 업데이트 확인' Label이
    둘 다 존재해야 하고, 링크 Label은 cursor='hand2'여야 한다(클릭 가능함을 시각적으로
    알리는 최소 조건)."""
    widgets = _j_run_widget_probe(tmp_path, "gui._ask_combined_input_popup(mass_type='adult')")
    version_labels = [w for w in widgets if w['class'] == 'Label' and str(w.get('text', '')).startswith('v')
                       and any(ch.isdigit() for ch in w.get('text', ''))]
    assert version_labels, widgets
    update_labels = [w for w in widgets if w['class'] == 'Label' and '업데이트' in str(w.get('text', ''))]
    assert len(update_labels) == 1, widgets
    assert update_labels[0]['cursor'] == 'hand2', update_labels[0]


def test_j10d_update_control_is_label_not_button_and_reuses_on_check_update():
    """'업데이트 확인'이 더 이상 tk.Button이 아니라 tk.Label이어야 하고(cursor+bind
    스타일), §J3에서 옮겨둔 `_on_check_update`를 재구현 없이 그대로 재사용해야 한다."""
    import inspect
    import missa_gui as gui
    src = inspect.getsource(gui._ask_combined_input_popup)
    assert "Button(frame, text='🔄" not in src, src
    assert "Button(_version_row" not in src or "text='🔄" not in src, src
    assert '_on_check_update' in src, src
    assert "cursor='hand2'" in src, src


def test_j10e_clicking_update_link_actually_invokes_on_check_update(tmp_path):
    """존재/커서 스타일만으로는 실제 클릭 배선까지 검증되지 않는다 — 위젯이 있어도
    바인딩이 빠지면 눌러도 아무 일도 안 일어나는 조용한 회귀가 가능하므로, 실제
    `<Button-1>` 이벤트를 발생시켜 monkeypatch한 `_on_check_update`가 호출되는지까지
    확인한다."""
    result = _j10_run_click_probe(tmp_path)
    assert result['clicked'] == [True], result


def test_j10f_version_row_not_counted_by_j4_gridded_label_filter(tmp_path):
    """§J4의 파일 행 순서 검증(`test_j4_*`)은 grid_column==0 && gridded인 Label만 센다.
    버전/링크 Label은 별도 Frame 안에 pack()으로 배치돼 gridded=False여야 하므로, J4의
    카운트(성인 5개/청년 4개)를 건드리지 않는다는 걸 이 배치 자체에서도 재확인한다."""
    widgets = _j_run_widget_probe(tmp_path, "gui._ask_combined_input_popup(mass_type='adult')")
    update_labels = [w for w in widgets if w['class'] == 'Label' and '업데이트' in str(w.get('text', ''))]
    assert len(update_labels) == 1
    assert update_labels[0].get('gridded') is False, update_labels[0]


# ═══════════════════════════════════════════════════════════════════════════
# K그룹 — 2026-09-27 OneDrive 브라우저 재귀 프리페치(K1) + 로컬 캐시 재사용(K2)
#
# K1: `_prefetch_onedrive_children()`(§J5)이 3개 허용 루트 폴더의 직계 자식만 채우던 것을
# 모든 하위 폴더까지 재귀적으로 채우도록 확장한다. 실측(코디네이터 요청, 실제 Graph API로
# 3개 루트 폴더 트리를 순차 재귀 조회): 총 API 호출 9회·하위폴더 6개(루트 3개 제외)·
# 파일 31개·최대 깊이 2·순차 총 22.04초(평균 2.449초/호출, 네트워크 왕복이 대부분) —
# 트리가 이 정도로 작으면 동시성 없이 순차 재귀만으로 429/타임아웃 위험 없이 충분히
# 빠르다(백그라운드 스레드에서 실행되므로 사용자가 날짜/미사유형을 입력하는 동안 끝난다).
# 트리가 실제로 훨씬 커지면 재검토 필요(코디네이터 지시).
#
# K2: OneDrive 브라우저로 받는 모든 파일(공지사항/시작기도/미사후기도/성가 PPT 등, 파일
# 종류로 분기하지 않음)에 대해, 날짜 폴더와 별개인 공유 캐시 폴더(`onedrive_cache/`)에
# 원본을 보관하고 lastModifiedDateTime이 캐시 기록 시점과 정확히 같을 때만 재다운로드를
# 스킵한다. 조금이라도 다르면(더 최신이든, 기록/캐시본이 없든, 메타데이터 조회 자체가
# 불확실하든) 무조건 재다운로드한다 — "이름이 같다"는 절대 재사용 근거가 아니다(코디네이터
# 명시 불변조건). 캐시 히트 시에도 날짜 폴더(`dest_dir`)에 복사해 기존 파이프라인이 기대하는
# 경로 가정을 깨지 않는다.
# ═══════════════════════════════════════════════════════════════════════════


# ═══════════════════════════════════════════════════════════════════════════
# L그룹 — 2026-09-27 OneDrive 프리페치 트리거 시점 앞당기기(L1) + 파일 선택 즉시 반환·
# 백그라운드 다운로드(L2)
#
# L1: K1의 재귀 프리페치가 "찾아보기"를 눌러 `_ask_onedrive_file_browser_popup()`이 열릴
# 때 트리거되던 것을, 첫 입력창(`_ask_combined_input_popup()`) 표시 직후(mass_type='youth'
# 일 때만)로 앞당긴다. 로그인/네트워크 실패는 조용히 삼키고 첫 입력창 자체는 항상 즉시
# 뜬다(백그라운드 스레드라 구조적으로 블로킹하지 않음).
# ═══════════════════════════════════════════════════════════════════════════


# ═══════════════════════════════════════════════════════════════════════════
# L2 — 2026-09-27 OneDrive 파일 선택 즉시 반환, 다운로드는 백그라운드
#
# `_download_with_retries()`(순수 함수, 자동 재시도 담당)는 GUI 없이 직접 단위 테스트한다.
# `_browse()`/`on_ok()`의 "즉시 반환·확인 시 대기·재선택 최신성 보장" 통합 동작은
# `_ask_onedrive_file_browser_popup()`과 `_download_with_retries()`를 모킹한 서브프로세스
# 팝업 프로브로 검증한다(실제 Treeview 탐색은 H/I/K1에서 이미 커버됨 — 이 배치의 관심사는
# "선택 이후 무슨 일이 일어나는가"이지 "어떻게 탐색해서 고르는가"가 아니다).
# ═══════════════════════════════════════════════════════════════════════════


# --- 2026-10-18 나주노 최근 생성물 결함 수정(479/519/692/468) -----------------------
# 근거: 흐린/기운 스캔의 오선 탈락(479 5줄 중 3줄), 시스템 경계가 가사 2째 줄을 자름(692),
# 시스템 아래 2절 가사 별도 세그먼트 누락(468), 저작권 캐시 미등록 곡의 무음 생략(479/519/692).
_FIX_EXPECTED_SYSTEMS = {479: 5, 519: 5, 692: 4, 468: 6, 603: 8}


@pytest.fixture(scope="module")
def fix_bands():
    import numpy as _np
    result = {}
    d = _fitz.open(hp.SOURCES["나주노"]["pdf"])
    try:
        for num in _FIX_EXPECTED_SYSTEMS:
            gray = _np.asarray(hp.render_region(
                d[hp.find_song_page("나주노", num)], None, hp.RENDER_DPI).convert("L"))
            result[num] = {"gray": gray, "bands": hp.detect_system_bands(gray)}
    finally:
        d.close()
    return result


@pytest.mark.parametrize("num", sorted(_FIX_EXPECTED_SYSTEMS))
@pytest.mark.skipif(not _YOUTH_PDFS_AVAILABLE, reason=_YOUTH_SKIP_REASON)
def test_F1_system_counts_faint_scans(fix_bands, num):
    """F1: 오선이 흐린 곡(479)·가사 보충줄이 있는 곡(468)도 원본 시스템 수와 같다."""
    assert len(fix_bands[num]["bands"]) == _FIX_EXPECTED_SYSTEMS[num]


@pytest.mark.parametrize("num", sorted(_FIX_EXPECTED_SYSTEMS))
@pytest.mark.skipif(not _YOUTH_PDFS_AVAILABLE, reason=_YOUTH_SKIP_REASON)
def test_F2_no_ink_left_between_bands(fix_bands, num):
    """F2: 인접한 두 밴드 사이(밴드 밖)에 잉크가 남지 않는다 — 남으면 그 부분이 슬라이드에서
    잘린 것이다(692 가사 2째 줄, 468 2절 가사)."""
    gray = fix_bands[num]["gray"]
    bands = fix_bands[num]["bands"]
    for (_t0, b0), (t1, _b1) in zip(bands, bands[1:]):
        # 스캔 얼룩(행당 ≤7px)은 무시, 가사/악보 줄은 행당 수십 px 이상
        between = (gray[b0 + 1:t1] < 128).sum(axis=1)
        assert not (between >= 15).any(), f"{num}: 밴드 사이 [{b0 + 1},{t1}) 잉크 잘림"


@pytest.mark.parametrize("num", [479, 519, 692])
@pytest.mark.skipif(not _YOUTH_PDFS_AVAILABLE, reason=_YOUTH_SKIP_REASON)
def test_F3_naju_copyright_present(num):
    """F3: 479/519/692는 원본 하단에 저작권이 있으므로 크롭이 있어야 한다."""
    d = _fitz.open(hp.SOURCES["나주노"]["pdf"])
    try:
        assert hp.resolve_copyright_crop("나주노", num, d[hp.find_song_page("나주노", num)]) is not None
    finally:
        d.close()


@pytest.mark.skipif(not _YOUTH_PDFS_AVAILABLE, reason=_YOUTH_SKIP_REASON)
def test_F4_naju_copyright_fallback_when_cache_misses(monkeypatch):
    """F4: 캐시에 없는 곡도 저작권을 조용히 생략하지 않고 하단 첫 줄 폴백 크롭을 쓴다."""
    monkeypatch.setattr(hp, "_load_naju_cache", lambda: {})
    d = _fitz.open(hp.SOURCES["나주노"]["pdf"])
    try:
        crop = hp.resolve_copyright_crop("나주노", 479, d[hp.find_song_page("나주노", 479)])
    finally:
        d.close()
    assert crop is not None and crop.width > 300 and 15 < crop.height < 120


# ═════════════════════════════════════════════════════════════════════════════
# OneDrive 폴더 구조 개편 대응(2026-10-18) — docs/SPECS/SPEC_OneDrive_폴더구조_개편.md
# 접두 test_fs_. 루트 3종(공용/성인/청년)·경로 규칙·찾아보기 기억·Graph 폴백 폐지.
# ═════════════════════════════════════════════════════════════════════════════
import json as _fs_json


@pytest.fixture
def fs_cfg(tmp_path, monkeypatch):
    import missa_gui as gui
    path = tmp_path / "config.json"
    monkeypatch.setattr(gui, "CONFIG_FILE", path)
    gui._reset_ppt_roots()
    return path


def _fs_write_cfg(cfg, **kw):
    cfg.write_text(_fs_json.dumps(kw, ensure_ascii=False), encoding="utf-8")


def _fs_read_cfg(cfg):
    return _fs_json.loads(cfg.read_text(encoding="utf-8"))


@pytest.fixture
def fs_tree(tmp_path):
    """새 구조의 'PPT 문서'(마스터 PC). 반환: (parent, common, adult, youth)."""
    parent = tmp_path / "PPT 문서"
    common = parent / "00.공용 PPT"
    adult = parent / "10.성인 미사"
    youth = parent / "20.청년 미사"
    for p in (common / "09.가톨릭 성가" / "성가-악보버전", common / "13.기도문",
              common / "11.공지사항 PPT문서",
              adult / "10.기타 PPT문서" / "PPT 자동화" / "1.Template", adult / "14.화답송 악보",
              youth / "1.Template", youth / "2.성가" / "나주노 성가", youth / "2.성가" / "야훼이레 성가"):
        p.mkdir(parents=True)
    return parent, common, adult, youth


def test_fs_root_validity_markers_and_space_insensitive(fs_tree, tmp_path):
    import missa_gui as gui
    _p, common, adult, youth = fs_tree
    assert gui.is_valid_ppt_root("common", common)
    assert gui.is_valid_ppt_root("adult", adult)
    assert gui.is_valid_ppt_root("youth", youth)
    # 역할이 맞지 않는 폴더는 무효
    assert not gui.is_valid_ppt_root("youth", adult)
    assert not gui.is_valid_ppt_root("common", tmp_path / "없음")
    # 공백 유무 차이 무시
    spaced = tmp_path / "x" / "20.청년 미사"
    (spaced / "2. 성가").mkdir(parents=True)
    assert gui.is_valid_ppt_root("youth", spaced)


def test_fs_youth_run_asks_only_common_and_youth_once_then_never_again(fs_cfg, fs_tree, monkeypatch):
    """인수 기준 1: 서로 다른 위치의 두 루트를 한 번씩만 선택, 성인 폴더는 묻지 않는다."""
    import missa_gui as gui
    _p, common, _a, youth = fs_tree
    asked = []
    answers = {"00.공용 PPT": str(common), "20.청년 미사": str(youth)}

    def fake(**k):
        asked.append(k["dialog_title"])
        for name, val in answers.items():
            if name in k["dialog_title"]:
                return val
        pytest.fail(f"예상 밖 선택창: {k['dialog_title']}")

    monkeypatch.setattr(gui, "_ask_onedrive_path_popup", fake)
    assert gui.ensure_ppt_roots("youth", ask_if_missing=True) is True
    assert len(asked) == 2 and not any("성인" in t for t in asked)
    saved = _fs_read_cfg(fs_cfg)
    assert Path(saved["onedrive_common_folder"]) == common
    assert Path(saved["onedrive_youth_folder"]) == youth
    assert "onedrive_adult_folder" not in saved
    # 재실행(캐시 초기화 = 새 프로세스): 다시 묻지 않는다
    gui._reset_ppt_roots()
    monkeypatch.setattr(gui, "_ask_onedrive_path_popup", lambda **k: pytest.fail("다시 묻지 않아야 함"))
    assert gui.ensure_ppt_roots("youth", ask_if_missing=True) is True
    assert gui.get_root("youth") == youth and gui.get_root("common") == common


def test_fs_adult_run_does_not_ask_youth_root(fs_cfg, fs_tree, monkeypatch):
    import missa_gui as gui
    _p, common, adult, _y = fs_tree
    _fs_write_cfg(fs_cfg, onedrive_common_folder=str(common), onedrive_adult_folder=str(adult))
    monkeypatch.setattr(gui, "_ask_onedrive_path_popup", lambda **k: pytest.fail("묻지 않아야 함"))
    assert gui.ensure_ppt_roots("adult", ask_if_missing=True) is True
    assert gui.get_root("youth") is None


def test_fs_master_picks_parent_fills_all_three_roots(fs_cfg, fs_tree, monkeypatch):
    """인수 기준 2: 'PPT 문서'를 고르면 세 루트가 한 번에 채워지고 다시 묻지 않는다."""
    import missa_gui as gui
    parent, common, adult, youth = fs_tree
    calls = []
    monkeypatch.setattr(gui, "_ask_onedrive_path_popup", lambda **k: calls.append(1) or str(parent))
    assert gui.ensure_ppt_roots("youth", ask_if_missing=True) is True
    assert len(calls) == 1
    saved = _fs_read_cfg(fs_cfg)
    assert (Path(saved["onedrive_common_folder"]), Path(saved["onedrive_adult_folder"]),
            Path(saved["onedrive_youth_folder"])) == (common, adult, youth)
    assert gui.ensure_ppt_roots("adult", ask_if_missing=True) is True
    assert len(calls) == 1


def test_fs_invalid_selection_reasks_with_marker_hint_and_cancel_returns_false(fs_cfg, tmp_path, monkeypatch):
    import missa_gui as gui
    bad = tmp_path / "bad"
    bad.mkdir()
    answers = iter([str(bad), ""])
    msgs = []

    def fake(**k):
        msgs.append(k["message"])
        return next(answers)

    monkeypatch.setattr(gui, "_ask_onedrive_path_popup", fake)
    assert gui.ensure_ppt_roots("youth", ask_if_missing=True) is False
    assert len(msgs) == 2 and "09.가톨릭 성가" in msgs[1]
    assert not fs_cfg.exists()


def test_fs_without_ask_never_opens_popup(fs_cfg, monkeypatch):
    import missa_gui as gui
    monkeypatch.setattr(gui, "_ask_onedrive_path_popup", lambda **k: pytest.fail("팝업 금지"))
    assert gui.ensure_ppt_roots("adult") is False
    assert gui.ensure_ppt_roots("youth") is False


def test_fs_legacy_ppt_folder_with_new_structure_is_migrated(fs_cfg, fs_tree, monkeypatch):
    """기존 설정 이관: 옛 키 onedrive_ppt_folder가 새 구조를 품고 있으면 세 루트를 자동 유도."""
    import missa_gui as gui
    parent, common, adult, youth = fs_tree
    _fs_write_cfg(fs_cfg, onedrive_ppt_folder=str(parent))
    monkeypatch.setattr(gui, "_ask_onedrive_path_popup", lambda **k: pytest.fail("묻지 않아야 함"))
    assert gui.ensure_ppt_roots("adult", ask_if_missing=True) is True
    saved = _fs_read_cfg(fs_cfg)
    assert Path(saved["onedrive_adult_folder"]) == adult and Path(saved["onedrive_common_folder"]) == common
    assert Path(saved["onedrive_youth_folder"]) == youth


def test_fs_old_structure_config_is_invalid_and_leads_to_picker(fs_cfg, tmp_path, monkeypatch):
    """인수 기준 4: 옛 구조(09/10/13/20이 'PPT 문서' 바로 아래)를 가리키는 config는 오류 없이 선택창으로."""
    import missa_gui as gui
    old = tmp_path / "PPT 문서"
    for p in ("09.가톨릭 성가/성가-악보버전", "10.기타 PPT문서", "13.기도문", "20.청년 미사/2.성가"):
        (old / p).mkdir(parents=True)
    _fs_write_cfg(fs_cfg, onedrive_ppt_folder=str(old),
                  onedrive_hymn_folder=str(old / "09.가톨릭 성가" / "성가-악보버전"),
                  onedrive_youth_hymn_folder=str(old / "20.청년 미사" / "2.성가"))
    asked = []
    monkeypatch.setattr(gui, "_ask_onedrive_path_popup", lambda **k: asked.append(k["dialog_title"]) or "")
    assert gui.ensure_ppt_roots("adult", ask_if_missing=True) is False
    assert asked  # 선택창으로 안내됨(예외 없음)
    assert gui.get_root("adult") is None and gui.get_root("common") is None


def test_fs_saved_root_that_vanished_is_asked_again(fs_cfg, fs_tree, tmp_path, monkeypatch):
    import missa_gui as gui
    _p, common, _a, youth = fs_tree
    _fs_write_cfg(fs_cfg, onedrive_common_folder=str(common), onedrive_youth_folder=str(tmp_path / "사라짐"))
    asked = []

    def fake(**k):
        asked.append(k["dialog_title"])
        return str(youth)

    monkeypatch.setattr(gui, "_ask_onedrive_path_popup", fake)
    assert gui.ensure_ppt_roots("youth", ask_if_missing=True) is True
    assert len(asked) == 1 and "20.청년 미사" in asked[0]


def test_fs_path_rules_follow_new_locations(fs_cfg, fs_tree):
    """인수 기준 3: 경로 규칙 표(성가 검색·신규 저장·결과 복사)."""
    import missa_gui as gui
    _p, common, adult, youth = fs_tree
    gui._ROOTS.update(common=common, adult=adult, youth=youth)
    assert gui.ppt_subdir("catholic_hymn") == common / "09.가톨릭 성가" / "성가-악보버전"
    assert gui.get_onedrive_hymn_folder() == common / "09.가톨릭 성가" / "성가-악보버전"
    assert gui.ppt_subdir("youth_hymn", "나주노 성가") == youth / "2.성가" / "나주노 성가"
    assert gui.result_copy_dir("adult", "2026") == adult / "2026"
    assert gui.result_copy_dir("youth", "2026") == youth / "2026"


def test_fs_hymn_folder_without_common_root_raises_runtime_error(fs_cfg):
    import missa_gui as gui
    with pytest.raises(RuntimeError):
        gui.get_onedrive_hymn_folder()


def test_fs_browse_defaults_per_row(fs_cfg, fs_tree):
    import missa_gui as gui
    _p, common, adult, youth = fs_tree
    gui._ROOTS.update(common=common, adult=adult, youth=youth)
    f = gui.browse_initial_dir
    for mt in ("adult", "youth"):
        assert f(mt, "시작기도") == common / "13.기도문"
        assert f(mt, "미사후기도") == common / "13.기도문"
        assert f(mt, "공지사항") == common / "11.공지사항 PPT문서"
    assert f("adult", "ref_pptx") == adult / "10.기타 PPT문서" / "PPT 자동화" / "1.Template"
    assert f("youth", "ref_pptx") == youth / "1.Template"
    assert f("adult", "화답송_pptx") == adult / "14.화답송 악보"


def test_fs_browse_default_falls_back_to_deepest_existing_then_root(fs_cfg, tmp_path):
    import missa_gui as gui
    adult = tmp_path / "10.성인 미사"
    (adult / "10.기타 PPT문서").mkdir(parents=True)
    gui._ROOTS.update(adult=adult)
    assert gui.browse_initial_dir("adult", "ref_pptx") == adult / "10.기타 PPT문서"
    assert gui.browse_initial_dir("adult", "화답송_pptx") == adult
    assert gui.browse_initial_dir("adult", "시작기도") is None  # 공용 루트 없음


def test_fs_browse_memory_is_per_row_and_ref_pptx_split_by_mass_type(fs_cfg, fs_tree, tmp_path):
    """인수 기준 5 + 결정 4: 행별 기억, 참조 PPT는 성인/청년 구분."""
    import missa_gui as gui
    _p, common, adult, youth = fs_tree
    gui._ROOTS.update(common=common, adult=adult, youth=youth)
    other = tmp_path / "다른 곳"
    other.mkdir()
    gui.remember_browse_dir("adult", "ref_pptx", other / "a.pptx")
    assert gui.browse_initial_dir("adult", "ref_pptx") == other
    assert gui.browse_initial_dir("youth", "ref_pptx") == youth / "1.Template"   # 청년은 영향 없음
    assert gui.browse_initial_dir("adult", "시작기도") == common / "13.기도문"      # 다른 행 영향 없음
    # 저장 폴더가 사라지면 기본 폴더로
    other.rmdir()
    assert gui.browse_initial_dir("adult", "ref_pptx") == adult / "10.기타 PPT문서" / "PPT 자동화" / "1.Template"


def test_fs_browse_memory_works_without_root_and_absorbs_legacy_psalm_dir(fs_cfg, tmp_path):
    import missa_gui as gui
    legacy = tmp_path / "예전 화답송"
    legacy.mkdir()
    _fs_write_cfg(fs_cfg, last_psalm_dir=str(legacy))
    assert gui.browse_initial_dir("adult", "화답송_pptx") == legacy   # 옛 키 흡수
    newer = tmp_path / "새 화답송"
    newer.mkdir()
    gui.remember_browse_dir("adult", "화답송_pptx", newer / "x.png")
    assert gui.browse_initial_dir("adult", "화답송_pptx") == newer     # 행별 기억이 우선
    assert "last_psalm_dir" not in _fs_read_cfg(fs_cfg)


def test_fs_graph_fallback_is_removed():
    """결정 1: OneDrive 로그인(Graph) 폴백·원격 브라우저·프리페치·캐시가 코드에서 사라졌다."""
    import missa_gui as gui
    for name in ("ensure_onedrive_login", "_mirror_onedrive_folder", "_resolve_onedrive_folder",
                 "get_onedrive_youth_hymn_root", "get_youth_ppt_mode", "_YOUTH_PPT_STATE",
                 "_ask_onedrive_file_browser_popup", "_start_onedrive_prefetch_thread",
                 "_ensure_onedrive_blob_fresh", "_find_onedrive_file", "find_youth_onedrive_hymn_file",
                 "_show_onedrive_device_code_window", "get_adult_ppt_folder"):
        assert not hasattr(gui, name), name
    assert not (BASE / "missa_onedrive.py").exists()


def test_fs_run_gui_mode_cancelled_root_selection_exits_without_popup(monkeypatch):
    """취소하면 안내만 하고 종료(입력창은 열리지 않는다)."""
    import missa_to_ppt as m
    shown = []
    monkeypatch.setattr(m, "ensure_ppt_roots", lambda mass_type, ask_if_missing=False: False)
    monkeypatch.setattr(m, "_show_root_required_message", lambda mt: shown.append(mt))
    monkeypatch.setattr(m, "_ask_combined_input_popup",
                        lambda mass_type: pytest.fail("입력창이 열리면 안 됨"))
    with pytest.raises(SystemExit):
        m._run_gui_mode("청년")
    assert shown == ["youth"]


def test_fs_new_youth_hymn_saved_under_youth_root(fs_cfg, fs_tree, monkeypatch):
    """인수 기준 3: 신규 성가 PPT 저장 위치 = {청년}/2.성가/{나주노 성가}."""
    import missa_content_updaters as cu
    import missa_gui as gui
    import missa_youth_hymn_pdf as yh
    from pptx import Presentation
    _p, common, adult, youth = fs_tree
    gui._ROOTS.update(common=common, adult=adult, youth=youth)
    monkeypatch.setattr(yh, "find_song_title", lambda src, n: "나의 고백")
    monkeypatch.setattr(yh, "build_hymn_pptx", lambda *a, **k: Presentation())
    monkeypatch.setattr(yh, "build_header_runs", lambda *a, **k: None)
    _prs, title = cu.resolve_youth_hymn_pptx({"출처": "나주노", "번호": 172, "제목": None}, "입당")
    assert title == "나의 고백"
    assert (youth / "2.성가" / "나주노 성가" / "나주노 성가 172 나의 고백.pptx").exists()


def test_fs_catholic_hymn_lookup_uses_common_root(fs_cfg, fs_tree):
    import missa_gui as gui
    _p, common, adult, youth = fs_tree
    gui._ROOTS.update(common=common, adult=adult, youth=youth)
    f = common / "09.가톨릭 성가" / "성가-악보버전" / "성가 96 하느님 약속하신.pptx"
    f.write_bytes(b"x")
    assert gui.find_youth_hymn_file("catholic_hymn", None, r"성가 96(?!\d)") == f
    assert gui.find_youth_hymn_file("catholic_hymn", None, r"성가 97(?!\d)") is None


# --- 2026-10-18 독립 리뷰 지적 반영 + 공유 페이지(곡 하나만 생성) -----------------------
@pytest.mark.parametrize("num,expected_top", [(362, 305), (173, 358), (146, 347), (267, 355), (447, 346)])
@pytest.mark.skipif(not _YOUTH_PDFS_AVAILABLE, reason=_YOUTH_SKIP_REASON)
def test_S7_first_band_excludes_composer_line(num, expected_top):
    """S7: 첫 시스템 상단이 작곡가명 아래(회귀 기준값 = 수정 전 HEAD 실측)에서 시작한다.
    `_system_top`이 오선 블록 앞 블록(작곡가)을 다시 포함하던 결함(독립 리뷰)의 회귀 방지."""
    import numpy as _np
    d = _fitz.open(hp.SOURCES["나주노"]["pdf"])
    try:
        gray = _np.asarray(hp.render_region(
            d[hp.find_song_page("나주노", num)], None, hp.RENDER_DPI).convert("L"))
    finally:
        d.close()
    assert hp.detect_system_bands(gray)[0][0] >= expected_top - 3


@pytest.mark.skipif(not _YOUTH_PDFS_AVAILABLE, reason=_YOUTH_SKIP_REASON)
def test_F4b_fallback_copyright_is_single_line(monkeypatch):
    """F4b: 폴백 크롭은 둘째 줄(All rights reserved…)을 포함하지 않는다(479·692: 줄이 붙어 있음)."""
    monkeypatch.setattr(hp, "_load_naju_cache", lambda: {})
    d = _fitz.open(hp.SOURCES["나주노"]["pdf"])
    try:
        for num in (479, 692):
            crop = hp.resolve_copyright_crop("나주노", num, d[hp.find_song_page("나주노", num)])
            assert crop is not None and crop.height <= 45, (num, crop.size)
    finally:
        d.close()


def _shared_song_numbers():
    d = _fitz.open(hp.SOURCES["나주노"]["pdf"])
    try:
        import re as _re
        pages = {}
        for _l, t, p in d.get_toc():
            m = _re.match(r"^\s*(\d+)\b", t)
            if m:
                pages.setdefault(p - 1, []).append(int(m.group(1)))
    finally:
        d.close()
    return [(pi, v) for pi, v in sorted(pages.items()) if len(v) > 1]


@pytest.mark.skipif(not _YOUTH_PDFS_AVAILABLE, reason=_YOUTH_SKIP_REASON)
def test_G1_all_shared_page_songs_have_own_region_and_partition_the_page():
    """G1: 곡이 둘 이상인 모든 페이지의 곡마다 y_range가 캐시에 있고, 곡 순서대로 서로 겹치지 않으며
    페이지 전체를 빈틈없이 나눈다(곡 하나만 생성하기 위한 전제)."""
    cache = hp._load_naju_cache()
    for _pi, nums in _shared_song_numbers():
        ranges = [cache[str(n)]["y_range"] for n in nums]
        assert ranges[0][0] == 0.0 and ranges[-1][1] == 1.0, nums
        for (a0, a1), (b0, b1) in zip(ranges, ranges[1:]):
            assert a1 == b0 and a0 < a1, nums


@pytest.mark.skipif(not _YOUTH_PDFS_AVAILABLE, reason=_YOUTH_SKIP_REASON)
def test_G2_shared_page_song_image_contains_only_that_song():
    """G2: 14번 영역 이미지는 y_range 높이만큼이고, 그 안의 시스템 밴드는 15번 영역과 겹치지 않는다."""
    cache = hp._load_naju_cache()
    d = _fitz.open(hp.SOURCES["나주노"]["pdf"])
    try:
        page = d[hp.find_song_page("나주노", 14)]
        full_h = hp.render_region(page, None, hp.RENDER_DPI).height
        img14 = hp.naju_song_region_image(page, 14)
        img15 = hp.naju_song_region_image(page, 15)
    finally:
        d.close()
    y14, y15 = cache["14"]["y_range"], cache["15"]["y_range"]
    assert img14.height <= round((y14[1] - y14[0]) * full_h) + 2
    assert img15.height <= round((y15[1] - y15[0]) * full_h) + 2
    assert img14.height + img15.height < full_h  # 저작권 줄 등은 제외돼 합이 전체보다 작다


@pytest.mark.skipif(not _YOUTH_PDFS_AVAILABLE, reason=_YOUTH_SKIP_REASON)
def test_G3_shared_page_without_registered_range_raises(monkeypatch):
    """G3: 공유 페이지 곡인데 y_range가 없으면 다른 곡 악보가 섞이므로 조용히 진행하지 않고 예외."""
    monkeypatch.setattr(hp, "_load_naju_cache", lambda: {})
    d = _fitz.open(hp.SOURCES["나주노"]["pdf"])
    try:
        with pytest.raises(RuntimeError):
            hp.naju_song_region_image(d[hp.find_song_page("나주노", 14)], 14)
    finally:
        d.close()


@pytest.mark.skipif(not _YOUTH_PDFS_AVAILABLE, reason=_YOUTH_SKIP_REASON)
def test_G4_copyright_bbox_lies_inside_own_song_region():
    """G4: 공유 페이지 곡의 저작권 bbox는 그 곡 자신의 y_range 안에 있고, 곡 영역 이미지는 그 위에서 끝난다."""
    cache = hp._load_naju_cache()
    for _pi, nums in _shared_song_numbers():
        for n in nums:
            e = cache[str(n)]
            if e.get("has_copyright"):
                y0, y1 = e["y_range"]
                assert y0 <= e["bbox"][1] < y1, n


_G5_SONGS = [479, 519, 692, 468, 603, 107, 332, 255, 35, 68, 78, 358, 359, 11, 14, 15, 531, 212]


@pytest.mark.parametrize("num", _G5_SONGS)
@pytest.mark.skipif(not _YOUTH_PDFS_AVAILABLE, reason=_YOUTH_SKIP_REASON)
def test_G5_no_music_or_lyric_ink_outside_bands(num):
    """G5: 곡 영역 안에서 첫 밴드 상단~마지막 밴드 하단 사이의 밴드 밖, 그리고 마지막 밴드 아래에
    글줄(행당 잉크 ≥15px)이 남지 않는다 — 남으면 슬라이드에서 잘린 악보/가사다(반복 종지 괄호 잘림 107·332·255,
    공유 페이지 마지막 가사줄 누락 68·78·358·359, 35번 2째 가사줄; 독립 재리뷰 지적의 부정 단언)."""
    d = _fitz.open(hp.SOURCES["나주노"]["pdf"])
    try:
        region, bands = hp.naju_song_bands(d[hp.find_song_page("나주노", num)], num)
    finally:
        d.close()
    import numpy as _np
    ink_rows = (_np.asarray(region.convert("L")) < 128).sum(axis=1) >= 15
    covered = _np.zeros(len(ink_rows), dtype=bool)
    for t, b in bands:
        covered[t:b + 1] = True
    stray = _np.where(ink_rows & ~covered)[0]
    stray = stray[stray > bands[0][0]]   # 첫 밴드 위(제목·작곡가)는 제외 대상이 맞다
    assert len(stray) < 3, f"{num}: 밴드 밖 글줄 {len(stray)}행 (예: {stray[:5].tolist()})"


@pytest.mark.skipif(not _YOUTH_PDFS_AVAILABLE, reason=_YOUTH_SKIP_REASON)
def test_G6_known_false_copyright_songs_now_have_copyright():
    """G6: 재리뷰에서 '저작권이 있는데 false'로 확인된 곡(19·24·29·34·348, 공유 페이지 11)은 저작권 bbox가 있다."""
    cache = hp._load_naju_cache()
    for n in (19, 24, 29, 34, 348, 11):
        assert cache[str(n)].get("has_copyright"), n


@pytest.mark.skipif(not _YOUTH_PDFS_AVAILABLE, reason=_YOUTH_SKIP_REASON)
def test_G7_region_image_ignores_degenerate_bbox(monkeypatch):
    """G7: bbox 상단이 영역 상단 근처(또는 위)여도 곡 영역이 1px로 무너지지 않는다."""
    cache = dict(hp._load_naju_cache())
    cache["479"] = {"has_copyright": True, "bbox": [0.1, 0.0, 0.2, 0.01], "trigger": ""}
    monkeypatch.setattr(hp, "_load_naju_cache", lambda: cache)
    d = _fitz.open(hp.SOURCES["나주노"]["pdf"])
    try:
        region = hp.naju_song_region_image(d[hp.find_song_page("나주노", 479)], 479)
    finally:
        d.close()
    assert region.height >= hp._MIN_REGION_H
