# -*- coding: utf-8 -*-
"""프로그레션 테스트 — 아직 안정화되지 않은 신규 기능을 구현 전에 명세한다.

현재 진행 중: 화답송 악보 이미지(PNG/JPG) 입력 지원.
- 1차 마일스톤(missa_psalm_score_image.py — 이미지 → 슬라이드 변환): 완료·사용자 승인.
- 2차 마일스톤(missa_to_ppt.py 배선 — find_files/CLI 오버라이드/입력창/update_화답송):
  이번 라운드. 설계서 §4(통합 지점) 그대로 구현.

설계서: _workspace/01_architect_design.md §4(통합 지점), §7(1차 마일스톤 16개 테스트 대상 행동).
기대값은 3개 예제 세트(20260809/20260816/20260823) 실측 정답이다(추측 아님).

안정화되면 regression-qa가 test_missa_regression.py로 승격하고 이 파일에서 제거한다.
"""
import io
from pathlib import Path

import numpy as np
import pytest
from PIL import Image

import missa_psalm_score_image as ps

BASE = Path(__file__).resolve().parent
REF = BASE / "reference" / "화답송악보"
TEMPLATE = BASE / "assets" / "화답송_악보_template.pptx"

# 세트별 원본 이미지 경로와 실측 기대값 (설계서 §0.1, §2.3, §7)
CASES = {
    "20260809": {
        "img": REF / "20260809" / "20260809_043842947.jpg",
        "crop_left": 372,
        "crop_right": 3906,
        "split": 2146,
        "top_ratio": 0.502,
    },
    "20260816": {
        "img": REF / "20260816" / "20260816_043822155.png",
        "crop_left": 223,
        "crop_right": 1953,
        "split": 1201,
        "top_ratio": 0.565,
    },
    "20260823": {
        "img": REF / "20260823" / "20260823_043736869.png",
        "crop_left": 372,
        "crop_right": 3906,
        # split=2304(top_ratio 0.547)는 이전 라운드의 오검출 버그가 낳은 값이었다 — 실제로는
        # 8분음표 기둥+깃발(스템+플래그)이었고 진짜 마디선이 아니었다(사용자가 실제 생성물을
        # 육안 확대 검토해 발견, _workspace/02c 라운드). detect_barlines()의 노트헤드/스템 인접
        # 배제 필터 추가 후 실측 재검증한 참값은 split=1971(진짜 마디선, 확대 이미지로 확인:
        # 좌우가 깨끗한 공백인 순수 수직선). "값이 그럴듯하게 보인다"만으로 설계서 수치를
        # 맹신하지 않고, 매 라운드 실제 이미지 확대 검토로 재검증해야 한다는 교훈.
        "split": 1971,
        "top_ratio": 0.452,
    },
}


@pytest.fixture(scope="module", params=list(CASES))
def case(request):
    c = CASES[request.param]
    if not c["img"].is_file():
        pytest.skip(
            f"화답송 악보 참조 이미지가 로컬에 없음(reference/화답송악보/, git 미포함): {c['img']}"
        )
    c = dict(c, name=request.param)
    return c


@pytest.fixture(scope="module")
def gray(case):
    return ps.load_gray_array(case["img"])


# ---------------------------------------------------------------------------
# 이미지 처리 (PowerPoint 불필요, 순수 함수 — 빠른 회귀 후보)
# ---------------------------------------------------------------------------

def test_01_watermark_crop_start(case, gray):
    """워터마크 크롭 시작 컬럼이 분리 공백 띠 우측에 온다 (±5px)."""
    x = ps.detect_watermark_crop_x(gray)
    assert abs(x - case["crop_left"]) <= 5, (case["name"], x)


def test_02_watermark_excluded_from_crop(case, gray):
    """워터마크가 크롭 결과에서 제외된다.

    (밀도 비교는 부적합 — 실측상 오선보 시작부 clef/조표가 워터마크 글자보다 잉크가
    더 촘촘하다.) 대신 두 사실로 강하게 단언한다: (1) 제거되는 좌측 영역에 실제
    워터마크 잉크가 존재하고(제거할 대상이 있음), (2) 크롭 경계가 분리 공백 띠 위에
    정확히 놓여 크롭 직전 컬럼이 완전 백색이다 → 워터마크 컬럼은 전부 크롭 좌측에 남아
    결과에서 빠진다.
    """
    crop_x = ps.detect_watermark_crop_x(gray)
    ink = gray < 128
    assert crop_x > 0, case["name"]
    assert ink[:, 10:crop_x].sum() > 0, ("제거 영역에 워터마크 잉크가 있어야 함", case["name"])
    assert ink[:, crop_x - 1].sum() == 0, ("크롭 경계가 공백 띠 위에 놓여야 함", case["name"])


def test_03_content_right_edge(case, gray):
    """오선보 유효 우측끝 검출 — 트레일링 공백 제거 (±5px)."""
    left, right, top, bottom = ps.detect_content_bounds(gray)
    assert abs(right - case["crop_right"]) <= 5, (case["name"], right)


def test_04_split_on_barline(case, gray):
    """분할점이 마디 경계(바라인) 위에 온다 — 중앙창 안 + 탐지 바라인 집합의 원소 (±20px)."""
    crop_left = ps.detect_watermark_crop_x(gray)
    left, right, top, bottom = ps.detect_content_bounds(gray)
    band = ps.staff_band(gray)
    barlines = ps.detect_barlines(gray, band)
    split = ps.choose_split_x(barlines, crop_left, right)
    assert abs(split - case["split"]) <= 20, (case["name"], split)
    # 중앙창 [0.30, 0.70] 안
    content_w = right - crop_left
    assert crop_left + 0.30 * content_w <= split <= crop_left + 0.70 * content_w
    # 폴백이 아닌 한 실제 탐지 바라인 중 하나여야 한다
    assert any(abs(split - b) <= 20 for b in barlines), (case["name"], barlines)


def test_05_balanced_two_lines(case, gray):
    """두 줄 폭 비율이 균형: |상단폭-하단폭|/content_w ≤ 0.15, 상단비율 실측 일치(±0.03)."""
    crop_left = ps.detect_watermark_crop_x(gray)
    left, right, top, bottom = ps.detect_content_bounds(gray)
    band = ps.staff_band(gray)
    split = ps.choose_split_x(ps.detect_barlines(gray, band), crop_left, right)
    content_w = right - crop_left
    up = split - crop_left
    low = right - split
    assert abs(up - low) / content_w <= 0.15, (case["name"], up, low)
    assert abs(up / content_w - case["top_ratio"]) <= 0.03, (case["name"], up / content_w)


def test_06_lossless_split(case, gray):
    """원본 픽셀 무손실 분할: 상단폭 + 하단폭 == crop_right - crop_left (겹침·누락 없음)."""
    crop_left = ps.detect_watermark_crop_x(gray)
    left, right, top, bottom = ps.detect_content_bounds(gray)
    band = ps.staff_band(gray)
    split = ps.choose_split_x(ps.detect_barlines(gray, band), crop_left, right)
    img = ps.load_flattened_image(CASES[case["name"]]["img"])
    up_img, low_img = ps.split_and_place(img, (crop_left, top, right, bottom + 1), split)
    assert up_img.width + low_img.width == right - crop_left, (case["name"],)
    assert up_img.height == low_img.height  # 같은 높이로 잘림


def test_07_rgba_flatten_safe():
    """RGBA 투명 입력이 검정 잉크로 오인되지 않는다 (알파 평탄화 검증).

    완전 투명 RGBA 이미지는 흰 배경 합성 후 그레이스케일에서 잉크(어두운 픽셀)가
    전혀 없어야 한다. 평탄화를 빠뜨리면 투명 픽셀이 0(검정)으로 변해 전부 잉크가 된다.
    """
    transparent = Image.new("RGBA", (100, 40), (0, 0, 0, 0))
    buf = io.BytesIO()
    transparent.save(buf, format="PNG")
    buf.seek(0)
    gray = ps.load_gray_array(buf)
    ink = gray < 128
    assert ink.mean() == 0.0, "투명 픽셀이 잉크로 오인됨 — 평탄화 누락"


def test_08_fallback_no_watermark_gap():
    """폴백: 워터마크 미검출 시 미크롭(crop_left==0) + 경고, 예외 없음.

    좌측 30%에 넓은 zero-run이 없는 합성 입력. 음표 손실보다 워터마크 잔존이 안전.
    """
    # 좌측 30% 안에 완전 백색 컬럼(zero-run)이 하나도 없도록 전 폭에 가로 잉크 바를 둔다
    arr = np.full((40, 300), 255, dtype=np.uint8)
    arr[10:30, :] = 0  # 모든 컬럼에 잉크 → 분리 공백 띠 없음
    x = ps.detect_watermark_crop_x(arr)
    assert x == 0


def test_09_fallback_no_barline_geometric_center():
    """폴백: 중앙창 바라인 0개 → 예외 없이 기하 중앙 근처 분할점 반환."""
    # 바라인(고밀도 세로선)이 전혀 없는 합성 이미지
    arr = np.full((40, 300), 255, dtype=np.uint8)
    arr[18:22, 20:280] = 0  # 얇은 가로 오선만
    band = ps.staff_band(arr)
    barlines = ps.detect_barlines(arr, band)
    split = ps.choose_split_x(barlines, 20, 280)
    assert 20 < split < 280
    assert abs(split - (20 + 280) // 2) <= 30  # 기하 중앙 근처


def test_09b_barline_rejects_notehead_fused_stem():
    """detect_barlines가 노트헤드/스템에 융합된 획을 마디선 후보에서 제외한다.

    실사용자가 20260823 미리보기에서 col=2304를 확대해 8분음표 기둥+깃발임을 확인한 사례의
    재발 방지 테스트(설계서 §2.4, 02c 라운드). 합성 이미지로 두 후보를 만든다:
    - col=100: 오선 전체를 관통하는 순수 수직선(좌우 공백) — 진짜 마디선.
    - col=150: 오선 전체를 관통하는 수직선 + 한쪽 행간에 노트헤드 모양 블록이 융합됨 — 가짜.
    둘 다 col_fill>0.85(밴드 전체 관통)라 1차 임계값만으로는 구분되지 않으므로, 노트헤드 인접
    배제 필터가 실제로 두 후보를 갈라내는지 확인한다.
    """
    w, h = 220, 60
    img = np.full((h, w), 255, dtype=np.uint8)
    for r in (10, 20, 30, 40, 50):
        img[r, :] = 0  # 오선 5줄, 폭 전체 관통
    img[10:51, 100:102] = 0  # 진짜 마디선: 순수 수직선, 좌우 공백
    img[10:51, 150:152] = 0  # 가짜 후보의 기둥(스템)
    img[24:37, 143:161] = 0  # 가짜 후보에 융합된 노트헤드(한쪽 행간에 폭넓게 번짐)

    band = ps.staff_band(img)
    barlines = ps.detect_barlines(img, band)
    assert any(abs(b - 100) <= 2 for b in barlines), ("진짜 마디선이 후보에서 빠짐", barlines)
    assert not any(abs(b - 150) <= 2 for b in barlines), (
        "노트헤드에 융합된 스템이 마디선으로 오검출됨", barlines
    )


# ---------------------------------------------------------------------------
# 슬라이드 생성 (템플릿 보존)
# ---------------------------------------------------------------------------

@pytest.fixture(scope="module")
def rendered(case):
    return ps.render_화답송_score_slide(CASES[case["name"]]["img"], "시편 00(00)", TEMPLATE)


def _iter_shapes(slide):
    return list(slide.shapes)


def test_10_exactly_one_slide(rendered):
    """출력 Presentation은 슬라이드 정확히 1장."""
    assert len(rendered.slides) == 1


def test_11_slide_size_preserved(rendered):
    """슬라이드 크기 12192000×6858000 유지 (copy_slide_from_prs 스케일 1:1 보장)."""
    assert rendered.slide_width == 12192000
    assert rendered.slide_height == 6858000


def test_12_template_text_preserved(rendered):
    """템플릿 라벨/저작권 텍스트박스가 그대로 보존된다."""
    texts = [
        sh.text_frame.text
        for sh in _iter_shapes(rendered.slides[0])
        if sh.has_text_frame
    ]
    joined = "\n".join(texts)
    assert "Responsorial" in joined
    assert "한국천주교중앙협의회" in joined
    assert "2022" in joined


def test_13_title_replaced(rendered):
    """제목박스 텍스트가 '화 답 송   {title}'로 교체, 다른 텍스트박스는 불변."""
    slide = rendered.slides[0]
    title_boxes = [
        sh for sh in _iter_shapes(slide)
        if sh.has_text_frame and "화 답 송" in sh.text_frame.text
    ]
    assert len(title_boxes) == 1
    assert title_boxes[0].text_frame.text.strip() == "화 답 송   시편 00(00)"
    # 저작권 박스는 제목 교체의 영향을 받지 않는다
    copyright_boxes = [
        sh for sh in _iter_shapes(slide)
        if sh.has_text_frame and "한국천주교중앙협의회" in sh.text_frame.text
    ]
    assert len(copyright_boxes) == 1


def test_13b_title_paragraph_child_order_schema_valid(rendered):
    """제목 교체 후 <a:p> 자식 순서가 CT_TextParagraph 스키마를 지킨다.

    스키마 시퀀스: pPr?, (run|br)*, endParaRPr?. run이 endParaRPr보다 뒤에 오면
    PowerPoint가 "손상된 파일"로 취급할 수 있는 스키마 위반이다. validate_pptx_structure()는
    XML 제어문자·끊어진 rId만 검사하고 요소 순서를 검사하지 않으므로(missa_to_ppt.py:6071),
    이 순서는 lxml로 직접 확인해야 한다(02b_review_report.md 리뷰 지적 — CLAUDE.md의
    "pPr append 금지" 함정이 <a:p> 자식 레벨(run vs endParaRPr)에서 재발했었음).
    """
    from pptx.oxml.ns import qn

    slide = rendered.slides[0]
    title_boxes = [
        sh for sh in _iter_shapes(slide)
        if sh.has_text_frame and "화 답 송" in sh.text_frame.text
    ]
    assert len(title_boxes) == 1
    p = title_boxes[0].text_frame.paragraphs[0]._p
    tags = [child.tag for child in p]
    eprp_tag = qn("a:endParaRPr")
    r_tag = qn("a:r")
    assert r_tag in tags, "제목 run이 존재해야 함"
    if eprp_tag in tags:
        r_indices = [i for i, t in enumerate(tags) if t == r_tag]
        eprp_index = tags.index(eprp_tag)
        assert max(r_indices) < eprp_index, (
            "run이 endParaRPr보다 앞에 와야 한다(스키마 위반)", tags
        )


def test_14_two_score_pictures_placed(rendered):
    """악보 이미지 PICTURE 2장 추가 (width>1M), 각각 상/하 밴드 안, 종횡비 유지."""
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    slide = rendered.slides[0]
    big_pics = [
        sh for sh in _iter_shapes(slide)
        if sh.shape_type == MSO_SHAPE_TYPE.PICTURE and sh.width > 1_000_000
    ]
    assert len(big_pics) == 2, [sh.width for sh in big_pics]
    tops = sorted(sh.top for sh in big_pics)
    # 상단 밴드 top≈1.0M, 하단 밴드 top≈3.65M (밴드 내 세로 중앙정렬이라 약간 아래로 이동)
    assert 1_000_000 <= tops[0] < 3_000_000, tops
    assert 3_000_000 <= tops[1] < 6_000_000, tops
    # 밴드 폭(11392000, left여백 400000) 안에 있고 왜곡 없이 배치
    for sh in big_pics:
        assert sh.left >= 400_000 - 1
        assert sh.left + sh.width <= 400_000 + 11_392_000 + 1
        assert sh.height <= 2_500_000 + 1


def test_14b_uniform_scale_across_both_lines():
    """상/하 두 단의 폭이 크게 다르면 공통(더 작은) 배율을 적용해 음표 크기를 통일한다.

    사용자 요구(02c 라운드): "일부 악보는 첫째 단과 둘째 단의 폭이 다를 수 있다 — 그 경우
    전체적인 비율을 조정해서... 첫째 단과 둘째 단의 높이가 더 작아지는 거지." 각자 독립적으로
    자기 밴드에 맞춰 스케일링하면(예전 방식) 폭이 좁은 쪽이 상대적으로 확대돼 두 줄의 음표
    크기가 달라 보인다. 원본 픽셀 높이가 같은 두 이미지(폭만 크게 다름)를 넣었을 때, 결과
    슬라이드의 두 PICTURE 높이(EMU)가 같아야 한다 — 배율이 공통이라는 강한 증거.
    """
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    wide_img = Image.new("RGB", (2000, 100), "white")
    narrow_img = Image.new("RGB", (300, 100), "white")
    prs = ps.build_slide(TEMPLATE, "시편 00(00)", wide_img, narrow_img)
    slide = prs.slides[0]
    big_pics = [
        sh for sh in slide.shapes
        if sh.shape_type == MSO_SHAPE_TYPE.PICTURE and sh.width > 1_000_000
    ]
    assert len(big_pics) == 2
    heights = sorted(sh.height for sh in big_pics)
    # 원본 픽셀 높이가 동일(100=100)하므로 배율이 같다면 결과 EMU 높이도 같아야 한다.
    assert abs(heights[0] - heights[1]) <= 2, (
        "두 단의 배율이 달라 높이가 다름 — 공통 배율 미적용", heights
    )
    widths = sorted(sh.width for sh in big_pics)
    # 폭 비율(2000:300 ≈ 6.67:1)이 결과에도 그대로 유지돼야 한다(공통 배율은 왜곡을 만들지 않음).
    assert abs(widths[1] / widths[0] - 2000 / 300) < 0.05


def test_15_residual_shapes_preserved(rendered):
    """잔재 4도형(12×7·41×17 EMU PICTURE) 그대로 — 템플릿 무손상."""
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    slide = rendered.slides[0]
    tiny_pics = [
        sh for sh in _iter_shapes(slide)
        if sh.shape_type == MSO_SHAPE_TYPE.PICTURE and sh.width < 1_000_000
    ]
    assert len(tiny_pics) == 4, [(sh.width, sh.height) for sh in tiny_pics]


def test_16_copy_slide_from_prs_compatible(rendered, tmp_path):
    """생성 Presentation을 타깃에 copy_slide_from_prs로 복사 시 손상 없이 rel까지 복사."""
    import missa_to_ppt as mtp
    import missa_ooxml_utils as ou
    from pptx import Presentation

    target = Presentation(str(TEMPLATE))
    ou.copy_slide_from_prs(target, len(target.slides), rendered, 0)
    out = tmp_path / "copied.pptx"
    target.save(str(out))
    problems = mtp.validate_pptx_structure(str(out))
    assert problems == [], problems


# ---------------------------------------------------------------------------
# 2차 마일스톤: missa_to_ppt.py 배선 (설계서 §4 통합 지점)
# ---------------------------------------------------------------------------

import missa_to_ppt as mtp  # noqa: E402  (모듈 상단 import 관례를 따르되, 1차 섹션과 분리해 배치)


def test_find_files_prefers_pptx_over_image(tmp_path):
    """화답송_pptx가 있으면 화답송_img 탐색 자체를 건너뛰고 pptx를 사용한다(기존 동작 보존)."""
    (tmp_path / "화답송 악보 시편 1.pptx").write_bytes(b"")
    (tmp_path / "photo.jpg").write_bytes(b"")
    files = mtp.find_files(str(tmp_path), {}, is_sunday=False)
    assert files["화답송_pptx"] is not None
    assert files["화답송_pptx"].name == "화답송 악보 시편 1.pptx"
    assert files["화답송_img"] is None


def test_find_files_named_image_when_no_pptx(tmp_path):
    """pptx가 없고 파일명에 '화답송'이 들어간 이미지가 있으면 그것을 화답송_img로 찾는다."""
    (tmp_path / "화답송 사진.jpg").write_bytes(b"")
    files = mtp.find_files(str(tmp_path), {}, is_sunday=False)
    assert files["화답송_pptx"] is None
    assert files["화답송_img"] is not None
    assert files["화답송_img"].name == "화답송 사진.jpg"


def test_find_files_unique_image_fallback_when_unnamed(tmp_path):
    """파일명에 '화답송'이 없어도(촬영 앱 타임스탬프 파일명) 폴더 내 유일 이미지면 폴백한다."""
    (tmp_path / "20260809_043842947.jpg").write_bytes(b"")
    files = mtp.find_files(str(tmp_path), {}, is_sunday=False)
    assert files["화답송_img"] is not None
    assert files["화답송_img"].name == "20260809_043842947.jpg"


def test_find_files_ambiguous_multiple_unnamed_images_no_fallback(tmp_path):
    """이름 매칭이 안 되고 이미지가 여러 개면 어떤 게 화답송 악보인지 모호하므로 폴백하지
    않는다 — 잘못된 이미지를 화답송으로 오인하는 것보다 미검출이 안전하다."""
    (tmp_path / "photo1.jpg").write_bytes(b"")
    (tmp_path / "photo2.png").write_bytes(b"")
    files = mtp.find_files(str(tmp_path), {}, is_sunday=False)
    assert files["화답송_img"] is None


def test_apply_hwadapsong_override_routes_image_extension():
    """--화답송 오버라이드가 이미지 확장자면 화답송_img로 가고, 반대쪽 키는 비워진다."""
    files = {"화답송_pptx": Path("auto_found.pptx"), "화답송_img": None}
    mtp.apply_화답송_override(files, "C:/mass/photo.jpg")
    assert files["화답송_img"] == Path("C:/mass/photo.jpg")
    assert files["화답송_pptx"] is None


def test_apply_hwadapsong_override_routes_pptx_extension():
    """--화답송 오버라이드가 pptx 확장자면 화답송_pptx로 가고, 반대쪽 키는 비워진다."""
    files = {"화답송_pptx": None, "화답송_img": Path("auto_found.jpg")}
    mtp.apply_화답송_override(files, "C:/mass/manual.pptx")
    assert files["화답송_pptx"] == Path("C:/mass/manual.pptx")
    assert files["화답송_img"] is None


def _화답송_json(title: str) -> dict:
    return {
        "화답송": {
            "title": title,
            "content": "◎ 후렴 테스트\n○ 절1\n○ 절2\n○ 절3",
        }
    }


def test_update_화답송_uses_image_when_no_pptx_end_to_end(tmp_path):
    """update_화답송이 화답송_img_path만 주어졌을 때 render_화답송_score_slide()를 거쳐
    만들어진 슬라이드가 실제 결과 PPT에 copy_slide_from_prs 경로로 정상 복사되는지 확인한다
    (실제 참조 예제 이미지로 end-to-end 검증, 설계서 §4.3)."""
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    prs = mtp.Presentation(
        str(BASE / "output" / "20260712" / "Template_20260628_연중 제13주일 (교황주일).pptx")
    )
    sections = mtp.find_sections(prs)
    json_data = _화답송_json("시편 99(98)")
    img_path = CASES["20260809"]["img"]

    mtp.update_화답송(prs, json_data, sections, None, 화답송_img_path=img_path, is_sunday=True)

    start = sections["화답송_start"]
    score_slide = prs.slides[start]  # i=0(짝수) => 악보 슬라이드
    big_pics = [
        sh for sh in score_slide.shapes
        if sh.shape_type == MSO_SHAPE_TYPE.PICTURE and sh.width > 1_000_000
    ]
    assert len(big_pics) == 2, "이미지 입력 경로로 악보 슬라이드에 상/하 악보 그림 2장이 있어야 함"

    texts = [sh.text_frame.text for sh in score_slide.shapes if sh.has_text_frame]
    assert any("시편 99(98)" in t for t in texts), texts

    out = tmp_path / "e2e_img.pptx"
    prs.save(str(out))
    assert mtp.validate_pptx_structure(str(out)) == []


def test_update_화답송_existing_pptx_path_unaffected(tmp_path):
    """기존 pptx 입력 경로가 이번 배선 변경으로 전혀 달라지지 않았는지 확인(회귀 방지)."""
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    prs = mtp.Presentation(
        str(BASE / "output" / "20260712" / "Template_20260628_연중 제13주일 (교황주일).pptx")
    )
    sections = mtp.find_sections(prs)
    json_data = _화답송_json("시편 67(66)")
    pptx_path = REF / "20260816" / "시편 67(66).pptx"

    mtp.update_화답송(prs, json_data, sections, pptx_path, 화답송_img_path=None, is_sunday=True)

    start = sections["화답송_start"]
    score_slide = prs.slides[start]
    big_pics = [
        sh for sh in score_slide.shapes
        if sh.shape_type == MSO_SHAPE_TYPE.PICTURE and sh.width > 1_000_000
    ]
    assert len(big_pics) == 2

    out = tmp_path / "e2e_pptx.pptx"
    prs.save(str(out))
    assert mtp.validate_pptx_structure(str(out)) == []


def test_update_화답송_prefers_pptx_over_image(monkeypatch):
    """화답송_pptx_path와 화답송_img_path가 둘 다 주어지면 pptx가 우선한다(설계서 §4.3 우선순위).

    render_화답송_score_slide가 절대 호출되지 않아야 함을 몽키패치로 직접 확인한다.
    """
    import missa_psalm_score_image

    def _must_not_be_called(*args, **kwargs):
        raise AssertionError("이미지 경로가 호출됨 — pptx가 우선이어야 하는데 위반됨")

    monkeypatch.setattr(
        missa_psalm_score_image, "render_화답송_score_slide", _must_not_be_called
    )

    prs = mtp.Presentation(
        str(BASE / "output" / "20260712" / "Template_20260628_연중 제13주일 (교황주일).pptx")
    )
    sections = mtp.find_sections(prs)
    json_data = _화답송_json("시편 1(1)")
    pptx_path = REF / "20260816" / "시편 67(66).pptx"
    img_path = CASES["20260809"]["img"]

    # 예외 없이 끝나야 함 = render_화답송_score_slide가 호출되지 않았다는 뜻
    mtp.update_화답송(
        prs, json_data, sections, pptx_path, 화답송_img_path=img_path, is_sunday=True
    )


# ---------------------------------------------------------------------------
# 공지사항 슬라이드 삽입 (설계서 _workspace/공지사항_슬라이드/01_architect_design.md)
# §6 테스트 대상 행동 1~9. 8번(GUI tkinter 확장자 검증)만 자동 테스트 제외.
# 기대 인덱스는 세 git 추적 픽스처의 실측 find_sections() 값이다(추측 아님):
#   20260712(주일) n=147, 2차봉헌_content_end=132(blank)
#   20260705(성수축복) n=144, 2차봉헌_content_end=129(blank)
#   20260624(평일) 2차봉헌_* 키 부재
# ---------------------------------------------------------------------------

from pptx import Presentation as _Presentation  # noqa: E402
from pptx.util import Emu as _Emu  # noqa: E402
from missa_content_updaters import insert_공지사항  # noqa: E402
from missa_ooxml_utils import all_slide_texts as _all_slide_texts  # noqa: E402

_FIXTURES = {
    "20260712": BASE / "output" / "20260712" / "20260712_연중 제15주일.pptx",
    "20260705": BASE / "output" / "20260705"
    / "20260705_한국 성직자들의 수호자 성 김대건 안드레아 사제 순교자 - 신심 미사.pptx",
    "20260624": BASE / "output" / "20260624" / "20260624_성 요한 세례자 탄생 대축일.pptx",
}


def _make_공지사항_pptx(path, n):
    """식별 가능한 텍스트를 가진 n장짜리 공지사항 PPT를 임시 생성한다.

    각 슬라이드에 '공지사항 슬라이드 {i}' 텍스트박스를 넣어 삽입 후 서식/순서 보존을
    강하게 단언할 수 있게 한다. n==0이면 슬라이드 없는 빈 PPT를 만든다.
    """
    prs = _Presentation()
    for i in range(n):
        slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank layout
        box = slide.shapes.add_textbox(_Emu(914400), _Emu(914400), _Emu(3657600), _Emu(914400))
        box.text_frame.text = f"공지사항 슬라이드 {i}"
    prs.save(str(path))
    return prs


def _src_texts(path):
    src = _Presentation(str(path))
    return [
        "\n".join(
            sh.text_frame.text for sh in s.shapes if sh.has_text_frame
        ).strip()
        for s in src.slides
    ]


def test_공지사항_01_insert_count_주일(tmp_path):
    """[삽입 총수] N장짜리 공지사항 → 반환 N+1, 전체 슬라이드 == 기존 + N + 1."""
    from missa_sections import find_sections

    n_src = 4
    src = tmp_path / "공지사항.pptx"
    _make_공지사항_pptx(src, n_src)

    prs = _Presentation(str(_FIXTURES["20260712"]))
    before = len(prs.slides)
    sec = find_sections(prs)
    added = insert_공지사항(prs, src, sec)

    assert added == n_src + 1
    assert len(prs.slides) == before + n_src + 1


def test_공지사항_02_one_divider_each_side(tmp_path):
    """[앞뒤 구분선 정확히 1장씩] 앞 구분선(기존 content_end)·뒤 구분선(신규) 모두 blank,
    그 뒤는 non-blank(마침 축복). 공지사항 경계에 2장 연속 blank가 없음."""
    from missa_sections import find_sections

    n_src = 3
    src = tmp_path / "공지사항.pptx"
    _make_공지사항_pptx(src, n_src)

    prs = _Presentation(str(_FIXTURES["20260712"]))
    sec = find_sections(prs)
    front = sec["2차봉헌_content_end"]  # 132 (실측)
    insert_공지사항(prs, src, sec)

    texts = _all_slide_texts(prs)
    is_blank = [not t.strip() for t in texts]

    # 앞 구분선: content_end 자리 그대로 blank
    assert is_blank[front], ("앞 구분선이 blank가 아님", front)
    # 공지사항 콘텐츠 n_src장: 모두 non-blank
    for i in range(n_src):
        assert not is_blank[front + 1 + i], ("공지사항 콘텐츠가 blank임", front + 1 + i)
    # 뒤 구분선: 콘텐츠 마지막 다음 자리 blank
    rear = front + 1 + n_src
    assert is_blank[rear], ("뒤 구분선이 blank가 아님", rear)
    # 2026-09-13 실측 버그의 end-to-end 재발 방지: 뒤 구분선도 앞 구분선과 동일하게
    # showMasterSp="0"이어야 한다(다르면 마스터 상속 요소가 노출돼 색이 달라 보임).
    front_show_master = prs.slides[front].element.get("showMasterSp")
    rear_show_master = prs.slides[rear].element.get("showMasterSp")
    assert rear_show_master == front_show_master, (
        "뒤 구분선의 showMasterSp가 앞 구분선과 다름", front_show_master, rear_show_master,
    )
    # 뒤 구분선 다음은 non-blank(마침 축복 기도)
    assert not is_blank[rear + 1], ("뒤 구분선 다음이 blank임 — 2장 연속 blank", rear + 1)
    # front-1(2차봉헌 마지막 콘텐츠)도 non-blank → front 앞뒤로 2장 연속 blank 없음
    assert not is_blank[front - 1], ("앞 구분선 앞이 blank임 — 2장 연속 blank", front - 1)


def test_공지사항_03_content_format_preserved(tmp_path):
    """[콘텐츠 서식 보존] 삽입된 공지사항 슬라이드 텍스트가 원본과 동일(순서 포함).
    뒤 구분선 슬라이드는 shapes 0개."""
    from missa_sections import find_sections

    n_src = 3
    src = tmp_path / "공지사항.pptx"
    _make_공지사항_pptx(src, n_src)
    src_texts = _src_texts(src)  # ['공지사항 슬라이드 0', '...1', '...2']

    prs = _Presentation(str(_FIXTURES["20260712"]))
    sec = find_sections(prs)
    front = sec["2차봉헌_content_end"]
    insert_공지사항(prs, src, sec)

    for i in range(n_src):
        slide = prs.slides[front + 1 + i]
        joined = "\n".join(
            sh.text_frame.text for sh in slide.shapes if sh.has_text_frame
        ).strip()
        assert joined == src_texts[i], (i, joined, src_texts[i])

    rear = front + 1 + n_src
    assert len(prs.slides[rear].shapes) == 0, "뒤 구분선은 원본 blank 복제라 shapes 0개여야 함"


def test_공지사항_04_no_2차봉헌_slot_skips_평일(tmp_path):
    """[2차봉헌 없는 평일 → 경고 후 skip] 예외 없이 반환 0, 슬라이드 수 불변."""
    from missa_sections import find_sections

    src = tmp_path / "공지사항.pptx"
    _make_공지사항_pptx(src, 2)

    prs = _Presentation(str(_FIXTURES["20260624"]))
    before = len(prs.slides)
    sec = find_sections(prs)
    assert sec.get("2차봉헌_content_end") is None  # 실측: 키 자체가 없음

    added = insert_공지사항(prs, src, sec)
    assert added == 0
    assert len(prs.slides) == before


def test_공지사항_05_insert_성수축복(tmp_path):
    """[성수축복(2차봉헌 있는 특수 주일)] content_end=129 기준 삽입 성립."""
    from missa_sections import find_sections

    n_src = 2
    src = tmp_path / "공지사항.pptx"
    _make_공지사항_pptx(src, n_src)

    prs = _Presentation(str(_FIXTURES["20260705"]))
    before = len(prs.slides)
    sec = find_sections(prs)
    front = sec["2차봉헌_content_end"]  # 129 (실측)
    added = insert_공지사항(prs, src, sec)

    assert added == n_src + 1
    assert len(prs.slides) == before + n_src + 1

    texts = _all_slide_texts(prs)
    is_blank = [not t.strip() for t in texts]
    assert is_blank[front]
    rear = front + 1 + n_src
    assert is_blank[rear]
    assert not is_blank[rear + 1]


def test_공지사항_06a_missing_path_returns_zero(tmp_path):
    """[로드 실패 방어] 존재하지 않는 경로 → 반환 0, 슬라이드 불변, 예외 없음."""
    from missa_sections import find_sections

    prs = _Presentation(str(_FIXTURES["20260712"]))
    before = len(prs.slides)
    sec = find_sections(prs)
    added = insert_공지사항(prs, tmp_path / "없는파일.pptx", sec)
    assert added == 0
    assert len(prs.slides) == before


def test_공지사항_06b_zero_slide_pptx_returns_zero(tmp_path):
    """[0장 방어] 슬라이드 0장짜리 공지사항 PPT → 반환 0, 슬라이드 불변."""
    from missa_sections import find_sections

    src = tmp_path / "empty.pptx"
    _make_공지사항_pptx(src, 0)

    prs = _Presentation(str(_FIXTURES["20260712"]))
    before = len(prs.slides)
    sec = find_sections(prs)
    added = insert_공지사항(prs, src, sec)
    assert added == 0
    assert len(prs.slides) == before


def test_공지사항_07_front_not_blank_skips(tmp_path):
    """[앞 구분선이 blank 아님 → skip] content_end가 non-blank를 가리키면 반환 0(방어 (c))."""
    from missa_sections import find_sections

    src = tmp_path / "공지사항.pptx"
    _make_공지사항_pptx(src, 2)

    prs = _Presentation(str(_FIXTURES["20260712"]))
    before = len(prs.slides)
    sec = find_sections(prs)
    # 132(blank) 대신 133(마침 축복, non-blank)을 가리키도록 조작
    sec = dict(sec, **{"2차봉헌_content_end": 133})
    added = insert_공지사항(prs, src, sec)
    assert added == 0
    assert len(prs.slides) == before


def test_공지사항_structure_valid_after_insert(tmp_path):
    """[구조 무손상] 삽입·저장 후 validate_pptx_structure()가 빈 리스트(끊어진 rId/제어문자 없음)."""
    from missa_sections import find_sections

    src = tmp_path / "공지사항.pptx"
    _make_공지사항_pptx(src, 3)

    prs = _Presentation(str(_FIXTURES["20260712"]))
    sec = find_sections(prs)
    insert_공지사항(prs, src, sec)

    out = tmp_path / "out.pptx"
    prs.save(str(out))
    assert mtp.validate_pptx_structure(str(out)) == []


def test_공지사항_08_cli_rejects_non_pptx(monkeypatch):
    """[CLI 확장자 검증] --공지사항에 .pptx 아닌 경로 → SystemExit(오류 종료)."""
    monkeypatch.setattr(
        "sys.argv", ["missa_to_ppt.py", "20260712", "--공지사항", "foo.docx", "--test"]
    )
    with pytest.raises(SystemExit):
        mtp.parse_args()


def test_공지사항_08b_cli_accepts_pptx_into_5tuple(monkeypatch):
    """[CLI 통과] .pptx면 parse_args가 5-튜플을 반환하고 5번째가 공지사항 경로다."""
    monkeypatch.setattr(
        "sys.argv", ["missa_to_ppt.py", "20260712", "--공지사항", "notice.pptx", "--test"]
    )
    result = mtp.parse_args()
    assert len(result) == 5
    assert result[4] == "notice.pptx"


def test_공지사항_08c_cli_default_none(monkeypatch):
    """[CLI 기본값] --공지사항 미지정이면 5번째가 None(선택 입력)."""
    monkeypatch.setattr("sys.argv", ["missa_to_ppt.py", "20260712", "--test"])
    result = mtp.parse_args()
    assert len(result) == 5
    assert result[4] is None


def test_공지사항_09_find_files_detects_by_name(tmp_path):
    """[자동 탐색] 파일명에 '공지사항'이 든 .pptx가 files['공지사항']에 잡힌다."""
    (tmp_path / "20260712_ref.pptx").write_bytes(b"")
    (tmp_path / "공지사항.pptx").write_bytes(b"")
    files = mtp.find_files(str(tmp_path), {}, is_sunday=False)
    assert files["공지사항"] is not None
    assert files["공지사항"].name == "공지사항.pptx"


def test_공지사항_09b_find_files_none_when_absent(tmp_path):
    """[미입력 무변경 가드 근거] 공지사항 파일이 없으면 files['공지사항']==None →
    main의 `if files.get('공지사항')` 가드가 False라 insert_공지사항이 호출되지 않는다."""
    (tmp_path / "20260712_ref.pptx").write_bytes(b"")
    files = mtp.find_files(str(tmp_path), {}, is_sunday=False)
    assert files.get("공지사항") is None
    assert not files.get("공지사항")  # 가드 falsy


# ---------------------------------------------------------------------------
# duplicate_slide()의 showMasterSp 속성 보존 (2026-09-13 실측 버그)
# insert_공지사항이 만든 뒤 구분 슬라이드(142, duplicate_slide 산출물)가 앞 구분 슬라이드
# (133, hand-authored)와 색이 달라 보였다. 두 슬라이드는 layout/배경이 100% 동일했으나
# 최상위 <p:sld>의 showMasterSp 속성이 앞쪽은 "0"(마스터 상속 요소 숨김), 뒤쪽은 부재
# (python-pptx add_slide 기본값 = 마스터 요소 노출)였던 것이 원인. duplicate_slide가
# spTree/이미지 rel/p:bg는 복사하면서 <p:sld> 자체의 showMasterSp 속성은 복사하지 않았다.
# ---------------------------------------------------------------------------

from missa_ooxml_utils import duplicate_slide as _duplicate_slide  # noqa: E402


def test_showmastersp_present_copied_to_duplicate():
    """원본 <p:sld showMasterSp="0">이면 복제본의 <p:sld>도 showMasterSp="0"이어야 한다.

    이 속성이 누락되면 layout/배경이 동일해도 마스터 배치 요소(장식·로고 등)가 복제본에서만
    노출돼 색이 달라 보인다. p:bg 복사만으로는 커버되지 않는 별도 속성이다.
    """
    prs = _Presentation()
    src = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    src.element.set("showMasterSp", "0")
    new_idx = _duplicate_slide(prs, 0)
    assert prs.slides[new_idx].element.get("showMasterSp") == "0"


def test_showmastersp_absent_stays_absent():
    """원본에 showMasterSp 속성이 없으면 복제본에도 추가하지 않는다(불필요한 속성 주입 금지).

    원본이 add_slide 기본 상태(속성 부재 = 마스터 요소 노출)면 복제본도 그 상태를 그대로
    유지해야 한다 — 원본에 없던 속성을 새로 만들어 넣으면 안 된다.
    """
    prs = _Presentation()
    src = prs.slides.add_slide(prs.slide_layouts[6])
    assert src.element.get("showMasterSp") is None  # 전제: add_slide 기본은 속성 부재
    new_idx = _duplicate_slide(prs, 0)
    assert prs.slides[new_idx].element.get("showMasterSp") is None


def test_showmastersp_real_fixture_handauthored_slide():
    """실측 재현: 20260712 idx132(hand-authored, showMasterSp="0")를 복제하면 복제본도 "0".

    2026-09-13 프로덕션 파일의 앞 구분 슬라이드(idx132)와 동일한 성격의 슬라이드다(git 추적
    픽스처 20260712 idx132도 showMasterSp="0"임을 실측). 이 슬라이드를 재사용해 뒤 구분
    슬라이드를 만들 때 속성이 보존돼야 두 슬라이드가 같은 색으로 렌더링된다.
    """
    prs = _Presentation(str(_FIXTURES["20260712"]))
    assert prs.slides[132].element.get("showMasterSp") == "0"  # 픽스처 실측 전제
    new_idx = _duplicate_slide(prs, 132)
    assert prs.slides[new_idx].element.get("showMasterSp") == "0"


# ===========================================================================
# 청년미사 1단계 — 모듈 1: 영문 복음 조회 (missa_youth_gospel.py)
# 설계서 _workspace/청년미사_1단계/01_architect_design.md §2, §9(G1~G7).
#
# 파싱 의존 테스트(G1/G4~G7)는 커밋된 실측 HTML 픽스처로 결정론적으로 돈다(universalis는
# 약 열흘치만 서빙하므로 날짜 기반 라이브 테스트는 시간이 지나면 깨진다). 리다이렉트 함정
# (G2/G3)만 라이브로 검증한다 — 20200101/20270101은 현재 날짜와 무관하게 항상 범위를 벗어나
# 리다이렉트되므로 시간에 안정적이다.
# ===========================================================================
import missa_youth_gospel as yg

GOSPEL_FIX = BASE / "reference" / "청년미사" / "gospel_fixtures"


def _fix_html(date_str):
    return (GOSPEL_FIX / f"{date_str}.html").read_text(encoding="utf-8")


def test_G1_parse_reference_and_body_start():
    """G1: 범위 내 날짜 HTML → reference/본문 정확 추출."""
    g = yg.parse_gospel(_fix_html("20260913"))
    assert g["reference"] == "Matthew 18:21-35"
    assert g["content"].startswith("Peter went up to Jesus")


def test_G2_far_past_redirect_raises():
    """G2(최우선): 과거로 먼 날짜(→/mass.htm 리다이렉트) → GospelFetchError.

    라이브. 조용히 오늘자 복음이 반환되는 가장 위험한 케이스를 막는지 확인."""
    with pytest.raises(yg.GospelFetchError):
        yg.get_youth_gospel("20200101")


def test_G3_far_future_redirect_raises():
    """G3: 미래로 먼 날짜(→/n-otherdates.htm) → GospelFetchError. 라이브."""
    with pytest.raises(yg.GospelFetchError):
        yg.get_youth_gospel("20270101")


def test_G4_p_pi_paragraph_boundary():
    """G4: p/pi 형식 본문 단락 경계가 \\n\\n으로 유지."""
    g = yg.parse_gospel(_fix_html("20260913"))
    assert "\n\n" in g["content"]
    # pi 단락("And so the kingdom…")이 p 단락 뒤 \n\n 경계로 이어짐
    parts = g["content"].split("\n\n")
    assert len(parts) >= 2
    assert any("kingdom of heaven" in p for p in parts)


def test_G5_v_vi_format_collected():
    """G5: v/vi 형식(족보 20260908)도 본문 수집."""
    g = yg.parse_gospel(_fix_html("20260908"))
    assert g["reference"] == "Matthew 1:1-16,18-23"
    assert "Abraham was the father of Isaac" in g["content"]
    # v 형식은 절마다 div → 여러 단락으로 수집됨
    assert g["content"].count("\n\n") >= 5


def test_G6_html_entities_decoded():
    """G6: &#8216; 등 HTML 엔티티가 유니코드 문자로 디코딩(&# 잔존 없음)."""
    g = yg.parse_gospel(_fix_html("20260913"))
    assert "‘" in g["content"]  # &#8216; 곡따옴표
    assert "&#" not in g["content"]
    assert "&#" not in g["reference"]


def test_G7_theme_and_date_in_full_json(monkeypatch):
    """G7: get_youth_gospel JSON에 theme·date 포함, 값 정확(fetch를 픽스처로 대체)."""
    monkeypatch.setattr(yg, "fetch_gospel_html", lambda d: _fix_html("20260913"))
    out = yg.get_youth_gospel("20260913")
    assert out["date"] == "20260913"
    assert out["복음"]["theme"] == "To be forgiven, you must forgive"
    assert out["복음"]["reference"] == "Matthew 18:21-35"
    assert out["복음"]["content"].startswith("Peter went up to Jesus")


# ===========================================================================
# 청년미사 1단계 — 모듈 2: PDF 성가집 → PPT (missa_youth_hymn_pdf.py)
# 설계서 §3, §9(H1~H4, S1~S6, Hd1~Hd4, C1~C5, P1~P4).
# 실측 자산(나주노/야훼 PDF, 샘플 PPT)을 기준으로 검증한다 — 하드코딩 기대값 금지.
# ===========================================================================
import missa_youth_hymn_pdf as hp
import fitz as _fitz
import os as _os

_YOUTH_PDFS_AVAILABLE = _os.path.isfile(hp.SOURCES["나주노"]["pdf"]) and _os.path.isfile(
    hp.SOURCES["야훼 이레"]["pdf"]
)
_YOUTH_SKIP_REASON = "나주노/야훼이레 PDF 원본이 로컬에 없음(reference/청년미사/, git 미포함)"


@pytest.fixture(scope="module")
def naju_doc():
    d = _fitz.open(hp.SOURCES["나주노"]["pdf"])
    yield d
    d.close()


@pytest.fixture(scope="module")
def yahwe_doc():
    d = _fitz.open(hp.SOURCES["야훼 이레"]["pdf"])
    yield d
    d.close()


@pytest.mark.skipif(not _YOUTH_PDFS_AVAILABLE, reason=_YOUTH_SKIP_REASON)
def test_H1_naju_page_and_title():
    """H1: 나주노 362 → TOC 기반 page index 426, 제목 '삼위일체'(번호 접두 제거)."""
    assert hp.find_song_page("나주노", 362) == 426
    assert hp.find_song_title("나주노", 362) == "삼위일체"


@pytest.mark.skipif(not _YOUTH_PDFS_AVAILABLE, reason=_YOUTH_SKIP_REASON)
def test_H2_yahwe_page():
    """H2: 야훼 이레 810 → search_for 기반 page index 572."""
    assert hp.find_song_page("야훼 이레", 810) == 572


@pytest.mark.skipif(not _YOUTH_PDFS_AVAILABLE, reason=_YOUTH_SKIP_REASON)
def test_H3_yahwe_y_range_excludes_neighbors(yahwe_doc):
    """H3: 야훼810 y범위 = (제목아래≈396.6, 저작권top≈743.7). 809(y≈65)·811(다음 페이지) 미포함."""
    page = yahwe_doc[572]
    top, bottom = hp.find_song_y_range(page, "야훼 이레", 810)
    assert 390 < top < 400          # 810 제목 아래
    assert 740 < bottom < 748       # 810 저작권 줄 top
    assert top > 84                 # 809 제목(y1≈84) 아래에서 시작 → 809 혼입 없음


@pytest.mark.skipif(not _YOUTH_PDFS_AVAILABLE, reason=_YOUTH_SKIP_REASON)
def test_H4_naju_scan_no_text_layer(naju_doc):
    """H4: 나주노 스캔 페이지 get_text() 길이 0이어도 크래시 없이 이미지 경로 진행."""
    page = naju_doc[426]
    assert len(page.get_text("text").strip()) == 0  # 스캔 = 텍스트 레이어 없음(전제)
    # 렌더가 예외 없이 이미지를 반환
    img = hp.render_region(page, None, hp.RENDER_DPI)
    assert img.width > 0 and img.height > 0


# --- 시스템 밴드 / 슬라이드 수 (S1~S6) --------------------------------------
# 5곡의 시스템 밴드는 렌더가 느리므로(82MB PDF) 모듈 스코프에서 1회 계산해 공유한다.
_SONGS = {362: "나주노", 173: "나주노", 146: "나주노", 810: "야훼 이레", 267: "나주노"}
_EXPECTED_SYSTEMS = {362: 4, 173: 7, 146: 6, 810: 4, 267: 5}
_EXPECTED_SLIDES = {362: 2, 173: 4, 146: 3, 810: 2, 267: 3}


@pytest.fixture(scope="module")
def hymn_bands():
    import numpy as _np
    import missa_psalm_score_image as _ps
    result = {}
    for num, src in _SONGS.items():
        pi = hp.find_song_page(src, num)
        d = _fitz.open(hp.SOURCES[src]["pdf"])
        page = d[pi]
        yr = None if src == "나주노" else hp.find_song_y_range(page, src, num)
        img = hp.render_region(page, yr, hp.RENDER_DPI)
        gray = _np.asarray(img.convert("L"))
        bands = hp.detect_system_bands(gray)
        content_top = int(_np.where((gray < 128).any(axis=1))[0].min())
        result[num] = {"bands": bands, "content_top": content_top,
                       "shape": gray.shape, "gray": gray}
        d.close()
    return result


@pytest.fixture(scope="module")
def render_447():
    """447 전용 렌더(gray+bands) — 항목 5/6 검증용(447은 _SONGS에 없음)."""
    import numpy as _np
    pi = hp.find_song_page("나주노", 447)
    d = _fitz.open(hp.SOURCES["나주노"]["pdf"])
    page = d[pi]
    img = hp.render_region(page, None, hp.RENDER_DPI)
    gray = _np.asarray(img.convert("L"))
    bands = hp.detect_system_bands(gray)
    d.close()
    return {"gray": gray, "bands": bands}


@pytest.mark.parametrize("num", [362, 173, 146, 810, 267])
@pytest.mark.skipif(not _YOUTH_PDFS_AVAILABLE, reason=_YOUTH_SKIP_REASON)
def test_S1_S5_system_and_slide_counts(hymn_bands, num):
    """S1~S5: 각 곡의 시스템 수와 슬라이드 수(2개/장)가 요구사항 표와 일치."""
    bands = hymn_bands[num]["bands"]
    assert len(bands) == _EXPECTED_SYSTEMS[num]
    assert len(hp.group_systems(bands, 2)) == _EXPECTED_SLIDES[num]


@pytest.mark.parametrize("num", [362, 173, 146, 810, 267])
@pytest.mark.skipif(not _YOUTH_PDFS_AVAILABLE, reason=_YOUTH_SKIP_REASON)
def test_S6_bands_ordered_non_overlapping(hymn_bands, num):
    """S6: 밴드가 오름차순·비겹침이고, 상단 텍스트(제목/작곡가)가 첫 밴드에 포함되지 않는다.

    나주노는 페이지 전체를 렌더하므로 제목/작곡가가 콘텐츠 최상단에 있다 — 첫 시스템 밴드는
    반드시 그 아래에서 시작해야 한다(작곡가 병합 트림 검증). 크롭 PNG 육안 검증은 구현 노트
    참고(설계서 §3.4 실측 검증 원칙)."""
    info = hymn_bands[num]
    bands = info["bands"]
    for (t, b) in bands:
        assert b > t
    for (t0, b0), (t1, b1) in zip(bands, bands[1:]):
        assert t1 > b0  # 비겹침·오름차순
    if _SONGS[num] == "나주노":
        # 제목/작곡가가 콘텐츠 최상단에 있고, 첫 시스템은 그보다 아래에서 시작
        assert bands[0][0] > info["content_top"]


# --- 저작권 (C1~C4; C5는 슬라이드 조립 섹션) --------------------------------
@pytest.mark.parametrize("num", [362, 173, 146])
@pytest.mark.skipif(not _YOUTH_PDFS_AVAILABLE, reason=_YOUTH_SKIP_REASON)
def test_C1_naju_no_copyright(num):
    """C1: 나주노 362/173/146 → 저작권 크롭 None(이미지 요소 생성 안 함)."""
    pi = hp.find_song_page("나주노", num)
    d = _fitz.open(hp.SOURCES["나주노"]["pdf"])
    try:
        assert hp.resolve_copyright_crop("나주노", num, d[pi]) is None
    finally:
        d.close()


@pytest.mark.skipif(not _YOUTH_PDFS_AVAILABLE, reason=_YOUTH_SKIP_REASON)
def test_C2_naju_267_copyright_before_trigger():
    """C2: 나주노267 → 저작권 크롭 1개, 캐시 bbox가 'Adm. by' 트리거 전까지(가로 긴 한 줄)."""
    cache = hp._load_naju_cache()
    assert cache["267"]["has_copyright"] is True
    assert cache["267"]["trigger"] == "Adm. by"
    pi = hp.find_song_page("나주노", 267)
    d = _fitz.open(hp.SOURCES["나주노"]["pdf"])
    try:
        crop = hp.resolve_copyright_crop("나주노", 267, d[pi])
    finally:
        d.close()
    assert crop is not None
    assert crop.width > crop.height * 3  # 한 줄 텍스트(가로로 김)


@pytest.mark.skipif(not _YOUTH_PDFS_AVAILABLE, reason=_YOUTH_SKIP_REASON)
def test_C3_yahwe_810_right_edge_before_trigger(yahwe_doc):
    """C3: 야훼810 저작권 크롭 우측 경계가 'Administered by' 직전(그 이후 제외). 결정론적."""
    page = yahwe_doc[572]
    left_x, right_x, y0, y1 = hp.yahwe_copyright_bounds(page, 810)
    adm_x0 = min(h.x0 for h in page.search_for("Administered by"))
    assert left_x < right_x <= adm_x0        # 'Administered by' 시작 이전에서 끝남
    cw_x0 = min(h.x0 for h in page.search_for("Copyright"))
    assert abs(left_x - cw_x0) < 1           # 좌측은 'Copyright'에서 시작
    crop = hp.resolve_copyright_crop("야훼 이레", 810, page)
    assert crop is not None and crop.width > crop.height * 3


@pytest.mark.skipif(not _YOUTH_PDFS_AVAILABLE, reason=_YOUTH_SKIP_REASON)
def test_C4_naju_cache_miss_returns_none(caplog):
    """C4: 캐시에 없는 나주노 번호 → 경고 로그 + None(크래시/오크롭 없음)."""
    d = _fitz.open(hp.SOURCES["나주노"]["pdf"])
    try:
        import logging
        with caplog.at_level(logging.WARNING):
            crop = hp.resolve_copyright_crop("나주노", 999, d[0])
    finally:
        d.close()
    assert crop is None
    assert any("999" in r.message or "수동 확인" in r.message for r in caplog.records)


# 마침표(".")는 이 스캔·300dpi에서 폭 3~4px, 한글/라틴 글자는 ≥10px. 8px가 둘 사이 안전
# 경계다(글자로 끝나면 last-run 폭 > 8, 꼬리 마침표만 남으면 ≤ 8). test_C6 회귀 판정 기준.
_PERIOD_MAX_W_PX = 8


def _crop_col_ink_runs(pil_img):
    """크롭 이미지의 세로 잉크 열(어두운 픽셀 존재) 연속 구간 [(x0, x1, width), …]."""
    col = (np.asarray(pil_img.convert("L")) < 128).sum(axis=0)
    runs = []
    i, n = 0, len(col)
    while i < n:
        if col[i] > 0:
            s = i
            while i < n and col[i] > 0:
                i += 1
            runs.append((s, i - 1, i - s))
        else:
            i += 1
    return runs


@pytest.mark.parametrize("num", [447, 267])
@pytest.mark.skipif(not _YOUTH_PDFS_AVAILABLE, reason=_YOUTH_SKIP_REASON)
def test_C6_naju_copyright_excludes_trailing_period(num):
    """C6: 나주노 저작권 크롭이 이름 마지막 '글자'에서 끝난다(꼬리 마침표 제외).

    요구사항(§Copyright, 177~183행)은 '마침표 자체도 제외'를 명시한다. 이전 고정 4px back-off는
    마침표+공백 폭보다 작아 크롭 우측에 마침표가 남았고(리뷰 확정 이슈), 크기 비율만 보던
    기존 C2/P2는 이를 못 잡았다. 여기서는 런타임 경로(resolve_copyright_crop, 캐시 bbox·OCR
    없음)의 실제 크롭 우측 끝단 픽셀을 검사한다 — 마지막 세로 잉크 런이 글자 폭(>8px)이면
    글자로 끝난 것, 마침표 폭(≤8px)만 남으면 회귀."""
    d = _fitz.open(hp.SOURCES["나주노"]["pdf"])
    try:
        crop = hp.resolve_copyright_crop("나주노", num, d[hp.find_song_page("나주노", num)])
    finally:
        d.close()
    assert crop is not None
    runs = _crop_col_ink_runs(crop)
    assert len(runs) >= 4  # 'Copyright ©…이름' — 내용이 온전(과트림 방지)
    assert runs[-1][2] > _PERIOD_MAX_W_PX, (
        f"성가 {num}: 크롭이 글자가 아니라 마침표 폭({runs[-1][2]}px)으로 끝남 — 꼬리 마침표 잔존")


@pytest.mark.skipif(not _YOUTH_PDFS_AVAILABLE, reason=_YOUTH_SKIP_REASON)
def test_C6b_naju_267_keeps_internal_period():
    """C6b: 267은 이름 자체에 마침표가 있다('천태혁. 진경'). 꼬리 마침표만 제거하고 내부
    마침표는 보존돼야 한다 — 크롭 중간(마지막 런 이전)에 마침표 폭 런이 남아 있어야 하고,
    '진경'까지 온전히 포함돼(마지막 두 글자가 글자 폭) 과트림이 아니어야 한다."""
    d = _fitz.open(hp.SOURCES["나주노"]["pdf"])
    try:
        crop = hp.resolve_copyright_crop("나주노", 267, d[hp.find_song_page("나주노", 267)])
    finally:
        d.close()
    runs = _crop_col_ink_runs(crop)
    assert any(w <= _PERIOD_MAX_W_PX for _, _, w in runs[:-1]), "내부 마침표(천태혁.)가 사라짐"
    assert runs[-1][2] > _PERIOD_MAX_W_PX and runs[-2][2] > _PERIOD_MAX_W_PX, "'진경' 과트림"


@pytest.mark.skipif(not _YOUTH_PDFS_AVAILABLE, reason=_YOUTH_SKIP_REASON)
def test_C6c_447_matches_sample_no_trailing_period():
    """C6c: 447 런타임 크롭이 제공 샘플(SAMPLE447)과 같은 '마침표 없는 글자 끝' 속성을 가진다.

    샘플·재생성 크롭 모두 마지막 세로 잉크 런이 글자 폭(>8px)으로, 저작권자 이름 뒤 마침표가
    없다. 리뷰어가 지적한 '샘플과의 직접 대조' 갭을 이 속성 일치로 닫는다(해상도가 달라 픽셀
    완전 일치는 부적합)."""
    d = _fitz.open(hp.SOURCES["나주노"]["pdf"])
    try:
        crop = hp.resolve_copyright_crop("나주노", 447, d[hp.find_song_page("나주노", 447)])
    finally:
        d.close()
    sample = Image.open(
        Path(__file__).resolve().parent / "_workspace" / "청년미사_1단계" / "crops"
        / "SAMPLE447_copyright.png")
    assert _crop_col_ink_runs(sample)[-1][2] > _PERIOD_MAX_W_PX   # 샘플도 글자로 끝남
    assert _crop_col_ink_runs(crop)[-1][2] > _PERIOD_MAX_W_PX     # 재생성도 동일 속성


# --- 헤더 / 산출물 구조 (Hd1~Hd4, C5, P1~P4) --------------------------------
from pptx.oxml.ns import qn as _qn


def _hdr_runs(slide):
    for sh in slide.shapes:
        if sh.name == "Rectangle 11":
            return sh.text_frame.paragraphs[0]._p.findall(_qn("a:r"))
    return []


def _run_text(r):
    t = r.find(_qn("a:t"))
    return t.text if t is not None else ""


def _run_color(r):
    rPr = r.find(_qn("a:rPr"))
    if rPr is None:
        return None
    fill = rPr.find(_qn("a:solidFill"))
    if fill is None:
        return None
    srgb = fill.find(_qn("a:srgbClr"))
    if srgb is not None:
        return ("srgb", srgb.get("val"))
    prst = fill.find(_qn("a:prstClr"))
    if prst is not None:
        return ("prst", prst.get("val"))
    return None


def _score_pics(slide):
    return [s for s in slide.shapes if s.shape_type == 13 and s.top < 4800000]


def _copyright_pics(slide):
    return [s for s in slide.shapes if s.shape_type == 13 and s.top >= 4800000]


@pytest.fixture(scope="module")
def built():
    return {
        362: hp.build_hymn_pptx("입당", "나주노", 362),
        267: hp.build_hymn_pptx("파견", "나주노", 267),
        447: hp.build_hymn_pptx("입당", "나주노", 447),
        810: hp.build_hymn_pptx("2차봉헌", "야훼 이레", 810),
    }


@pytest.mark.skipif(not _YOUTH_PDFS_AVAILABLE, reason=_YOUTH_SKIP_REASON)
def test_Hd1_header_on_every_slide(built):
    """Hd1: 'Rectangle 11' 헤더가 모든 슬라이드에 반복."""
    for num, prs in built.items():
        for slide in prs.slides:
            assert any(sh.name == "Rectangle 11" for sh in slide.shapes), num


@pytest.mark.skipif(not _YOUTH_PDFS_AVAILABLE, reason=_YOUTH_SKIP_REASON)
def test_Hd2_run_order_and_colors(built):
    """Hd2: run 순서·색 — 구분=흰색(prstClr white), 출처·번호·제목=FFC000."""
    runs = _hdr_runs(built[362].slides[0])
    texts = [_run_text(r) for r in runs]
    assert texts[0] == "입당 "
    assert _run_color(runs[0]) == ("prst", "white")
    # 공백 아닌 세그먼트: 출처/번호/제목
    seg = [(_run_text(r), _run_color(r)) for r in runs[1:] if _run_text(r).strip()]
    assert seg[0] == ("나주노", ("srgb", "FFC000"))
    assert seg[1] == ("362", ("srgb", "FFC000"))
    assert seg[2] == ("삼위일체", ("srgb", "FFC000"))


@pytest.mark.skipif(not _YOUTH_PDFS_AVAILABLE, reason=_YOUTH_SKIP_REASON)
def test_Hd3_font_size_and_spacing_preserved(built):
    """Hd3: 헤더 모든 run이 28pt(sz=2800)·자간 spc=300 유지(템플릿 서식 보존)."""
    runs = _hdr_runs(built[362].slides[0])
    for r in runs:
        rPr = r.find(_qn("a:rPr"))
        assert rPr.get("sz") == "2800"
        assert rPr.get("spc") == "300"


@pytest.mark.skipif(not _YOUTH_PDFS_AVAILABLE, reason=_YOUTH_SKIP_REASON)
def test_Hd4_unknown_구분_raises():
    """Hd4: 구분이 HYMN_TYPES에 없으면 에러."""
    with pytest.raises(ValueError):
        hp.build_hymn_pptx("존재하지않는구분", "나주노", 362)


@pytest.mark.skipif(not _YOUTH_PDFS_AVAILABLE, reason=_YOUTH_SKIP_REASON)
def test_C1_C5_copyright_presence_and_position(built):
    """C1/C5: 362=저작권 0개; 267/810/447=저작권 1개씩, 우하단 샘플 계열 위치·높이."""
    for slide in built[362].slides:
        assert len(_copyright_pics(slide)) == 0
    for num in (267, 810, 447):
        for slide in built[num].slides:
            pics = _copyright_pics(slide)
            assert len(pics) == 1, num
            p = pics[0]
            assert abs(p.top - 4855468) < 30000            # 샘플 top
            assert abs(p.height - 288032) < 30000          # 샘플 높이 계열
            assert abs((p.left + p.width) - 8510525) < 30000  # 샘플 우측 끝


@pytest.mark.skipif(not _YOUTH_PDFS_AVAILABLE, reason=_YOUTH_SKIP_REASON)
def test_P1_slide_size(built):
    """P1: 슬라이드 크기 9144000×5143500."""
    for prs in built.values():
        assert prs.slide_width == 9144000
        assert prs.slide_height == 5143500


@pytest.mark.skipif(not _YOUTH_PDFS_AVAILABLE, reason=_YOUTH_SKIP_REASON)
def test_P2_447_matches_sample(built):
    """P2: 나주노447 재생성 → 4장, 헤더 텍스트·색, 저작권 유무 샘플과 일치."""
    prs = built[447]
    assert len(prs.slides) == 4
    runs = _hdr_runs(prs.slides[0])
    seg = [(_run_text(r), _run_color(r)) for r in runs[1:] if _run_text(r).strip()]
    assert _run_text(runs[0]) == "입당 " and _run_color(runs[0]) == ("prst", "white")
    assert seg[0] == ("나주노", ("srgb", "FFC000"))
    assert seg[1] == ("447", ("srgb", "FFC000"))
    assert seg[2] == ("영원도 하시어라 그 사랑이여", ("srgb", "FFC000"))
    for slide in prs.slides:
        assert len(_copyright_pics(slide)) == 1  # 447은 저작권 있음


@pytest.mark.skipif(not _YOUTH_PDFS_AVAILABLE, reason=_YOUTH_SKIP_REASON)
def test_P3_validate_structure(built, tmp_path):
    """P3: validate_pptx_structure 통과(구조 손상 없음). PowerPoint 육안은 구현 노트."""
    from missa_sections import validate_pptx_structure
    for num, prs in built.items():
        f = tmp_path / f"out_{num}.pptx"
        prs.save(str(f))
        assert validate_pptx_structure(str(f)) == []


@pytest.mark.skipif(not _YOUTH_PDFS_AVAILABLE, reason=_YOUTH_SKIP_REASON)
def test_P4_one_score_per_slide_below_header(built):
    """P4: 슬라이드별 악보 크롭 1장, 헤더(523220) 아래 배치."""
    for num, prs in built.items():
        for slide in prs.slides:
            score = _score_pics(slide)
            assert len(score) == 1, num
            assert score[0].top >= 523220


@pytest.mark.skipif(not _YOUTH_PDFS_AVAILABLE, reason=_YOUTH_SKIP_REASON)
def test_Hd_header_child_order_valid(built):
    """헤더 <a:p> 자식 순서 검증 — 모든 run이 endParaRPr **앞**에 온다.

    validate_pptx_structure()는 요소 순서를 검사하지 않으므로(CLAUDE.md 2026-08-29,
    ooxml-pitfalls) run/endParaRPr 순서를 lxml로 직접 단언한다. build_header_runs가
    _insert_run_before_end_para_rpr로 삽입하지 않으면 run이 endParaRPr 뒤로 가 스키마 위반이
    되지만 검사기·다른 테스트를 통과해버린다."""
    for num, prs in built.items():
        for slide in prs.slides:
            for sh in slide.shapes:
                if sh.name != "Rectangle 11":
                    continue
                p = sh.text_frame.paragraphs[0]._p
                children = [c.tag for c in p]
                eprp_tag = _qn("a:endParaRPr")
                r_tag = _qn("a:r")
                if eprp_tag in children:
                    eprp_i = children.index(eprp_tag)
                    last_r = max((i for i, t in enumerate(children) if t == r_tag), default=-1)
                    assert last_r < eprp_i, (num, "run이 endParaRPr 뒤에 옴")


# ===========================================================================
# 청년미사 1단계 후속(v2) — 설계서 01b_architect_design_v2.md
# 항목 6(P5~P7) per-group 크롭 · 항목 5(S7~S11) 동적 패킹 · 항목 3(Hd8) 야훼이레 표시명
# 항목 4(C7 폴백 마침표 · C8 높이0.9 · C9 447 샘플 종횡비) · 항목 2(Hf1~Hf4 헤더 축소)
# 항목 1(M1~M9 다중 섹션 파싱)
# ===========================================================================


# --- 항목 6: per-group 크롭 bbox (P5~P7) -----------------------------------
def _tail_has_ink(gray, box, frac=0.98):
    """크롭 박스의 우측 끝(frac~1.0 열 범위)에 잉크가 있으면 True(우측 흰 여백 없음)."""
    left, top, right, bottom = box
    sub = (gray[top:bottom, left:right] < hp._INK)
    w = sub.shape[1]
    return bool(sub[:, int(w * frac):].any())


# (P5~P7은 test_missa_regression.py로 승격됨 — 2026-09-11 QA.)


# --- 항목 5: 동적 패킹 (S7~S11) --------------------------------------------
def _sizes(groups):
    return [len(g) for g in groups]


def _synthetic_gray(bands, width):
    """밴드 tuple 목록에 맞춰 전 폭에 잉크가 찬 합성 gray(그 외 흰색)를 만든다."""
    h = max(b for _, b in bands) + 5
    g = np.full((h, width), 255, np.uint8)
    for t, b in bands:
        g[t:b + 1, :] = 0
    return g


_PACK_EXPECTED = {362: [2, 2], 173: [2, 2, 2, 1], 146: [2, 2, 2],
                  810: [2, 2], 267: [2, 2, 1]}


@pytest.mark.parametrize("num", [362, 173, 146, 810, 267])
@pytest.mark.skipif(not _YOUTH_PDFS_AVAILABLE, reason=_YOUTH_SKIP_REASON)
def test_S7_pack_systems_no_regression_5songs(hymn_bands, num):
    """S7: min-slide 균형 패커가 5곡 각각 현재와 동일한 그룹 크기를 반환(무회귀)."""
    groups = hp.pack_systems(hymn_bands[num]["bands"], hymn_bands[num]["gray"])
    assert _sizes(groups) == _PACK_EXPECTED[num]


@pytest.mark.skipif(not _YOUTH_PDFS_AVAILABLE, reason=_YOUTH_SKIP_REASON)
def test_S7_pack_systems_447(render_447):
    """S7: 447도 현재와 동일 [2,2,2,1](무회귀)."""
    groups = hp.pack_systems(render_447["bands"], render_447["gray"])
    assert _sizes(groups) == [2, 2, 2, 1]


def test_S8_three_pack_reduces_slides():
    """S8: 짧은 시스템 5개, 첫 3개가 종횡비·safety 통과 → [3,2] 2장(현재 규칙이면 [2,2,1] 3장).
    3-pack이 슬라이드를 줄인다. (합성 데이터라 PDF 불필요.)"""
    bands = [(0, 39), (60, 99), (120, 159), (180, 219), (240, 279)]
    gray = _synthetic_gray(bands, width=2000)
    assert _sizes(hp.pack_systems(bands, gray)) == [3, 2]


def test_S9_tight_three_pack_rejected_by_safety():
    """S9: 3개가 종횡비는 통과하나 vfill이 safety 초과(810형 97%) → 3-pack 탈락, [2,2] 유지."""
    # 3-band span/width = 449/1000 = 0.449 > band_aspect*safety(0.463*0.95=0.440),
    # 하지만 ≤ band_aspect(0.463) — 즉 겹치진 않으나 너무 빡빡. 2-band=289/1000=0.289 통과.
    bands = [(0, 129), (160, 289), (320, 449), (480, 609)]
    gray = _synthetic_gray(bands, width=1000)
    assert _sizes(hp.pack_systems(bands, gray)) == [2, 2]


def test_S10_balance_prefers_2_2_over_3_1():
    """S10: 3-pack이 기하적으로 들어가도 슬라이드 수가 같으면 [3,1]이 아니라 균형 [2,2]."""
    bands = [(0, 39), (60, 99), (120, 159), (180, 219)]  # 4개 다 짧음(3-pack 여유 통과)
    gray = _synthetic_gray(bands, width=2000)
    assert _sizes(hp.pack_systems(bands, gray)) == [2, 2]


# S11(3-pack 발동 슬라이드 PNG 겹침 없음)은 육안 검증 — 구현 노트 참고.


# --- 항목 3: 야훼이레 표시명 (Hd8) -----------------------------------------
@pytest.mark.skipif(not _YOUTH_PDFS_AVAILABLE, reason=_YOUTH_SKIP_REASON)
def test_Hd8_yahwe_display_label_no_space(built):
    """Hd8: 야훼 810 헤더 출처 run == '야훼이레'(공백 없음). SOURCES 키/조회는 '야훼 이레' 유지."""
    runs = _hdr_runs(built[810].slides[0])
    seg = [(_run_text(r), _run_color(r)) for r in runs[1:] if _run_text(r).strip()]
    assert seg[0] == ("야훼이레", ("srgb", "FFC000"))
    assert hp.find_song_page("야훼 이레", 810) == 572  # 내부 조회는 원본 키 유지


def test_Hd8_source_display_mapping():
    """Hd8: SOURCE_DISPLAY는 표시명만 매핑, SOURCES 키는 원본 유지."""
    assert hp.SOURCE_DISPLAY.get("야훼 이레") == "야훼이레"
    assert "야훼 이레" in hp.SOURCES and "야훼이레" not in hp.SOURCES


# --- 항목 4-1: 폴백 꼬리 마침표 트림 일반화 (C7) ---------------------------
def test_C7_trailing_period_trim_generalized():
    """C7: _trim_trailing_period_px가 트리거 유무와 무관하게 우측 꼬리 마침표를 제거하고,
    이름 내부 마침표(267식)는 보존한다. (합성 이미지라 PDF 불필요.)"""
    from PIL import Image as _Im
    # 글자(20px) + 공백 + 꼬리 마침표(4px)
    a = np.full((30, 60), 255, np.uint8)
    a[10:20, 5:25] = 0
    a[16:20, 40:44] = 0
    out = hp._trim_trailing_period_px(_Im.fromarray(a))
    runs = _crop_col_ink_runs(out)
    assert runs[-1][2] > _PERIOD_MAX_W_PX  # 마침표 제거됨(마지막 런=글자)
    # 글자 + 내부 마침표 + 글자 → 내부 마침표 보존, 트림 없음
    b = np.full((30, 80), 255, np.uint8)
    b[10:20, 5:25] = 0
    b[16:20, 30:34] = 0
    b[10:20, 45:65] = 0
    out2 = hp._trim_trailing_period_px(_Im.fromarray(b))
    runs2 = _crop_col_ink_runs(out2)
    assert any(w <= _PERIOD_MAX_W_PX for _, _, w in runs2[:-1])  # 내부 마침표 보존
    assert runs2[-1][2] > _PERIOD_MAX_W_PX


# --- 항목 4-2: 저작권 높이 0.9배 (C8) --------------------------------------
@pytest.mark.skipif(not _YOUTH_PDFS_AVAILABLE, reason=_YOUTH_SKIP_REASON)
def test_C8_copyright_height_fixed_09(built):
    """C8: 저작권 이미지 높이 == round(288032*0.9)=259229(전곡), 폭 == round(높이×크롭종횡비),
    우측끝 8510525·바닥 5143500 앵커 유지."""
    assert hp._COPY_H == 259229
    for num in (267, 810, 447):
        for slide in built[num].slides:
            p = _copyright_pics(slide)[0]
            assert p.height == 259229
            assert abs((p.left + p.width) - 8510525) < 3        # 우측끝 앵커
            assert abs((p.top + p.height) - 5143500) < 3        # 바닥 앵커


# --- 항목 4-3: 447 캐시 크롭이 샘플 종횡비에 근접 (C9) ----------------------
@pytest.mark.skipif(not _YOUTH_PDFS_AVAILABLE, reason=_YOUTH_SKIP_REASON)
def test_C9_447_crop_aspect_matches_sample():
    """C9: 447 저작권 크롭 종횡비가 손제작 샘플(≈5.17)에 근접(±0.5). 재조정 전 6.03은 red.

    샘플보다 과도하게 넓던 447 크롭을 샘플 비율로 재조정(항목 4-3). 크롭은 여전히 마침표 없이
    글자로 끝나야 한다(C6c 속성 유지)."""
    d = _fitz.open(hp.SOURCES["나주노"]["pdf"])
    try:
        crop = hp.resolve_copyright_crop("나주노", 447, d[hp.find_song_page("나주노", 447)])
    finally:
        d.close()
    aspect = crop.width / crop.height
    sample = Image.open(
        Path(__file__).resolve().parent / "_workspace" / "청년미사_1단계" / "crops"
        / "SAMPLE447_copyright.png")
    s_aspect = sample.width / sample.height
    assert abs(aspect - s_aspect) < 0.5, (aspect, s_aspect)
    assert _crop_col_ink_runs(crop)[-1][2] > _PERIOD_MAX_W_PX  # 여전히 글자로 끝남


# --- 항목 2: 헤더 자동 축소 (Hf1~Hf4) --------------------------------------
def _template_header_runs(구분, 출처, number, title):
    prs = hp.Presentation(str(hp.TEMPLATE))
    hp.build_header_runs(prs.slides[0], 구분, 출처, number, title)
    return _hdr_runs(prs.slides[0])


def _run_sz(r):
    rPr = r.find(_qn("a:rPr"))
    return rPr.get("sz") if rPr is not None else None


def _run_spc(r):
    rPr = r.find(_qn("a:rPr"))
    return rPr.get("spc") if rPr is not None else None


@pytest.mark.skipif(not _YOUTH_PDFS_AVAILABLE, reason=_YOUTH_SKIP_REASON)
def test_Hf1_short_title_no_shrink(built):
    """Hf1: 짧은 제목(362 '삼위일체')은 spc=300·28pt 그대로(축소 없음)."""
    for r in _hdr_runs(built[362].slides[0]):
        assert _run_spc(r) == "300"
        assert _run_sz(r) == "2800"


def test_Hf2_overflow_reduces_spacing_only():
    """Hf2: spc=300에서 넘치나 spc=0이면 들어가는 제목 → 전 run spc→0, 폰트는 28pt 유지."""
    runs = _template_header_runs("입당", "나주노", 447, "가" * 15)
    for r in runs:
        assert _run_spc(r) == "0"
        assert _run_sz(r) == "2800"


def test_Hf3_still_overflow_reduces_title_font_only():
    """Hf3: spc=0으로도 넘치는 제목 → 제목 run sz만 축소(<2800), 구분/출처/번호는 2800 유지."""
    runs = _template_header_runs("입당", "나주노", 447, "가" * 20)
    title_r = runs[-1]
    assert int(_run_sz(title_r)) < 2800
    assert int(_run_sz(title_r)) >= hp._MIN_TITLE_PT * 100
    # 제목 아닌 run(구분/출처/번호/공백)은 2800 유지
    for r in runs[:-1]:
        assert _run_sz(r) == "2800"


def test_Hf4_after_shrink_fits_usable():
    """Hf4: 축소 후 최종 측정폭 ≤ usable(안 넘침)."""
    prs = hp.Presentation(str(hp.TEMPLATE))
    hp.build_header_runs(prs.slides[0], "입당", "나주노", 447, "가" * 20)
    hdr = hp._find_header(prs.slides[0])
    runs = _hdr_runs(prs.slides[0])
    total = hp._header_total_emu(runs)
    assert total <= hp._usable_header_emu(hdr)


# --- 항목 1: 다중 섹션 파싱 (M1~M9) ----------------------------------------
def test_M1_parse_mass_all_sections():
    """M1: parse_mass가 5개 섹션 키 반환, Gospel reference 정확."""
    m = yg.parse_mass(_fix_html("20260913"))
    for k in ("First reading", "Responsorial Psalm", "Second reading",
              "Gospel Acclamation", "Gospel"):
        assert k in m, k
    assert m["Gospel"]["reference"] == "Matthew 18:21-35"


def test_M2_first_reading():
    m = yg.parse_mass(_fix_html("20260913"))
    assert m["First reading"]["reference"] == "Ecclesiasticus 27:33-28:9"
    assert m["First reading"]["content"]


def test_M3_psalm_reference_only():
    """M3: Responsorial Psalm은 reference만(content/theme 키 없음)."""
    psalm = yg.parse_mass(_fix_html("20260913"))["Responsorial Psalm"]
    assert psalm["reference"] == "Psalm 102(103):1-4,9-12"
    assert "content" not in psalm and "theme" not in psalm


def test_M4_acclamation_content_excludes_or_body():
    """M4: Acclamation content가 'Alleluia'로 시작, Or: 본문이 섞이지 않음(다음 섹션 종료 경계)."""
    accl = yg.parse_mass(_fix_html("20260913"))["Gospel Acclamation"]
    assert accl["content"].startswith("Alleluia")
    assert "I give you a new commandment" not in accl["content"]
    assert accl["content"].count("\n\n") >= 3


def test_M5_acclamation_or_nested():
    """M5: Gospel Acclamation의 'or' 중첩 = {reference 'Jn13:34', content '…love one another…'}."""
    accl = yg.parse_mass(_fix_html("20260913"))["Gospel Acclamation"]
    assert accl["or"]["reference"] == "Jn13:34"
    assert "love one another" in accl["or"]["content"]


def test_M6_no_korean_keys():
    """M6: 모든 키가 ASCII(한글 '복음' 부재)."""
    m = yg.parse_mass(_fix_html("20260913"))
    def _all_keys(d):
        for k, v in d.items():
            yield k
            if isinstance(v, dict):
                yield from _all_keys(v)
    for k in _all_keys(m):
        assert k.isascii(), k


def test_M7_theme_presence_by_section():
    """M7: First/Second/Gospel엔 theme, Psalm/Acclamation엔 theme 키 없음."""
    m = yg.parse_mass(_fix_html("20260913"))
    assert m["First reading"]["theme"] == "Forgive your neighbour the hurt he does you; and when you pray, your sins will be forgiven"
    assert m["Second reading"]["theme"] == "Alive or dead, we belong to the Lord"
    assert m["Gospel"]["theme"] == "To be forgiven, you must forgive"
    assert "theme" not in m["Responsorial Psalm"]
    assert "theme" not in m["Gospel Acclamation"]


def test_M8_main_default_filename(monkeypatch, tmp_path):
    """M8: main 인자 미지정 시 missa_en_YYYYMMDD.json으로 저장, 영어 섹션 포함."""
    monkeypatch.setattr(yg, "fetch_gospel_html", lambda d: _fix_html("20260913"))
    monkeypatch.chdir(tmp_path)
    rc = yg.main(["20260913"])
    assert rc == 0
    out = tmp_path / "missa_en_20260913.json"
    assert out.exists()
    import json as _json
    data = _json.loads(out.read_text(encoding="utf-8"))
    assert data["date"] == "20260913"
    assert data["Gospel"]["reference"] == "Matthew 18:21-35"
    assert "복음" not in data


def test_M9_get_youth_mass_schema(monkeypatch):
    """M9: get_youth_mass가 date + 영어 섹션 반환, Psalm 빈 본문은 오류 아님."""
    monkeypatch.setattr(yg, "fetch_gospel_html", lambda d: _fix_html("20260913"))
    out = yg.get_youth_mass("20260913")
    assert out["date"] == "20260913"
    assert out["First reading"]["reference"] == "Ecclesiasticus 27:33-28:9"
    assert "복음" not in out
