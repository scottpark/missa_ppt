"""
missa_to_ppt.py 회귀 테스트 스위트

목적
----
주일미사·평일미사 처리 로직의 회귀 테스트. 향후 토요일미사 지원이 추가되거나
다른 기능이 변경됐을 때, 기존 주일/평일 로직이 깨지지 않았는지 확인하는 기준선.

실행
----
    pip install pytest
    pytest test_missa_regression.py -v

테스트 데이터
------------
- 주일: 20260712 (2026-07-12, 연중 제15주일)
- 평일: 20260624 (2026-06-24, 성 요한 세례자 탄생 대축일, 수요일)

두 폴더 모두 참조 PPT·JSON·성가 파일이 준비되어 있어야 한다.

새 케이스(예: 토요일미사) 추가 시 MASS_CASES 리스트에 항목을 추가하면 된다.
"""
from __future__ import annotations

import copy
import json
import re
import subprocess
import sys
import types
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

import missa_to_ppt as m

REPO_ROOT = Path(__file__).resolve().parent

MASS_CASES = [
    {
        "id": "sunday",
        "date": "20260712",
        "is_sunday": True,
        "hymns": {"입당": "329", "봉헌": "221", "성체": "156", "2차봉헌": "221", "파견": "25"},
    },
    {
        "id": "weekday",
        "date": "20260624",
        "is_sunday": False,
        "hymns": {"입당": "329", "봉헌": "221", "성체": "156", "파견": "25"},
    },
]

READING_SECTIONS = [
    ("제1독서_start", "제1독서_end"),
    ("제2독서_start", "제2독서_end"),
    ("복음_start", "복음_end"),
]

HYMN_TYPES = ["입당", "봉헌", "성체", "2차봉헌", "파견"]


# ─────────────────────────────────────────────────────────────────
# 1. 단위 테스트 — 순수 함수
# ─────────────────────────────────────────────────────────────────

class TestIsSundayMass:
    """향후 토요일미사(전야미사) 판단 로직 추가 시 이 클래스에
    '토요일을 평일미사/주일 전야 중 무엇으로 볼 것인가' 케이스를 추가해서
    회귀를 잡는다."""

    def test_sunday(self):
        assert m.is_sunday_mass("20260712") is True  # 일요일

    def test_weekday_friday(self):
        assert m.is_sunday_mass("20260612") is False  # 금요일

    def test_weekday_wednesday(self):
        assert m.is_sunday_mass("20260624") is False  # 수요일

    def test_saturday_currently_treated_as_weekday(self):
        # 2026-07-11은 토요일. 현재 로직(주일 여부만 판단)에서는 '평일미사'로 처리된다.
        # 토요일미사 지원이 추가되면 이 테스트는 새 기대값으로 업데이트해야 한다.
        assert m.is_sunday_mass("20260711") is False


class TestWrapLineCount:
    def test_empty_text(self):
        assert m._wrap_line_count("") == 0

    def test_short_line_is_one_line(self):
        assert m._wrap_line_count("짧은 문장입니다.") == 1

    def test_long_text_wraps_to_multiple_lines(self):
        text = "가" * 60  # CHARS_PER_LINE(27) 초과
        assert m._wrap_line_count(text) >= 2


class TestParseIntoVerseUnits:
    def test_verse_numbers_are_detected(self):
        content = (
            "1 태초에 하느님께서 하늘과 땅을 창조하셨다. "
            "2 땅은 아직 모양을 갖추지 않고 비어 있었다."
        )
        units = m.parse_into_verse_units(content)
        assert len(units) >= 2
        assert units[0]["verse_num"] == "1"


class TestSplitParaPreservesVerseColors:
    """2026-07-26 발견: 여러 절이 하나의 논리 단락으로 병합된 경우(예: continuation
    절), _split_para_at_lines()가 문단을 두 슬라이드로 분리하면서 단락 중간에 있는
    절 번호 run의 오렌지색이 사라지는 회귀를 방지한다. 당시 원인은 이 함수가 단락에
    run이 정확히 2개(절 번호+본문)라고 가정하고 3개 이상이면 뒤쪽 run들의 서식을
    버렸기 때문이다."""

    @staticmethod
    def _build_two_verse_para():
        from PIL import ImageFont
        from pptx.oxml import parse_xml as pptx_parse_xml

        A = 'http://schemas.openxmlformats.org/drawingml/2006/main'
        xml = (
            f'<a:p xmlns:a="{A}">'
            f'<a:r><a:rPr sz="3200"><a:solidFill><a:srgbClr val="FFC000"/></a:solidFill></a:rPr><a:t>51 </a:t></a:r>'
            f'<a:r><a:rPr sz="3200"><a:solidFill><a:schemeClr val="bg1"/></a:solidFill></a:rPr><a:t>many words here for wrapping test purposes today </a:t></a:r>'
            f'<a:r><a:rPr sz="3200"><a:solidFill><a:srgbClr val="FFC000"/></a:solidFill></a:rPr><a:t>52 </a:t></a:r>'
            f'<a:r><a:rPr sz="3200"><a:solidFill><a:schemeClr val="bg1"/></a:solidFill></a:rPr><a:t>more words after the second verse marker appear here</a:t></a:r>'
            f'</a:p>'
        )
        p_elem = pptx_parse_xml(xml)
        font = ImageFont.truetype(r'C:\Windows\Fonts\arial.ttf', 32)
        return p_elem, font

    @staticmethod
    def _orange_texts(p_elem):
        A = 'http://schemas.openxmlformats.org/drawingml/2006/main'
        out = []
        for r in p_elem.findall(f'{{{A}}}r'):
            rPr = r.find(f'{{{A}}}rPr')
            t = r.find(f'{{{A}}}t')
            if rPr is None or t is None:
                continue
            sf = rPr.find(f'{{{A}}}solidFill')
            if sf is None:
                continue
            sc = sf.find(f'{{{A}}}srgbClr')
            if sc is not None and sc.get('val', '').upper() == 'FFC000':
                out.append((t.text or '').strip())
        return out

    def test_verse_numbers_survive_paragraph_split(self):
        p_elem, font = self._build_two_verse_para()
        box_px = 220  # 좁게 잡아 여러 줄로 강제 wrap
        rest_p = m._split_para_at_lines(p_elem, keep_lines=1, pil_font=font, box_px=box_px)
        assert rest_p is not None, "테스트 문단이 1줄에 다 들어가 분리가 일어나지 않음 — box_px를 줄일 것"
        combined = set(self._orange_texts(p_elem)) | set(self._orange_texts(rest_p))
        assert combined == {"51", "52"}, f"절 번호 오렌지색 유실: {combined}"


class TestSplitAndAdjustViaCom:
    """_split_and_adjust_via_com()이 COM 실측과 Pillow 추정이 어긋날 때
    분리 지점(keep)을 올바른 방향으로 조정하는지, 최대 조정 횟수를 지키는지,
    조정 과정에서도 run 서식(오렌지 절 번호 등)이 유지되는지 확인한다.
    _split_para_at_lines() 자체는 "몇 줄인지 판단"에만 관여하고 분리 지점
    계산은 여전히 Pillow 기반이므로, COM이 최종 결과가 어긋났다고 보고하면
    keep을 ±1 조정해 재분리해야 한다(2026-08-04 발견: COM 검증이 도입된 뒤에도
    분리 지점 계산 자체는 Pillow 편향을 그대로 물려받아 재시도 캡 내에서
    수렴하지 못하는 사례가 실측으로 확인됨)."""

    @staticmethod
    def _build_para(text, sz=3200):
        from pptx.oxml import parse_xml as pptx_parse_xml
        A = 'http://schemas.openxmlformats.org/drawingml/2006/main'
        xml = f'<a:p xmlns:a="{A}"><a:r><a:rPr sz="{sz}"/><a:t>{text}</a:t></a:r></a:p>'
        return pptx_parse_xml(xml)

    @staticmethod
    def _text(p_elem):
        A = 'http://schemas.openxmlformats.org/drawingml/2006/main'
        return ''.join(
            (r.find(f'{{{A}}}t').text or '') for r in p_elem.findall(f'{{{A}}}r')
        )

    @staticmethod
    def _font():
        from PIL import ImageFont
        return ImageFont.truetype(r'C:\Windows\Fonts\arial.ttf', 32)

    def _long_text(self):
        return (
            "alpha beta gamma delta epsilon zeta eta theta iota kappa "
            "lambda mu nu xi omicron pi rho sigma tau upsilon phi chi psi omega"
        )

    def test_restore_para_from_backup_reverts_split(self):
        p = self._build_para(self._long_text())
        backup = copy.deepcopy(p)
        rest = m._split_para_at_lines(p, keep_lines=1, pil_font=self._font(), box_px=150)
        assert rest is not None, "테스트 문단이 1줄에 다 들어가 분리가 일어나지 않음"
        assert self._text(p) != self._text(backup)
        m._restore_para_from_backup(p, backup)
        assert self._text(p) == self._text(backup)

    def _run_adjust(self, monkeypatch, responses, **kwargs):
        calls = []

        def fake_verified(prs, slide):
            val = responses[len(calls)] if len(calls) < len(responses) else responses[-1]
            calls.append(val)
            return val

        monkeypatch.setattr(m, "_count_slide_lines_verified", fake_verified)
        placed = []

        def place_rest(rp):
            placed.append(rp)

        def remove_rest(rp):
            placed.remove(rp)

        defaults = dict(
            prs=None, cur_slide=object(), pil_font=self._font(), box_px=150,
            place_rest=place_rest, remove_rest=remove_rest, max_adjust=2,
        )
        defaults.update(kwargs)
        result = m._split_and_adjust_via_com(**defaults)
        return result, calls, placed

    def test_keep_increases_when_cur_slide_too_short(self, monkeypatch):
        p = self._build_para(self._long_text())
        (rest_p, final_lines), calls, placed = self._run_adjust(
            monkeypatch,
            responses=[m.LINES_PER_SLIDE - 1, m.LINES_PER_SLIDE],
            p_elem=p, keep=2,
        )
        assert final_lines == m.LINES_PER_SLIDE
        assert len(calls) == 2
        assert len(placed) == 1  # 마지막으로 배치된 rest_p 하나만 남아 있어야 함

    def test_keep_decreases_when_cur_slide_too_long(self, monkeypatch):
        p = self._build_para(self._long_text())
        (rest_p, final_lines), calls, placed = self._run_adjust(
            monkeypatch,
            responses=[m.LINES_PER_SLIDE + 1, m.LINES_PER_SLIDE],
            p_elem=p, keep=3,
        )
        assert final_lines == m.LINES_PER_SLIDE
        assert len(calls) == 2
        assert len(placed) == 1

    def test_gives_up_after_max_adjust_without_crashing(self, monkeypatch):
        p = self._build_para(self._long_text())
        original_text = self._text(p)
        (rest_p, final_lines), calls, placed = self._run_adjust(
            monkeypatch,
            responses=[m.LINES_PER_SLIDE + 1] * 5,  # 절대 수렴하지 않는 상황(항상 초과)
            p_elem=p, keep=3, max_adjust=2,
        )
        # 초기 1회 + 조정 최대 2회 = 최대 3회 호출로 멈춰야 함(무한 루프 금지)
        assert len(calls) <= 3
        # 2026-08-04 발견 버그의 회귀 방지: 재시도 캡을 다 써도 여전히
        # LINES_PER_SLIDE를 초과하면(오버플로가 남으면) 이 시도를 통째로
        # 되돌려야 한다 — "일부만 고쳐진 채로 오버플로가 남은" 상태를
        # 성공(rest_p is not None)으로 잘못 보고하면 안 된다.
        assert rest_p is None
        assert len(placed) == 0
        assert self._text(p) == original_text

    def test_never_reports_success_while_still_overflowing(self, monkeypatch):
        """조정을 거듭해도 실측이 계속 LINES_PER_SLIDE를 넘으면(수렴 실패),
        절대 rest_p를 성공으로 반환하면 안 된다 — 반환하면 호출부가 "고쳐졌다"고
        오인해 실제로는 넘치는 슬라이드를 그대로 최종본에 남기게 된다."""
        p = self._build_para(self._long_text())
        (rest_p, final_lines), calls, placed = self._run_adjust(
            monkeypatch,
            responses=[m.LINES_PER_SLIDE + 1, m.LINES_PER_SLIDE + 1],
            p_elem=p, keep=2, max_adjust=1,
        )
        assert rest_p is None
        assert len(placed) == 0

    def test_verse_colors_survive_adjustment_retry(self, monkeypatch):
        p_elem, font = TestSplitParaPreservesVerseColors._build_two_verse_para()
        (rest_p, final_lines), calls, placed = self._run_adjust(
            monkeypatch,
            responses=[m.LINES_PER_SLIDE - 1, m.LINES_PER_SLIDE],
            p_elem=p_elem, keep=1, pil_font=font, box_px=220,
        )
        assert rest_p is not None
        combined = set(TestSplitParaPreservesVerseColors._orange_texts(p_elem)) | set(
            TestSplitParaPreservesVerseColors._orange_texts(rest_p)
        )
        assert combined == {"51", "52"}, f"조정 재시도 후 절 번호 오렌지색 유실: {combined}"


class TestComVerificationEnabledConfig:
    """config.json의 com_verification_enabled 플래그를 읽는
    _com_verification_enabled()가 true/false/키 없음(기본값 True)을 올바르게
    반영하는지 확인한다. PowerPoint COM 검증을 끄고 싶은 머신을 위한 opt-out이므로
    실패 시에도 항상 안전하게(Pillow 전용으로) 동작해야 한다."""

    def _reset_cache_and_point_config(self, monkeypatch, tmp_path, config_dict):
        config_path = tmp_path / "config.json"
        config_path.write_text(
            json.dumps(config_dict, ensure_ascii=False), encoding="utf-8"
        )
        monkeypatch.setattr(m, "CONFIG_FILE", config_path)
        m._COM_VERIFY_ENABLED_CACHE[0] = None

    def test_true_when_explicitly_enabled(self, monkeypatch, tmp_path):
        self._reset_cache_and_point_config(
            monkeypatch, tmp_path, {"com_verification_enabled": True}
        )
        assert m._com_verification_enabled() is True

    def test_false_when_explicitly_disabled(self, monkeypatch, tmp_path):
        self._reset_cache_and_point_config(
            monkeypatch, tmp_path, {"com_verification_enabled": False}
        )
        assert m._com_verification_enabled() is False

    def test_defaults_true_when_key_absent(self, monkeypatch, tmp_path):
        self._reset_cache_and_point_config(monkeypatch, tmp_path, {})
        assert m._com_verification_enabled() is True

    def test_result_is_memoized(self, monkeypatch, tmp_path):
        self._reset_cache_and_point_config(
            monkeypatch, tmp_path, {"com_verification_enabled": False}
        )
        assert m._com_verification_enabled() is False
        # 캐시된 이후에는 config.json 내용이 바뀌어도 재확인하지 않는다
        (tmp_path / "config.json").write_text(
            json.dumps({"com_verification_enabled": True}), encoding="utf-8"
        )
        assert m._com_verification_enabled() is False


class TestCountSlideLinesVerified:
    """_count_slide_lines_verified()의 제어 흐름(경계값 최적화, COM 값 채택,
    실패 시 영구 비활성화, config opt-out)을 실제 PowerPoint COM 없이
    monkeypatch로 결정적으로 검증한다."""

    class _FakeSlide:
        def __init__(self, slide_id):
            self.slide_id = slide_id

    @pytest.fixture(autouse=True)
    def _reset_state(self, monkeypatch):
        monkeypatch.setattr(m, "_COM_DISABLED", [False], raising=False)
        monkeypatch.setattr(m, "_COM_MISMATCH_COUNT", {}, raising=False)
        monkeypatch.setattr(m, "_COM_VERIFY_ENABLED_CACHE", [True], raising=False)
        monkeypatch.setattr(m, "_build_com_probe_pptx", lambda prs, slide: (Path("dummy.pptx"), 1), raising=False)
        monkeypatch.setattr(m, "_COM_ATEXIT_REGISTERED", [False], raising=False)

    @staticmethod
    def _install_fake_com(monkeypatch, real_lines=None, raises=False):
        fake = types.ModuleType("ppt_com_verify")

        class _FakeComUnavailable(Exception):
            pass

        fake.ComVerificationUnavailable = _FakeComUnavailable
        calls = []

        def _count_slide_lines(path, shape_index):
            calls.append((path, shape_index))
            if raises:
                raise _FakeComUnavailable("simulated COM failure")
            return real_lines

        fake.count_slide_lines = _count_slide_lines
        fake.shutdown = lambda: None
        monkeypatch.setitem(sys.modules, "ppt_com_verify", fake)
        return calls

    def test_com_not_called_when_pillow_estimate_is_not_boundary(self, monkeypatch):
        monkeypatch.setattr(m, "_count_slide_lines_rendered", lambda slide: 8)
        calls = self._install_fake_com(monkeypatch, real_lines=10)
        result = m._count_slide_lines_verified(object(), self._FakeSlide(1))
        assert result == 8
        assert calls == []

    def test_com_confirms_boundary_estimate(self, monkeypatch, capsys):
        monkeypatch.setattr(m, "_count_slide_lines_rendered", lambda slide: m.LINES_PER_SLIDE)
        calls = self._install_fake_com(monkeypatch, real_lines=m.LINES_PER_SLIDE)
        result = m._count_slide_lines_verified(object(), self._FakeSlide(2))
        assert result == m.LINES_PER_SLIDE
        assert len(calls) == 1
        assert "불일치" not in capsys.readouterr().out

    def test_com_mismatch_is_adopted_and_logged(self, monkeypatch, capsys):
        monkeypatch.setattr(m, "_count_slide_lines_rendered", lambda slide: m.LINES_PER_SLIDE)
        self._install_fake_com(monkeypatch, real_lines=m.LINES_PER_SLIDE + 1)
        result = m._count_slide_lines_verified(object(), self._FakeSlide(3))
        assert result == m.LINES_PER_SLIDE + 1
        assert "불일치" in capsys.readouterr().out
        assert m._COM_MISMATCH_COUNT[3] == 1

    def test_com_is_always_consulted_even_after_repeated_mismatches(self, monkeypatch):
        """2026-08-04 발견 버그의 회귀 방지: 예전에는 같은 슬라이드에 대해 불일치가
        누적되면(예전 캡=3회) 그 슬라이드에 한해 이후 영구히 COM을 건너뛰고 Pillow
        값을 실측인 것처럼 반환했다. 이로 인해 실제로는 여전히 오버플로인 슬라이드가
        "성공"으로 잘못 보고된 사례(제1독서 슬라이드 19, 실제 10줄인데 9줄로 보고)가
        실측으로 확인됐다. 이제는 슬라이드별 이력과 무관하게 경계값(9)일 때마다
        매번 실제로 COM에 묻어야 한다."""
        monkeypatch.setattr(m, "_count_slide_lines_rendered", lambda slide: m.LINES_PER_SLIDE)
        calls = self._install_fake_com(monkeypatch, real_lines=m.LINES_PER_SLIDE + 1)
        slide = self._FakeSlide(4)
        for _ in range(5):
            result = m._count_slide_lines_verified(object(), slide)
            assert result == m.LINES_PER_SLIDE + 1  # 매번 실제 COM 값을 그대로 반환
        assert len(calls) == 5  # 호출 이력과 무관하게 COM이 매번 실제로 호출됨

    def test_com_failure_permanently_disables_for_rest_of_run(self, monkeypatch, capsys):
        monkeypatch.setattr(m, "_count_slide_lines_rendered", lambda slide: m.LINES_PER_SLIDE)
        calls = self._install_fake_com(monkeypatch, raises=True)
        result_1 = m._count_slide_lines_verified(object(), self._FakeSlide(5))
        assert result_1 == m.LINES_PER_SLIDE
        assert m._COM_DISABLED[0] is True
        assert "[경고]" in capsys.readouterr().out

        # 이후 같은 프로세스 내에서는(경계값이어도) COM을 다시 시도하지 않는다
        result_2 = m._count_slide_lines_verified(object(), self._FakeSlide(6))
        assert result_2 == m.LINES_PER_SLIDE
        assert len(calls) == 1

    def test_config_disabled_skips_com_entirely(self, monkeypatch):
        monkeypatch.setattr(m, "_count_slide_lines_rendered", lambda slide: m.LINES_PER_SLIDE)
        monkeypatch.setattr(m, "_COM_VERIFY_ENABLED_CACHE", [False])
        calls = self._install_fake_com(monkeypatch, real_lines=m.LINES_PER_SLIDE + 1)
        result = m._count_slide_lines_verified(object(), self._FakeSlide(7))
        assert result == m.LINES_PER_SLIDE
        assert calls == []

    def test_probe_build_failure_falls_back_without_disabling_com_globally(self, monkeypatch, capsys):
        """코드 리뷰 발견: _build_com_probe_pptx()는 python-pptx만 쓰는 순수
        파이썬 코드라 ComVerificationUnavailable이 아닌 예외(예: content shape가
        없어 ValueError)를 낼 수 있다. 예전에는 이 예외가 try 블록을 빠져나가
        함수 자체가 통째로 실패했다 — "어떤 실패 경로도 예외를 밖으로 내보내지
        않는다"는 문서화된 계약을 어겼다. 이제는 이 슬라이드만 Pillow로
        폴백하고, COM 자체는 다른 슬라이드에 대해 계속 사용 가능해야 한다
        (probe 생성 실패는 COM 전체의 문제가 아니라 그 슬라이드만의 문제)."""
        monkeypatch.setattr(m, "_count_slide_lines_rendered", lambda slide: m.LINES_PER_SLIDE)

        def _raise_probe_build(prs, slide):
            raise ValueError("simulated: content shape 없음")

        monkeypatch.setattr(m, "_build_com_probe_pptx", _raise_probe_build)
        calls = self._install_fake_com(monkeypatch, real_lines=m.LINES_PER_SLIDE)

        result = m._count_slide_lines_verified(object(), self._FakeSlide(8))

        assert result == m.LINES_PER_SLIDE  # 예외 없이 Pillow 값으로 폴백
        assert calls == []  # COM 자체는 호출되지도 않음(probe 생성 단계에서 실패)
        assert m._COM_DISABLED[0] is False  # COM을 전역적으로 비활성화하지 않음
        assert "[경고]" in capsys.readouterr().out


# ─────────────────────────────────────────────────────────────────
# 2. 통합 테스트 — 실제 PPT 생성 + 구조 검증
# ─────────────────────────────────────────────────────────────────

def _find_output(date_str: str) -> Path:
    folder = REPO_ROOT / date_str
    candidates = [
        f for f in folder.glob(f"{date_str}_*.pptx")
        if not f.name.startswith("~$")
    ]
    assert candidates, f"{date_str} 출력 PPT를 찾을 수 없음"
    return max(candidates, key=lambda f: f.stat().st_mtime)


def _generate(case: dict) -> str:
    date_str = case["date"]
    folder = REPO_ROOT / date_str

    locked = list(folder.glob("~$*.pptx"))
    if locked:
        pytest.skip(f"{locked[0].name} 이(가) PowerPoint에서 열려 있어 생성을 건너뜀")

    args = [sys.executable, str(REPO_ROOT / "missa_to_ppt.py"), date_str]
    for k, v in case["hymns"].items():
        args += [f"--{k}", v]

    result = subprocess.run(
        args, cwd=REPO_ROOT, capture_output=True, text=True, encoding="utf-8",
    )
    assert result.returncode == 0, (
        f"{date_str} 생성 실패 (exit={result.returncode})\n"
        f"stdout:\n{result.stdout}\nstderr:\n{result.stderr}"
    )
    assert "모든 검증 통과!" in result.stdout, (
        f"{date_str} 내부 검증(validate) 실패\nstdout:\n{result.stdout}"
    )
    return result.stdout


@pytest.fixture(scope="module", params=MASS_CASES, ids=lambda c: c["id"])
def generated_case(request):
    case = request.param
    log = _generate(case)
    path = _find_output(case["date"])
    prs = Presentation(str(path))
    sections = m.find_sections(prs)
    json_data = m.get_json_data(case["date"])
    return {
        "case": case, "prs": prs, "sections": sections, "log": log,
        "path": path, "json_data": json_data,
    }


def test_output_opens_without_corruption(generated_case):
    # Presentation()이 예외 없이 로드되면 OOXML 구조 자체는 유효함(fixture에서 이미 검증됨).
    assert len(generated_case["prs"].slides) > 0


def test_liturgy_type_detected_correctly(generated_case):
    log = generated_case["log"]
    case = generated_case["case"]
    expected = "미사 유형: 주일미사" if case["is_sunday"] else "미사 유형: 평일미사"
    assert expected in log


def test_required_sections_present(generated_case):
    sections = generated_case["sections"]
    required = [
        "입당송", "제1독서_start", "제1독서_end",
        "화답송_start", "화답송_end",
        "복음환호송", "복음_start", "복음_end",
        "영성체송",
    ]
    missing = [k for k in required if k not in sections]
    assert not missing, f"누락된 섹션: {missing}"


class TestBuildComProbePptx:
    """_build_com_probe_pptx()가 실제 독서 슬라이드를 손실 없이 임시 단일 슬라이드
    pptx로 복사하는지 확인한다 (PowerPoint COM 실측을 위한 probe 생성 단계)."""

    def test_probe_slide_dimensions_match_original(self, generated_case):
        prs, sections = generated_case["prs"], generated_case["sections"]
        slide = prs.slides[sections["제1독서_start"]]
        probe_path, _ = m._build_com_probe_pptx(prs, slide)
        probe = Presentation(str(probe_path))
        assert probe.slide_width == prs.slide_width
        assert probe.slide_height == prs.slide_height

    def test_probe_content_text_matches_original(self, generated_case):
        prs, sections = generated_case["prs"], generated_case["sections"]
        slide = prs.slides[sections["제1독서_start"]]
        original_text = m._find_content_shape(slide).text_frame.text
        probe_path, _ = m._build_com_probe_pptx(prs, slide)
        probe = Presentation(str(probe_path))
        probe_text = m._find_content_shape(probe.slides[0]).text_frame.text
        assert probe_text == original_text

    def test_probe_shape_index_is_1_based_and_points_to_content_shape(self, generated_case):
        prs, sections = generated_case["prs"], generated_case["sections"]
        slide = prs.slides[sections["제1독서_start"]]
        original_text = m._find_content_shape(slide).text_frame.text
        probe_path, shape_index = m._build_com_probe_pptx(prs, slide)
        probe = Presentation(str(probe_path))
        shapes = list(probe.slides[0].shapes)
        assert 1 <= shape_index <= len(shapes)
        pointed_shape = shapes[shape_index - 1]  # COM Shapes()는 1-based
        assert pointed_shape.has_text_frame
        assert pointed_shape.text_frame.text == original_text

    def test_probe_path_reused_across_calls(self, generated_case):
        prs, sections = generated_case["prs"], generated_case["sections"]
        slide_a = prs.slides[sections["제1독서_start"]]
        slide_b = prs.slides[sections["복음_start"]]
        path_a, _ = m._build_com_probe_pptx(prs, slide_a)
        path_b, _ = m._build_com_probe_pptx(prs, slide_b)
        assert path_a == path_b, "프로세스당 하나의 임시 경로를 재사용해야 함(누적 생성 금지)"
        # 두 번째 호출 후 파일 내용은 slide_b 기준으로 덮어써져 있어야 한다
        probe = Presentation(str(path_b))
        probe_text = m._find_content_shape(probe.slides[0]).text_frame.text
        assert probe_text == m._find_content_shape(slide_b).text_frame.text


def test_no_reading_slide_line_overflow(generated_case):
    """독서·복음 본문 슬라이드가 LINES_PER_SLIDE(9줄)를 넘으면 안 된다.
    2026-07-16 발견된 post-write 재조정 실패(연쇄 초과) 회귀 방지용."""
    prs, sections = generated_case["prs"], generated_case["sections"]
    problems = []
    for start_key, end_key in READING_SECTIONS:
        if start_key not in sections or end_key not in sections:
            continue
        s, e = sections[start_key], sections[end_key]
        for idx in range(s, e):
            lines = m._count_slide_lines_rendered(prs.slides[idx])
            if lines > m.LINES_PER_SLIDE:
                problems.append(f"{start_key} 슬라이드 {idx + 1}: {lines}줄 (>{m.LINES_PER_SLIDE})")
    assert not problems, "\n".join(problems)


def test_no_missing_orange_verse_numbers(generated_case):
    """독서·복음의 모든 절 번호가 실제로 오렌지색 run으로 렌더링됐는지 확인한다.
    텍스트 자체는 남아 있어도 색상만 사라지는 경우(2026-07-26)가 있어, 슬라이드
    텍스트 존재 여부가 아니라 절 번호 단위로 오렌지색 run을 직접 대조한다."""
    prs, sections, json_data = (
        generated_case["prs"], generated_case["sections"], generated_case["json_data"]
    )
    problems = []
    for label, start_key, end_key in [
        ("제1독서", "제1독서_start", "제1독서_end"),
        ("제2독서", "제2독서_start", "제2독서_end"),
        ("복음", "복음_start", "복음_end"),
    ]:
        reading = json_data.get(label)
        if not reading or not reading.get("content"):
            continue
        if start_key not in sections or end_key not in sections:
            continue
        missing = m._missing_orange_verse_numbers(
            prs, sections[start_key], sections[end_key], reading["content"]
        )
        if missing:
            problems.append(f"{label} 절 번호 오렌지색 누락: {', '.join(missing)}")
    assert not problems, "\n".join(problems)


def test_reading_verse_numbers_appear_in_ascending_order(generated_case):
    """2026-08-04 발견 회귀 방지: _rebalance_reading_slides_post_write()의
    COM 조정 give-up 롤백 경로 중 일부가 되돌릴 위치를 잘못 계산해(맨 앞이
    아니라 append) 다음 슬라이드에 이미 남아 있던 뒤쪽 단락보다 앞에 있어야 할
    단락이 뒤로 밀려 절 순서가 뒤바뀌는 버그가 있었다. 오렌지색 절 번호가
    슬라이드/단락 순서대로 오름차순으로 나타나는지 확인해 이런 단락 순서
    뒤바뀜을 감지한다(존재 여부만 확인하는 test_no_missing_orange_verse_numbers
    와 달리 순서 자체를 확인)."""
    prs, sections, json_data = (
        generated_case["prs"], generated_case["sections"], generated_case["json_data"]
    )
    problems = []
    for label, start_key, end_key in [
        ("제1독서", "제1독서_start", "제1독서_end"),
        ("제2독서", "제2독서_start", "제2독서_end"),
        ("복음", "복음_start", "복음_end"),
    ]:
        reading = json_data.get(label)
        if not reading or not reading.get("content"):
            continue
        if start_key not in sections or end_key not in sections:
            continue
        s, e = sections[start_key], sections[end_key]
        seq = []
        for idx in range(s, e):
            shape = m._find_content_shape(prs.slides[idx])
            if shape is None:
                continue
            for para in shape.text_frame.paragraphs:
                for run in para.runs:
                    try:
                        if run.font.color.rgb != m.ORANGE:
                            continue
                    except Exception:
                        continue
                    match = re.match(r'^(\d+)', run.text.strip())
                    if match:
                        seq.append(int(match.group(1)))
        for a, b in zip(seq, seq[1:]):
            if b < a:
                problems.append(f"{label}: 절 번호 순서 역전 {a} → {b} (전체 순서: {seq})")
    assert not problems, "\n".join(problems)


def test_blackbg_covers_full_slide(generated_case):
    """BlackBg 도형이 존재하면 항상 슬라이드 전체 크기를 덮어야 한다.
    2026-07-16 발견된 _align_ending_slides_to_제2독서의 BlackBg 오분류 회귀 방지용."""
    prs = generated_case["prs"]
    problems = []
    for i, slide in enumerate(prs.slides):
        for shape in slide.shapes:
            if shape.name == "BlackBg":
                if not (
                    shape.left == 0 and shape.top == 0
                    and shape.width == prs.slide_width
                    and shape.height == prs.slide_height
                ):
                    problems.append(
                        f"슬라이드 {i + 1}: BlackBg 크기 이상 "
                        f"({shape.left},{shape.top},{shape.width},{shape.height})"
                    )
    assert not problems, "\n".join(problems)


def test_화답송_pattern_matches_mass_type(generated_case):
    """주일: 악보/텍스트 슬라이드 교대(짝수 인덱스=악보). 평일: 텍스트만, 악보 이미지 없음."""
    prs, sections, case = (
        generated_case["prs"], generated_case["sections"], generated_case["case"]
    )
    if "화답송_start" not in sections:
        pytest.skip("화답송 섹션 없음")
    s, e = sections["화답송_start"], sections["화답송_end"]
    problems = []
    for i in range(s, e):
        slide = prs.slides[i]
        has_picture = any(sh.shape_type == MSO_SHAPE_TYPE.PICTURE for sh in slide.shapes)
        rel_i = i - s
        if case["is_sunday"]:
            expected_picture = (rel_i % 2 == 0)
            if has_picture != expected_picture:
                problems.append(
                    f"화답송 슬라이드 {i + 1}: 악보 유무 불일치 "
                    f"(기대={expected_picture}, 실제={has_picture})"
                )
        else:
            if has_picture:
                problems.append(f"화답송 슬라이드 {i + 1}: 평일미사인데 악보 이미지 존재")
    assert not problems, "\n".join(problems)


def test_hymn_score_copy_matches_mass_type(generated_case):
    """주일: 성가 악보 슬라이드 복사됨. 평일: 번호만 갱신, 악보 슬라이드 없음."""
    sections, case = generated_case["sections"], generated_case["case"]
    problems = []
    for hymn in HYMN_TYPES:
        s_key, e_key = f"{hymn}_content_start", f"{hymn}_content_end"
        if s_key not in sections or e_key not in sections:
            continue
        n_content = sections[e_key] - sections[s_key]
        if case["is_sunday"]:
            if n_content == 0:
                problems.append(f"{hymn}: 주일미사인데 악보 슬라이드 없음")
        else:
            if n_content != 0:
                problems.append(f"{hymn}: 평일미사인데 악보 슬라이드 {n_content}장 존재")
    assert not problems, "\n".join(problems)


# ─────────────────────────────────────────────────────────────────
# 3. PowerPoint COM 실측 통합 테스트 (환경 조건부 skip)
# ─────────────────────────────────────────────────────────────────

def _com_available() -> bool:
    try:
        import ppt_com_verify as com
    except ImportError:
        return False
    return com.is_available()


class TestComVerify:
    """PowerPoint COM 실측 경로 전용 테스트. pywin32 미설치 또는 PowerPoint COM
    연결 불가 환경에서는 skip한다(이 저장소의 기존 pytest.skip 전례를 따름)."""

    def test_com_available_or_skips(self):
        if not _com_available():
            pytest.skip("pywin32 미설치 또는 PowerPoint COM 연결 불가")
        assert True

    def test_com_verification_resolves_known_overflow_case(self):
        """2026-06-24 케이스는 과거 Pillow 추정으로는 9줄이지만 실제 PowerPoint
        렌더링으로는 10줄인 슬라이드가 존재했던 회귀였다(제1독서 슬라이드 19,
        복음 슬라이드 38). com_verification_enabled가 기본 활성화된 상태로
        재생성한 뒤, Pillow가 아니라 실제 COM 실측으로 모든 독서/복음 슬라이드가
        9줄 이하인지 확인한다 — 원래 버그를 실제로 잡아낼 수 있는 유일한 테스트."""
        if not _com_available():
            pytest.skip("pywin32 미설치 또는 PowerPoint COM 연결 불가")
        import ppt_com_verify as com

        case = next(c for c in MASS_CASES if c["date"] == "20260624")
        _generate(case)
        path = _find_output(case["date"])
        prs = Presentation(str(path))
        sections = m.find_sections(prs)

        problems = []
        for start_key, end_key in READING_SECTIONS:
            if start_key not in sections or end_key not in sections:
                continue
            s, e = sections[start_key], sections[end_key]
            for idx in range(s, e):
                slide = prs.slides[idx]
                if m._find_content_shape(slide) is None:
                    continue  # 종료 텍스트 병합 슬라이드 등 본문 텍스트박스가 없는 슬라이드는 대상 제외
                probe_path, shape_idx = m._build_com_probe_pptx(prs, slide)
                real_lines = com.count_slide_lines(str(probe_path), shape_idx)
                if real_lines > m.LINES_PER_SLIDE:
                    problems.append(
                        f"{start_key} 슬라이드 {idx + 1}: COM 실측 {real_lines}줄 (>{m.LINES_PER_SLIDE})"
                    )
        assert not problems, "\n".join(problems)
