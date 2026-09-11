# -*- coding: utf-8 -*-
"""청년미사 성가 슬라이드 템플릿 자산 생성(개발 1회성).

샘플 PPT(reference/청년미사/나주노 성가 447 ….pptx)에서 슬라이드 1장만 남기고, 헤더
'Rectangle 11'의 run을 [흰색 라벨 run][FFC000 run] 2개로 줄이고(나머지 run 제거·endParaRPr
보존), 저작권/악보 그림(Picture 2/3)을 제거해 assets/청년미사_성가_template.pptx로 저장한다.

이 템플릿을 duplicate_slide로 복제해 build_hymn_pptx가 헤더 텍스트·악보·저작권을 채운다.
헤더 rPr(28pt/spc300/라벨 prstClr white·본문 srgbClr FFC000/굴서 폰트)을 실측 XML 그대로
보존하는 것이 목적이라, 프로그램 생성이 아니라 실측 템플릿 복제 방식을 택했다(화답송 모듈과
동일 판단, 설계서 §3.6)."""
from pathlib import Path

from pptx import Presentation
from pptx.oxml.ns import qn

BASE = Path(__file__).resolve().parent.parent
SAMPLE = BASE / "reference" / "청년미사" / "나주노 성가 447 영원도하시어라 그 사랑이여.pptx"
OUT = BASE / "assets" / "청년미사_성가_template.pptx"


def _trim_header_runs(slide):
    for sh in slide.shapes:
        if sh.name != "Rectangle 11":
            continue
        p = sh.text_frame.paragraphs[0]._p
        runs = p.findall(qn("a:r"))
        # run0(흰색 라벨), run1(FFC000)만 남기고 나머지 제거. endParaRPr는 보존.
        for r in runs[2:]:
            p.remove(r)
        return
    raise RuntimeError("Rectangle 11 헤더를 찾지 못함")


def _remove_pictures(slide):
    for sh in list(slide.shapes):
        if sh.shape_type == 13:  # PICTURE
            sh._element.getparent().remove(sh._element)


def _keep_only_first_slide(prs):
    sldIdLst = prs.slides._sldIdLst
    for sldId in list(sldIdLst)[1:]:
        rId = sldId.get(qn("r:id"))
        prs.part.drop_rel(rId)
        sldIdLst.remove(sldId)


def main():
    prs = Presentation(str(SAMPLE))
    _keep_only_first_slide(prs)
    slide = prs.slides[0]
    _remove_pictures(slide)
    _trim_header_runs(slide)
    OUT.parent.mkdir(exist_ok=True)
    prs.save(str(OUT))
    print("저장:", OUT)
    # 검증
    chk = Presentation(str(OUT))
    assert len(chk.slides) == 1
    hdr = [s for s in chk.slides[0].shapes if s.name == "Rectangle 11"][0]
    runs = hdr.text_frame.paragraphs[0].runs
    print("헤더 run 수:", len(runs), "텍스트:", [r.text for r in runs])
    print("그림 수:", sum(1 for s in chk.slides[0].shapes if s.shape_type == 13))


if __name__ == "__main__":
    main()
