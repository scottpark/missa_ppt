"""화답송 악보 슬라이드용 고정 템플릿 자산 빌더 (일회성).

베이스: reference/화답송악보/20260816/시편 67(66).pptx (설계서 §5.1 근거 — 상/하 악보가
슬라이드 정상 범위 안이라 제거 후 잔여 XML이 깔끔).

이 스크립트는 자산을 만들어 커밋하기 위한 일회성 도구다. 런타임 코드
(missa_psalm_score_image.py)는 이 스크립트를 호출하지 않고 완성된 .pptx만 로드한다.

악보 PICTURE 식별은 텍스트가 아니라 크기/타입 기준(shape_type==PICTURE and
width>1_000_000 and height>1_000_000)으로 한다 — 잔재 4개(12×7·41×17 EMU)와
제목/라벨/저작권(TEXT_BOX/AUTO_SHAPE)은 자동 배제된다(설계서 §5.2, §8, CLAUDE.md
"도형 분류 시 빈 텍스트 falsy" 함정 회피: 텍스트 키워드로 분류하지 않음).
"""
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

BASE = Path(__file__).resolve().parent.parent
SOURCE = BASE / "reference" / "화답송악보" / "20260816" / "시편 67(66).pptx"
DEST = BASE / "assets" / "화답송_악보_template.pptx"

# 악보 그림으로 간주할 최소 크기(EMU). 잔재 도형(12×7·41×17)은 이 아래라 자동 배제.
_SCORE_MIN_EMU = 1_000_000


def build() -> Path:
    prs = Presentation(str(SOURCE))
    slide = prs.slides[0]
    removed = 0
    for shape in list(slide.shapes):
        if (
            shape.shape_type == MSO_SHAPE_TYPE.PICTURE
            and shape.width > _SCORE_MIN_EMU
            and shape.height > _SCORE_MIN_EMU
        ):
            sp = shape._element
            sp.getparent().remove(sp)
            removed += 1
    if removed != 2:
        raise RuntimeError(
            f"악보 PICTURE 2개를 제거해야 하는데 {removed}개 제거됨 — 베이스 PPT 구조 변경 의심"
        )
    DEST.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(DEST))
    return DEST


if __name__ == "__main__":
    out = build()
    print(f"템플릿 생성 완료: {out}")
